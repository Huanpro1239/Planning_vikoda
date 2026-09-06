import hashlib
import json
import math
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from planning_resource_timeline import ensure_resource_timelines


REPORT_PATH = Path("planning_schedule_report.json")
REPORT_SCHEMA_VERSION = 3
PLANNING_SHEET = "Ke_hoach_SX"
MASS_BALANCE_EPS = 1e-5


def _json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(key): _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    return value


def json_safe(value):
    """Return a recursively JSON-serializable copy of report metadata."""
    return _json_value(value)


def _read_uom_by_code(workbook_bytes):
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        if PLANNING_SHEET not in workbook.sheetnames:
            return {}
        sheet = workbook[PLANNING_SHEET]
        result = {}
        for row in range(2, sheet.max_row + 1):
            raw_code = sheet.cell(row=row, column=1).value
            if raw_code in (None, ""):
                continue
            if isinstance(raw_code, float) and raw_code.is_integer():
                raw_code = int(raw_code)
            code = str(raw_code).strip()
            result[code] = str(sheet.cell(row=row, column=3).value or "").strip()
        return result
    finally:
        workbook.close()


def _validate_mass_balance_values(code, planned, scheduled, carryover):
    values = {
        "planned_qty": float(planned),
        "scheduled_qty": float(scheduled),
        "carryover_qty": float(carryover),
    }
    for label, value in values.items():
        if not math.isfinite(value):
            raise RuntimeError(
                f"Schedule report mã {code} có {label} không hữu hạn: {value!r}."
            )
        if value < -MASS_BALANCE_EPS:
            raise RuntimeError(
                f"Schedule report mã {code} có {label} âm: {value}."
            )

    planned = values["planned_qty"]
    scheduled = values["scheduled_qty"]
    carryover = values["carryover_qty"]
    if scheduled > planned + MASS_BALANCE_EPS:
        raise RuntimeError(
            f"Schedule report mã {code} scheduled={scheduled} vượt P={planned}."
        )
    if carryover > planned + MASS_BALANCE_EPS:
        raise RuntimeError(
            f"Schedule report mã {code} carryover={carryover} vượt P={planned}."
        )
    if not math.isclose(
        scheduled + carryover,
        planned,
        rel_tol=1e-9,
        abs_tol=MASS_BALANCE_EPS,
    ):
        raise RuntimeError(
            f"Schedule report mã {code} mất cân đối: scheduled {scheduled} + "
            f"carryover {carryover} != P {planned}."
        )
    return planned, scheduled, carryover


def _enrich_resource_meta(meta, uom_by_code):
    """Carry deadline provenance from production timeline into late records."""
    meta = dict(meta or {})
    timeline_by_unit = {}
    for raw in meta.get("timeline") or []:
        if not isinstance(raw, dict) or raw.get("kind") != "production":
            continue
        code = str(raw.get("code") or "")
        unit_no = raw.get("unit_no")
        if code and unit_no is not None:
            timeline_by_unit[(code, str(unit_no))] = raw

    for field in (
        "unserved_due",
        "late_completed",
        "deadline_misses",
        "shortage_by_day",
        "timeline",
    ):
        records = meta.get(field)
        if not isinstance(records, list):
            continue
        enriched = []
        for record in records:
            if not isinstance(record, dict):
                enriched.append(record)
                continue
            item = dict(record)
            code = str(item.get("code") or "")
            if code and not item.get("uom"):
                item["uom"] = uom_by_code.get(code, "")

            if field in {"late_completed", "deadline_misses"}:
                source = timeline_by_unit.get((code, str(item.get("unit_no"))))
                if source:
                    for key in (
                        "demand_deadline_shift",
                        "production_deadline_shift",
                        "demand_due_date",
                        "production_deadline_date",
                    ):
                        if item.get(key) is None and source.get(key) is not None:
                            item[key] = source[key]
            enriched.append(item)
        meta[field] = enriched
    return meta


def _timeline_utilization(timeline, capacity, day_count):
    if capacity <= 0 or day_count <= 0:
        return 0.0
    total = 0.0
    for event in timeline or []:
        if not isinstance(event, dict):
            continue
        try:
            start = float(event.get("start_shift"))
            end = float(event.get("end_shift"))
        except (TypeError, ValueError):
            continue
        if math.isfinite(start) and math.isfinite(end) and end > start:
            total += end - start
    return total / (float(capacity) * day_count)


def _build_status(carryover, inventory, resource_proofs):
    carryover_codes = sorted(
        code for code, qty in carryover.items() if float(qty or 0) > MASS_BALANCE_EPS
    )
    stockout_codes = sorted(
        code
        for code, values in inventory.items()
        if isinstance(values, dict) and values.get("first_stockout") is not None
    )
    resource_failures = sorted(
        resource
        for resource, proof in resource_proofs.items()
        if not bool((proof.get("validation") or {}).get("validated"))
    )

    monthly_ok = not carryover_codes
    resource_ok = not resource_failures
    service_ok = not stockout_codes
    return {
        "monthly_quantity": {
            "ok": monthly_ok,
            "state": "complete" if monthly_ok else "carryover",
            "carryover_skus": carryover_codes,
        },
        "resource_validation": {
            "ok": resource_ok,
            "state": "passed" if resource_ok else "failed",
            "validated_resources": sorted(resource_proofs),
            "failed_resources": resource_failures,
        },
        "service": {
            "ok": service_ok,
            "state": "no_stockout" if service_ok else "stockout_risk",
            "stockout_skus": stockout_codes,
        },
    }


def build_schedule_report(
    workbook_bytes,
    info,
    *,
    input_revision=None,
    algorithm="priority_v8",
):
    """Build an operational report with mass-balance and physical-resource proof."""
    headers = list(info["headers"])
    products = list(info["products"])
    schedule = info["schedule"]
    carryover = dict(info.get("carryover") or {})
    utilization = dict(info.get("utilization") or {})
    uom_by_code = _read_uom_by_code(workbook_bytes)

    mass_balance = {}
    for product in products:
        code = product["code"]
        planned = float(product.get("planned_qty", 0) or 0)
        scheduled = sum(float(schedule[code].get(day, 0) or 0) for day in headers)
        missing = float(carryover.get(code, 0) or 0)
        planned, scheduled, missing = _validate_mass_balance_values(
            code,
            planned,
            scheduled,
            missing,
        )
        mass_balance[code] = {
            "uom": str(product.get("uom") or uom_by_code.get(code, "")),
            "planned_qty": planned,
            "scheduled_qty": scheduled,
            "carryover_qty": missing,
            "balanced_qty": scheduled + missing,
        }

    # Independent proof for every serial resource. KHS/PET re-validates the
    # scheduler-native timeline. RGB/Galon build a physically feasible daily
    # order from the exact workbook quantities and reject the report if setup
    # cannot fit capacity.
    resource_proofs = ensure_resource_timelines(info)

    resources = {}
    for resource, capacity in sorted((info.get("line_capacity") or {}).items()):
        raw_meta = dict((info.get("optimizer_meta") or {}).get(resource, {}) or {})
        proof = resource_proofs.get(resource)
        if proof:
            raw_meta["timeline"] = proof.get("timeline") or []
            raw_meta["scheduler_setup_shifts"] = proof.get("scheduler_setup_shifts", 0.0)
            raw_meta["setup_shifts"] = proof.get("setup_shifts", 0.0)
            raw_meta["setup_delta_shifts"] = proof.get("setup_delta_shifts", 0.0)
            raw_meta["resource_validation"] = proof.get("validation") or {}

        meta = _enrich_resource_meta(raw_meta, uom_by_code)
        timeline = meta.get("timeline") or []
        reported_utilization = float(utilization.get(resource, 0) or 0)
        if proof:
            reported_utilization = _timeline_utilization(
                timeline,
                float(capacity),
                len(headers),
            )
        resources[resource] = {
            "capacity_shifts_per_day": float(capacity),
            "utilization": reported_utilization,
            "meta": meta,
        }

    inventory = {}
    for code, values in (info.get("inventory") or {}).items():
        inventory[code] = dict(values)

    status = _build_status(carryover, inventory, resource_proofs)
    ready = all(section.get("ok") for section in status.values())
    report = {
        "schema_version": REPORT_SCHEMA_VERSION,
        "algorithm": algorithm,
        "plan_month": headers[0].strftime("%Y-%m") if headers else None,
        "input_revision": dict(input_revision or {}),
        "input_sha256": hashlib.sha256(workbook_bytes).hexdigest(),
        "mass_balance": mass_balance,
        "resources": resources,
        "inventory": inventory,
        "status": status,
        # Deliberately no longer means "P fits in month == feasible".
        # Any stockout/service issue keeps the plan in review_required.
        "publish_status": "ready_for_publish" if ready else "review_required",
    }
    return json_safe(report)


def attach_output_hash(report, workbook_bytes):
    result = dict(report)
    result["output_sha256"] = hashlib.sha256(workbook_bytes).hexdigest()
    return result


def save_schedule_report(report, path=REPORT_PATH):
    path = Path(path)
    path.write_text(
        json.dumps(json_safe(report), ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return path


def load_schedule_report(path=REPORT_PATH):
    path = Path(path)
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def print_operational_report(report):
    print(
        f"[PLAN_REPORT] algorithm={report.get('algorithm')}; "
        f"plan={report.get('plan_month')}; status={report.get('publish_status')}."
    )
    split_status = report.get("status") or {}
    if split_status:
        print(
            "[PLAN_REPORT][STATUS] "
            f"monthly={((split_status.get('monthly_quantity') or {}).get('state'))}; "
            f"resource={((split_status.get('resource_validation') or {}).get('state'))}; "
            f"service={((split_status.get('service') or {}).get('state'))}."
        )

    carryovers = [
        (code, values)
        for code, values in sorted((report.get("mass_balance") or {}).items())
        if float(values.get("carryover_qty", 0) or 0) > 1e-9
    ]
    if carryovers:
        print("[PLAN_REPORT][CARRYOVER] Sản lượng hợp lệ nhưng chưa xếp được:")
        for code, values in carryovers:
            print(
                f"  - {code}: {values['carryover_qty']:g} {values.get('uom') or ''}; "
                f"scheduled={values['scheduled_qty']:g}; P={values['planned_qty']:g}"
            )

    for resource, resource_info in sorted((report.get("resources") or {}).items()):
        meta = resource_info.get("meta") or {}
        for event in meta.get("unserved_due") or []:
            print(
                f"[PLAN_REPORT][UNSERVED_DUE] {resource} {event.get('code')}: "
                f"{event.get('qty')} {event.get('uom') or ''}; "
                f"need={event.get('demand_due_date')}; "
                f"production_deadline={event.get('production_deadline_date')}"
            )
        for event in meta.get("late_completed") or []:
            print(
                f"[PLAN_REPORT][LATE_COMPLETED] {resource} {event.get('code')}: "
                f"{event.get('qty')} {event.get('uom') or ''}; "
                f"need={event.get('demand_due_date')}; "
                f"production_deadline={event.get('production_deadline_date')}"
            )
