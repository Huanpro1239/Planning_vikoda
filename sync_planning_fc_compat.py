from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import sync_planning_fc as fc


def read_planning_fc_targets_robust(workbook_bytes):
    values_workbook = load_workbook(
        BytesIO(workbook_bytes),
        data_only=True,
        read_only=True,
    )

    try:
        for sheet_name in (fc.FC_SHEET, fc.PLANNING_SHEET):
            if sheet_name not in values_workbook.sheetnames:
                raise RuntimeError(
                    f"Không tìm thấy sheet {sheet_name!r} trong Sắp kế hoạch.xlsx."
                )

        fc_sheet = values_workbook[fc.FC_SHEET]
        selector = fc_sheet[fc.FC_SELECTOR_CELL].value
        selector_key = fc._normalize_header(selector)
        if not selector_key:
            raise RuntimeError(f"{fc.FC_SHEET}!{fc.FC_SELECTOR_CELL} đang trống.")

        header = next(
            fc_sheet.iter_rows(
                min_row=fc.FC_HEADER_ROW,
                max_row=fc.FC_HEADER_ROW,
                min_col=fc.FC_HEADER_MIN_COL,
                max_col=fc.FC_HEADER_MAX_COL,
                values_only=True,
            ),
            (),
        )
        matched_columns = [
            fc.FC_HEADER_MIN_COL + offset
            for offset, value in enumerate(header)
            if fc._normalize_header(value) == selector_key
        ]
        if not matched_columns:
            raise RuntimeError(
                f"{fc.FC_SHEET}!{fc.FC_SELECTOR_CELL}={selector!r} không khớp "
                f"tiêu đề nào trong {fc.FC_SHEET}!E1:P1."
            )
        if len(matched_columns) > 1:
            letters = ", ".join(get_column_letter(c) for c in matched_columns)
            raise RuntimeError(
                f"Tiêu đề {selector!r} bị lặp trong {fc.FC_SHEET}!E1:P1 "
                f"tại các cột {letters}."
            )

        source_column = matched_columns[0]
        source_column_letter = get_column_letter(source_column)
        fc_values = {}

        for row_number, values in enumerate(
            fc_sheet.iter_rows(
                min_row=2,
                max_col=max(fc.FC_CODE_COL, source_column),
                values_only=True,
            ),
            start=2,
        ):
            code = fc.normalize_code(values[fc.FC_CODE_COL - 1])
            if not code:
                continue
            if code in fc_values:
                raise RuntimeError(f"Mã {code} bị lặp trong {fc.FC_SHEET}!B.")
            fc_values[code] = fc._forecast_value(
                values[source_column - 1],
                f"{fc.FC_SHEET}!{source_column_letter}{row_number}",
            )

        if not fc_values:
            raise RuntimeError(f"Không đọc được mã sản phẩm trong {fc.FC_SHEET}!B.")

        planning = values_workbook[fc.PLANNING_SHEET]
        targets = {}
        seen_codes = set()
        missing_codes = []
        changed_count = 0

        for row_number, values in enumerate(
            planning.iter_rows(
                min_row=2,
                max_col=fc.PLANNING_TARGET_COL,
                values_only=True,
            ),
            start=2,
        ):
            code = fc.normalize_code(values[fc.PLANNING_CODE_COL - 1])
            if not code:
                continue
            if code in seen_codes:
                raise RuntimeError(
                    f"Mã {code} bị lặp trong {fc.PLANNING_SHEET}!A."
                )
            seen_codes.add(code)
            if code not in fc_values:
                missing_codes.append(code)
                continue

            target_value = fc_values[code]
            targets[code] = target_value
            current_value = values[fc.PLANNING_TARGET_COL - 1]
            if not fc._values_equal(current_value, target_value):
                changed_count += 1

        if missing_codes:
            preview = ", ".join(missing_codes[:10])
            suffix = "..." if len(missing_codes) > 10 else ""
            raise RuntimeError(
                f"Có {len(missing_codes)} mã trong {fc.PLANNING_SHEET}!A "
                f"không tồn tại trong {fc.FC_SHEET}!B: {preview}{suffix}"
            )
        if not targets:
            raise RuntimeError(
                f"Không đọc được mã sản phẩm trong {fc.PLANNING_SHEET}!A."
            )

        print(
            f"[FC] Đọc {len(fc_values)} mã bằng chế độ tương thích sheet "
            "không có dimension."
        )
        return {
            "selector": selector,
            "source_column": source_column,
            "source_column_letter": source_column_letter,
            "targets": targets,
            "changed_count": changed_count,
            "fc_hash": fc.compute_fc_hash(workbook_bytes),
        }
    finally:
        values_workbook.close()


fc.read_planning_fc_targets = read_planning_fc_targets_robust


if __name__ == "__main__":
    fc.main()
