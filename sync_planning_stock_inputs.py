import math
import zipfile
from io import BytesIO

from lxml import etree
from openpyxl import load_workbook

from sync_planning_fc import (
    PLANNING_SHEET,
    _find_sheet_xml_path,
    _get_or_create_cell,
    _load_shared_strings,
    _read_cell_text,
)
from sync_stock import (
    DEST_PATH,
    GraphClient,
    clean_number,
    get_access_token,
    normalize_code,
    to_number,
)


STOCK_SHEET = "Ton_kho"
STOCK_CODE_COL = 1          # A - Mã sản phẩm
STOCK_ACTUAL_COL = 4        # D - Tồn thực tế
STOCK_BOOK_COLS = (5, 6, 7, 8)  # E:H - toàn bộ tồn sổ sách/hệ thống

PLANNING_CODE_COL = 1       # A
PLANNING_ACTUAL_COL = 10    # J - Tồn đầu thực tế
PLANNING_BOOK_COL = 11      # K - Tồn đầu sổ sách


def _values_equal(current, target):
    if current in (None, "") and target in (None, ""):
        return True
    if isinstance(current, (int, float)) and isinstance(target, (int, float)):
        return math.isclose(float(current), float(target), rel_tol=1e-12, abs_tol=1e-9)
    return current == target


def _set_numeric_cell(row_element, row_number, column_letter, value):
    cell = _get_or_create_cell(row_element, row_number, column_letter)
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)

    namespace = etree.QName(cell).namespace
    value_node = etree.SubElement(cell, f"{{{namespace}}}v")
    value = clean_number(value)
    if isinstance(value, int):
        value_node.text = str(value)
    else:
        value_node.text = format(float(value), ".15g")


def read_stock_targets(workbook_bytes):
    values_wb = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    formulas_wb = load_workbook(BytesIO(workbook_bytes), data_only=False, read_only=True)
    try:
        for sheet_name in (STOCK_SHEET, PLANNING_SHEET):
            if sheet_name not in values_wb.sheetnames:
                raise RuntimeError(f"Không tìm thấy sheet {sheet_name!r} trong Sắp kế hoạch.xlsx.")

        stock = values_wb[STOCK_SHEET]
        stock_values = {}
        for row in range(2, stock.max_row + 1):
            code = normalize_code(stock.cell(row=row, column=STOCK_CODE_COL).value)
            if not code:
                continue
            if code in stock_values:
                raise RuntimeError(f"Mã {code} bị lặp trong {STOCK_SHEET}!A.")

            actual = to_number(
                stock.cell(row=row, column=STOCK_ACTUAL_COL).value,
                f"{STOCK_SHEET}!D{row}",
            )
            book = sum(
                to_number(
                    stock.cell(row=row, column=column).value,
                    f"{STOCK_SHEET}!{chr(64 + column)}{row}",
                )
                for column in STOCK_BOOK_COLS
            )
            stock_values[code] = {
                "actual_stock": clean_number(actual),
                "book_stock": clean_number(book),
            }

        if not stock_values:
            raise RuntimeError(f"Không đọc được mã sản phẩm trong {STOCK_SHEET}!A.")

        planning_values = values_wb[PLANNING_SHEET]
        planning_formulas = formulas_wb[PLANNING_SHEET]
        targets = {}
        missing_codes = []
        changed_count = 0
        seen_codes = set()

        for row in range(2, planning_values.max_row + 1):
            code = normalize_code(planning_values.cell(row=row, column=PLANNING_CODE_COL).value)
            if not code:
                continue
            if code in seen_codes:
                raise RuntimeError(f"Mã {code} bị lặp trong {PLANNING_SHEET}!A.")
            seen_codes.add(code)

            if code not in stock_values:
                missing_codes.append(code)
                continue

            target = stock_values[code]
            targets[code] = target

            current_actual = planning_values.cell(row=row, column=PLANNING_ACTUAL_COL).value
            current_book = planning_values.cell(row=row, column=PLANNING_BOOK_COL).value
            formula_actual = planning_formulas.cell(row=row, column=PLANNING_ACTUAL_COL)
            formula_book = planning_formulas.cell(row=row, column=PLANNING_BOOK_COL)

            if (
                formula_actual.data_type == "f"
                or formula_book.data_type == "f"
                or not _values_equal(current_actual, target["actual_stock"])
                or not _values_equal(current_book, target["book_stock"])
            ):
                changed_count += 1

        if missing_codes:
            preview = ", ".join(missing_codes[:10])
            suffix = "..." if len(missing_codes) > 10 else ""
            raise RuntimeError(
                f"Có {len(missing_codes)} mã trong {PLANNING_SHEET}!A không tồn tại trong "
                f"{STOCK_SHEET}!A: {preview}{suffix}"
            )

        if not targets:
            raise RuntimeError(f"Không đọc được mã sản phẩm trong {PLANNING_SHEET}!A.")

        return {
            "targets": targets,
            "changed_count": changed_count,
        }
    finally:
        values_wb.close()
        formulas_wb.close()


def patch_planning_stock_workbook(workbook_bytes, targets):
    source_buffer = BytesIO(workbook_bytes)
    output_buffer = BytesIO()

    with zipfile.ZipFile(source_buffer, "r") as source_zip:
        sheet_path = _find_sheet_xml_path(source_zip, PLANNING_SHEET)
        shared_strings = _load_shared_strings(source_zip)
        root = etree.fromstring(source_zip.read(sheet_path))

        patched = 0
        seen_codes = set()
        rows = root.xpath('//*[local-name()="sheetData"]/*[local-name()="row"]')
        for row_element in rows:
            row_number_text = row_element.get("r")
            if not row_number_text:
                continue
            row_number = int(row_number_text)
            code = None
            for cell in row_element.xpath('./*[local-name()="c"]'):
                if cell.get("r") == f"A{row_number}":
                    code = normalize_code(_read_cell_text(cell, shared_strings))
                    break
            if not code or code not in targets:
                continue
            if code in seen_codes:
                raise RuntimeError(f"Mã {code} bị lặp trong XML {PLANNING_SHEET}!A.")
            seen_codes.add(code)

            target = targets[code]
            _set_numeric_cell(row_element, row_number, "J", target["actual_stock"])
            _set_numeric_cell(row_element, row_number, "K", target["book_stock"])
            patched += 1

        if patched != len(targets):
            missing = sorted(set(targets) - seen_codes)
            preview = ", ".join(missing[:10])
            suffix = "..." if len(missing) > 10 else ""
            raise RuntimeError(
                f"Chỉ cập nhật được {patched}/{len(targets)} mã vào {PLANNING_SHEET}!J:K. "
                f"Thiếu: {preview}{suffix}"
            )

        new_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        with zipfile.ZipFile(output_buffer, "w", compression=zipfile.ZIP_DEFLATED) as output_zip:
            for item in source_zip.infolist():
                data = new_xml if item.filename == sheet_path else source_zip.read(item.filename)
                output_zip.writestr(item, data)

    return output_buffer.getvalue(), patched


def prepare_stock_input_update(workbook_bytes):
    info = read_stock_targets(workbook_bytes)
    if info["changed_count"] == 0:
        return workbook_bytes, info
    updated, patched = patch_planning_stock_workbook(workbook_bytes, info["targets"])
    info["patched_count"] = patched
    return updated, info


def main():
    token = get_access_token()
    graph = GraphClient(token)
    site_id = graph.get_site_id()
    drive_id = graph.get_default_drive_id(site_id)

    dest_item = graph.get_item_by_path(drive_id, DEST_PATH)
    dest_bytes = graph.download_file(drive_id, dest_item["id"])
    updated, info = prepare_stock_input_update(dest_bytes)

    print(
        f"[{STOCK_SHEET} -> {PLANNING_SHEET}] J = Ton_kho!D; "
        f"K = SUM(Ton_kho!E:H); {len(info['targets'])} mã."
    )

    if info["changed_count"] == 0:
        print(f"[{PLANNING_SHEET}] J:K đã đúng theo {STOCK_SHEET}; không cần upload lại.")
        return

    result = graph.upload_file(
        drive_id,
        dest_item["id"],
        updated,
        expected_etag=dest_item["eTag"],
    )
    print(
        f"[{PLANNING_SHEET}] Đã cập nhật {info.get('patched_count', 0)} dòng vào J:K."
    )
    print("Upload thành công:", result.get("name", "Sắp kế hoạch.xlsx"))


if __name__ == "__main__":
    main()
