import hashlib
import json
import os
import posixpath
import re
import zipfile
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

import msal
import requests
from lxml import etree
from openpyxl import load_workbook

GRAPH = "https://graph.microsoft.com/v1.0"
HOSTNAME = "vikodacomvn.sharepoint.com"
SITE_PATH = "/sites/Planning"

SOURCE_ACTUAL_PATH = (
    "Tinh san xuat Mua hang 2027/"
    "Ton thuc te/"
    "Bao cao ton thuc te hien tai.xlsx"
)

SOURCE_FACTORY_VIKODA_PATH = (
    "Tinh san xuat Mua hang 2027/"
    "Ton He thong/"
    "Ton Nha May/"
    "NXT_Vikoda.xlsm"
)

SOURCE_FACTORY_VKD_PATH = (
    "Tinh san xuat Mua hang 2027/"
    "Ton He thong/"
    "Ton Nha May/"
    "NXT_VKD.xlsm"
)

SOURCE_ACCOUNTING_VIKODA_PATH = (
    "Tinh san xuat Mua hang 2027/"
    "Ton He thong/"
    "Ton Ke Toan/"
    "XNT_ketoan_Vikoda.xlsm"
)

SOURCE_ACCOUNTING_VKD_PATH = (
    "Tinh san xuat Mua hang 2027/"
    "Ton He thong/"
    "Ton Ke Toan/"
    "XNT_ketoan_VKD.xlsm"
)

DEST_PATH = "Tinh san xuat Mua hang 2027/Sắp kế hoạch.xlsx"
DEST_SHEET = "Ton_kho"
MASTER_SHEET = "Danh_muc"

STATE_FILE = Path("state.json")
CODE_PATTERN = re.compile(r"^\d{6,}$")
SYNC_VERSION = 3


class GraphRequestError(RuntimeError):
    """Structured Microsoft Graph failure used by retry policies."""

    def __init__(self, message, *, status_code=None, error_code=None, detail=None):
        super().__init__(message)
        self.status_code = status_code
        self.error_code = error_code
        self.detail = detail


def is_retryable_graph_error(exc):
    """Return True only for transient/concurrency/network failures."""
    if isinstance(exc, GraphRequestError):
        if exc.status_code in {412, 423, 429, 500, 502, 503, 504}:
            return True
        if str(exc.error_code or "").casefold() in {
            "resourcelocked",
            "preconditionfailed",
            "toomanyrequests",
            "timeout",
            "serviceunavailable",
        }:
            return True
        return False

    if isinstance(exc, (requests.Timeout, requests.ConnectionError)):
        return True

    # Backward compatibility for existing fakes/tests while callers migrate.
    if isinstance(exc, RuntimeError):
        message = str(exc)
        return any(
            marker in message
            for marker in ("412", "423", "429", "resourceLocked", "500", "502", "503", "504")
        )
    return False


def get_access_token():
    app = msal.ConfidentialClientApplication(
        client_id=os.environ["MS_CLIENT_ID"],
        authority=(
            "https://login.microsoftonline.com/"
            + os.environ["MS_TENANT_ID"]
        ),
        client_credential=os.environ["MS_CLIENT_SECRET"],
    )

    result = app.acquire_token_for_client(
        scopes=["https://graph.microsoft.com/.default"]
    )

    token = result.get("access_token")
    if not token:
        raise RuntimeError(
            "Không lấy được Microsoft Graph access token: "
            + json.dumps(result, ensure_ascii=False)
        )

    return token


class GraphClient:
    def __init__(self, token):
        self.session = requests.Session()
        self.session.headers.update(
            {"Authorization": f"Bearer {token}"}
        )

    @staticmethod
    def _raise(response):
        if response.ok:
            return

        try:
            detail = response.json()
        except Exception:
            detail = getattr(response, "text", "")

        error_code = None
        if isinstance(detail, dict):
            error = detail.get("error")
            if isinstance(error, dict):
                error_code = error.get("code")

        raise GraphRequestError(
            f"Microsoft Graph lỗi {response.status_code}: {detail}",
            status_code=response.status_code,
            error_code=error_code,
            detail=detail,
        )

    def get_json(self, url, params=None):
        response = self.session.get(
            url,
            params=params,
            timeout=60,
        )
        self._raise(response)
        return response.json()

    def get_site_id(self):
        url = f"{GRAPH}/sites/{HOSTNAME}:{SITE_PATH}"
        return self.get_json(url, {"$select": "id"})["id"]

    def get_default_drive_id(self, site_id):
        url = f"{GRAPH}/sites/{site_id}/drive"
        return self.get_json(url, {"$select": "id"})["id"]

    def get_item_by_path(self, drive_id, file_path):
        encoded = quote(file_path, safe="/")
        url = f"{GRAPH}/drives/{drive_id}/root:/{encoded}"
        return self.get_json(
            url,
            {"$select": "id,name,eTag,size,lastModifiedDateTime"},
        )

    def download_file(self, drive_id, item_id):
        url = f"{GRAPH}/drives/{drive_id}/items/{item_id}/content"
        response = self.session.get(
            url,
            timeout=120,
            allow_redirects=True,
        )
        self._raise(response)
        return response.content

    def upload_file(self, drive_id, item_id, content, expected_etag):
        url = f"{GRAPH}/drives/{drive_id}/items/{item_id}/content"
        headers = {
            "Content-Type": (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            "If-Match": expected_etag,
        }

        response = self.session.put(
            url,
            headers=headers,
            data=content,
            timeout=180,
        )

        if response.status_code == 412:
            raise GraphRequestError(
                "Microsoft Graph HTTP 412 preconditionFailed: file đích vừa thay đổi; "
                "phải tải lại workbook, tính lại patch và dùng ETag mới.",
                status_code=412,
                error_code="preconditionFailed",
            )

        self._raise(response)
        return response.json()


def load_state():
    if not STATE_FILE.exists():
        return {}

    try:
        data = json.loads(STATE_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}

    if "sources" not in data and data.get("source_etag"):
        return {
            "sources": {
                "actual_stock": data["source_etag"],
            }
        }

    return data


def save_state(source_etags, conversion_hash, fc_hash=None):
    payload = {
        "sync_version": SYNC_VERSION,
        "conversion_hash": conversion_hash,
        "sources": source_etags,
    }
    if fc_hash is not None:
        payload["fc_hash"] = fc_hash
    STATE_FILE.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )


def normalize_code(value, vkd_to_vikoda=False):
    if value is None:
        return None

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    text = str(value).strip()
    if not CODE_PATTERN.fullmatch(text):
        return None

    if vkd_to_vikoda and text.startswith("2"):
        text = "1" + text[1:]

    return text


def to_number(value, cell_name):
    if value is None or value == "":
        return 0

    if isinstance(value, bool):
        raise ValueError(
            f"{cell_name} chứa TRUE/FALSE, không phải số."
        )

    if isinstance(value, (int, float)):
        return value

    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(
            f"{cell_name} có giá trị không phải số: {value!r}"
        ) from exc


def clean_number(value):
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def read_actual_stock(source_bytes):
    workbook = load_workbook(
        BytesIO(source_bytes),
        data_only=True,
        read_only=True,
    )

    try:
        worksheet = workbook.worksheets[0]
        print(f"[Tồn thực tế] Sheet nguồn: {worksheet.title}")

        result = {}

        for row in range(1, worksheet.max_row + 1):
            code = normalize_code(
                worksheet.cell(row=row, column=3).value
            )
            if not code:
                continue

            if code in result:
                raise RuntimeError(
                    f"[Tồn thực tế] Mã {code} bị lặp trong cột C."
                )

            value_n = to_number(
                worksheet.cell(row=row, column=14).value,
                f"N{row}",
            )
            value_o = to_number(
                worksheet.cell(row=row, column=15).value,
                f"O{row}",
            )

            result[code] = clean_number(value_n + value_o)

        if not result:
            raise RuntimeError(
                "[Tồn thực tế] Không đọc được mã từ cột C."
            )

        print(
            f"[Tồn thực tế] Đọc {len(result)} mã; "
            "Ton_kho!D = N + O."
        )
        return result
    finally:
        workbook.close()


def read_single_value_source(
    source_bytes,
    *,
    label,
    source_name,
    sheet_name,
    code_column,
    value_column,
    value_column_letter,
    vkd_to_vikoda=False,
):
    workbook = load_workbook(
        BytesIO(source_bytes),
        data_only=True,
        read_only=True,
        keep_vba=True,
    )

    try:
        if sheet_name not in workbook.sheetnames:
            raise RuntimeError(
                f"[{label}] Không tìm thấy {sheet_name} "
                f"trong {source_name}."
            )

        worksheet = workbook[sheet_name]
        result = {}

        for row in range(1, worksheet.max_row + 1):
            raw_code = worksheet.cell(
                row=row,
                column=code_column,
            ).value

            code = normalize_code(
                raw_code,
                vkd_to_vikoda=vkd_to_vikoda,
            )

            if not code:
                continue

            if code in result:
                raise RuntimeError(
                    f"[{label}] Mã {code} bị lặp sau chuẩn hóa."
                )

            value = to_number(
                worksheet.cell(
                    row=row,
                    column=value_column,
                ).value,
                f"{value_column_letter}{row}",
            )

            result[code] = clean_number(value)

        if not result:
            raise RuntimeError(
                f"[{label}] Không đọc được mã sản phẩm."
            )

        mode = (
            " (chuẩn hóa 2xxxxxxxx → 1xxxxxxxx)"
            if vkd_to_vikoda
            else ""
        )

        print(
            f"[{label}] Đọc {len(result)} mã từ "
            f"{source_name}!{sheet_name}{mode}."
        )

        return result
    finally:
        workbook.close()


def read_conversion_factors(dest_bytes):
    workbook = load_workbook(
        BytesIO(dest_bytes),
        data_only=True,
        read_only=True,
    )

    try:
        if MASTER_SHEET not in workbook.sheetnames:
            raise RuntimeError(
                f"Không tìm thấy sheet {MASTER_SHEET!r} trong file đích."
            )

        worksheet = workbook[MASTER_SHEET]
        factors = {}

        for row in range(2, worksheet.max_row + 1):
            code = normalize_code(
                worksheet.cell(row=row, column=1).value
            )
            if not code:
                continue

            if code in factors:
                raise RuntimeError(
                    f"[{MASTER_SHEET}] Mã {code} bị lặp trong cột A."
                )

            factor = to_number(
                worksheet.cell(row=row, column=9).value,
                f"{MASTER_SHEET}!I{row}",
            )

            if factor <= 0:
                raise RuntimeError(
                    f"[{MASTER_SHEET}] Quy cách của mã {code} "
                    f"phải > 0, hiện là {factor!r}."
                )

            factors[code] = factor

        if not factors:
            raise RuntimeError(
                f"[{MASTER_SHEET}] Không đọc được mã/quy cách từ A:I."
            )

        payload = json.dumps(
            {k: factors[k] for k in sorted(factors)},
            ensure_ascii=False,
            separators=(",", ":"),
        ).encode("utf-8")

        conversion_hash = hashlib.sha256(payload).hexdigest()

        print(
            f"[{MASTER_SHEET}] Đọc {len(factors)} quy cách từ cột I."
        )

        return factors, conversion_hash
    finally:
        workbook.close()


def load_shared_strings(archive):
    path = "xl/sharedStrings.xml"
    if path not in archive.namelist():
        return []

    root = etree.fromstring(archive.read(path))
    result = []

    for si in root.xpath('//*[local-name()="si"]'):
        result.append("".join(si.itertext()))

    return result


def read_cell_text(cell, shared_strings):
    cell_type = cell.get("t")

    if cell_type == "inlineStr":
        return "".join(cell.itertext()).strip()

    value_nodes = cell.xpath('./*[local-name()="v"]')
    if not value_nodes:
        return ""

    raw = value_nodes[0].text or ""

    if cell_type == "s":
        try:
            return shared_strings[int(raw)]
        except Exception:
            return ""

    return raw


def find_sheet_xml_path(archive, sheet_name):
    workbook_root = etree.fromstring(
        archive.read("xl/workbook.xml")
    )

    relationship_ns = (
        "http://schemas.openxmlformats.org/"
        "officeDocument/2006/relationships"
    )

    relationship_id = None

    for sheet in workbook_root.xpath('//*[local-name()="sheet"]'):
        if sheet.get("name") == sheet_name:
            relationship_id = sheet.get(
                f"{{{relationship_ns}}}id"
            )
            break

    if not relationship_id:
        raise RuntimeError(
            f"Không tìm thấy sheet {sheet_name!r} trong file đích."
        )

    rels_root = etree.fromstring(
        archive.read("xl/_rels/workbook.xml.rels")
    )

    target = None

    for rel in rels_root.xpath(
        '//*[local-name()="Relationship"]'
    ):
        if rel.get("Id") == relationship_id:
            target = rel.get("Target")
            break

    if not target:
        raise RuntimeError(
            f"Không xác định được XML của sheet {sheet_name!r}."
        )

    if target.startswith("/"):
        return target.lstrip("/")

    return posixpath.normpath(
        posixpath.join("xl", target)
    )


def column_number(cell_reference):
    letters = re.match(r"([A-Z]+)", cell_reference)
    if not letters:
        return 10**9

    number = 0
    for char in letters.group(1):
        number = number * 26 + (ord(char) - 64)

    return number


def set_numeric_cell(row_element, row_number, column_letter, value):
    target_ref = f"{column_letter}{row_number}"
    target_cell = None

    cells = row_element.xpath('./*[local-name()="c"]')

    for cell in cells:
        if cell.get("r") == target_ref:
            target_cell = cell
            break

    if target_cell is None:
        namespace = etree.QName(row_element).namespace
        target_cell = etree.Element(
            f"{{{namespace}}}c",
            r=target_ref,
        )

        target_col = column_number(target_ref)
        inserted = False

        for existing in cells:
            if column_number(existing.get("r", "")) > target_col:
                existing.addprevious(target_cell)
                inserted = True
                break

        if not inserted:
            row_element.append(target_cell)

    for child in list(target_cell):
        target_cell.remove(child)

    target_cell.attrib.pop("t", None)

    namespace = etree.QName(target_cell).namespace
    value_node = etree.SubElement(
        target_cell,
        f"{{{namespace}}}v",
    )

    value = clean_number(value)

    if isinstance(value, int):
        value_node.text = str(value)
    else:
        value_node.text = format(float(value), ".15g")


def patch_destination_workbook(
    dest_bytes,
    *,
    actual_stock,
    factory_vikoda,
    factory_vkd,
    accounting_vikoda,
    accounting_vkd,
    conversion_factors,
):
    source_buffer = BytesIO(dest_bytes)
    output_buffer = BytesIO()

    counters = {
        "D": 0,
        "E": 0,
        "F": 0,
        "G": 0,
        "H": 0,
    }

    with zipfile.ZipFile(source_buffer, "r") as source_zip:
        sheet_path = find_sheet_xml_path(
            source_zip,
            DEST_SHEET,
        )

        shared_strings = load_shared_strings(source_zip)
        sheet_root = etree.fromstring(
            source_zip.read(sheet_path)
        )

        seen_dest_codes = set()

        rows = sheet_root.xpath(
            '//*[local-name()="sheetData"]/*[local-name()="row"]'
        )

        for row_element in rows:
            row_number_text = row_element.get("r")
            if not row_number_text:
                continue

            row_number = int(row_number_text)
            code = None

            for cell in row_element.xpath('./*[local-name()="c"]'):
                if cell.get("r") == f"A{row_number}":
                    code = normalize_code(
                        read_cell_text(cell, shared_strings)
                    )
                    break

            if not code:
                continue

            if code in seen_dest_codes:
                raise RuntimeError(
                    f"Mã {code} bị lặp trong {DEST_SHEET}!A."
                )

            seen_dest_codes.add(code)

            if code not in conversion_factors:
                raise RuntimeError(
                    f"Không tìm thấy quy cách cho mã {code} "
                    f"trong {MASTER_SHEET}!A:I."
                )

            factor = conversion_factors[code]

            if code in actual_stock:
                value_d = actual_stock[code]
                set_numeric_cell(
                    row_element,
                    row_number,
                    "D",
                    value_d,
                )
                counters["D"] += 1

            value_e = factory_vikoda.get(code, 0)
            set_numeric_cell(
                row_element,
                row_number,
                "E",
                value_e,
            )
            counters["E"] += 1

            value_f = factory_vkd.get(code, 0)
            set_numeric_cell(
                row_element,
                row_number,
                "F",
                value_f,
            )
            counters["F"] += 1

            if code in accounting_vikoda:
                value_g = clean_number(
                    accounting_vikoda[code] / factor - value_e
                )
                set_numeric_cell(
                    row_element,
                    row_number,
                    "G",
                    value_g,
                )
                counters["G"] += 1
            else:
                value_g = "giữ cũ"

            if code in accounting_vkd:
                value_h = clean_number(
                    accounting_vkd[code] / factor - value_f
                )
                set_numeric_cell(
                    row_element,
                    row_number,
                    "H",
                    value_h,
                )
                counters["H"] += 1
            else:
                value_h = "giữ cũ"

            print(
                f"{code}: Q={factor}; "
                f"D={actual_stock.get(code, 'giữ cũ')}; "
                f"E={value_e}; F={value_f}; "
                f"G={value_g}; H={value_h}"
            )

        if not seen_dest_codes:
            raise RuntimeError(
                f"Không đọc được mã sản phẩm trong {DEST_SHEET}!A."
            )

        for column in ("D", "E", "F", "G", "H"):
            if counters[column] == 0:
                raise RuntimeError(
                    f"Không cập nhật được cột {column} của {DEST_SHEET}."
                )

        print(
            f"[{DEST_SHEET}] Số dòng cập nhật: "
            + ", ".join(
                f"{col}={count}"
                for col, count in counters.items()
            )
        )

        new_sheet_xml = etree.tostring(
            sheet_root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        with zipfile.ZipFile(
            output_buffer,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as output_zip:
            for item in source_zip.infolist():
                data = (
                    new_sheet_xml
                    if item.filename == sheet_path
                    else source_zip.read(item.filename)
                )
                output_zip.writestr(item, data)

    return output_buffer.getvalue()


def main():
    token = get_access_token()
    graph = GraphClient(token)

    site_id = graph.get_site_id()
    drive_id = graph.get_default_drive_id(site_id)

    print("Đã kết nối SharePoint Planning.")

    sources = {
        "actual_stock": {
            "path": SOURCE_ACTUAL_PATH,
            "label": "Tồn thực tế",
        },
        "factory_vikoda": {
            "path": SOURCE_FACTORY_VIKODA_PATH,
            "label": "Tồn nhà máy Vikoda",
        },
        "factory_vkd": {
            "path": SOURCE_FACTORY_VKD_PATH,
            "label": "Tồn nhà máy VKD",
        },
        "accounting_vikoda": {
            "path": SOURCE_ACCOUNTING_VIKODA_PATH,
            "label": "Tồn kế toán Vikoda",
        },
        "accounting_vkd": {
            "path": SOURCE_ACCOUNTING_VKD_PATH,
            "label": "Tồn kế toán VKD",
        },
    }

    for source in sources.values():
        source["item"] = graph.get_item_by_path(
            drive_id,
            source["path"],
        )

    current_etags = {
        key: source["item"]["eTag"]
        for key, source in sources.items()
    }

    dest_item = graph.get_item_by_path(
        drive_id,
        DEST_PATH,
    )
    dest_bytes = graph.download_file(
        drive_id,
        dest_item["id"],
    )

    conversion_factors, conversion_hash = read_conversion_factors(
        dest_bytes
    )

    old_state = load_state()
    old_etags = old_state.get("sources", {})

    changed_sources = [
        key
        for key, etag in current_etags.items()
        if old_etags.get(key) != etag
    ]

    if old_state.get("sync_version") != SYNC_VERSION:
        changed_sources.append("logic_version")

    if old_state.get("conversion_hash") != conversion_hash:
        changed_sources.append("Danh_muc!I")

    for source in sources.values():
        print(
            f"[{source['label']}] sửa lần cuối:",
            source["item"].get("lastModifiedDateTime"),
        )

    if not changed_sources:
        print(
            "Không có file nguồn/quy cách nào thay đổi. Kết thúc."
        )
        return

    print(
        "Nguồn hoặc logic thay đổi: "
        + ", ".join(changed_sources)
    )

    source_bytes = {
        key: graph.download_file(
            drive_id,
            source["item"]["id"],
        )
        for key, source in sources.items()
    }

    actual_stock = read_actual_stock(
        source_bytes["actual_stock"]
    )

    factory_vikoda = read_single_value_source(
        source_bytes["factory_vikoda"],
        label="Tồn nhà máy Vikoda",
        source_name="NXT_Vikoda.xlsm",
        sheet_name="Sheet1",
        code_column=2,
        value_column=12,
        value_column_letter="L",
    )

    factory_vkd = read_single_value_source(
        source_bytes["factory_vkd"],
        label="Tồn nhà máy VKD",
        source_name="NXT_VKD.xlsm",
        sheet_name="Sheet1",
        code_column=2,
        value_column=12,
        value_column_letter="L",
        vkd_to_vikoda=True,
    )

    accounting_vikoda = read_single_value_source(
        source_bytes["accounting_vikoda"],
        label="Tồn kế toán Vikoda",
        source_name="XNT_ketoan_Vikoda.xlsm",
        sheet_name="Sheet1",
        code_column=2,
        value_column=13,
        value_column_letter="M",
    )

    accounting_vkd = read_single_value_source(
        source_bytes["accounting_vkd"],
        label="Tồn kế toán VKD",
        source_name="XNT_ketoan_VKD.xlsm",
        sheet_name="Sheet1",
        code_column=2,
        value_column=13,
        value_column_letter="M",
        vkd_to_vikoda=True,
    )

    updated_dest_bytes = patch_destination_workbook(
        dest_bytes,
        actual_stock=actual_stock,
        factory_vikoda=factory_vikoda,
        factory_vkd=factory_vkd,
        accounting_vikoda=accounting_vikoda,
        accounting_vkd=accounting_vkd,
        conversion_factors=conversion_factors,
    )

    result = graph.upload_file(
        drive_id,
        dest_item["id"],
        updated_dest_bytes,
        expected_etag=dest_item["eTag"],
    )

    print(
        "Upload thành công:",
        result.get("name", "Sắp kế hoạch.xlsx"),
    )

    save_state(current_etags, conversion_hash)
    print("SYNC THÀNH CÔNG.")


if __name__ == "__main__":
    main()
