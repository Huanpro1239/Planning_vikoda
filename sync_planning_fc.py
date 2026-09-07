import hashlib
import json
import math
import posixpath
import re
import zipfile
from io import BytesIO

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from sync_stock import DEST_PATH, GraphClient, get_access_token, normalize_code


FC_SHEET = "FC"
PLANNING_SHEET = "Ke_hoach_SX"
FC_SELECTOR_CELL = "R1"
FC_HEADER_ROW = 1
FC_HEADER_MIN_COL = 5  # E
FC_HEADER_MAX_COL = 16  # P
FC_CODE_COL = 2  # B
PLANNING_CODE_COL = 1  # A
PLANNING_TARGET_COL = 12  # L
PLANNING_TARGET_LETTER = "L"


def _normalize_header(value):
    if value is None:
        return ""
    return str(value).strip().casefold()


def _clean_number(value):
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _forecast_value(value, cell_name):
    if value is None or value == "":
        return None

    if isinstance(value, bool):
        raise ValueError(
            f"{cell_name} chứa TRUE/FALSE, không phải số FC."
        )

    if isinstance(value, (int, float)):
        return _clean_number(value)

    text = str(value).strip().replace(",", "")
    try:
        return _clean_number(float(text))
    except ValueError as exc:
        raise ValueError(
            f"{cell_name} có giá trị FC không phải số: {value!r}"
        ) from exc


def _values_equal(current, target):
    if current in (None, "") and target is None:
        return True

    if isinstance(current, (int, float)) and isinstance(
        target, (int, float)
    ):
        return math.isclose(
            float(current),
            float(target),
            rel_tol=1e-12,
            abs_tol=1e-9,
        )

    return current == target


def compute_fc_hash(workbook_bytes):
    """Compute deterministic SHA-256 hash of all data in sheet FC."""
    values_workbook = load_workbook(
        BytesIO(workbook_bytes),
        data_only=True,
        read_only=True,
    )
    try:
        if FC_SHEET not in values_workbook.sheetnames:
            return ""
        fc_sheet = values_workbook[FC_SHEET]
        selector = str(fc_sheet[FC_SELECTOR_CELL].value or "").strip()
        data = [("selector", selector)]
        for row in fc_sheet.iter_rows(values_only=True):
            if any(cell not in (None, "") for cell in row):
                data.append(tuple(
                    _clean_number(c) if isinstance(c, (int, float)) else str(c or "").strip()
                    for c in row
                ))
        canonical = json.dumps(data, ensure_ascii=False, sort_keys=True).encode("utf-8")
        return hashlib.sha256(canonical).hexdigest()
    finally:
        values_workbook.close()


def _find_sheet_xml_path(archive, sheet_name):
    workbook_root = etree.fromstring(
        archive.read("xl/workbook.xml")
    )

    relationship_ns = (
        "http://schemas.openxmlformats.org/"
        "officeDocument/2006/relationships"
    )

    relationship_id = None
    for sheet in workbook_root.xpath('//*[local-name()="sheet"]'):
        if sheet.get("name") == sheet_name:
            relationship_id = sheet.get(
                f"{{{relationship_ns}}}id"
            )
            break

    if not relationship_id:
        raise RuntimeError(
            f"Không tìm thấy sheet {sheet_name!r} trong file đích."
        )

    rels_root = etree.fromstring(
        archive.read("xl/_rels/workbook.xml.rels")
    )

    target = None
    for rel in rels_root.xpath(
        '//*[local-name()="Relationship"]'
    ):
        if rel.get("Id") == relationship_id:
            target = rel.get("Target")
            break

    if not target:
        raise RuntimeError(
            f"Không xác định được XML của sheet {sheet_name!r}."
        )

    if target.startswith("/"):
        return target.lstrip("/")

    return posixpath.normpath(posixpath.join("xl", target))


def _load_shared_strings(archive):
    path = "xl/sharedStrings.xml"
    if path not in archive.namelist():
        return []

    root = etree.fromstring(archive.read(path))
    return [
        "".join(si.itertext())
        for si in root.xpath('//*[local-name()="si"]')
    ]


def _read_cell_text(cell, shared_strings):
    cell_type = cell.get("t")

    if cell_type == "inlineStr":
        return "".join(cell.itertext()).strip()

    value_nodes = cell.xpath('./*[local-name()="v"]')
    if not value_nodes:
        return ""

    raw = value_nodes[0].text or ""
    if cell_type == "s":
        try:
            return shared_strings[int(raw)]
        except Exception:
            return ""

    return raw


def _column_number(cell_reference):
    letters = re.match(r"([A-Z]+)", cell_reference)
    if not letters:
        return 10**9

    number = 0
    for char in letters.group(1):
        number = number * 26 + (ord(char) - 64)

    return number


def _get_or_create_cell(row_element, row_number, column_letter):
    target_ref = f"{column_letter}{row_number}"
    cells = row_element.xpath('./*[local-name()="c"]')

    for cell in cells:
        if cell.get("r") == target_ref:
            return cell

    namespace = etree.QName(row_element).namespace
    target_cell = etree.Element(
        f"{{{namespace}}}c",
        r=target_ref,
    )

    target_col = _column_number(target_ref)
    for existing in cells:
        if _column_number(existing.get("r", "")) > target_col:
            existing.addprevious(target_cell)
            return target_cell

    row_element.append(target_cell)
    return target_cell


def _set_planning_value(row_element, row_number, value):
    target_cell = _get_or_create_cell(
        row_element,
        row_number,
        PLANNING_TARGET_LETTER,
    )

    for child in list(target_cell):
        target_cell.remove(child)

    target_cell.attrib.pop("t", None)

    if value is None:
        return

    namespace = etree.QName(target_cell).namespace
    value_node = etree.SubElement(
        target_cell,
        f"{{{namespace}}}v",
    )

    if isinstance(value, int):
        value_node.text = str(value)
    else:
        value_node.text = format(float(value), ".15g")


def read_planning_fc_targets(workbook_bytes):
    values_workbook = load_workbook(
        BytesIO(workbook_bytes),
        data_only=True,
        read_only=True,
    )
    formulas_workbook = load_workbook(
        BytesIO(workbook_bytes),
        data_only=False,
        read_only=True,
    )

    try:
        for sheet_name in (FC_SHEET, PLANNING_SHEET):
            if sheet_name not in values_workbook.sheetnames:
                raise RuntimeError(
                    f"Không tìm thấy sheet {sheet_name!r} "
                    "trong Sắp kế hoạch.xlsx."
                )

        fc_sheet = values_workbook[FC_SHEET]
        selector = fc_sheet[FC_SELECTOR_CELL].value
        selector_key = _normalize_header(selector)

        if not selector_key:
            raise RuntimeError(
                f"{FC_SHEET}!{FC_SELECTOR_CELL} đang trống."
            )

        matched_columns = [
            column
            for column in range(
                FC_HEADER_MIN_COL,
                FC_HEADER_MAX_COL + 1,
            )
            if _normalize_header(
                fc_sheet.cell(
                    row=FC_HEADER_ROW,
                    column=column,
                ).value
            )
            == selector_key
        ]

        if not matched_columns:
            raise RuntimeError(
                f"{FC_SHEET}!{FC_SELECTOR_CELL}={selector!r} "
                f"không khớp tiêu đề nào trong {FC_SHEET}!E1:P1."
            )

        if len(matched_columns) > 1:
            letters = ", ".join(
                get_column_letter(column)
                for column in matched_columns
            )
            raise RuntimeError(
                f"Tiêu đề {selector!r} bị lặp trong {FC_SHEET}!E1:P1 "
                f"tại các cột {letters}."
            )

        source_column = matched_columns[0]
        source_column_letter = get_column_letter(source_column)

        fc_values = {}
        for row in range(2, fc_sheet.max_row + 1):
            code = normalize_code(
                fc_sheet.cell(row=row, column=FC_CODE_COL).value
            )
            if not code:
                continue

            if code in fc_values:
                raise RuntimeError(
                    f"Mã {code} bị lặp trong {FC_SHEET}!B."
                )

            fc_values[code] = _forecast_value(
                fc_sheet.cell(
                    row=row,
                    column=source_column,
                ).value,
                f"{FC_SHEET}!{source_column_letter}{row}",
            )

        if not fc_values:
            raise RuntimeError(
                f"Không đọc được mã sản phẩm trong {FC_SHEET}!B."
            )

        planning_values = values_workbook[PLANNING_SHEET]
        planning_formulas = formulas_workbook[PLANNING_SHEET]

        targets = {}
        seen_planning_codes = set()
        missing_codes = []
        changed_count = 0

        for row in range(2, planning_values.max_row + 1):
            code = normalize_code(
                planning_values.cell(
                    row=row,
                    column=PLANNING_CODE_COL,
                ).value
            )
            if not code:
                continue

            if code in seen_planning_codes:
                raise RuntimeError(
                    f"Mã {code} bị lặp trong {PLANNING_SHEET}!A."
                )
            seen_planning_codes.add(code)

            if code not in fc_values:
                missing_codes.append(code)
                continue

            target_value = fc_values[code]
            targets[code] = target_value

            current_value = planning_values.cell(
                row=row,
                column=PLANNING_TARGET_COL,
            ).value
            current_formula_cell = planning_formulas.cell(
                row=row,
                column=PLANNING_TARGET_COL,
            )

            if (
                current_formula_cell.data_type == "f"
                or not _values_equal(current_value, target_value)
            ):
                changed_count += 1

        if missing_codes:
            preview = ", ".join(missing_codes[:10])
            suffix = "..." if len(missing_codes) > 10 else ""
            raise RuntimeError(
                f"Có {len(missing_codes)} mã trong {PLANNING_SHEET}!A "
                f"không tồn tại trong {FC_SHEET}!B: {preview}{suffix}"
            )

        if not targets:
            raise RuntimeError(
                f"Không đọc được mã sản phẩm trong {PLANNING_SHEET}!A."
            )

        return {
            "selector": selector,
            "source_column": source_column,
            "source_column_letter": source_column_letter,
            "targets": targets,
            "changed_count": changed_count,
            "fc_hash": compute_fc_hash(workbook_bytes),
        }
    finally:
        values_workbook.close()
        formulas_workbook.close()


def patch_planning_fc_workbook(workbook_bytes, targets):
    source_buffer = BytesIO(workbook_bytes)
    output_buffer = BytesIO()

    with zipfile.ZipFile(source_buffer, "r") as source_zip:
        sheet_path = _find_sheet_xml_path(
            source_zip,
            PLANNING_SHEET,
        )
        shared_strings = _load_shared_strings(source_zip)
        sheet_root = etree.fromstring(source_zip.read(sheet_path))

        patched = 0
        seen_codes = set()

        rows = sheet_root.xpath(
            '//*[local-name()="sheetData"]/*[local-name()="row"]'
        )

        for row_element in rows:
            row_number_text = row_element.get("r")
            if not row_number_text:
                continue

            row_number = int(row_number_text)
            code = None

            for cell in row_element.xpath('./*[local-name()="c"]'):
                if cell.get("r") == f"A{row_number}":
                    code = normalize_code(
                        _read_cell_text(cell, shared_strings)
                    )
                    break

            if not code or code not in targets:
                continue

            if code in seen_codes:
                raise RuntimeError(
                    f"Mã {code} bị lặp trong XML {PLANNING_SHEET}!A."
                )
            seen_codes.add(code)

            _set_planning_value(
                row_element,
                row_number,
                targets[code],
            )
            patched += 1

        if patched != len(targets):
            missing = sorted(set(targets) - seen_codes)
            preview = ", ".join(missing[:10])
            suffix = "..." if len(missing) > 10 else ""
            raise RuntimeError(
                f"Chỉ cập nhật được {patched}/{len(targets)} mã vào "
                f"{PLANNING_SHEET}!L. Thiếu: {preview}{suffix}"
            )

        new_sheet_xml = etree.tostring(
            sheet_root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        with zipfile.ZipFile(
            output_buffer,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as output_zip:
            for item in source_zip.infolist():
                data = (
                    new_sheet_xml
                    if item.filename == sheet_path
                    else source_zip.read(item.filename)
                )
                output_zip.writestr(item, data)

    return output_buffer.getvalue(), patched


def prepare_planning_fc_update(workbook_bytes):
    info = read_planning_fc_targets(workbook_bytes)

    if info["changed_count"] == 0:
        return workbook_bytes, info

    updated_bytes, patched = patch_planning_fc_workbook(
        workbook_bytes,
        info["targets"],
    )
    info["patched_count"] = patched
    return updated_bytes, info


def main():
    token = get_access_token()
    graph = GraphClient(token)

    site_id = graph.get_site_id()
    drive_id = graph.get_default_drive_id(site_id)

    dest_item = graph.get_item_by_path(drive_id, DEST_PATH)
    dest_bytes = graph.download_file(
        drive_id,
        dest_item["id"],
    )

    updated_bytes, info = prepare_planning_fc_update(dest_bytes)

    print(
        f"[{FC_SHEET} -> {PLANNING_SHEET}] "
        f"{FC_SELECTOR_CELL}={info['selector']!r} -> "
        f"{FC_SHEET}!{info['source_column_letter']}; "
        f"{len(info['targets'])} mã sản phẩm."
    )

    if info["changed_count"] == 0:
        print(
            f"[{PLANNING_SHEET}] Cột L đã đúng theo FC; "
            "không cần upload lại file."
        )
        return

    result = graph.upload_file(
        drive_id,
        dest_item["id"],
        updated_bytes,
        expected_etag=dest_item["eTag"],
    )

    print(
        f"[{PLANNING_SHEET}] Đã cập nhật "
        f"{info.get('patched_count', 0)} dòng vào cột L."
    )
    print(
        "Upload thành công:",
        result.get("name", "Sắp kế hoạch.xlsx"),
    )


if __name__ == "__main__":
    main()
