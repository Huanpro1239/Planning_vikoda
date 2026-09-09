"""Module tổng hợp kế hoạch sản xuất theo tuần cho sheet KHSX_ki.

Quy tắc nghiệp vụ:
- Sheet Ke_hoach_SX là nguồn kế hoạch sản xuất chi tiết đã được engine phân bổ theo ngày
  và kiểm tra công suất.
- Sheet KHSX_ki là kết quả tổng hợp phân bổ theo tuần từ lịch ngày đó.
- Không chạy thuật toán lập lịch độc lập cho KHSX_ki.
- Chu kỳ tuần tuân thủ nghiêm ngặt chuẩn Monday-Sunday (Thứ Hai đến Chủ Nhật).
  Tuần 1: từ ngày 01 đến Chủ Nhật đầu tiên của tháng.
  Các tuần tiếp theo: từ Thứ Hai đến Chủ Nhật.
  Tuần cuối cùng: kết thúc vào ngày cuối tháng (days_in_month).
  Không tuần nào vượt quá 7 ngày; không có khoảng trống (gap), không chồng lấn (overlap).
  Số tuần thực tế là 4, 5, hoặc 6 tuần tùy theo lịch của từng tháng.
- Định dạng và bố cục KHSX_ki:
  + Tháng 6 tuần (ví dụ Tháng 11/2026, Tháng 3/2026, Tháng 8/2026, Tháng 5/2027, Tháng 8/2027):
    Cột tuần chiếm các cột E..J (cột 5..10); Cột Tổng cộng chuyển sang Cột K (cột 11).
    Tiêu đề hợp nhất A1:K1 .. A4:K4 và khối chữ ký I31:K31 được mở rộng.
  + Tháng 5 tuần (ví dụ Tháng 9/2026):
    Cột tuần chiếm các cột E..I (cột 5..9); Cột Tổng cộng là Cột J (cột 10).
    Cột K (cột 11) được dọn sạch hoàn toàn nếu trước đó chạy tháng 6 tuần.
  + Tháng 4 tuần (ví dụ Tháng 2/2027):
    Cột tuần 1..4 chiếm các cột E..H (cột 5..8); Cột I (cột 9) là Tuần 5 rỗng (giá trị 0);
    Cột Tổng cộng là Cột J (cột 10). Cột K được dọn sạch.
- Tính toàn vẹn SKU:
  + Cả hai sheet Ke_hoach_SX và KHSX_ki phải có tập mã SKU duy nhất và khớp 100% (set(kh) == set(ki)).
  + Bắt lỗi và chặn ngay nếu có mã SKU trùng, mã SKU bị thiếu, hoặc mã SKU lạ xuất hiện.
- Tính độc lập của Header và Parser:
  + Tiêu đề tuần trong KHSX_ki là kết quả sinh tất định từ lịch, không phải cấu hình đầu vào.
  + Parser và Verifier thuần đọc (read-only), không ghi đè ô khi xác minh.
  + Verifier bắt buộc tiêu đề dòng 6 khớp chính xác phân hoạch lịch của tháng.
"""
from __future__ import annotations

import calendar
import datetime
import math
from collections import defaultdict
from copy import copy
from io import BytesIO
from typing import Any, Collection

from openpyxl import load_workbook
from openpyxl.styles import Border, Font, PatternFill
from openpyxl.utils.cell import get_column_letter

import sync_stock

KHSX_KI_SHEET = "KHSX_ki"
PLANNING_SHEET = "Ke_hoach_SX"
START_DAILY_COLUMN = 19  # Cột S trong Ke_hoach_SX
P_COLUMN = 16            # Cột P trong Ke_hoach_SX
EPS = 1e-6


def has_khsx_ki_sheet(workbook_bytes: bytes) -> bool:
    """Kiểm tra sự tồn tại của sheet KHSX_ki trong workbook mà không nạp toàn bộ dữ liệu."""
    wb = load_workbook(BytesIO(workbook_bytes), read_only=True)
    try:
        return KHSX_KI_SHEET in wb.sheetnames
    finally:
        wb.close()


def compute_standard_calendar_weeks(
    year: int,
    month: int,
) -> list[dict[str, Any]]:
    """Phân bổ các ngày trong tháng thành các tuần theo chuẩn Monday-Sunday.

    Quy tắc:
    - Tuần 1: từ ngày 1 đến Chủ Nhật đầu tiên của tháng (hoặc ngày cuối tháng nếu tháng kết thúc trước).
    - Các tuần tiếp theo: từ Thứ Hai đến Chủ Nhật (hoặc ngày cuối tháng nếu tháng kết thúc trước).
    - Độ dài mỗi tuần từ 1 đến 7 ngày. Không tuần nào dài quá 7 ngày.
    - Không tuần nào vượt ra ngoài tháng (1 <= start_day <= end_day <= days_in_month).
    - Tập hợp các ngày của tất cả các tuần là phân hoạch chính xác của [1, days_in_month]
      (không có ngày trùng lặp, không có ngày bị bỏ sót).
    Số tuần W sẽ là 4, 5, hoặc 6 tùy theo tháng.
    """
    days_in_month = calendar.monthrange(year, month)[1]
    weeks: list[dict[str, Any]] = []
    w_start = 1
    w_idx = 1
    while w_start <= days_in_month:
        curr = w_start
        # Monday is 0, Sunday is 6
        while curr < days_in_month and datetime.date(year, month, curr).weekday() != 6:
            curr += 1
        w_end = curr
        days = list(range(w_start, w_end + 1))
        label = f"Tuần {w_idx}\n{w_start:02d}/{month:02d}-{w_end:02d}/{month:02d}"
        weeks.append({
            "week_num": w_idx,
            "start_day": w_start,
            "end_day": w_end,
            "days": days,
            "label": label,
        })
        w_start = w_end + 1
        w_idx += 1
    return weeks


def compute_month_weeks(
    year: int,
    month: int,
    *,
    num_week_cols: int | None = None,
) -> list[dict[str, Any]]:
    """Trả về danh sách tuần theo chuẩn Monday-Sunday.

    Nếu num_week_cols được chỉ định và lớn hơn số tuần lịch thực tế (ví dụ tháng 4 tuần
    nhưng template có 5 cột), các tuần bổ sung sẽ là tuần trống (start_day=0, end_day=0, days=[]).
    Nếu số tuần lịch thực tế lớn hơn num_week_cols (ví dụ tháng 6 tuần nhưng num_week_cols=5),
    vẫn trả về đầy đủ các tuần thực tế để không bị gộp/ép ngày vi phạm ranh giới tuần.
    """
    std_weeks = compute_standard_calendar_weeks(year, month)
    if num_week_cols is None or num_week_cols <= len(std_weeks):
        return std_weeks

    # Padding thêm tuần rỗng nếu num_week_cols > len(std_weeks) (ví dụ tháng 4 tuần trong template 5 cột)
    padded = list(std_weeks)
    for w_idx in range(len(std_weeks) + 1, num_week_cols + 1):
        padded.append({
            "week_num": w_idx,
            "start_day": 0,
            "end_day": 0,
            "days": [],
            "label": f"Tuần {w_idx}\n-",
        })
    return padded


def get_layout_spec(
    plan_year: int,
    plan_month: int,
) -> tuple[list[dict[str, Any]], int]:
    """Xác định cấu trúc cột tuần và cột Tổng cộng chuẩn theo lịch.

    Trả về (weeks, total_col):
    - Tháng 6 tuần: 6 cột tuần (E..J, col 5..10), cột Tổng cộng là K (col 11).
    - Tháng 5 tuần: 5 cột tuần (E..I, col 5..9), cột Tổng cộng là J (col 10).
    - Tháng 4 tuần: 4 cột tuần active (E..H, col 5..8), cột I (col 9) là tuần 5 rỗng,
                    cột Tổng cộng là J (col 10).
    """
    std_weeks = compute_standard_calendar_weeks(plan_year, plan_month)
    w_count = len(std_weeks)
    if w_count >= 6:
        weeks = [dict(w, col=5 + i) for i, w in enumerate(std_weeks)]
        total_col = 5 + w_count  # Col 11 cho 6 tuần
    elif w_count == 4:
        weeks = [dict(w, col=5 + i) for i, w in enumerate(std_weeks)]
        weeks.append({
            "week_num": 5,
            "start_day": 0,
            "end_day": 0,
            "days": [],
            "label": "Tuần 5\n-",
            "col": 9,
        })
        total_col = 10
    else:  # 5 tuần
        weeks = [dict(w, col=5 + i) for i, w in enumerate(std_weeks)]
        total_col = 10

    return weeks, total_col


def parse_week_columns(
    ws,
    *,
    plan_year: int,
    plan_month: int,
    header_row: int = 6,
) -> tuple[list[dict[str, Any]], int]:
    """Xác định danh sách các cột tuần và cột Tổng cộng thuần đọc (read-only).

    Hàm này TUYỆT ĐỐI KHÔNG ghi/sửa ô trên ws để an toàn 100% với read-only workbook.
    """
    return get_layout_spec(plan_year, plan_month)


def _snapshot_cell_style(cell) -> dict[str, Any] | None:
    """Lưu snapshot toàn bộ định dạng hiển thị của một ô."""
    if not getattr(cell, "has_style", False):
        return None
    return {
        "font": copy(cell.font) if cell.font else None,
        "border": copy(cell.border) if cell.border else None,
        "fill": copy(cell.fill) if cell.fill else None,
        "number_format": copy(cell.number_format) if cell.number_format else None,
        "protection": copy(cell.protection) if cell.protection else None,
        "alignment": copy(cell.alignment) if cell.alignment else None,
    }


def _apply_cell_style(dst_cell, style_snapshot: dict[str, Any] | None) -> None:
    """Áp dụng snapshot định dạng vào ô đích."""
    if not style_snapshot:
        return
    if style_snapshot.get("font") is not None:
        dst_cell.font = copy(style_snapshot["font"])
    if style_snapshot.get("border") is not None:
        dst_cell.border = copy(style_snapshot["border"])
    if style_snapshot.get("fill") is not None:
        dst_cell.fill = copy(style_snapshot["fill"])
    if style_snapshot.get("number_format") is not None:
        dst_cell.number_format = copy(style_snapshot["number_format"])
    if style_snapshot.get("protection") is not None:
        dst_cell.protection = copy(style_snapshot["protection"])
    if style_snapshot.get("alignment") is not None:
        dst_cell.alignment = copy(style_snapshot["alignment"])


def _copy_cell_style(src_cell, dst_cell) -> None:
    """Sao chép toàn bộ định dạng hiển thị từ một ô sang ô khác."""
    _apply_cell_style(dst_cell, _snapshot_cell_style(src_cell))


def _clear_cell(cell) -> None:
    """Xóa giá trị và đặt lại định dạng rỗng cho một ô."""
    cell.value = None
    cell.border = Border()
    cell.fill = PatternFill(fill_type=None)
    cell.font = Font()
    cell.number_format = "General"


def _detect_current_layout(
    ws_ki,
    header_row: int = 6,
    sku_rows: Collection[int] | None = None,
    total_row: int | None = None,
) -> int:
    """Xác định cột tổng hiện tại trong sheet (10 hoặc 11) dựa trên tiêu đề bảng và dữ liệu SKU.

    Quy tắc nghiệp vụ:
    - Nhận diện bố cục từ tiêu đề tổng và các cột tuần hợp lệ trong vùng bảng.
    - Cột 10 (J) có 'Tổng cộng' và Cột 11 không có 'Tổng cộng' -> trả về 10 (bố cục 4 hoặc 5 tuần).
    - Cột 11 (K) có 'Tổng cộng' và Cột 10 không có 'Tổng cộng' -> trả về 11 (bố cục 6 tuần).
    - Cột 10 (J) có 'Tuần 6' và không có 'Tổng cộng' -> trả về 11 (bố cục 6 tuần).
    - Tuyệt đối không dùng độ rộng cột, font.bold hay merged ranges ngoài bảng làm căn cứ.
    - Với tiêu đề thiếu, hỏng hoặc mâu thuẫn: đối soát các dòng dữ liệu SKU được parser nhận diện.
      Loại bỏ hoàn toàn việc quét cố định 30 dòng; không dùng ghi chú/công thức ngoài bảng làm bằng chứng.
    - Dữ liệu SKU bằng 0 (0, 0.0) phải được phân biệt với ô trống (None / chuỗi rỗng).
    - Nếu không đủ bằng chứng hoặc bằng chứng mâu thuẫn: báo lỗi rõ ràng kèm tên sheet và vị trí tiêu đề.
    """
    def _norm(val) -> str:
        return str(val or "").strip().lower()

    def _is_total(val) -> bool:
        v = _norm(val)
        return "tổng" in v or "total" in v or "tong cong" in v

    def _is_w6(val) -> bool:
        v = _norm(val)
        return any(k in v for k in ("tuần 6", "tuan 6", "t6", "week 6", "w6"))

    c10_val = ws_ki.cell(header_row, 10).value
    c11_val = ws_ki.cell(header_row, 11).value

    c10_is_total = _is_total(c10_val)
    c11_is_total = _is_total(c11_val)
    c10_is_w6 = _is_w6(c10_val)

    # 1. Tín hiệu rõ ràng và không mâu thuẫn từ hàng tiêu đề hợp lệ trong bảng
    if c10_is_total and not c11_is_total and not c10_is_w6:
        return 10
    if c11_is_total and not c10_is_total:
        return 11
    if c10_is_w6 and not c10_is_total and not c11_is_total:
        return 11

    # 2. Xử lý tiêu đề mâu thuẫn hoặc thiếu/hỏng:
    # Tái sử dụng kết quả parser hiện có; chỉ đối soát các dòng SKU hợp lệ
    sheet_name = getattr(ws_ki, "title", KHSX_KI_SHEET)
    if sku_rows is None or total_row is None:
        parsed_seen, parsed_total = _read_khsx_ki_skus(ws_ki, header_row=header_row)
        if sku_rows is None:
            sku_rows = list(parsed_seen.values())
        if total_row is None:
            total_row = parsed_total

    if not sku_rows:
        raise RuntimeError(
            f"Sheet '{sheet_name}': Không đủ bằng chứng để xác định bố cục cột tổng tại hàng {header_row} "
            f"(J{header_row}={c10_val!r}, K{header_row}={c11_val!r}) do không tìm thấy dòng SKU nào trong bảng."
        )

    # 3. Phân tích bằng chứng nội bộ trong các dòng SKU
    def _is_sum_formula(val) -> bool:
        s = str(val or "").strip().upper()
        return s.startswith("=") and "SUM" in s

    def _has_data(val) -> bool:
        if val is None:
            return False
        if isinstance(val, (int, float)):
            return True  # 0 và 0.0 là dữ liệu SKU thực sự, phân biệt với ô trống
        s = str(val).strip()
        return len(s) > 0

    def _is_numeric_or_formula(val) -> bool:
        if val is None:
            return False
        if isinstance(val, (int, float)):
            return True
        s = str(val).strip()
        if s.startswith("="):
            return True
        try:
            float(s.replace(",", ""))
            return True
        except ValueError:
            return False

    sum_count_10 = sum(1 for r in sku_rows if _is_sum_formula(ws_ki.cell(r, 10).value))
    sum_count_11 = sum(1 for r in sku_rows if _is_sum_formula(ws_ki.cell(r, 11).value))

    # Nếu dòng tổng cộng có công thức =SUM
    if total_row is not None:
        if _is_sum_formula(ws_ki.cell(total_row, 10).value):
            sum_count_10 += 1
        if _is_sum_formula(ws_ki.cell(total_row, 11).value):
            sum_count_11 += 1

    if sum_count_11 > 0 and sum_count_10 == 0:
        return 11
    if sum_count_10 > 0 and sum_count_11 == 0:
        return 10

    # 4. Khi không có công thức SUM hoặc cả hai cột đều có: Đối soát giá trị dữ liệu SKU
    sku_data_count_10 = sum(1 for r in sku_rows if _has_data(ws_ki.cell(r, 10).value))
    sku_data_count_11 = sum(1 for r in sku_rows if _has_data(ws_ki.cell(r, 11).value))
    sku_numeric_count_11 = sum(1 for r in sku_rows if _is_numeric_or_formula(ws_ki.cell(r, 11).value))

    # Nếu cả hai tiêu đề đều ghi 'Tổng cộng' (mâu thuẫn tiêu đề):
    if c10_is_total and c11_is_total:
        if sum_count_11 > 0 or sku_data_count_11 > 0:
            return 11
        return 10

    # Trường hợp tiêu đề bị thiếu/hỏng:
    # Trường hợp A: Cột 10 có dữ liệu SKU nhưng Cột 11 hoàn toàn rỗng trên toàn bộ các dòng SKU
    # (Đặc trưng của bố cục 5 tuần / 4 tuần khi tiêu đề bị thiếu/hỏng)
    if sku_data_count_10 > 0 and sku_data_count_11 == 0:
        return 10

    # Trường hợp B: Cột 11 có dữ liệu SKU hợp lệ (số hoặc công thức) trên các dòng SKU
    # (Đặc trưng của bố cục 6 tuần: Cột 11 là cột tổng sản lượng của SKU)
    if sku_numeric_count_11 > 0 and sku_numeric_count_11 == sku_data_count_11:
        return 11

    # Trường hợp C: Không đủ bằng chứng hoặc dữ liệu mâu thuẫn
    raise RuntimeError(
        f"Sheet '{sheet_name}': Tiêu đề cột tổng tại hàng {header_row} không xác định hoặc mâu thuẫn "
        f"(J{header_row}={c10_val!r}, K{header_row}={c11_val!r}) và không đủ bằng chứng rõ ràng trong {len(sku_rows)} dòng SKU "
        f"(Cột 10 có {sku_data_count_10} dòng dữ liệu, Cột 11 có {sku_data_count_11} dòng dữ liệu) để xác định bố cục."
    )



def _is_cell_occupied(ws, row: int, col: int, exclude_rng=None) -> bool:
    """Kiểm tra xem ô (row, col) có chứa giá trị, công thức, chú thích, định dạng người dùng
    hoặc thuộc một merged range khác không."""
    cell = ws.cell(row, col)
    # Giá trị hoặc công thức
    if cell.value is not None:
        return True
    # Chú thích
    if getattr(cell, "comment", None) is not None:
        return True
    # Định dạng tùy biến người dùng (fill, border không rỗng)
    if cell.fill and getattr(cell.fill, "fill_type", None) is not None:
        return True
    if cell.border:
        for side in ("left", "right", "top", "bottom"):
            border_side = getattr(cell.border, side, None)
            if border_side and getattr(border_side, "style", None) is not None:
                return True
    # Thuộc một merged range khác
    for rng in ws.merged_cells.ranges:
        if exclude_rng is not None and rng == exclude_rng:
            continue
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return True
    return False


def _safe_update_merge(ws, rng, target_end_col: int) -> bool:
    """Cập nhật an toàn một vùng merged cell hiện tại tới target_end_col.

    Quy tắc an toàn:
    - Không bao giờ tạo dải ô có cột kết thúc nhỏ hơn cột bắt đầu (target_end_col < rng.min_col).
    - Nếu đã đúng kích thước (target_end_col == rng.max_col) -> không làm gì, trả về True.
    - Khi mở rộng (target_end_col > rng.max_col): kiểm tra mọi ô trong phạm vi mở rộng.
      Nếu có ô chứa giá trị, công thức, ghi chú, định dạng hoặc thuộc vùng merge khác:
      -> Bảo toàn vùng người dùng: KHÔNG mở rộng, giữ nguyên vùng gộp hiện tại và trả về False.
    - Khi thu hẹp (target_end_col < rng.max_col): giải phóng vùng gộp an toàn.
    """
    if target_end_col < rng.min_col:
        return False

    if target_end_col == rng.max_col:
        return True

    min_r = rng.min_row
    max_r = rng.max_row
    start_col = rng.min_col
    current_end_col = rng.max_col

    # Kiểm tra xung đột khi mở rộng
    if target_end_col > current_end_col:
        for r in range(min_r, max_r + 1):
            for c in range(current_end_col + 1, target_end_col + 1):
                if _is_cell_occupied(ws, r, c, exclude_rng=rng):
                    # Phát hiện xung đột với nội dung người dùng: giữ nguyên vùng ngoài bảng
                    return False

    start_letter = get_column_letter(start_col)
    target_letter = get_column_letter(target_end_col)
    target_ref = f"{start_letter}{min_r}:{target_letter}{max_r}"

    ws.unmerge_cells(rng.coord)
    ws.merge_cells(target_ref)
    return True


def _is_planning_signature_block(ws, rng, total_row: int | None) -> bool:
    """Xác định xem vùng gộp bên dưới bảng có phải là khối chữ ký kế hoạch/người lập do module quản lý không."""
    min_footer_row = (total_row + 1) if total_row is not None else 7
    if rng.min_row < min_footer_row:
        return False
    # Khối chữ ký kế hoạch/người lập nằm ở góc dưới bên phải bảng:
    # Bắt đầu tại cột H (8) hoặc I (9), hiện kết thúc tại J (10) hoặc K (11)
    if rng.min_col not in (8, 9):
        return False
    if rng.max_col not in (10, 11):
        return False

    val_start = str(ws.cell(rng.min_row, rng.min_col).value or "").strip().lower()
    val_above = str(ws.cell(rng.min_row - 1, rng.min_col).value or "").strip().lower()
    keywords = (
        "kế hoạch", "ke hoach", "người lập", "nguoi lap",
        "lập biểu", "lap bieu", "lập bảng", "lap bang",
        "planning", "tp.kh", "tp. kh"
    )
    if any(k in val_start or k in val_above for k in keywords):
        return True

    # Kiểm tra nếu trên cùng dòng có các chức danh duyệt/ký khác (CEO, Giám đốc, Sản xuất, Mua hàng...)
    row_text = " ".join(
        str(ws.cell(rng.min_row, c).value or "").lower()
        for c in range(1, 12)
    )
    sig_titles = ("ceo", "giám đốc", "giam doc", "sản xuất", "san xuat", "mua hàng", "mua hang")
    if any(t in row_text for t in sig_titles):
        return True

    return False


def _read_planning_sheet_data(
    ws_kh,
    *,
    plan_year: int,
    plan_month: int,
) -> tuple[dict[int, int], dict[int, float], dict[tuple[int, int], float]]:
    """Đọc dữ liệu từ sheet Ke_hoach_SX: (seen_kh, kh_p, kh_daily).

    seen_kh: {code: row_idx} (phát hiện trùng mã ngay lập tức)
    kh_p: {code: committed_qty}
    kh_daily: {(code, day): qty}
    """
    days_in_month = calendar.monthrange(plan_year, plan_month)[1]
    seen_kh: dict[int, int] = {}
    kh_p: dict[int, float] = {}
    kh_daily: dict[tuple[int, int], float] = defaultdict(float)

    for r in range(2, ws_kh.max_row + 1):
        raw_a = str(ws_kh.cell(r, 1).value or "").strip().casefold()
        if "tổng cộng" in raw_a or "tong cong" in raw_a:
            break
        raw_code = ws_kh.cell(r, 1).value
        code_str = sync_stock.normalize_code(raw_code)
        if not code_str:
            continue
        code = int(code_str)
        if code in seen_kh:
            raise RuntimeError(
                f"Trùng mã SKU {code} trong sheet {PLANNING_SHEET} tại dòng {r} "
                f"(đã xuất hiện tại dòng {seen_kh[code]})."
            )
        seen_kh[code] = r
        kh_p[code] = float(ws_kh.cell(r, P_COLUMN).value or 0.0)
        for d in range(1, days_in_month + 1):
            col = START_DAILY_COLUMN + d - 1
            qty = float(ws_kh.cell(r, col).value or 0.0)
            if qty > EPS:
                kh_daily[(code, d)] += qty

    return seen_kh, kh_p, kh_daily


def _read_khsx_ki_skus(
    ws_ki,
    header_row: int = 6,
) -> tuple[dict[int, int], int | None]:
    """Đọc danh sách mã SKU và vị trí dòng TỔNG CỘNG từ sheet KHSX_ki.

    seen_ki: {code: row_idx} (phát hiện trùng mã ngay lập tức)
    total_row: index dòng TỔNG CỘNG (nếu có)
    """
    seen_ki: dict[int, int] = {}
    total_row: int | None = None

    for r in range(header_row + 1, ws_ki.max_row + 1):
        raw_a = str(ws_ki.cell(r, 1).value or "").strip().casefold()
        raw_b = str(ws_ki.cell(r, 2).value or "").strip().casefold()
        if "tổng cộng" in raw_a or "tong cong" in raw_a or "tổng cộng" in raw_b or "tong cong" in raw_b:
            total_row = r
            break
        raw_code = ws_ki.cell(r, 2).value
        code_str = sync_stock.normalize_code(raw_code)
        if not code_str:
            continue
        code = int(code_str)
        if code in seen_ki:
            raise RuntimeError(
                f"Trùng mã SKU {code} trong sheet {KHSX_KI_SHEET} tại dòng {r} "
                f"(đã xuất hiện tại dòng {seen_ki[code]})."
            )
        seen_ki[code] = r

    return seen_ki, total_row


def reconcile_skus(
    seen_kh: dict[int, int],
    seen_ki: dict[int, int],
) -> None:
    """Đối chiếu tập mã SKU giữa Ke_hoach_SX và KHSX_ki.

    Yêu cầu khớp 100% chính xác (set(kh_skus) == set(ki_skus)).
    Chặn đứng mọi trường hợp thiếu mã, thừa mã lạ, hoặc trùng mã.
    """
    kh_skus = set(seen_kh.keys())
    ki_skus = set(seen_ki.keys())

    if not kh_skus:
        raise RuntimeError(f"Không tìm thấy mã SKU nào trong sheet {PLANNING_SHEET}.")
    if not ki_skus:
        raise RuntimeError(f"Không tìm thấy mã SKU nào trong sheet {KHSX_KI_SHEET}.")

    missing_in_ki = kh_skus - ki_skus
    if missing_in_ki:
        details = [f"{c} (dòng {seen_kh[c]} {PLANNING_SHEET})" for c in sorted(missing_in_ki)]
        raise RuntimeError(
            f"Sheet {KHSX_KI_SHEET} thiếu {len(missing_in_ki)} mã SKU từ {PLANNING_SHEET}: "
            f"{', '.join(details)}"
        )

    extra_in_ki = ki_skus - kh_skus
    if extra_in_ki:
        details = [f"{c} (dòng {seen_ki[c]} {KHSX_KI_SHEET})" for c in sorted(extra_in_ki)]
        raise RuntimeError(
            f"Sheet {KHSX_KI_SHEET} chứa {len(extra_in_ki)} mã SKU không có trong {PLANNING_SHEET}: "
            f"{', '.join(details)}"
        )


def _read_daily_plan_from_analysis_or_sheet(
    workbook_bytes: bytes,
    weekly_analysis: Any | None,
    *,
    plan_year: int,
    plan_month: int,
) -> tuple[dict[tuple[int, int], float], dict[int, float]]:
    """Hàm tương thích đọc kế hoạch ngày và sản lượng cam kết."""
    wb = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        ws_kh = wb[PLANNING_SHEET]
        _, kh_p, kh_daily = _read_planning_sheet_data(ws_kh, plan_year=plan_year, plan_month=plan_month)
        if weekly_analysis is not None and getattr(weekly_analysis, "daily_plan", None):
            for item in weekly_analysis.daily_plan:
                code = int(item.ma_sp)
                day = item.date.day
                kh_daily[(code, day)] = max(kh_daily[(code, day)], float(item.qty))
            for calc in getattr(weekly_analysis, "calculated", []):
                code = int(calc.input.ma_sp)
                kh_p[code] = float(calc.schedulable_qty)
        return kh_daily, kh_p
    finally:
        wb.close()


def patch_khsx_ki_workbook(
    workbook_bytes: bytes,
    *,
    weekly_analysis: Any | None = None,
    plan_year: int,
    plan_month: int,
) -> tuple[bytes, dict[str, Any]]:
    """Tổng hợp lịch ngày từ Ke_hoach_SX vào sheet KHSX_ki, bảo toàn định dạng và cấu trúc."""
    wb = load_workbook(BytesIO(workbook_bytes), data_only=False)
    try:
        if KHSX_KI_SHEET not in wb.sheetnames:
            return workbook_bytes, {"ok": False, "reason": f"Thiếu sheet {KHSX_KI_SHEET}"}
        if PLANNING_SHEET not in wb.sheetnames:
            return workbook_bytes, {"ok": False, "reason": f"Thiếu sheet {PLANNING_SHEET}"}

        ws_ki = wb[KHSX_KI_SHEET]
        ws_kh = wb[PLANNING_SHEET]

        # 1. Đọc dữ liệu Ke_hoach_SX và KHSX_ki, kiểm tra tính duy nhất
        seen_kh, kh_p, kh_daily = _read_planning_sheet_data(
            ws_kh,
            plan_year=plan_year,
            plan_month=plan_month,
        )
        if weekly_analysis is not None and getattr(weekly_analysis, "daily_plan", None):
            for item in weekly_analysis.daily_plan:
                code = int(item.ma_sp)
                day = item.date.day
                kh_daily[(code, day)] = max(kh_daily[(code, day)], float(item.qty))
            for calc in getattr(weekly_analysis, "calculated", []):
                code = int(calc.input.ma_sp)
                if code not in seen_kh:
                    seen_kh[code] = len(seen_kh) + 2
                kh_p[code] = float(calc.schedulable_qty)

        seen_ki, total_row = _read_khsx_ki_skus(ws_ki, header_row=6)

        # 2. Đối chiếu toàn vẹn 100% SKU
        reconcile_skus(seen_kh, seen_ki)

        # 3. Lấy cấu trúc tuần chuẩn lịch Monday-Sunday
        std_weeks = compute_standard_calendar_weeks(plan_year, plan_month)
        weeks, total_col = get_layout_spec(plan_year, plan_month)
        is_6_weeks = (len(std_weeks) >= 6)
        target_end_col = 11 if is_6_weeks else 10

        # 4. Snapshot bố cục và định dạng hiện tại TRƯỚC KHI chỉnh sửa bất kỳ ô nào
        input_total_col = _detect_current_layout(
            ws_ki,
            header_row=6,
            sku_rows=list(seen_ki.values()),
            total_row=total_row,
        )
        input_is_6_weeks = (input_total_col == 11)

        # Cột mẫu tuần ổn định: Cột 8 (H - Tuần 4) luôn là cột tuần thực sự trong mọi tháng 4/5/6 tuần
        week_sample_col = 8

        # Snapshot styles cho header (dòng 6)
        total_header_style = _snapshot_cell_style(ws_ki.cell(6, input_total_col))
        week_header_style = _snapshot_cell_style(ws_ki.cell(6, week_sample_col))

        # Snapshot styles cho từng dòng SKU
        total_sku_styles = {
            r: _snapshot_cell_style(ws_ki.cell(r, input_total_col))
            for r in seen_ki.values()
        }
        week_sku_styles = {
            r: _snapshot_cell_style(ws_ki.cell(r, week_sample_col))
            for r in seen_ki.values()
        }

        # Snapshot styles cho dòng TỔNG CỘNG
        total_summary_style = None
        week_summary_style = None
        if total_row is not None:
            total_summary_style = _snapshot_cell_style(ws_ki.cell(total_row, input_total_col))
            week_summary_style = _snapshot_cell_style(ws_ki.cell(total_row, week_sample_col))

        # Đảm bảo cột tổng luôn giữ định dạng in đậm
        if total_header_style is not None and total_header_style.get("font") is not None:
            f = total_header_style["font"]
            total_header_style["font"] = Font(
                name=f.name, size=f.size, bold=True, italic=f.italic,
                vertAlign=f.vertAlign, underline=f.underline, strike=f.strike, color=f.color
            )
        if total_summary_style is not None and total_summary_style.get("font") is not None:
            f = total_summary_style["font"]
            total_summary_style["font"] = Font(
                name=f.name, size=f.size, bold=True, italic=f.italic,
                vertAlign=f.vertAlign, underline=f.underline, strike=f.strike, color=f.color
            )
        for r, s in total_sku_styles.items():
            if s is not None and s.get("font") is not None:
                f = s["font"]
                s["font"] = Font(
                    name=f.name, size=f.size, bold=True, italic=f.italic,
                    vertAlign=f.vertAlign, underline=f.underline, strike=f.strike, color=f.color
                )

        # Độ rộng cột ổn định
        if input_is_6_weeks:
            total_width = ws_ki.column_dimensions["K"].width or 17.0
            week_width = (
                ws_ki.column_dimensions["J"].width
                or ws_ki.column_dimensions["I"].width
                or 19.44140625
            )
        else:
            total_width = ws_ki.column_dimensions["J"].width or 17.0
            week_width = (
                ws_ki.column_dimensions["I"].width
                or ws_ki.column_dimensions["H"].width
                or 19.44140625
            )

        # 5. Cập nhật tiêu đề kỳ kế hoạch và ngày lập
        period_text = f"Kỳ kế hoạch: Tháng {plan_month}"
        date_text = f"Ngày lập: 01/{plan_month:02d}/{plan_year}"
        if ws_ki.cell(3, 1).value is not None:
            ws_ki.cell(3, 1).value = period_text
        if ws_ki.cell(4, 1).value is not None:
            ws_ki.cell(4, 1).value = date_text

        # Kiểm tra trước nếu có vùng gộp ô xung đột trong phạm vi dữ liệu bảng (dòng 6..total_row)
        for rng in list(ws_ki.merged_cells.ranges):
            if 6 <= rng.min_row <= (total_row or 6) and (rng.min_col <= 11 and rng.max_col >= 10):
                if total_row is not None and rng.min_row == total_row and rng.min_col == 1 and rng.max_col == 4:
                    continue
                raise ValueError(
                    f"Sheet '{ws_ki.title}' ô '{rng.coord}' có vùng gộp ô xung đột trong bảng kế hoạch sản xuất."
                )

        # 6. Cập nhật các ô hợp nhất cho tiêu đề và khối chữ ký/chân trang
        for rng in list(ws_ki.merged_cells.ranges):
            # Khối tiêu đề (dòng 1..5, bắt đầu ở cột 1)
            if rng.min_row <= 5 and rng.max_row <= 5 and rng.min_col == 1 and rng.max_col in (10, 11):
                _safe_update_merge(ws_ki, rng, target_end_col)
            # Khối chữ ký / chân trang dưới dòng tổng
            elif _is_planning_signature_block(ws_ki, rng, total_row):
                _safe_update_merge(ws_ki, rng, target_end_col)

        # 7. Thiết lập tiêu đề dòng 6 (Header) và định dạng cột tuần / tổng
        if is_6_weeks:
            _apply_cell_style(ws_ki.cell(6, 10), week_header_style)
            _apply_cell_style(ws_ki.cell(6, 11), total_header_style)
            ws_ki.cell(6, 11).value = "Tổng cộng"
            for w in weeks:
                ws_ki.cell(6, w["col"]).value = w["label"]
            ws_ki.column_dimensions["J"].width = week_width
            ws_ki.column_dimensions["K"].width = total_width
        else:
            _apply_cell_style(ws_ki.cell(6, 10), total_header_style)
            ws_ki.cell(6, 10).value = "Tổng cộng"
            for w in weeks:
                ws_ki.cell(6, w["col"]).value = w["label"]
            _clear_cell(ws_ki.cell(6, 11))
            ws_ki.column_dimensions["J"].width = total_width
            has_k_content = any(
                ws_ki.cell(r, 11).value is not None
                for r in range((total_row + 1) if total_row else 7, ws_ki.max_row + 1)
            )
            if not has_k_content and "K" in ws_ki.column_dimensions:
                del ws_ki.column_dimensions["K"]

        col_totals: dict[int, float] = defaultdict(float)
        sku_reports = []

        # 8. Ghi dữ liệu từng dòng SKU
        for code, r in seen_ki.items():
            if is_6_weeks:
                _apply_cell_style(ws_ki.cell(r, 10), week_sku_styles[r])
                _apply_cell_style(ws_ki.cell(r, 11), total_sku_styles[r])
            else:
                _apply_cell_style(ws_ki.cell(r, 10), total_sku_styles[r])
                _clear_cell(ws_ki.cell(r, 11))

            row_total = 0.0
            sku_week_vals = {}
            for w in weeks:
                col = w["col"]
                if w["days"]:
                    w_qty = sum(kh_daily.get((code, d), 0.0) for d in w["days"])
                    val = float(w_qty) if w_qty > EPS else 0.0
                else:
                    val = 0.0
                ws_ki.cell(r, col).value = val
                col_totals[col] += val
                row_total += val
                sku_week_vals[f"week_{w['week_num']}"] = val

            ws_ki.cell(r, total_col).value = float(row_total) if row_total > EPS else 0.0
            col_totals[total_col] += row_total
            sku_reports.append({
                "code": code,
                "weeks": sku_week_vals,
                "total": row_total,
                "committed_p": kh_p.get(code, 0.0),
            })

        # 9. Ghi dòng TỔNG CỘNG nếu có
        if total_row is not None:
            if is_6_weeks:
                _apply_cell_style(ws_ki.cell(total_row, 10), week_summary_style)
                _apply_cell_style(ws_ki.cell(total_row, 11), total_summary_style)
            else:
                _apply_cell_style(ws_ki.cell(total_row, 10), total_summary_style)
                _clear_cell(ws_ki.cell(total_row, 11))

            for w in weeks:
                col = w["col"]
                ws_ki.cell(total_row, col).value = float(col_totals[col])
            ws_ki.cell(total_row, total_col).value = float(col_totals[total_col])

        # Chú ý: Không xóa bất kỳ ô nào sau total_row để bảo toàn 100% ghi chú và dữ liệu ngoài bảng.

        out = BytesIO()
        wb.save(out)
        patched_bytes = out.getvalue()

        report = {
            "ok": True,
            "sheet": KHSX_KI_SHEET,
            "period": f"{plan_year:04d}-{plan_month:02d}",
            "checked_skus": len(seen_ki),
            "week_columns": [w["label"].replace("\n", " ") for w in weeks],
            "total_production": col_totals[total_col],
            "grand_total": col_totals[total_col],
            "num_weeks": len(std_weeks),
        }
        return patched_bytes, report
    finally:
        wb.close()


def verify_khsx_ki(
    workbook_bytes: bytes,
    *,
    plan_year: int,
    plan_month: int,
) -> dict[str, Any]:
    """Kiểm tra độc lập tính toàn vẹn số liệu và cấu trúc của sheet KHSX_ki.

    Các bất biến bắt buộc:
    1. Sheet KHSX_ki tồn tại.
    2. Toàn bộ ô sản lượng tuần, tổng dòng, tổng cột là số thực hữu hạn >= 0.
    3. Phân hoạch tuần khớp 100% lịch Monday-Sunday (không gap, không overlap).
    4. Tiêu đề dòng 6 khớp chính xác phân hoạch chuẩn theo lịch.
    5. Tập mã SKU duy nhất và khớp 100% giữa Ke_hoach_SX và KHSX_ki (chặn trùng, thiếu, thừa).
    6. Với từng SKU: Tổng các tuần == Cột Tổng cộng == Tổng lịch ngày Ke_hoach_SX == Cột P.
    7. Dòng TỔNG CỘNG khớp tổng các dòng của từng cột tuần.
    8. Ô grand total khớp tổng các dòng và tổng các cột.
    9. Với tháng <= 5 tuần, cột 11 (K) trong vùng quản lý không được chứa dữ liệu dư thừa.
    """
    wb = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        if KHSX_KI_SHEET not in wb.sheetnames:
            return {"ok": True, "present": False, "message": "Sheet KHSX_ki không có trong workbook"}

        ws_ki = wb[KHSX_KI_SHEET]
        if PLANNING_SHEET not in wb.sheetnames:
            raise RuntimeError(f"Không tìm thấy sheet {PLANNING_SHEET} để đối chiếu KHSX_ki.")
        ws_kh = wb[PLANNING_SHEET]

        days_in_month = calendar.monthrange(plan_year, plan_month)[1]
        std_weeks = compute_standard_calendar_weeks(plan_year, plan_month)
        weeks, total_col = get_layout_spec(plan_year, plan_month)

        # 1. Xác minh phân hoạch tuần phủ 100% các ngày trong tháng không trùng/khuyết
        active_days = [d for w in weeks for d in w["days"]]
        if sorted(active_days) != list(range(1, days_in_month + 1)):
            raise RuntimeError(
                f"Phân hoạch tuần KHSX_ki không phủ đúng dải ngày [1, {days_in_month}] của tháng {plan_month}."
            )
        if len(active_days) != len(set(active_days)):
            raise RuntimeError("Phân hoạch tuần KHSX_ki có ngày bị chồng lấn (trùng lặp).")

        # 2. Kiểm tra tiêu đề dòng 6 (Header)
        for w in weeks:
            col = w["col"]
            val = str(ws_ki.cell(6, col).value or "").strip()
            if w["days"]:
                expected_range = f"{w['start_day']:02d}/{plan_month:02d}-{w['end_day']:02d}/{plan_month:02d}"
                clean_val = val.replace(" ", "")
                if expected_range not in clean_val:
                    raise RuntimeError(
                        f"KHSX_ki tiêu đề cột {col} ({val!r}) không khớp dải ngày chuẩn {expected_range}."
                    )
                if f"Tuần{w['week_num']}" not in clean_val and f"tuần{w['week_num']}" not in clean_val.lower():
                    raise RuntimeError(
                        f"KHSX_ki tiêu đề cột {col} ({val!r}) không đúng tên Tuần {w['week_num']}."
                    )
            else:
                if "5" not in val or ("-" not in val and val != ""):
                    raise RuntimeError(
                        f"KHSX_ki tiêu đề cột {col} ({val!r}) không đúng định dạng tuần rỗng ('Tuần 5\n-')."
                    )

        val_total = str(ws_ki.cell(6, total_col).value or "").strip().casefold()
        if "tổng cộng" not in val_total and "tong cong" not in val_total:
            raise RuntimeError(
                f"KHSX_ki tiêu đề cột {total_col} ({val_total!r}) không phải cột Tổng cộng."
            )

        if len(std_weeks) <= 5:
            val_k = ws_ki.cell(6, 11).value
            if val_k is not None and str(val_k).strip() != "":
                raise RuntimeError(
                    f"KHSX_ki tiêu đề cột 11 (K) vẫn còn dữ liệu ({val_k!r}) dù tháng {plan_month} chỉ có {len(std_weeks)} tuần."
                )

        # 3. Đọc dữ liệu Ke_hoach_SX và KHSX_ki, đối chiếu 100% SKU
        seen_kh, kh_p, kh_daily = _read_planning_sheet_data(
            ws_kh,
            plan_year=plan_year,
            plan_month=plan_month,
        )
        seen_ki, total_row = _read_khsx_ki_skus(ws_ki, header_row=6)
        reconcile_skus(seen_kh, seen_ki)

        col_sums: dict[int, float] = defaultdict(float)
        checked_count = 0

        # 4. Kiểm tra từng dòng SKU
        for code, r in seen_ki.items():
            week_vals = []
            for w in weeks:
                val = ws_ki.cell(r, w["col"]).value
                if val is None:
                    val = 0.0
                if not isinstance(val, (int, float)) or not math.isfinite(val) or val < -EPS:
                    raise RuntimeError(f"KHSX_ki!R{r}C{w['col']} có giá trị không hợp lệ: {val!r}")
                val = float(val)
                if not w["days"] and abs(val) > EPS:
                    raise RuntimeError(f"KHSX_ki!R{r}C{w['col']} thuộc tuần không có ngày nhưng có giá trị: {val}")
                week_vals.append(val)
                col_sums[w["col"]] += val

            row_total = ws_ki.cell(r, total_col).value
            if row_total is None:
                row_total = 0.0
            if not isinstance(row_total, (int, float)) or not math.isfinite(row_total) or row_total < -EPS:
                raise RuntimeError(f"KHSX_ki!R{r}C{total_col} (Tổng cộng) không hợp lệ: {row_total!r}")
            row_total = float(row_total)
            col_sums[total_col] += row_total

            # Đối chiếu 1: Tổng các tuần phải bằng cột Tổng cộng
            sum_weeks = sum(week_vals)
            if abs(sum_weeks - row_total) > EPS:
                raise RuntimeError(
                    f"KHSX_ki mã {code} (dòng {r}): Tổng tuần ({sum_weeks}) != Cột tổng ({row_total})"
                )

            # Đối chiếu 2: Cột Tổng cộng phải khớp tổng lịch ngày trong Ke_hoach_SX
            daily_sum = sum(kh_daily.get((code, d), 0.0) for d in range(1, days_in_month + 1))
            if abs(row_total - daily_sum) > EPS:
                raise RuntimeError(
                    f"KHSX_ki mã {code} (dòng {r}): Cột tổng ({row_total}) != Tổng lịch ngày ({daily_sum})"
                )

            # Đối chiếu 3: Cột Tổng cộng phải khớp cột P cam kết trong Ke_hoach_SX
            p_qty = kh_p.get(code, 0.0)
            if abs(row_total - p_qty) > EPS:
                raise RuntimeError(
                    f"KHSX_ki mã {code} (dòng {r}): Cột tổng ({row_total}) != Cột P cam kết ({p_qty})"
                )

            # Đối chiếu 4: Từng tuần phải khớp chính xác tổng các ngày tương ứng trong lịch
            for idx, w in enumerate(weeks):
                expected_w = sum(kh_daily.get((code, d), 0.0) for d in w["days"])
                actual_w = week_vals[idx]
                if abs(actual_w - expected_w) > EPS:
                    raise RuntimeError(
                        f"KHSX_ki mã {code} Tuần {w['week_num']} (dòng {r}): "
                        f"Thực tế {actual_w} != Kỳ vọng {expected_w}"
                    )

            if len(std_weeks) <= 5:
                val_k = ws_ki.cell(r, 11).value
                if val_k is not None and str(val_k).strip() != "":
                    raise RuntimeError(f"KHSX_ki!R{r}C11 vẫn có dữ liệu ({val_k!r}) dù tháng chỉ có {len(std_weeks)} tuần.")

            checked_count += 1

        # 5. Kiểm tra dòng TỔNG CỘNG
        if total_row is not None:
            for w in weeks:
                col = w["col"]
                actual_col_tot = float(ws_ki.cell(total_row, col).value or 0.0)
                expected_col_tot = col_sums[col]
                if abs(actual_col_tot - expected_col_tot) > EPS:
                    raise RuntimeError(
                        f"KHSX_ki dòng TỔNG CỘNG cột {col}: {actual_col_tot} != Tổng dòng ({expected_col_tot})"
                    )

            actual_grand = float(ws_ki.cell(total_row, total_col).value or 0.0)
            expected_grand = col_sums[total_col]
            if abs(actual_grand - expected_grand) > EPS:
                raise RuntimeError(
                    f"KHSX_ki grand total: {actual_grand} != Tổng cột {expected_grand}"
                )

            if len(std_weeks) <= 5:
                val_k = ws_ki.cell(total_row, 11).value
                if val_k is not None and str(val_k).strip() != "":
                    raise RuntimeError(f"KHSX_ki!R{total_row}C11 vẫn có dữ liệu ({val_k!r}) dù tháng chỉ có {len(std_weeks)} tuần.")

        return {
            "ok": True,
            "present": True,
            "checked_skus": checked_count,
            "grand_total": col_sums[total_col],
            "period": f"{plan_year:04d}-{plan_month:02d}",
            "num_weeks": len(std_weeks),
        }
    finally:
        wb.close()
