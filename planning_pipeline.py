import copy
import math
from io import BytesIO

from openpyxl import load_workbook

from planning_cleanup import remove_sheet_formulas
from planning_schedule_report import attach_output_hash, json_safe
from sync_planning_calendar_all_months import prepare_calendar_update_all_months
import sync_planning_fc_compat  # noqa: F401 - installs dimension-tolerant FC reader
from sync_planning_fc import prepare_planning_fc_update
import sync_planning_metrics as metrics
import sync_planning_metrics_all_months  # noqa: F401 - installs all-month/direct hooks
import sync_planning_metrics_compat as metrics_compat
from sync_planning_stock_inputs import prepare_stock_input_update
from sync_planning_weekly_model import (
    prepare_weekly_schedule_update,
    verify_weekly_workbook,
)
import sync_stock
from sync_stock_compat import read_conversion_factors_robust


SOURCE_KEYS = (
    "actual_stock",
    "factory_vikoda",
    "factory_vkd",
    "accounting_vikoda",
    "accounting_vkd",
)


def _planning_snapshot(workbook_bytes):
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        sheet = workbook["Ke_hoach_SX"]
        result = {}
        for row in range(2, sheet.max_row + 1):
            raw_code = sheet.cell(row=row, column=1).value
            code = sync_stock.normalize_code(raw_code)
            if not code:
                continue
            result[code] = {
                "J": sheet.cell(row=row, column=10).value,
                "K": sheet.cell(row=row, column=11).value,
                "L": sheet.cell(row=row, column=12).value,
                "M": sheet.cell(row=row, column=13).value,
                "N": sheet.cell(row=row, column=14).value,
                "O": sheet.cell(row=row, column=15).value,
                "P": sheet.cell(row=row, column=16).value,
                "Q": sheet.cell(row=row, column=17).value,
                "R": sheet.cell(row=row, column=18).value,
            }
        return result
    finally:
        workbook.close()


def _diff_snapshot(before, after):
    changed = {}
    for code in sorted(set(before) | set(after)):
        fields = {}
        for field in ("J", "K", "L", "M", "N", "O", "P", "Q", "R"):
            old = before.get(code, {}).get(field)
            new = after.get(code, {}).get(field)
            same = old == new
            if isinstance(old, (int, float)) and isinstance(new, (int, float)):
                same = math.isclose(float(old), float(new), rel_tol=1e-10, abs_tol=1e-7)
            if not same:
                fields[field] = {"before": old, "after": new}
        if fields:
            changed[code] = fields
    return changed


def prepare_pipeline_output(
    target_workbook_bytes,
    source_bytes,
    *,
    runtime_state,
    input_revision=None,
):
    """Compute one validated proposal from one complete SharePoint snapshot.

    Network I/O remains outside this function. The planning engine is the
    data-driven implementation of the ``Ke hoach SX tuan`` model with Service First:
    sales/debt are mandatory and safety stock uses only remaining capacity.
    """
    missing_sources = [key for key in SOURCE_KEYS if key not in source_bytes]
    if missing_sources:
        raise RuntimeError("Thiếu nguồn dry-run: " + ", ".join(missing_sources))

    initial_bytes = target_workbook_bytes
    before_snapshot = _planning_snapshot(initial_bytes)
    work = initial_bytes
    steps = []

    # 1) Recompute Ton_kho from the same raw SharePoint snapshots used by production.
    conversion_factors, conversion_hash = read_conversion_factors_robust(work)
    actual_stock = sync_stock.read_actual_stock(source_bytes["actual_stock"])
    factory_vikoda = sync_stock.read_single_value_source(
        source_bytes["factory_vikoda"],
        label="Tồn nhà máy Vikoda",
        source_name="NXT_Vikoda.xlsm",
        sheet_name="Sheet1",
        code_column=2,
        value_column=12,
        value_column_letter="L",
    )
    factory_vkd = sync_stock.read_single_value_source(
        source_bytes["factory_vkd"],
        label="Tồn nhà máy VKD",
        source_name="NXT_VKD.xlsm",
        sheet_name="Sheet1",
        code_column=2,
        value_column=12,
        value_column_letter="L",
        vkd_to_vikoda=True,
    )
    accounting_vikoda = sync_stock.read_single_value_source(
        source_bytes["accounting_vikoda"],
        label="Tồn kế toán Vikoda",
        source_name="XNT_ketoan_Vikoda.xlsm",
        sheet_name="Sheet1",
        code_column=2,
        value_column=13,
        value_column_letter="M",
    )
    accounting_vkd = sync_stock.read_single_value_source(
        source_bytes["accounting_vkd"],
        label="Tồn kế toán VKD",
        source_name="XNT_ketoan_VKD.xlsm",
        sheet_name="Sheet1",
        code_column=2,
        value_column=13,
        value_column_letter="M",
        vkd_to_vikoda=True,
    )
    work = sync_stock.patch_destination_workbook(
        work,
        actual_stock=actual_stock,
        factory_vikoda=factory_vikoda,
        factory_vkd=factory_vkd,
        accounting_vikoda=accounting_vikoda,
        accounting_vkd=accounting_vkd,
        conversion_factors=conversion_factors,
    )
    steps.append("Ton_kho")

    # 2) J/K come from Ton_kho and L from the selected FC month.
    work, stock_info = prepare_stock_input_update(work)
    work, fc_info = prepare_planning_fc_update(work)
    steps.extend(["J_K", "FC_L"])

    # 3) Resolve period from the current source snapshot, never wall clock.
    report_date, actual_receipts, consignments = metrics.read_actual_inputs(
        source_bytes["actual_stock"]
    )
    selector, plan_month, planning_rows = metrics.read_planning_rows(work)
    plan_year = metrics._resolve_plan_year(report_date, plan_month)

    # 4) Calendar must match the same resolved period.
    work, calendar_info = prepare_calendar_update_all_months(
        work,
        plan_year=plan_year,
    )
    steps.append("calendar")

    # Re-read after layout/calendar migration. This also installs Leadtime from
    # Danh_muc!J as the single source of truth for the M/N base calculation.
    selector, plan_month, planning_rows = metrics.read_planning_rows(work)
    conversion_factors, _ = metrics_compat.read_conversion_factors_and_leadtime(work)
    system_receipts = metrics.read_system_receipts(
        source_bytes["factory_vikoda"],
        conversion_factors,
        set(planning_rows),
    )

    # 5) M/N base inputs remain driven by SharePoint/state. The weekly engine
    # then recalculates O:R and the daily schedule from these runtime values.
    next_state = copy.deepcopy(runtime_state or {})
    metric_values, state_changed, source_key, plan_key = metrics.calculate_metrics(
        report_date=report_date,
        plan_month=plan_month,
        planning_rows=planning_rows,
        actual_receipts=actual_receipts,
        system_receipts=system_receipts,
        current_consignments=consignments,
        state=next_state,
    )
    metric_changes = metrics.count_changes(planning_rows, metric_values)
    if metric_changes:
        work, metric_patched = metrics.patch_workbook(work, metric_values)
    else:
        metric_patched = 0
    steps.append("M_N_base")

    # 6) Production planning uses only the Ke hoach SX tuan algorithm.
    scheduled_bytes, report, weekly_analysis = prepare_weekly_schedule_update(
        work,
        plan_year=plan_year,
        plan_month=plan_month,
        input_revision=dict(input_revision or {}),
    )
    steps.append("ke_hoach_sx_tuan")

    # 7) Clean formulas on the proposal copy and validate the exact output bytes.
    final_bytes, formulas_removed = remove_sheet_formulas(
        scheduled_bytes,
        "Ke_hoach_SX",
    )
    steps.append("cleanup")
    report = attach_output_hash(report, final_bytes)
    verify_info = verify_weekly_workbook(
        final_bytes,
        schedule_report=report,
        plan_year=plan_year,
        plan_month=plan_month,
    )
    steps.append("verify")

    after_snapshot = _planning_snapshot(final_bytes)
    report["pipeline"] = {
        "mode": "offline_single_snapshot",
        "engine": "ke_hoach_sx_tuan_v2_service_first",
        "steps": steps,
        "source_period": source_key,
        "plan_period": plan_key,
        "selector": selector,
        "conversion_hash": conversion_hash,
        "fc_hash": fc_info.get("fc_hash", ""),
        "state_would_change": bool(state_changed),
        "stock_rows_changed": stock_info.get("changed_count", 0),
        "fc_rows_changed": fc_info.get("changed_count", 0),
        "calendar_cells_changed": calendar_info.get("changed_count", 0),
        "metric_rows_changed": metric_changes,
        "metric_rows_patched": metric_patched,
        "schedule_cells_changed": weekly_analysis.changed_cells,
        "policy_warning_count": len(weekly_analysis.policy_warnings),
        "formulas_removed": formulas_removed,
        "planning_changes": _diff_snapshot(before_snapshot, after_snapshot),
        "verify": verify_info,
    }

    return final_bytes, json_safe(report), next_state
