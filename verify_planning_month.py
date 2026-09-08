import math
import re
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import sync_planning_fc as fc
from planning_schedule_report import load_schedule_report
from sync_planning_calendar import build_date_headers, parse_plan_month
from sync_planning_calendar_all_months import resolve_plan_year
from sync_stock import DEST_PATH, GraphClient, get_access_token, normalize_code, to_number
from weekly_planning_engine import is_sugar_classification


PLANNING_SHEET = "Ke_hoach_SX"
STOCK_SHEET = "Ton_kho"
START_COLUMN = 19  # S
SHARED_RESOURCE = "KHS + PET 9000"
# Tên resource máy chung mà sync_planning_weekly_model ghi vào report. Verifier
# chấp nhận cả hai tên để độc lập với thay đổi nhãn của report.
SHARED_RESOURCE_ALIASES = (SHARED_RESOURCE, "KHS/PET 9000 shared machine")
SHARED_LINES = {"KHS", "PET 9000"}
SETUP_SHIFTS = 0.5


def _find_shared_resource_info(schedule_report):
    """Tìm block resource máy chung trong report theo nhiều tên khóa.

    Trả về ``(resource_info, meta)`` nếu tìm thấy, ngược lại ``(None, None)``.
    Ưu tiên các alias đã biết, sau đó tới bất kỳ resource nào có ``meta.timeline``.
    """
    resources = (schedule_report or {}).get("resources") or {}
    for name in SHARED_RESOURCE_ALIASES:
        info = resources.get(name)
        if isinstance(info, dict):
            return info, (info.get("meta") or {})
    for info in resources.values():
        if isinstance(info, dict) and isinstance((info.get("meta") or {}).get("timeline"), list):
            return info, info.get("meta") or {}
    return None, None


def _header_day(value):
    text = str(value or "").strip()
    return text.splitlines()[0] if text else "?"


def _finite_number(value, label):
    if value in (None, ""):
        return 0.0
    if isinstance(value, bool):
        raise RuntimeError(f"{label} chứa TRUE/FALSE, không phải số.")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise RuntimeError(f"{label} không phải số: {value!r}") from exc
    if not math.isfinite(result):
        raise RuntimeError(f"{label} phải là số hữu hạn, hiện {value!r}.")
    return result


def _validate_report_provenance(schedule_report, plan_year, plan_month):
    if not isinstance(schedule_report, dict):
        raise RuntimeError("Thiếu schedule report có provenance để kiểm carryover/setup timeline.")
    revision = schedule_report.get("input_revision")
    if not isinstance(revision, dict) or not revision:
        raise RuntimeError("Schedule report thiếu input_revision/provenance.")
    reported_month = schedule_report.get("plan_month")
    expected_month = f"{plan_year:04d}-{plan_month:02d}"
    if reported_month and reported_month != expected_month:
        raise RuntimeError(
            f"Schedule report thuộc {reported_month}, không phải kỳ {expected_month}."
        )


def _validate_shared_timeline(
    schedule_report,
    headers,
    planning_data,
    capacity,
):
    _validate_report_provenance(
        schedule_report,
        planning_data["__plan_year"],
        planning_data["__plan_month"],
    )
    resource_info, meta = _find_shared_resource_info(schedule_report)
    if not isinstance(resource_info, dict):
        raise RuntimeError(
            f"Schedule report thiếu resource {SHARED_RESOURCE!r}; không chứng minh được setup/timeline."
        )

    # capacity có thể nằm ở cấp resource hoặc trong meta (shared_machine).
    capacity_source = (
        resource_info.get("capacity_shifts_per_day")
        if resource_info.get("capacity_shifts_per_day") is not None
        else meta.get("capacity_shifts_per_day")
    )
    report_capacity = _finite_number(
        capacity_source,
        f"report {SHARED_RESOURCE} capacity",
    )
    if not math.isclose(report_capacity, capacity, rel_tol=1e-9, abs_tol=1e-6):
        raise RuntimeError(
            f"Schedule report capacity {report_capacity} khác workbook capacity {capacity} cho {SHARED_RESOURCE}."
        )

    timeline = (meta or {}).get("timeline")
    if not isinstance(timeline, list) or not timeline:
        raise RuntimeError(
            f"Schedule report thiếu timeline production/setup cho {SHARED_RESOURCE}."
        )

    events = []
    horizon = len(headers) * capacity
    for index, raw in enumerate(timeline):
        if not isinstance(raw, dict):
            raise RuntimeError(f"Timeline event #{index + 1} không phải object.")
        event = dict(raw)
        kind = str(event.get("kind") or "").strip()
        if kind not in {"production", "setup"}:
            raise RuntimeError(f"Timeline event #{index + 1} có kind={kind!r} không hợp lệ.")
        start = _finite_number(event.get("start_shift"), f"timeline[{index}].start_shift")
        end = _finite_number(event.get("end_shift"), f"timeline[{index}].end_shift")
        if start < -1e-7 or end < start - 1e-7 or end > horizon + 1e-7:
            raise RuntimeError(
                f"Timeline event #{index + 1} ngoài horizon: start={start}, end={end}, horizon={horizon}."
            )
        event["_start"] = start
        event["_end"] = end
        events.append(event)

    events.sort(key=lambda item: (item["_start"], item["_end"]))
    previous_end = 0.0
    for event in events:
        if event["_start"] < previous_end - 1e-7:
            raise RuntimeError(
                f"Timeline {SHARED_RESOURCE} bị chồng lấn tại shift {event['_start']:.3f}."
            )
        previous_end = max(previous_end, event["_end"])

    daily_usage = [0.0] * len(headers)
    reconstructed = {
        code: [0.0] * len(headers)
        for code, item in planning_data.items()
        if not code.startswith("__") and item["line"] in SHARED_LINES
    }

    last_production_code = None
    setup_since_last_production = 0.0
    for event in events:
        duration = event["_end"] - event["_start"]
        kind = event["kind"]

        for day_index in range(len(headers)):
            day_start = day_index * capacity
            day_end = day_start + capacity
            overlap = max(
                0.0,
                min(event["_end"], day_end) - max(event["_start"], day_start),
            )
            if overlap > 1e-9:
                daily_usage[day_index] += overlap

        if kind == "setup":
            setup_since_last_production += duration
            continue

        code = normalize_code(event.get("code")) or str(event.get("code") or "").strip()
        if code not in reconstructed:
            raise RuntimeError(f"Timeline production chứa mã {code!r} không thuộc máy chung trong workbook.")

        if last_production_code is not None and code != last_production_code:
            if setup_since_last_production < SETUP_SHIFTS - 1e-7:
                raise RuntimeError(
                    f"Timeline đổi mã {last_production_code} -> {code} thiếu setup {SETUP_SHIFTS:g} ca; "
                    f"chỉ có {setup_since_last_production:g} ca."
                )
        setup_since_last_production = 0.0
        last_production_code = code

        per_shift = planning_data[code]["per_shift"]
        expected_qty = duration * per_shift
        event_qty = _finite_number(event.get("qty"), f"timeline production {code} qty")
        if not math.isclose(event_qty, expected_qty, rel_tol=1e-9, abs_tol=1e-5):
            raise RuntimeError(
                f"Timeline mã {code} qty={event_qty} không khớp duration*E={expected_qty}."
            )

        for day_index in range(len(headers)):
            day_start = day_index * capacity
            day_end = day_start + capacity
            overlap = max(
                0.0,
                min(event["_end"], day_end) - max(event["_start"], day_start),
            )
            if overlap > 1e-9:
                reconstructed[code][day_index] += overlap * per_shift

    for day_index, used in enumerate(daily_usage):
        if used > capacity + 1e-6:
            raise RuntimeError(
                f"Timeline {SHARED_RESOURCE} ngày {_header_day(headers[day_index])} dùng "
                f"{used:.3f} ca gồm production/setup > capacity {capacity:.3f}."
            )

    for code, values in reconstructed.items():
        workbook_values = planning_data[code]["daily_values"]
        for index, (expected, actual) in enumerate(zip(values, workbook_values)):
            if not math.isclose(expected, actual, rel_tol=1e-9, abs_tol=1e-5):
                raise RuntimeError(
                    f"Timeline mã {code} ngày {_header_day(headers[index])}={expected} "
                    f"khác workbook={actual}."
                )


def _validate_shared_from_workbook(headers, planning_data, capacity):
    """Kiểm tra khả thi máy chung KHS/PET 9000 chỉ từ dữ liệu workbook.

    Dùng khi report không kèm timeline production/setup. Ràng buộc capacity
    theo ngày đã được kiểm ở vòng lặp chính; ở đây kiểm thêm chặn dưới theo
    tháng có tính setup:

        tổng_ca_sản_xuất + (số_SKU_được_SX - 1) * setup <= capacity * số_ngày

    Đây là điều kiện cần: mọi lịch hợp lệ do engine sinh ra đều thỏa (vì đã xếp
    được trên một timeline chung có setup), nên không tạo lỗi giả; đồng thời bắt
    được trường hợp workbook bị chỉnh tay vượt tổng năng lực máy chung.
    """
    if capacity <= 0:
        raise RuntimeError(
            f"{SHARED_RESOURCE} có capacity {capacity} không hợp lệ."
        )

    production_shifts = 0.0
    produced_codes = 0
    for code, item in planning_data.items():
        if code.startswith("__") or item["line"] not in SHARED_LINES:
            continue
        scheduled = float(item["scheduled"])
        per_shift = float(item["per_shift"])
        if scheduled <= 1e-7:
            continue
        if per_shift <= 0:
            raise RuntimeError(
                f"Mã {code} máy chung có SL/ca <= 0 nhưng có lịch SX."
            )
        production_shifts += scheduled / per_shift
        produced_codes += 1

    if produced_codes <= 1:
        return

    setup_shifts = max(0, produced_codes - 1) * SETUP_SHIFTS
    total_capacity = capacity * len(headers)
    required = production_shifts + setup_shifts
    if required > total_capacity + 1e-6:
        raise RuntimeError(
            f"{SHARED_RESOURCE}: tổng nhu cầu {required:.3f} ca "
            f"(SX {production_shifts:.3f} + setup {setup_shifts:.3f}) vượt "
            f"năng lực tháng {total_capacity:.3f} ca ({capacity:g} ca/ngày x "
            f"{len(headers)} ngày)."
        )


def verify_workbook(workbook_bytes, schedule_report=None, plan_year=None):
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        for sheet_name in (fc.FC_SHEET, STOCK_SHEET, PLANNING_SHEET):
            if sheet_name not in workbook.sheetnames:
                raise RuntimeError(f"Không tìm thấy sheet {sheet_name!r}.")

        fc_sheet = workbook[fc.FC_SHEET]
        stock_sheet = workbook[STOCK_SHEET]
        planning = workbook[PLANNING_SHEET]

        # 1) J/K phải lấy trực tiếp từ Ton_kho theo mã sản phẩm.
        stock_values = {}
        for row in range(2, stock_sheet.max_row + 1):
            code = normalize_code(stock_sheet.cell(row=row, column=1).value)
            if not code:
                continue
            if code in stock_values:
                raise RuntimeError(f"Mã {code} bị lặp trong {STOCK_SHEET}!A.")
            actual = to_number(stock_sheet.cell(row=row, column=4).value, f"{STOCK_SHEET}!D{row}")
            book = sum(
                to_number(stock_sheet.cell(row=row, column=column).value, f"{STOCK_SHEET}!{get_column_letter(column)}{row}")
                for column in range(5, 9)
            )
            stock_values[code] = (actual, book)

        stock_mismatches = []
        stock_checked = 0
        for row_number, values in enumerate(
            planning.iter_rows(min_row=2, max_col=11, values_only=True),
            start=2,
        ):
            code = normalize_code(values[0])
            if not code:
                continue
            if code not in stock_values:
                raise RuntimeError(f"Mã {code} ở {PLANNING_SHEET}!A{row_number} không có trong {STOCK_SHEET}!A.")
            expected_actual, expected_book = stock_values[code]
            actual_j = values[9]
            actual_k = values[10]
            stock_checked += 1
            if not fc._values_equal(actual_j, expected_actual) or not fc._values_equal(actual_k, expected_book):
                stock_mismatches.append((code, actual_j, expected_actual, actual_k, expected_book))

        if stock_mismatches:
            preview = "; ".join(
                f"{code}: J={actual_j!r}/Ton_kho!D={expected_j!r}, K={actual_k!r}/SUM(E:H)={expected_k!r}"
                for code, actual_j, expected_j, actual_k, expected_k in stock_mismatches[:10]
            )
            raise RuntimeError(
                f"{PLANNING_SHEET}!J:K chưa theo {STOCK_SHEET} ({len(stock_mismatches)}/{stock_checked} mã sai): {preview}"
            )

        # 2) L phải lấy đúng cột tháng được chọn tại FC!R1.
        selector = fc_sheet[fc.FC_SELECTOR_CELL].value
        selector_key = fc._normalize_header(selector)
        plan_month = parse_plan_month(selector)
        # Ưu tiên năm được truyền vào; nếu không có thì lấy từ report
        # (plan_month dạng "YYYY-MM"); chỉ dùng đồng hồ hệ thống khi không còn
        # nguồn nào khác. Điều này tránh verify sai quanh mốc giao năm và giúp
        # test không phụ thuộc thời điểm chạy.
        if plan_year is None and isinstance(schedule_report, dict):
            reported = str(schedule_report.get("plan_month") or "")
            match = re.fullmatch(r"(\d{4})-(\d{2})", reported)
            if match and int(match.group(2)) == plan_month:
                plan_year = int(match.group(1))
        if plan_year is None:
            plan_year = resolve_plan_year(plan_month)

        matched_columns = []
        for column in range(fc.FC_HEADER_MIN_COL, fc.FC_HEADER_MAX_COL + 1):
            if fc._normalize_header(fc_sheet.cell(row=1, column=column).value) == selector_key:
                matched_columns.append(column)
        if len(matched_columns) != 1:
            raise RuntimeError(
                f"{fc.FC_SHEET}!{fc.FC_SELECTOR_CELL}={selector!r} phải khớp đúng 1 cột E:P, "
                f"hiện khớp {len(matched_columns)} cột."
            )
        source_column = matched_columns[0]

        fc_values = {}
        for values in fc_sheet.iter_rows(min_row=2, max_col=max(fc.FC_CODE_COL, source_column), values_only=True):
            code = normalize_code(values[fc.FC_CODE_COL - 1])
            if not code:
                continue
            fc_values[code] = values[source_column - 1]

        fc_mismatches = []
        fc_checked = 0
        for row_number, values in enumerate(
            planning.iter_rows(min_row=2, max_col=fc.PLANNING_TARGET_COL, values_only=True),
            start=2,
        ):
            code = normalize_code(values[fc.PLANNING_CODE_COL - 1])
            if not code:
                continue
            if code not in fc_values:
                raise RuntimeError(f"Mã {code} ở {PLANNING_SHEET}!A{row_number} không có trong FC!B.")
            expected = fc_values[code]
            actual = values[fc.PLANNING_TARGET_COL - 1]
            fc_checked += 1
            if not fc._values_equal(actual, expected):
                fc_mismatches.append((code, actual, expected))

        if fc_mismatches:
            preview = "; ".join(
                f"{code}: L={actual!r}, FC={expected!r}"
                for code, actual, expected in fc_mismatches[:10]
            )
            raise RuntimeError(
                f"{PLANNING_SHEET}!L chưa theo {selector!r} ({len(fc_mismatches)}/{fc_checked} mã sai): {preview}"
            )

        # 3) Header lịch phải đúng tháng và không còn cột legacy/dữ liệu dư.
        headers = build_date_headers(plan_year, plan_month)
        expected_start = headers[0]
        expected_end = headers[-1]
        end_column = START_COLUMN + len(headers) - 1
        after_end_column = end_column + 1

        actual_start = planning.cell(row=1, column=START_COLUMN).value
        actual_end = planning.cell(row=1, column=end_column).value
        if actual_start != expected_start or actual_end != expected_end:
            raise RuntimeError(
                f"Dải ngày sai cho {plan_month:02d}/{plan_year}: "
                f"S1={actual_start!r}, {get_column_letter(end_column)}1={actual_end!r}; "
                f"cần {expected_start!r} ... {expected_end!r}."
            )

        for offset, expected_header in enumerate(headers):
            column = START_COLUMN + offset
            actual_header = planning.cell(row=1, column=column).value
            if actual_header != expected_header:
                raise RuntimeError(
                    f"Tiêu đề ngày sai tại {get_column_letter(column)}1: "
                    f"{actual_header!r}; cần {expected_header!r}."
                )

        legacy_labels = {"Kỳ kế hoạch", "Tổng SX", "Chênh lệch (SX-P)"}
        for column in range(START_COLUMN, max(after_end_column + 4, 53)):
            value = planning.cell(row=1, column=column).value
            if value in legacy_labels:
                raise RuntimeError(
                    f"Còn cột legacy {value!r} tại {get_column_letter(column)}1."
                )

        stale_cells = []
        for row_number, values in enumerate(
            planning.iter_rows(
                min_row=1,
                min_col=after_end_column,
                max_col=after_end_column + 4,
                values_only=True,
            ),
            start=1,
        ):
            for offset, value in enumerate(values):
                if value not in (None, ""):
                    stale_cells.append(
                        f"{get_column_letter(after_end_column + offset)}{row_number}={value!r}"
                    )
                    if len(stale_cells) >= 10:
                        break
            if len(stale_cells) >= 10:
                break
        if stale_cells:
            raise RuntimeError(
                "Còn dữ liệu sau ngày cuối tháng: " + "; ".join(stale_cells)
            )

        # 4) Hậu kiểm độc lập dữ liệu kế hoạch và lịch ngày.
        EPS = 1e-7
        MASS_EPS = 1e-5
        resource_capacity = {}
        daily_resource_usage = {}
        checked_schedule_rows = 0
        planning_data = {
            "__plan_year": plan_year,
            "__plan_month": plan_month,
        }

        for row in range(2, planning.max_row + 1):
            code = normalize_code(planning.cell(row=row, column=1).value)
            if not code:
                continue

            batch = _finite_number(planning.cell(row=row, column=4).value, f"D{row}")
            per_shift = _finite_number(planning.cell(row=row, column=5).value, f"E{row}")
            line = str(planning.cell(row=row, column=6).value or "").strip()
            classification = str(planning.cell(row=row, column=8).value or "").strip()
            shifts_per_day = _finite_number(planning.cell(row=row, column=9).value, f"I{row}")
            _finite_number(planning.cell(row=row, column=10).value, f"J{row}")
            book_stock = _finite_number(planning.cell(row=row, column=11).value, f"K{row}")
            forecast = _finite_number(planning.cell(row=row, column=12).value, f"L{row}")
            target_stock = _finite_number(planning.cell(row=row, column=13).value, f"M{row}")
            debt = _finite_number(planning.cell(row=row, column=14).value, f"N{row}")
            required = _finite_number(planning.cell(row=row, column=15).value, f"O{row}")
            planned = _finite_number(planning.cell(row=row, column=16).value, f"P{row}")
            production_days = _finite_number(planning.cell(row=row, column=17).value, f"Q{row}")

            # M/P/Q không được âm. O (=p_need của engine) CÓ THỂ âm khi tồn
            # vượt nhu cầu (thừa hàng) — engine ghi p_need thô, không kẹp 0.
            for label, value in (("M", target_stock), ("P", planned), ("Q", production_days)):
                if value < -EPS:
                    raise RuntimeError(f"{label}{row} của mã {code} không được âm: {value}.")

            if per_shift <= 0 or shifts_per_day <= 0:
                if planned > EPS:
                    raise RuntimeError(
                        f"Mã {code} có P>0 nhưng E/I không hợp lệ: E={per_shift}, I={shifts_per_day}."
                    )
            else:
                # O = p_need của weekly_planning_engine (nguồn sự thật ghi workbook):
                #   - Không nợ: O = FC - tồn_sổ(K) + tồn_cuối(M).
                #   - Có nợ: theo Debt mode, KHÔNG cộng M:
                #       SUBTRACT_BOOK_ON_DEBT: O = FC + nợ - tồn_sổ(K)
                #       IGNORE_BOOK_ON_DEBT:   O = FC + nợ
                # Verifier tự lập không đọc Debt mode nên chấp nhận cả hai nhánh nợ.
                if debt > EPS:
                    candidates = (
                        forecast + debt - book_stock,
                        forecast + debt,
                    )
                    if not any(
                        math.isclose(required, cand, rel_tol=1e-9, abs_tol=1e-5)
                        for cand in candidates
                    ):
                        raise RuntimeError(
                            f"O{row} mã {code} sai: {required}; cần "
                            f"{candidates[0]} (trừ tồn sổ) hoặc {candidates[1]} (bỏ qua tồn sổ)."
                        )
                else:
                    expected_required = forecast - book_stock + target_stock
                    if not math.isclose(required, expected_required, rel_tol=1e-9, abs_tol=1e-5):
                        raise RuntimeError(
                            f"O{row} mã {code} sai: {required}; cần {expected_required}."
                        )

                # P là sản lượng CAM KẾT = min(ROUNDUP(O/mẻ hoặc /ca), đã xếp lịch).
                # Do đó P nằm trong [0, ROUNDUP(O)] chứ không nhất thiết bằng ROUNDUP(O)
                # (thiếu capacity thì P < ROUNDUP và phần thiếu là carryover). Ràng buộc
                # P = tổng SX ngày và cân bằng khối lượng được kiểm ở phần report bên dưới.
                base_qty = batch if is_sugar_classification(classification) else per_shift
                if planned > EPS and base_qty <= 0:
                    raise RuntimeError(f"Mã {code} có quantum <=0 nhưng P={planned}.")
                if base_qty > 0:
                    rounded_upper = (
                        0.0
                        if required <= EPS
                        else math.ceil(required / base_qty - 1e-12) * base_qty
                    )
                    if planned > rounded_upper + 1e-5:
                        raise RuntimeError(
                            f"P{row} mã {code}={planned} vượt ROUNDUP(O)={rounded_upper}."
                        )

                expected_q = planned / per_shift / shifts_per_day
                if not math.isclose(production_days, expected_q, rel_tol=1e-9, abs_tol=1e-6):
                    raise RuntimeError(
                        f"Q{row} mã {code} sai: {production_days}; cần {expected_q}."
                    )

            daily_values = []
            for offset, current_day in enumerate(headers):
                column = START_COLUMN + offset
                qty = _finite_number(
                    planning.cell(row=row, column=column).value,
                    f"{get_column_letter(column)}{row}",
                )
                if qty < -EPS:
                    raise RuntimeError(
                        f"Lịch mã {code} ngày {_header_day(current_day)} âm: {qty}."
                    )
                daily_values.append(qty)

            total_scheduled = sum(daily_values)
            report_balance = None
            if isinstance(schedule_report, dict):
                report_balance = (schedule_report.get("mass_balance") or {}).get(code)

            if report_balance is not None:
                report_planned = _finite_number(
                    report_balance.get("planned_qty"),
                    f"report {code}.planned_qty",
                )
                report_scheduled = _finite_number(
                    report_balance.get("scheduled_qty"),
                    f"report {code}.scheduled_qty",
                )
                carryover_qty = _finite_number(
                    report_balance.get("carryover_qty"),
                    f"report {code}.carryover_qty",
                )
                for label, value in (
                    ("planned_qty", report_planned),
                    ("scheduled_qty", report_scheduled),
                    ("carryover_qty", carryover_qty),
                ):
                    if value < -MASS_EPS:
                        raise RuntimeError(
                            f"Report mã {code} có {label} âm: {value}."
                        )
                if total_scheduled > planned + MASS_EPS:
                    raise RuntimeError(
                        f"Mã {code} scheduled={total_scheduled} vượt P={planned}; "
                        "carryover âm không được phép che sản xuất vượt kế hoạch."
                    )
                if report_scheduled > report_planned + MASS_EPS:
                    raise RuntimeError(
                        f"Report mã {code} scheduled={report_scheduled} vượt P={report_planned}."
                    )
                if carryover_qty > report_planned + MASS_EPS:
                    raise RuntimeError(
                        f"Report mã {code} carryover={carryover_qty} vượt P={report_planned}."
                    )
                if not math.isclose(report_planned, planned, rel_tol=1e-9, abs_tol=MASS_EPS):
                    raise RuntimeError(f"Report P mã {code}={report_planned} khác workbook P={planned}.")
                if not math.isclose(report_scheduled, total_scheduled, rel_tol=1e-9, abs_tol=MASS_EPS):
                    raise RuntimeError(
                        f"Report scheduled mã {code}={report_scheduled} khác workbook={total_scheduled}."
                    )
                if not math.isclose(total_scheduled + carryover_qty, planned, rel_tol=1e-9, abs_tol=MASS_EPS):
                    raise RuntimeError(
                        f"Mass balance mã {code}: scheduled {total_scheduled} + carryover {carryover_qty} != P {planned}."
                    )
            elif not math.isclose(total_scheduled, planned, rel_tol=1e-9, abs_tol=MASS_EPS):
                raise RuntimeError(
                    f"Tổng SX ngày của mã {code} = {total_scheduled} khác P={planned}. "
                    "Cần schedule report có provenance để xác nhận carryover hợp lệ."
                )

            resource = SHARED_RESOURCE if line in SHARED_LINES else line
            if resource:
                existing_capacity = resource_capacity.get(resource)
                if existing_capacity is None:
                    resource_capacity[resource] = shifts_per_day
                elif resource == SHARED_RESOURCE or line not in {"RGB", "Galon"}:
                    resource_capacity[resource] = min(existing_capacity, shifts_per_day)
                else:
                    resource_capacity[resource] = max(existing_capacity, shifts_per_day)

                if per_shift > EPS:
                    for current_day, qty in zip(headers, daily_values):
                        key = (resource, current_day)
                        daily_resource_usage[key] = daily_resource_usage.get(key, 0.0) + qty / per_shift

            planning_data[code] = {
                "line": line,
                "per_shift": per_shift,
                "daily_values": daily_values,
                "planned": planned,
                "scheduled": total_scheduled,
            }
            checked_schedule_rows += 1

        for (resource, current_day), used_shifts in daily_resource_usage.items():
            capacity = resource_capacity.get(resource, 0.0)
            if used_shifts > capacity + 1e-6:
                raise RuntimeError(
                    f"Resource {resource} ngày {_header_day(current_day)} có ít nhất "
                    f"{used_shifts:.3f} ca sản xuất > capacity {capacity:.3f}; chưa tính setup."
                )

        shared_codes = [
            code
            for code, item in planning_data.items()
            if not code.startswith("__")
            and item["line"] in SHARED_LINES
            and item["scheduled"] > EPS
        ]
        if len(shared_codes) > 1:
            shared_capacity = resource_capacity[SHARED_RESOURCE]
            _, shared_meta = _find_shared_resource_info(schedule_report)
            has_timeline = isinstance((shared_meta or {}).get("timeline"), list) and (
                shared_meta or {}
            ).get("timeline")
            if has_timeline:
                # Report có timeline production/setup: kiểm chứng nghiêm ngặt
                # không chồng lấn, đủ 0,5 ca setup khi đổi mã, và khớp workbook.
                _validate_shared_timeline(
                    schedule_report,
                    headers,
                    planning_data,
                    shared_capacity,
                )
            else:
                # Report hiện hành chưa xuất timeline: hậu kiểm tự lập từ workbook
                # (ràng buộc capacity/ngày đã kiểm ở trên; ở đây kiểm chặn dưới
                # theo tháng có tính setup).
                _validate_shared_from_workbook(
                    headers,
                    planning_data,
                    shared_capacity,
                )

        if schedule_report is not None:
            _validate_report_provenance(schedule_report, plan_year, plan_month)

        print(
            f"[VERIFY] {stock_checked} mã J=Ton_kho!D, K=SUM(Ton_kho!E:H) đúng; "
            f"{selector!r} -> FC!{get_column_letter(source_column)}, {fc_checked} mã L đúng; "
            f"lịch {expected_start.splitlines()[0]} -> {expected_end.splitlines()[0]} đúng."
        )
        return {
            "selector": selector,
            "plan_month": plan_month,
            "plan_year": plan_year,
            "source_column": source_column,
            "stock_checked": stock_checked,
            "fc_checked": fc_checked,
            "schedule_checked": checked_schedule_rows,
            "publish_status": (
                schedule_report.get("publish_status")
                if isinstance(schedule_report, dict)
                else "verified_without_carryover_report"
            ),
        }
    finally:
        workbook.close()


def main():
    token = get_access_token()
    graph = GraphClient(token)
    site_id = graph.get_site_id()
    drive_id = graph.get_default_drive_id(site_id)
    dest_item = graph.get_item_by_path(drive_id, DEST_PATH)
    dest_bytes = graph.download_file(drive_id, dest_item["id"])
    verify_workbook(dest_bytes, schedule_report=load_schedule_report())


if __name__ == "__main__":
    main()
