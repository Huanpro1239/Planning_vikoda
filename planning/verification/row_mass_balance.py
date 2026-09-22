import math

from openpyxl.utils import get_column_letter

from stock import normalize_code
from planning.weekly_engine import is_sugar_classification

from .shared_machine import _finite_number, _header_day
from .workbook import START_COLUMN


def validate_planning_rows(planning, headers, schedule_report=None):
    EPS = 1e-7
    MASS_EPS = 1e-5
    validated_rows = []

    for row in range(2, planning.max_row + 1):
        code = normalize_code(planning.cell(row=row, column=1).value)
        if not code:
            continue
    
        batch = _finite_number(planning.cell(row=row, column=4).value, f"D{row}")
        per_shift = _finite_number(planning.cell(row=row, column=5).value, f"E{row}")
        line = str(planning.cell(row=row, column=6).value or "").strip()
        classification = str(planning.cell(row=row, column=8).value or "").strip()
        shifts_per_day = _finite_number(planning.cell(row=row, column=9).value, f"I{row}")
        _finite_number(planning.cell(row=row, column=10).value, f"J{row}")
        book_stock = _finite_number(planning.cell(row=row, column=11).value, f"K{row}")
        forecast = _finite_number(planning.cell(row=row, column=12).value, f"L{row}")
        target_stock = _finite_number(planning.cell(row=row, column=13).value, f"M{row}")
        debt = _finite_number(planning.cell(row=row, column=14).value, f"N{row}")
        required = _finite_number(planning.cell(row=row, column=15).value, f"O{row}")
        planned = _finite_number(planning.cell(row=row, column=16).value, f"P{row}")
        production_days = _finite_number(planning.cell(row=row, column=17).value, f"Q{row}")
    
        # M/P/Q không được âm. O (=p_need của engine) CÓ THỂ âm khi tồn
        # vượt nhu cầu (thừa hàng) — engine ghi p_need thô, không kẹp 0.
        for label, value in (("M", target_stock), ("P", planned), ("Q", production_days)):
            if value < -EPS:
                raise RuntimeError(f"{label}{row} của mã {code} không được âm: {value}.")
    
        if per_shift <= 0 or shifts_per_day <= 0:
            if planned > EPS:
                raise RuntimeError(
                    f"Mã {code} có P>0 nhưng E/I không hợp lệ: E={per_shift}, I={shifts_per_day}."
                )
        else:
            # O = p_need của planning.weekly_engine (nguồn sự thật ghi workbook):
            #   - Không nợ: O = FC - tồn_sổ(K) + tồn_cuối(M).
            #   - Có nợ: theo Debt mode, KHÔNG cộng M:
            #       SUBTRACT_BOOK_ON_DEBT: O = FC + nợ - tồn_sổ(K)
            #       IGNORE_BOOK_ON_DEBT:   O = FC + nợ
            # Verifier tự lập không đọc Debt mode nên chấp nhận cả hai nhánh nợ.
            if debt > EPS:
                candidates = (
                    forecast + debt - book_stock,
                    forecast + debt,
                )
                if not any(
                    math.isclose(required, cand, rel_tol=1e-9, abs_tol=1e-5)
                    for cand in candidates
                ):
                    raise RuntimeError(
                        f"O{row} mã {code} sai: {required}; cần "
                        f"{candidates[0]} (trừ tồn sổ) hoặc {candidates[1]} (bỏ qua tồn sổ)."
                    )
            else:
                expected_required = forecast - book_stock + target_stock
                if not math.isclose(required, expected_required, rel_tol=1e-9, abs_tol=1e-5):
                    raise RuntimeError(
                        f"O{row} mã {code} sai: {required}; cần {expected_required}."
                    )
    
            # P là sản lượng CAM KẾT = min(ROUNDUP(O/mẻ hoặc /ca), đã xếp lịch).
            # Do đó P nằm trong [0, ROUNDUP(O)] chứ không nhất thiết bằng ROUNDUP(O)
            # (thiếu capacity thì P < ROUNDUP và phần thiếu là carryover). Ràng buộc
            # P = tổng SX ngày và cân bằng khối lượng được kiểm ở phần report bên dưới.
            base_qty = batch if is_sugar_classification(classification) else per_shift
            if planned > EPS and base_qty <= 0:
                raise RuntimeError(f"Mã {code} có quantum <=0 nhưng P={planned}.")
            if base_qty > 0:
                rounded_upper = (
                    0.0
                    if required <= EPS
                    else math.ceil(required / base_qty - 1e-12) * base_qty
                )
                if planned > rounded_upper + 1e-5:
                    raise RuntimeError(
                        f"P{row} mã {code}={planned} vượt ROUNDUP(O)={rounded_upper}."
                    )
    
            expected_q = planned / per_shift / shifts_per_day
            if not math.isclose(production_days, expected_q, rel_tol=1e-9, abs_tol=1e-6):
                raise RuntimeError(
                    f"Q{row} mã {code} sai: {production_days}; cần {expected_q}."
                )
    
        daily_values = []
        for offset, current_day in enumerate(headers):
            column = START_COLUMN + offset
            qty = _finite_number(
                planning.cell(row=row, column=column).value,
                f"{get_column_letter(column)}{row}",
            )
            if qty < -EPS:
                raise RuntimeError(
                    f"Lịch mã {code} ngày {_header_day(current_day)} âm: {qty}."
                )
            daily_values.append(qty)
    
        total_scheduled = sum(daily_values)
        report_balance = None
        if isinstance(schedule_report, dict):
            report_balance = (schedule_report.get("mass_balance") or {}).get(code)
    
        if report_balance is not None:
            report_planned = _finite_number(
                report_balance.get("planned_qty"),
                f"report {code}.planned_qty",
            )
            report_scheduled = _finite_number(
                report_balance.get("scheduled_qty"),
                f"report {code}.scheduled_qty",
            )
            carryover_qty = _finite_number(
                report_balance.get("carryover_qty"),
                f"report {code}.carryover_qty",
            )
            for label, value in (
                ("planned_qty", report_planned),
                ("scheduled_qty", report_scheduled),
                ("carryover_qty", carryover_qty),
            ):
                if value < -MASS_EPS:
                    raise RuntimeError(
                        f"Report mã {code} có {label} âm: {value}."
                    )
            if total_scheduled > planned + MASS_EPS:
                raise RuntimeError(
                    f"Mã {code} scheduled={total_scheduled} vượt P={planned}; "
                    "carryover âm không được phép che sản xuất vượt kế hoạch."
                )
            if report_scheduled > report_planned + MASS_EPS:
                raise RuntimeError(
                    f"Report mã {code} scheduled={report_scheduled} vượt P={report_planned}."
                )
            if carryover_qty > report_planned + MASS_EPS:
                raise RuntimeError(
                    f"Report mã {code} carryover={carryover_qty} vượt P={report_planned}."
                )
            if not math.isclose(report_planned, planned, rel_tol=1e-9, abs_tol=MASS_EPS):
                raise RuntimeError(f"Report P mã {code}={report_planned} khác workbook P={planned}.")
            if not math.isclose(report_scheduled, total_scheduled, rel_tol=1e-9, abs_tol=MASS_EPS):
                raise RuntimeError(
                    f"Report scheduled mã {code}={report_scheduled} khác workbook={total_scheduled}."
                )
            if not math.isclose(total_scheduled + carryover_qty, planned, rel_tol=1e-9, abs_tol=MASS_EPS):
                raise RuntimeError(
                    f"Mass balance mã {code}: scheduled {total_scheduled} + carryover {carryover_qty} != P {planned}."
                )
        elif not math.isclose(total_scheduled, planned, rel_tol=1e-9, abs_tol=MASS_EPS):
            raise RuntimeError(
                f"Tổng SX ngày của mã {code} = {total_scheduled} khác P={planned}. "
                "Cần schedule report có provenance để xác nhận carryover hợp lệ."
            )
    
        validated_rows.append({
            "code": code,
            "line": line,
            "per_shift": per_shift,
            "shifts_per_day": shifts_per_day,
            "daily_values": daily_values,
            "planned": planned,
            "scheduled": total_scheduled,
        })

    return validated_rows
