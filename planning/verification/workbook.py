import re

from openpyxl.utils import get_column_letter

import sync_planning_fc as fc
from sync_planning_calendar import build_date_headers, parse_plan_month
from sync_planning_calendar_all_months import resolve_plan_year
from stock import normalize_code, to_number


PLANNING_SHEET = "Ke_hoach_SX"
STOCK_SHEET = "Ton_kho"
START_COLUMN = 19  # S


def validate_workbook_context(workbook, schedule_report=None, plan_year=None):
    for sheet_name in (fc.FC_SHEET, STOCK_SHEET, PLANNING_SHEET):
        if sheet_name not in workbook.sheetnames:
            raise RuntimeError(f"Không tìm thấy sheet {sheet_name!r}.")

    fc_sheet = workbook[fc.FC_SHEET]
    stock_sheet = workbook[STOCK_SHEET]
    planning = workbook[PLANNING_SHEET]

    # 1) J/K phải lấy trực tiếp từ Ton_kho theo mã sản phẩm.
    stock_values = {}
    for row in range(2, stock_sheet.max_row + 1):
        code = normalize_code(stock_sheet.cell(row=row, column=1).value)
        if not code:
            continue
        if code in stock_values:
            raise RuntimeError(f"Mã {code} bị lặp trong {STOCK_SHEET}!A.")
        actual = to_number(stock_sheet.cell(row=row, column=4).value, f"{STOCK_SHEET}!D{row}")
        book = sum(
            to_number(stock_sheet.cell(row=row, column=column).value, f"{STOCK_SHEET}!{get_column_letter(column)}{row}")
            for column in range(5, 9)
        )
        stock_values[code] = (actual, book)

    stock_mismatches = []
    stock_checked = 0
    for row_number, values in enumerate(
        planning.iter_rows(min_row=2, max_col=11, values_only=True),
        start=2,
    ):
        code = normalize_code(values[0])
        if not code:
            continue
        if code not in stock_values:
            raise RuntimeError(f"Mã {code} ở {PLANNING_SHEET}!A{row_number} không có trong {STOCK_SHEET}!A.")
        expected_actual, expected_book = stock_values[code]
        actual_j = values[9]
        actual_k = values[10]
        stock_checked += 1
        if not fc._values_equal(actual_j, expected_actual) or not fc._values_equal(actual_k, expected_book):
            stock_mismatches.append((code, actual_j, expected_actual, actual_k, expected_book))

    if stock_mismatches:
        preview = "; ".join(
            f"{code}: J={actual_j!r}/Ton_kho!D={expected_j!r}, K={actual_k!r}/SUM(E:H)={expected_k!r}"
            for code, actual_j, expected_j, actual_k, expected_k in stock_mismatches[:10]
        )
        raise RuntimeError(
            f"{PLANNING_SHEET}!J:K chưa theo {STOCK_SHEET} ({len(stock_mismatches)}/{stock_checked} mã sai): {preview}"
        )

    # 2) L phải lấy đúng cột tháng được chọn tại FC!R1.
    selector = fc_sheet[fc.FC_SELECTOR_CELL].value
    selector_key = fc._normalize_header(selector)
    plan_month = parse_plan_month(selector)
    # Ưu tiên năm được truyền vào; nếu không có thì lấy từ report
    # (plan_month dạng "YYYY-MM"); chỉ dùng đồng hồ hệ thống khi không còn
    # nguồn nào khác. Điều này tránh verify sai quanh mốc giao năm và giúp
    # test không phụ thuộc thời điểm chạy.
    if plan_year is None and isinstance(schedule_report, dict):
        reported = str(schedule_report.get("plan_month") or "")
        match = re.fullmatch(r"(\d{4})-(\d{2})", reported)
        if match and int(match.group(2)) == plan_month:
            plan_year = int(match.group(1))
    if plan_year is None:
        plan_year = resolve_plan_year(plan_month)

    matched_columns = []
    for column in range(fc.FC_HEADER_MIN_COL, fc.FC_HEADER_MAX_COL + 1):
        if fc._normalize_header(fc_sheet.cell(row=1, column=column).value) == selector_key:
            matched_columns.append(column)
    if len(matched_columns) != 1:
        raise RuntimeError(
            f"{fc.FC_SHEET}!{fc.FC_SELECTOR_CELL}={selector!r} phải khớp đúng 1 cột E:P, "
            f"hiện khớp {len(matched_columns)} cột."
        )
    source_column = matched_columns[0]

    fc_values = {}
    for values in fc_sheet.iter_rows(min_row=2, max_col=max(fc.FC_CODE_COL, source_column), values_only=True):
        code = normalize_code(values[fc.FC_CODE_COL - 1])
        if not code:
            continue
        fc_values[code] = values[source_column - 1]

    fc_mismatches = []
    fc_checked = 0
    for row_number, values in enumerate(
        planning.iter_rows(min_row=2, max_col=fc.PLANNING_TARGET_COL, values_only=True),
        start=2,
    ):
        code = normalize_code(values[fc.PLANNING_CODE_COL - 1])
        if not code:
            continue
        if code not in fc_values:
            raise RuntimeError(f"Mã {code} ở {PLANNING_SHEET}!A{row_number} không có trong FC!B.")
        expected = fc_values[code]
        actual = values[fc.PLANNING_TARGET_COL - 1]
        fc_checked += 1
        if not fc._values_equal(actual, expected):
            fc_mismatches.append((code, actual, expected))

    if fc_mismatches:
        preview = "; ".join(
            f"{code}: L={actual!r}, FC={expected!r}"
            for code, actual, expected in fc_mismatches[:10]
        )
        raise RuntimeError(
            f"{PLANNING_SHEET}!L chưa theo {selector!r} ({len(fc_mismatches)}/{fc_checked} mã sai): {preview}"
        )

    # 3) Header lịch phải đúng tháng và không còn cột legacy/dữ liệu dư.
    headers = build_date_headers(plan_year, plan_month)
    expected_start = headers[0]
    expected_end = headers[-1]
    end_column = START_COLUMN + len(headers) - 1
    after_end_column = end_column + 1

    actual_start = planning.cell(row=1, column=START_COLUMN).value
    actual_end = planning.cell(row=1, column=end_column).value
    if actual_start != expected_start or actual_end != expected_end:
        raise RuntimeError(
            f"Dải ngày sai cho {plan_month:02d}/{plan_year}: "
            f"S1={actual_start!r}, {get_column_letter(end_column)}1={actual_end!r}; "
            f"cần {expected_start!r} ... {expected_end!r}."
        )

    for offset, expected_header in enumerate(headers):
        column = START_COLUMN + offset
        actual_header = planning.cell(row=1, column=column).value
        if actual_header != expected_header:
            raise RuntimeError(
                f"Tiêu đề ngày sai tại {get_column_letter(column)}1: "
                f"{actual_header!r}; cần {expected_header!r}."
            )

    legacy_labels = {"Kỳ kế hoạch", "Tổng SX", "Chênh lệch (SX-P)"}
    for column in range(START_COLUMN, max(after_end_column + 4, 53)):
        value = planning.cell(row=1, column=column).value
        if value in legacy_labels:
            raise RuntimeError(
                f"Còn cột legacy {value!r} tại {get_column_letter(column)}1."
            )

    stale_cells = []
    for row_number, values in enumerate(
        planning.iter_rows(
            min_row=1,
            min_col=after_end_column,
            max_col=after_end_column + 4,
            values_only=True,
        ),
        start=1,
    ):
        for offset, value in enumerate(values):
            if value not in (None, ""):
                stale_cells.append(
                    f"{get_column_letter(after_end_column + offset)}{row_number}={value!r}"
                )
                if len(stale_cells) >= 10:
                    break
        if len(stale_cells) >= 10:
            break
    if stale_cells:
        raise RuntimeError(
            "Còn dữ liệu sau ngày cuối tháng: " + "; ".join(stale_cells)
        )

    return {
        "planning": planning,
        "headers": headers,
        "plan_year": plan_year,
        "plan_month": plan_month,
        "stock_checked": stock_checked,
        "selector": selector,
        "source_column": source_column,
        "fc_checked": fc_checked,
        "expected_start": expected_start,
        "expected_end": expected_end,
    }
