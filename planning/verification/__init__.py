from io import BytesIO
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from planning.schedule_report import load_schedule_report
from sharepoint.client import GraphClient, get_access_token
from stock import DEST_PATH

from .workbook import (
    PLANNING_SHEET,
    START_COLUMN,
    STOCK_SHEET,
    validate_workbook_context,
)

from .row_mass_balance import validate_planning_rows
from .resources import validate_schedule_resources


def verify_workbook(workbook_bytes, schedule_report=None, plan_year=None):
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        workbook_context = validate_workbook_context(
            workbook,
            schedule_report=schedule_report,
            plan_year=plan_year,
        )
        planning = workbook_context["planning"]
        headers = workbook_context["headers"]
        plan_year = workbook_context["plan_year"]
        plan_month = workbook_context["plan_month"]
        stock_checked = workbook_context["stock_checked"]
        selector = workbook_context["selector"]
        source_column = workbook_context["source_column"]
        fc_checked = workbook_context["fc_checked"]
        expected_start = workbook_context["expected_start"]
        expected_end = workbook_context["expected_end"]

        # 4) Hậu kiểm độc lập dữ liệu kế hoạch và lịch ngày.
        validated_rows = validate_planning_rows(
            planning,
            headers,
            schedule_report=schedule_report,
        )
        checked_schedule_rows = validate_schedule_resources(
            validated_rows,
            headers,
            plan_year=plan_year,
            plan_month=plan_month,
            schedule_report=schedule_report,
        )

        print(
            f"[VERIFY] {stock_checked} mã J=Ton_kho!D, K=SUM(Ton_kho!E:H) đúng; "
            f"{selector!r} -> FC!{get_column_letter(source_column)}, {fc_checked} mã L đúng; "
            f"lịch {expected_start.splitlines()[0]} -> {expected_end.splitlines()[0]} đúng."
        )
        return {
            "selector": selector,
            "plan_month": plan_month,
            "plan_year": plan_year,
            "source_column": source_column,
            "stock_checked": stock_checked,
            "fc_checked": fc_checked,
            "schedule_checked": checked_schedule_rows,
            "publish_status": (
                schedule_report.get("publish_status")
                if isinstance(schedule_report, dict)
                else "verified_without_carryover_report"
            ),
        }
    finally:
        workbook.close()


def main():
    token = get_access_token()
    graph = GraphClient(token)
    site_id = graph.get_site_id()
    drive_id = graph.get_default_drive_id(site_id)
    dest_item = graph.get_item_by_path(drive_id, DEST_PATH)
    dest_bytes = graph.download_file(drive_id, dest_item["id"])
    verify_workbook(dest_bytes, schedule_report=load_schedule_report())


if __name__ == "__main__":
    main()
