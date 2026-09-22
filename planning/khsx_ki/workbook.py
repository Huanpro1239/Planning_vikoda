"""KHSX_ki workbook readers and patching."""

from __future__ import annotations

import calendar
from collections import defaultdict
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font

from stock import normalize_code

from .calendar import compute_standard_calendar_weeks
from .layout import (
    _apply_cell_style,
    _clear_cell,
    _detect_current_layout,
    _is_planning_signature_block,
    _safe_update_merge,
    _snapshot_cell_style,
    get_layout_spec,
)


KHSX_KI_SHEET = "KHSX_ki"
PLANNING_SHEET = "Ke_hoach_SX"
START_DAILY_COLUMN = 19
P_COLUMN = 16
EPS = 1e-6


def has_khsx_ki_sheet(workbook_bytes: bytes) -> bool:
    """Kiểm tra sự tồn tại của sheet KHSX_ki trong workbook mà không nạp toàn bộ dữ liệu."""
    wb = load_workbook(BytesIO(workbook_bytes), read_only=True)
    try:
        return KHSX_KI_SHEET in wb.sheetnames
    finally:
        wb.close()

def _read_planning_sheet_data(
    ws_kh,
    *,
    plan_year: int,
    plan_month: int,
) -> tuple[dict[int, int], dict[int, float], dict[tuple[int, int], float]]:
    """Đọc dữ liệu từ sheet Ke_hoach_SX: (seen_kh, kh_p, kh_daily).

    seen_kh: {code: row_idx} (phát hiện trùng mã ngay lập tức)
    kh_p: {code: committed_qty}
    kh_daily: {(code, day): qty}
    """
    days_in_month = calendar.monthrange(plan_year, plan_month)[1]
    seen_kh: dict[int, int] = {}
    kh_p: dict[int, float] = {}
    kh_daily: dict[tuple[int, int], float] = defaultdict(float)

    for r in range(2, ws_kh.max_row + 1):
        raw_a = str(ws_kh.cell(r, 1).value or "").strip().casefold()
        if "tổng cộng" in raw_a or "tong cong" in raw_a:
            break
        raw_code = ws_kh.cell(r, 1).value
        code_str = normalize_code(raw_code)
        if not code_str:
            continue
        code = int(code_str)
        if code in seen_kh:
            raise RuntimeError(
                f"Trùng mã SKU {code} trong sheet {PLANNING_SHEET} tại dòng {r} "
                f"(đã xuất hiện tại dòng {seen_kh[code]})."
            )
        seen_kh[code] = r
        kh_p[code] = float(ws_kh.cell(r, P_COLUMN).value or 0.0)
        for d in range(1, days_in_month + 1):
            col = START_DAILY_COLUMN + d - 1
            qty = float(ws_kh.cell(r, col).value or 0.0)
            if qty > EPS:
                kh_daily[(code, d)] += qty

    return seen_kh, kh_p, kh_daily

def _read_khsx_ki_skus(
    ws_ki,
    header_row: int = 6,
) -> tuple[dict[int, int], int | None]:
    """Đọc danh sách mã SKU và vị trí dòng TỔNG CỘNG từ sheet KHSX_ki.

    seen_ki: {code: row_idx} (phát hiện trùng mã ngay lập tức)
    total_row: index dòng TỔNG CỘNG (nếu có)
    """
    seen_ki: dict[int, int] = {}
    total_row: int | None = None

    for r in range(header_row + 1, ws_ki.max_row + 1):
        raw_a = str(ws_ki.cell(r, 1).value or "").strip().casefold()
        raw_b = str(ws_ki.cell(r, 2).value or "").strip().casefold()
        if "tổng cộng" in raw_a or "tong cong" in raw_a or "tổng cộng" in raw_b or "tong cong" in raw_b:
            total_row = r
            break
        raw_code = ws_ki.cell(r, 2).value
        code_str = normalize_code(raw_code)
        if not code_str:
            continue
        code = int(code_str)
        if code in seen_ki:
            raise RuntimeError(
                f"Trùng mã SKU {code} trong sheet {KHSX_KI_SHEET} tại dòng {r} "
                f"(đã xuất hiện tại dòng {seen_ki[code]})."
            )
        seen_ki[code] = r

    return seen_ki, total_row

def reconcile_skus(
    seen_kh: dict[int, int],
    seen_ki: dict[int, int],
) -> None:
    """Đối chiếu tập mã SKU giữa Ke_hoach_SX và KHSX_ki.

    Yêu cầu khớp 100% chính xác (set(kh_skus) == set(ki_skus)).
    Chặn đứng mọi trường hợp thiếu mã, thừa mã lạ, hoặc trùng mã.
    """
    kh_skus = set(seen_kh.keys())
    ki_skus = set(seen_ki.keys())

    if not kh_skus:
        raise RuntimeError(f"Không tìm thấy mã SKU nào trong sheet {PLANNING_SHEET}.")
    if not ki_skus:
        raise RuntimeError(f"Không tìm thấy mã SKU nào trong sheet {KHSX_KI_SHEET}.")

    missing_in_ki = kh_skus - ki_skus
    if missing_in_ki:
        details = [f"{c} (dòng {seen_kh[c]} {PLANNING_SHEET})" for c in sorted(missing_in_ki)]
        raise RuntimeError(
            f"Sheet {KHSX_KI_SHEET} thiếu {len(missing_in_ki)} mã SKU từ {PLANNING_SHEET}: "
            f"{', '.join(details)}"
        )

    extra_in_ki = ki_skus - kh_skus
    if extra_in_ki:
        details = [f"{c} (dòng {seen_ki[c]} {KHSX_KI_SHEET})" for c in sorted(extra_in_ki)]
        raise RuntimeError(
            f"Sheet {KHSX_KI_SHEET} chứa {len(extra_in_ki)} mã SKU không có trong {PLANNING_SHEET}: "
            f"{', '.join(details)}"
        )

def _read_daily_plan_from_analysis_or_sheet(
    workbook_bytes: bytes,
    weekly_analysis: Any | None,
    *,
    plan_year: int,
    plan_month: int,
) -> tuple[dict[tuple[int, int], float], dict[int, float]]:
    """Hàm tương thích đọc kế hoạch ngày và sản lượng cam kết."""
    wb = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        ws_kh = wb[PLANNING_SHEET]
        _, kh_p, kh_daily = _read_planning_sheet_data(ws_kh, plan_year=plan_year, plan_month=plan_month)
        if weekly_analysis is not None and getattr(weekly_analysis, "daily_plan", None):
            for item in weekly_analysis.daily_plan:
                code = int(item.ma_sp)
                day = item.date.day
                kh_daily[(code, day)] = max(kh_daily[(code, day)], float(item.qty))
            for calc in getattr(weekly_analysis, "calculated", []):
                code = int(calc.input.ma_sp)
                kh_p[code] = float(calc.schedulable_qty)
        return kh_daily, kh_p
    finally:
        wb.close()

def patch_khsx_ki_workbook(
    workbook_bytes: bytes,
    *,
    weekly_analysis: Any | None = None,
    plan_year: int,
    plan_month: int,
) -> tuple[bytes, dict[str, Any]]:
    """Tổng hợp lịch ngày từ Ke_hoach_SX vào sheet KHSX_ki, bảo toàn định dạng và cấu trúc."""
    wb = load_workbook(BytesIO(workbook_bytes), data_only=False)
    try:
        if KHSX_KI_SHEET not in wb.sheetnames:
            return workbook_bytes, {"ok": False, "reason": f"Thiếu sheet {KHSX_KI_SHEET}"}
        if PLANNING_SHEET not in wb.sheetnames:
            return workbook_bytes, {"ok": False, "reason": f"Thiếu sheet {PLANNING_SHEET}"}

        ws_ki = wb[KHSX_KI_SHEET]
        ws_kh = wb[PLANNING_SHEET]

        # 1. Đọc dữ liệu Ke_hoach_SX và KHSX_ki, kiểm tra tính duy nhất
        seen_kh, kh_p, kh_daily = _read_planning_sheet_data(
            ws_kh,
            plan_year=plan_year,
            plan_month=plan_month,
        )
        if weekly_analysis is not None and getattr(weekly_analysis, "daily_plan", None):
            for item in weekly_analysis.daily_plan:
                code = int(item.ma_sp)
                day = item.date.day
                kh_daily[(code, day)] = max(kh_daily[(code, day)], float(item.qty))
            for calc in getattr(weekly_analysis, "calculated", []):
                code = int(calc.input.ma_sp)
                if code not in seen_kh:
                    seen_kh[code] = len(seen_kh) + 2
                kh_p[code] = float(calc.schedulable_qty)

        seen_ki, total_row = _read_khsx_ki_skus(ws_ki, header_row=6)

        # 2. Đối chiếu toàn vẹn 100% SKU
        reconcile_skus(seen_kh, seen_ki)

        # 3. Lấy cấu trúc tuần chuẩn lịch Monday-Sunday
        std_weeks = compute_standard_calendar_weeks(plan_year, plan_month)
        weeks, total_col = get_layout_spec(plan_year, plan_month)
        is_6_weeks = (len(std_weeks) >= 6)
        target_end_col = 11 if is_6_weeks else 10

        # 4. Snapshot bố cục và định dạng hiện tại TRƯỚC KHI chỉnh sửa bất kỳ ô nào
        input_total_col = _detect_current_layout(ws_ki, header_row=6)
        input_is_6_weeks = (input_total_col == 11)

        # Cột mẫu tuần ổn định: Cột 8 (H - Tuần 4) luôn là cột tuần thực sự trong mọi tháng 4/5/6 tuần
        week_sample_col = 8

        # Snapshot styles cho header (dòng 6)
        total_header_style = _snapshot_cell_style(ws_ki.cell(6, input_total_col))
        week_header_style = _snapshot_cell_style(ws_ki.cell(6, week_sample_col))

        # Snapshot styles cho từng dòng SKU
        total_sku_styles = {
            r: _snapshot_cell_style(ws_ki.cell(r, input_total_col))
            for r in seen_ki.values()
        }
        week_sku_styles = {
            r: _snapshot_cell_style(ws_ki.cell(r, week_sample_col))
            for r in seen_ki.values()
        }

        # Snapshot styles cho dòng TỔNG CỘNG
        total_summary_style = None
        week_summary_style = None
        if total_row is not None:
            total_summary_style = _snapshot_cell_style(ws_ki.cell(total_row, input_total_col))
            week_summary_style = _snapshot_cell_style(ws_ki.cell(total_row, week_sample_col))

        # Đảm bảo cột tổng luôn giữ định dạng in đậm
        if total_header_style is not None and total_header_style.get("font") is not None:
            f = total_header_style["font"]
            total_header_style["font"] = Font(
                name=f.name, size=f.size, bold=True, italic=f.italic,
                vertAlign=f.vertAlign, underline=f.underline, strike=f.strike, color=f.color
            )
        if total_summary_style is not None and total_summary_style.get("font") is not None:
            f = total_summary_style["font"]
            total_summary_style["font"] = Font(
                name=f.name, size=f.size, bold=True, italic=f.italic,
                vertAlign=f.vertAlign, underline=f.underline, strike=f.strike, color=f.color
            )
        for r, s in total_sku_styles.items():
            if s is not None and s.get("font") is not None:
                f = s["font"]
                s["font"] = Font(
                    name=f.name, size=f.size, bold=True, italic=f.italic,
                    vertAlign=f.vertAlign, underline=f.underline, strike=f.strike, color=f.color
                )

        # Độ rộng cột ổn định
        if input_is_6_weeks:
            total_width = ws_ki.column_dimensions["K"].width or 17.0
            week_width = (
                ws_ki.column_dimensions["J"].width
                or ws_ki.column_dimensions["I"].width
                or 19.44140625
            )
        else:
            total_width = ws_ki.column_dimensions["J"].width or 17.0
            week_width = (
                ws_ki.column_dimensions["I"].width
                or ws_ki.column_dimensions["H"].width
                or 19.44140625
            )

        # 5. Cập nhật tiêu đề kỳ kế hoạch và ngày lập
        period_text = f"Kỳ kế hoạch: Tháng {plan_month}"
        date_text = f"Ngày lập: 01/{plan_month:02d}/{plan_year}"
        if ws_ki.cell(3, 1).value is not None:
            ws_ki.cell(3, 1).value = period_text
        if ws_ki.cell(4, 1).value is not None:
            ws_ki.cell(4, 1).value = date_text

        # Kiểm tra trước nếu có vùng gộp ô xung đột trong phạm vi dữ liệu bảng (dòng 6..total_row)
        for rng in list(ws_ki.merged_cells.ranges):
            if 6 <= rng.min_row <= (total_row or 6) and (rng.min_col <= 11 and rng.max_col >= 10):
                if total_row is not None and rng.min_row == total_row and rng.min_col == 1 and rng.max_col == 4:
                    continue
                raise ValueError(
                    f"Sheet '{ws_ki.title}' ô '{rng.coord}' có vùng gộp ô xung đột trong bảng kế hoạch sản xuất."
                )

        # 6. Cập nhật các ô hợp nhất cho tiêu đề và khối chữ ký/chân trang
        for rng in list(ws_ki.merged_cells.ranges):
            # Khối tiêu đề (dòng 1..5, bắt đầu ở cột 1)
            if rng.min_row <= 5 and rng.max_row <= 5 and rng.min_col == 1 and rng.max_col in (10, 11):
                _safe_update_merge(ws_ki, rng, target_end_col)
            # Khối chữ ký / chân trang dưới dòng tổng
            elif _is_planning_signature_block(ws_ki, rng, total_row):
                _safe_update_merge(ws_ki, rng, target_end_col)

        # 7. Thiết lập tiêu đề dòng 6 (Header) và định dạng cột tuần / tổng
        if is_6_weeks:
            _apply_cell_style(ws_ki.cell(6, 10), week_header_style)
            _apply_cell_style(ws_ki.cell(6, 11), total_header_style)
            ws_ki.cell(6, 11).value = "Tổng cộng"
            for w in weeks:
                ws_ki.cell(6, w["col"]).value = w["label"]
            ws_ki.column_dimensions["J"].width = week_width
            ws_ki.column_dimensions["K"].width = total_width
        else:
            _apply_cell_style(ws_ki.cell(6, 10), total_header_style)
            ws_ki.cell(6, 10).value = "Tổng cộng"
            for w in weeks:
                ws_ki.cell(6, w["col"]).value = w["label"]
            _clear_cell(ws_ki.cell(6, 11))
            ws_ki.column_dimensions["J"].width = total_width
            has_k_content = any(
                ws_ki.cell(r, 11).value is not None
                for r in range((total_row + 1) if total_row else 7, ws_ki.max_row + 1)
            )
            if not has_k_content and "K" in ws_ki.column_dimensions:
                del ws_ki.column_dimensions["K"]

        col_totals: dict[int, float] = defaultdict(float)
        sku_reports = []

        # 8. Ghi dữ liệu từng dòng SKU
        for code, r in seen_ki.items():
            if is_6_weeks:
                _apply_cell_style(ws_ki.cell(r, 10), week_sku_styles[r])
                _apply_cell_style(ws_ki.cell(r, 11), total_sku_styles[r])
            else:
                _apply_cell_style(ws_ki.cell(r, 10), total_sku_styles[r])
                _clear_cell(ws_ki.cell(r, 11))

            row_total = 0.0
            sku_week_vals = {}
            for w in weeks:
                col = w["col"]
                if w["days"]:
                    w_qty = sum(kh_daily.get((code, d), 0.0) for d in w["days"])
                    val = float(w_qty) if w_qty > EPS else 0.0
                else:
                    val = 0.0
                ws_ki.cell(r, col).value = val
                col_totals[col] += val
                row_total += val
                sku_week_vals[f"week_{w['week_num']}"] = val

            ws_ki.cell(r, total_col).value = float(row_total) if row_total > EPS else 0.0
            col_totals[total_col] += row_total
            sku_reports.append({
                "code": code,
                "weeks": sku_week_vals,
                "total": row_total,
                "committed_p": kh_p.get(code, 0.0),
            })

        # 9. Ghi dòng TỔNG CỘNG nếu có
        if total_row is not None:
            if is_6_weeks:
                _apply_cell_style(ws_ki.cell(total_row, 10), week_summary_style)
                _apply_cell_style(ws_ki.cell(total_row, 11), total_summary_style)
            else:
                _apply_cell_style(ws_ki.cell(total_row, 10), total_summary_style)
                _clear_cell(ws_ki.cell(total_row, 11))

            for w in weeks:
                col = w["col"]
                ws_ki.cell(total_row, col).value = float(col_totals[col])
            ws_ki.cell(total_row, total_col).value = float(col_totals[total_col])

        # Chú ý: Không xóa bất kỳ ô nào sau total_row để bảo toàn 100% ghi chú và dữ liệu ngoài bảng.

        out = BytesIO()
        wb.save(out)
        patched_bytes = out.getvalue()

        report = {
            "ok": True,
            "sheet": KHSX_KI_SHEET,
            "period": f"{plan_year:04d}-{plan_month:02d}",
            "checked_skus": len(seen_ki),
            "week_columns": [w["label"].replace("\n", " ") for w in weeks],
            "total_production": col_totals[total_col],
            "grand_total": col_totals[total_col],
            "num_weeks": len(std_weeks),
        }
        return patched_bytes, report
    finally:
        wb.close()

