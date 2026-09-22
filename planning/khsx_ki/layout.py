"""KHSX_ki column layout, style and merge helpers."""

from __future__ import annotations

from copy import copy
from typing import Any

from openpyxl.styles import Border, Font, PatternFill
from openpyxl.utils.cell import get_column_letter

from .calendar import compute_standard_calendar_weeks


def get_layout_spec(
    plan_year: int,
    plan_month: int,
) -> tuple[list[dict[str, Any]], int]:
    """Xác định cấu trúc cột tuần và cột Tổng cộng chuẩn theo lịch.

    Trả về (weeks, total_col):
    - Tháng 6 tuần: 6 cột tuần (E..J, col 5..10), cột Tổng cộng là K (col 11).
    - Tháng 5 tuần: 5 cột tuần (E..I, col 5..9), cột Tổng cộng là J (col 10).
    - Tháng 4 tuần: 4 cột tuần active (E..H, col 5..8), cột I (col 9) là tuần 5 rỗng,
                    cột Tổng cộng là J (col 10).
    """
    std_weeks = compute_standard_calendar_weeks(plan_year, plan_month)
    w_count = len(std_weeks)
    if w_count >= 6:
        weeks = [dict(w, col=5 + i) for i, w in enumerate(std_weeks)]
        total_col = 5 + w_count  # Col 11 cho 6 tuần
    elif w_count == 4:
        weeks = [dict(w, col=5 + i) for i, w in enumerate(std_weeks)]
        weeks.append({
            "week_num": 5,
            "start_day": 0,
            "end_day": 0,
            "days": [],
            "label": "Tuần 5\n-",
            "col": 9,
        })
        total_col = 10
    else:  # 5 tuần
        weeks = [dict(w, col=5 + i) for i, w in enumerate(std_weeks)]
        total_col = 10

    return weeks, total_col

def parse_week_columns(
    ws,
    *,
    plan_year: int,
    plan_month: int,
    header_row: int = 6,
) -> tuple[list[dict[str, Any]], int]:
    """Xác định danh sách các cột tuần và cột Tổng cộng thuần đọc (read-only).

    Hàm này TUYỆT ĐỐI KHÔNG ghi/sửa ô trên ws để an toàn 100% với read-only workbook.
    """
    return get_layout_spec(plan_year, plan_month)

def _snapshot_cell_style(cell) -> dict[str, Any] | None:
    """Lưu snapshot toàn bộ định dạng hiển thị của một ô."""
    if not getattr(cell, "has_style", False):
        return None
    return {
        "font": copy(cell.font) if cell.font else None,
        "border": copy(cell.border) if cell.border else None,
        "fill": copy(cell.fill) if cell.fill else None,
        "number_format": copy(cell.number_format) if cell.number_format else None,
        "protection": copy(cell.protection) if cell.protection else None,
        "alignment": copy(cell.alignment) if cell.alignment else None,
    }

def _apply_cell_style(dst_cell, style_snapshot: dict[str, Any] | None) -> None:
    """Áp dụng snapshot định dạng vào ô đích."""
    if not style_snapshot:
        return
    if style_snapshot.get("font") is not None:
        dst_cell.font = copy(style_snapshot["font"])
    if style_snapshot.get("border") is not None:
        dst_cell.border = copy(style_snapshot["border"])
    if style_snapshot.get("fill") is not None:
        dst_cell.fill = copy(style_snapshot["fill"])
    if style_snapshot.get("number_format") is not None:
        dst_cell.number_format = copy(style_snapshot["number_format"])
    if style_snapshot.get("protection") is not None:
        dst_cell.protection = copy(style_snapshot["protection"])
    if style_snapshot.get("alignment") is not None:
        dst_cell.alignment = copy(style_snapshot["alignment"])

def _copy_cell_style(src_cell, dst_cell) -> None:
    """Sao chép toàn bộ định dạng hiển thị từ một ô sang ô khác."""
    _apply_cell_style(dst_cell, _snapshot_cell_style(src_cell))

def _clear_cell(cell) -> None:
    """Xóa giá trị và đặt lại định dạng rỗng cho một ô."""
    cell.value = None
    cell.border = Border()
    cell.fill = PatternFill(fill_type=None)
    cell.font = Font()
    cell.number_format = "General"

def _detect_current_layout(ws_ki, header_row: int = 6) -> int:
    """Xác định cột tổng hiện tại trong sheet (10 hoặc 11) dựa trên tiêu đề bảng và dữ liệu bảng.

    Quy tắc nghiệp vụ:
    - Nhận diện bố cục từ tiêu đề tổng và các cột tuần hợp lệ trong vùng bảng.
    - Cột 10 (J) có 'Tổng cộng' -> trả về 10 (bố cục 4 hoặc 5 tuần).
    - Cột 11 (K) có 'Tổng cộng' hoặc Cột 10 (J) có tiêu đề 'Tuần 6' -> trả về 11 (bố cục 6 tuần).
    - Tuyệt đối không dùng độ rộng cột, font.bold hay merged ranges ngoài bảng làm căn cứ.
    - Với tiêu đề thiếu, hỏng hoặc mâu thuẫn: đối soát các dòng dữ liệu SKU (cột 11 có =SUM/dữ liệu hay rỗng hoàn toàn).
    """
    def _norm(val) -> str:
        return str(val or "").strip().lower()

    def _is_total(val) -> bool:
        v = _norm(val)
        return "tổng" in v or "total" in v

    def _is_w6(val) -> bool:
        v = _norm(val)
        return any(k in v for k in ("tuần 6", "tuan 6", "t6", "week 6", "w6"))

    c10_val = ws_ki.cell(header_row, 10).value
    c11_val = ws_ki.cell(header_row, 11).value

    # 1. Tín hiệu rõ ràng và không mâu thuẫn từ hàng tiêu đề
    if _is_total(c10_val) and not _is_total(c11_val):
        return 10
    if _is_total(c11_val) and not _is_total(c10_val):
        return 11
    if _is_w6(c10_val) and not _is_total(c10_val):
        return 11

    # 2. Xử lý tiêu đề mâu thuẫn (cả 10 và 11 đều là tổng) hoặc tiêu đề bị hỏng/thiếu:
    # Đối soát cấu trúc nội bộ của bảng qua các dòng SKU
    sku_start = header_row + 1
    sample_limit = min(ws_ki.max_row or (header_row + 10), header_row + 30)

    has_sum_formula_11 = False
    has_sum_formula_10 = False
    has_sku_data_11 = False

    for r in range(sku_start, sample_limit + 1):
        v10 = ws_ki.cell(r, 10).value
        v11 = ws_ki.cell(r, 11).value
        s10 = str(v10 or "").strip().upper()
        s11 = str(v11 or "").strip().upper()

        if "=SUM" in s11:
            has_sum_formula_11 = True
        if "=SUM" in s10:
            has_sum_formula_10 = True
        if v11 is not None and str(v11).strip() != "":
            has_sku_data_11 = True

    if has_sum_formula_11 and not has_sum_formula_10:
        return 11
    if has_sum_formula_10 and not has_sum_formula_11:
        return 10

    # Nếu cả hai đều ghi 'Tổng cộng': cột 11 có dữ liệu SKU chứng tỏ cột 11 là cột tổng
    if _is_total(c10_val) and _is_total(c11_val):
        if has_sum_formula_11 or has_sku_data_11:
            return 11
        return 10

    # Nếu tiêu đề bị hỏng/xóa sạch ở cả hai cột:
    if has_sku_data_11 or has_sum_formula_11:
        return 11

    # Mặc định an toàn cho mẫu chuẩn Vikoda (5 tuần)
    return 10

def _is_cell_occupied(ws, row: int, col: int, exclude_rng=None) -> bool:
    """Kiểm tra xem ô (row, col) có chứa giá trị, công thức, chú thích, định dạng người dùng
    hoặc thuộc một merged range khác không."""
    cell = ws.cell(row, col)
    # Giá trị hoặc công thức
    if cell.value is not None:
        return True
    # Chú thích
    if getattr(cell, "comment", None) is not None:
        return True
    # Định dạng tùy biến người dùng (fill, border không rỗng)
    if cell.fill and getattr(cell.fill, "fill_type", None) is not None:
        return True
    if cell.border:
        for side in ("left", "right", "top", "bottom"):
            border_side = getattr(cell.border, side, None)
            if border_side and getattr(border_side, "style", None) is not None:
                return True
    # Thuộc một merged range khác
    for rng in ws.merged_cells.ranges:
        if exclude_rng is not None and rng == exclude_rng:
            continue
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return True
    return False

def _safe_update_merge(ws, rng, target_end_col: int) -> bool:
    """Cập nhật an toàn một vùng merged cell hiện tại tới target_end_col.

    Quy tắc an toàn:
    - Không bao giờ tạo dải ô có cột kết thúc nhỏ hơn cột bắt đầu (target_end_col < rng.min_col).
    - Nếu đã đúng kích thước (target_end_col == rng.max_col) -> không làm gì, trả về True.
    - Khi mở rộng (target_end_col > rng.max_col): kiểm tra mọi ô trong phạm vi mở rộng.
      Nếu có ô chứa giá trị, công thức, ghi chú, định dạng hoặc thuộc vùng merge khác:
      -> Bảo toàn vùng người dùng: KHÔNG mở rộng, giữ nguyên vùng gộp hiện tại và trả về False.
    - Khi thu hẹp (target_end_col < rng.max_col): giải phóng vùng gộp an toàn.
    """
    if target_end_col < rng.min_col:
        return False

    if target_end_col == rng.max_col:
        return True

    min_r = rng.min_row
    max_r = rng.max_row
    start_col = rng.min_col
    current_end_col = rng.max_col

    # Kiểm tra xung đột khi mở rộng
    if target_end_col > current_end_col:
        for r in range(min_r, max_r + 1):
            for c in range(current_end_col + 1, target_end_col + 1):
                if _is_cell_occupied(ws, r, c, exclude_rng=rng):
                    # Phát hiện xung đột với nội dung người dùng: giữ nguyên vùng ngoài bảng
                    return False

    start_letter = get_column_letter(start_col)
    target_letter = get_column_letter(target_end_col)
    target_ref = f"{start_letter}{min_r}:{target_letter}{max_r}"

    ws.unmerge_cells(rng.coord)
    ws.merge_cells(target_ref)
    return True

def _is_planning_signature_block(ws, rng, total_row: int | None) -> bool:
    """Xác định xem vùng gộp bên dưới bảng có phải là khối chữ ký kế hoạch/người lập do module quản lý không."""
    min_footer_row = (total_row + 1) if total_row is not None else 7
    if rng.min_row < min_footer_row:
        return False
    # Khối chữ ký kế hoạch/người lập nằm ở góc dưới bên phải bảng:
    # Bắt đầu tại cột H (8) hoặc I (9), hiện kết thúc tại J (10) hoặc K (11)
    if rng.min_col not in (8, 9):
        return False
    if rng.max_col not in (10, 11):
        return False

    val_start = str(ws.cell(rng.min_row, rng.min_col).value or "").strip().lower()
    val_above = str(ws.cell(rng.min_row - 1, rng.min_col).value or "").strip().lower()
    keywords = (
        "kế hoạch", "ke hoach", "người lập", "nguoi lap",
        "lập biểu", "lap bieu", "lập bảng", "lap bang",
        "planning", "tp.kh", "tp. kh"
    )
    if any(k in val_start or k in val_above for k in keywords):
        return True

    # Kiểm tra nếu trên cùng dòng có các chức danh duyệt/ký khác (CEO, Giám đốc, Sản xuất, Mua hàng...)
    row_text = " ".join(
        str(ws.cell(rng.min_row, c).value or "").lower()
        for c in range(1, 12)
    )
    sig_titles = ("ceo", "giám đốc", "giam doc", "sản xuất", "san xuat", "mua hàng", "mua hang")
    if any(t in row_text for t in sig_titles):
        return True

    return False

