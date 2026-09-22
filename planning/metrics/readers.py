"""Workbook/source readers for Planning metrics."""

import hashlib
import json
import re
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from excel.openpyxl_io import safe_close_workbook
from stock import clean_number, normalize_code, read_conversion_factors, to_number

from .constants import (
    ACTUAL_CODE_COL,
    ACTUAL_CONSIGNMENT_COL,
    ACTUAL_RECEIPT_COL,
    COL_ACTUAL_STOCK,
    COL_BATCH,
    COL_BOOK_STOCK,
    COL_CLASSIFICATION,
    COL_CURRENT_DEBT,
    COL_FC,
    COL_PER_SHIFT,
    COL_SHIFTS_PER_DAY,
    DEBT_CODE_COL,
    DEBT_SHEET,
    DEBT_VALUE_COL,
    FC_SELECTOR_CELL,
    FC_SHEET,
    MASTER_LEADTIME_COL,
    MASTER_SHEET,
    PLANNING_SHEET,
    SYSTEM_CODE_COL,
    SYSTEM_RECEIPT_COL,
)


def find_report_date(worksheet):
    for row in range(1, min(worksheet.max_row, 6) + 1):
        for column in range(1, min(worksheet.max_column, 20) + 1):
            value = worksheet.cell(row=row, column=column).value
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value

    raise RuntimeError("Không tìm thấy ngày báo cáo trong file tồn thực tế.")


def read_actual_inputs(source_bytes):
    workbook = load_workbook(
        BytesIO(source_bytes),
        data_only=True,
        read_only=True,
    )
    try:
        worksheet = workbook.worksheets[0]
        report_date = find_report_date(worksheet)
        receipts = {}
        consignments = {}

        for row in range(1, worksheet.max_row + 1):
            code = normalize_code(
                worksheet.cell(row=row, column=ACTUAL_CODE_COL).value
            )
            if not code:
                continue
            if code in receipts:
                raise RuntimeError(
                    f"[Tồn thực tế] Mã {code} bị lặp trong cột C."
                )

            receipts[code] = clean_number(
                to_number(
                    worksheet.cell(row=row, column=ACTUAL_RECEIPT_COL).value,
                    f"G{row}",
                )
            )
            consignments[code] = clean_number(
                to_number(
                    worksheet.cell(row=row, column=ACTUAL_CONSIGNMENT_COL).value,
                    f"Q{row}",
                )
            )

        if not receipts:
            raise RuntimeError(
                "[Tồn thực tế] Không đọc được dữ liệu nhập tháng."
            )
        return report_date, receipts, consignments
    finally:
        safe_close_workbook(workbook)


def read_system_receipts(source_bytes, conversion_factors, planning_codes):
    workbook = load_workbook(
        BytesIO(source_bytes),
        data_only=True,
        read_only=True,
        keep_vba=True,
    )
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise RuntimeError("Không tìm thấy Sheet1 trong NXT_Vikoda.xlsm.")

        worksheet = workbook["Sheet1"]
        receipts = {}
        for row in range(1, worksheet.max_row + 1):
            code = normalize_code(
                worksheet.cell(row=row, column=SYSTEM_CODE_COL).value
            )
            if not code or code not in planning_codes:
                continue
            if code in receipts:
                raise RuntimeError(
                    f"[NXT_Vikoda] Mã {code} bị lặp trong cột B."
                )

            factor = conversion_factors.get(code)
            if factor is None or factor <= 0:
                raise RuntimeError(f"Không có quy cách hợp lệ cho mã {code}.")

            raw_value = to_number(
                worksheet.cell(row=row, column=SYSTEM_RECEIPT_COL).value,
                f"I{row}",
            )
            receipts[code] = clean_number(raw_value / factor)

        for code in planning_codes:
            receipts.setdefault(code, 0)
        return receipts
    finally:
        safe_close_workbook(workbook)


def parse_plan_month(selector):
    match = re.search(r"(\d{1,2})", str(selector or ""))
    if not match:
        raise RuntimeError(
            f"Không xác định được tháng từ {FC_SHEET}!{FC_SELECTOR_CELL}="
            f"{selector!r}."
        )
    month = int(match.group(1))
    if not 1 <= month <= 12:
        raise RuntimeError(f"Tháng kế hoạch không hợp lệ: {month}.")
    return month


def read_planning_rows_robust(dest_bytes):
    workbook = load_workbook(
        BytesIO(dest_bytes),
        data_only=True,
        read_only=True,
    )
    try:
        if PLANNING_SHEET not in workbook.sheetnames:
            raise RuntimeError(f"Không tìm thấy sheet {PLANNING_SHEET!r}.")
        if FC_SHEET not in workbook.sheetnames:
            raise RuntimeError(f"Không tìm thấy sheet {FC_SHEET!r}.")

        worksheet = workbook[PLANNING_SHEET]
        selector = workbook[FC_SHEET][FC_SELECTOR_CELL].value
        plan_month = parse_plan_month(selector)
        rows = {}

        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, max_col=18, values_only=True),
            start=2,
        ):
            code = normalize_code(values[0] if values else None)
            if not code:
                continue
            if code in rows:
                raise RuntimeError(
                    f"Mã {code} bị lặp trong {PLANNING_SHEET}!A."
                )

            rows[code] = {
                "row": row_number,
                "batch": to_number(
                    values[COL_BATCH - 1],
                    f"{PLANNING_SHEET}!D{row_number}",
                ),
                "per_shift": to_number(
                    values[COL_PER_SHIFT - 1],
                    f"{PLANNING_SHEET}!E{row_number}",
                ),
                "classification": str(
                    values[COL_CLASSIFICATION - 1] or ""
                ).strip(),
                "shifts_per_day": to_number(
                    values[COL_SHIFTS_PER_DAY - 1],
                    f"{PLANNING_SHEET}!I{row_number}",
                ),
                "actual_stock": to_number(
                    values[COL_ACTUAL_STOCK - 1],
                    f"{PLANNING_SHEET}!J{row_number}",
                ),
                "book_stock": to_number(
                    values[COL_BOOK_STOCK - 1],
                    f"{PLANNING_SHEET}!K{row_number}",
                ),
                "fc": to_number(
                    values[COL_FC - 1],
                    f"{PLANNING_SHEET}!L{row_number}",
                ),
                "current_debt": to_number(
                    values[COL_CURRENT_DEBT - 1],
                    f"{PLANNING_SHEET}!N{row_number}",
                ),
                "current_outputs": {
                    "expected_end_stock": values[12],
                    "warehouse_debt": values[13],
                    "required_production": values[14],
                    "rounded_production": values[15],
                    "production_days": values[16],
                    "production_start": values[17],
                },
            }

        if not rows:
            raise RuntimeError(
                f"Không đọc được dữ liệu trong {PLANNING_SHEET}."
            )

        print(
            f"[{PLANNING_SHEET}] Đọc {len(rows)} dòng bằng "
            "chế độ tương thích sheet không có dimension."
        )
        return selector, plan_month, rows
    finally:
        safe_close_workbook(workbook)


def read_debt_from_no_kho(dest_bytes):
    workbook = load_workbook(
        BytesIO(dest_bytes),
        data_only=True,
        read_only=True,
    )
    try:
        if DEBT_SHEET not in workbook.sheetnames:
            raise RuntimeError(f"Không tìm thấy sheet {DEBT_SHEET!r}.")

        worksheet = workbook[DEBT_SHEET]
        code_header = worksheet.cell(row=1, column=DEBT_CODE_COL).value
        debt_header = worksheet.cell(row=1, column=DEBT_VALUE_COL).value

        if str(code_header or "").strip().casefold() != "mã sản phẩm".casefold():
            raise RuntimeError(
                f"{DEBT_SHEET}!A1 phải là 'Mã Sản Phẩm', hiện là {code_header!r}."
            )
        if str(debt_header or "").strip().casefold() != "số lượng nợ".casefold():
            raise RuntimeError(
                f"{DEBT_SHEET}!D1 phải là 'Số lượng nợ', hiện là {debt_header!r}."
            )

        debts = {}
        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=2,
                max_col=DEBT_VALUE_COL,
                values_only=True,
            ),
            start=2,
        ):
            code = normalize_code(
                values[DEBT_CODE_COL - 1] if values else None
            )
            if not code:
                continue
            if code in debts:
                raise RuntimeError(
                    f"Mã {code} bị lặp trong {DEBT_SHEET}!A."
                )

            raw_debt = values[DEBT_VALUE_COL - 1]
            debt = 0 if raw_debt in (None, "") else to_number(
                raw_debt,
                f"{DEBT_SHEET}!D{row_number}",
            )
            debts[code] = clean_number(debt)

        if not debts:
            raise RuntimeError(
                f"Không đọc được dữ liệu từ {DEBT_SHEET}!A:D."
            )

        print(
            f"[{DEBT_SHEET}] Đọc {len(debts)} mã; "
            "Nợ kho lấy trực tiếp từ cột D."
        )
        return debts
    finally:
        safe_close_workbook(workbook)


def hash_debt_sheet(dest_bytes):
    workbook = load_workbook(
        BytesIO(dest_bytes),
        data_only=True,
        read_only=True,
    )
    try:
        if DEBT_SHEET not in workbook.sheetnames:
            return {}, ""
    finally:
        safe_close_workbook(workbook)

    debts = read_debt_from_no_kho(dest_bytes)
    payload = json.dumps(
        {key: debts[key] for key in sorted(debts)},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return debts, hashlib.sha256(payload).hexdigest()


def read_planning_rows(dest_bytes):
    selector, plan_month, rows = read_planning_rows_robust(dest_bytes)
    debts = read_debt_from_no_kho(dest_bytes)

    missing = sorted(set(rows) - set(debts))
    if missing:
        raise RuntimeError(
            f"Thiếu {len(missing)} mã Ke_hoach_SX trong {DEBT_SHEET}!A: "
            + ", ".join(missing[:10])
        )

    for code, row in rows.items():
        row["current_debt"] = debts[code]

    return selector, plan_month, rows


def read_leadtime_from_master(dest_bytes):
    workbook = load_workbook(
        BytesIO(dest_bytes),
        data_only=True,
        read_only=True,
    )
    try:
        if MASTER_SHEET not in workbook.sheetnames:
            raise RuntimeError(f"Không tìm thấy sheet {MASTER_SHEET!r}.")

        worksheet = workbook[MASTER_SHEET]
        header = worksheet.cell(row=1, column=MASTER_LEADTIME_COL).value
        if str(header or "").strip().casefold() != "leadtime":
            raise RuntimeError(
                f"{MASTER_SHEET}!J1 phải là 'Leadtime', hiện là {header!r}."
            )

        leadtimes = {}
        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=2,
                max_col=MASTER_LEADTIME_COL,
                values_only=True,
            ),
            start=2,
        ):
            code = normalize_code(values[0] if values else None)
            if not code:
                continue
            if code in leadtimes:
                raise RuntimeError(
                    f"Mã {code} bị lặp trong {MASTER_SHEET}!A."
                )

            raw_leadtime = values[MASTER_LEADTIME_COL - 1]
            if raw_leadtime in (None, ""):
                continue

            leadtime = to_number(
                raw_leadtime,
                f"{MASTER_SHEET}!J{row_number}",
            )
            if leadtime < 0:
                raise RuntimeError(
                    f"{MASTER_SHEET}!J{row_number} của mã {code} "
                    f"phải >= 0, hiện là {leadtime!r}."
                )
            leadtimes[code] = clean_number(leadtime)

        if not leadtimes:
            raise RuntimeError(
                f"Không đọc được Leadtime từ {MASTER_SHEET}!J:J."
            )

        print(
            f"[{MASTER_SHEET}] Đọc {len(leadtimes)} Leadtime từ cột J."
        )
        return leadtimes
    finally:
        safe_close_workbook(workbook)


def read_conversion_factors_and_leadtime(dest_bytes):
    conversion_factors, conversion_hash = read_conversion_factors(
        dest_bytes
    )
    leadtimes = read_leadtime_from_master(dest_bytes)
    return conversion_factors, conversion_hash, leadtimes
