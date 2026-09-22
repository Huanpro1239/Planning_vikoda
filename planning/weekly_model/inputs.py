"""Workbook input extraction and planning-input fingerprinting."""

from __future__ import annotations

import hashlib
import json
import math
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from planning.weekly_engine import WeeklyInputRow

from .policy import (
    DEBT_HEADER_NAMES,
    PROFILE_HEADER_NAMES,
    debt_mode,
    header_col,
    schedule_profile,
)


PLANNING_SHEET = "Ke_hoach_SX"
MASTER_SHEET = "Danh_muc"
START_COLUMN = 19
EPS = 1e-6


def code(value: Any) -> int:
    return int(float(value))


def num(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, bool):
        raise RuntimeError(
            "TRUE/FALSE không phải số kế hoạch."
        )
    result = float(value)
    if not math.isfinite(result):
        raise RuntimeError(
            f"Giá trị không hữu hạn: {value!r}"
        )
    return result


def read_inputs(
    workbook_bytes: bytes,
    *,
    plan_year: int,
    plan_month: int,
):
    workbook = load_workbook(
        BytesIO(workbook_bytes),
        data_only=True,
        read_only=True,
    )
    try:
        for name in (PLANNING_SHEET, MASTER_SHEET):
            if name not in workbook.sheetnames:
                raise RuntimeError(
                    f"Không tìm thấy sheet {name!r}."
                )

        planning = workbook[PLANNING_SHEET]
        master = workbook[MASTER_SHEET]

        debt_col = header_col(
            master,
            DEBT_HEADER_NAMES,
        )
        profile_col = header_col(
            master,
            PROFILE_HEADER_NAMES,
        )

        master_data: dict[int, dict[str, Any]] = {}
        for row in range(2, master.max_row + 1):
            raw = master.cell(row, 1).value
            if raw in (None, ""):
                continue

            product_code = code(raw)
            if product_code in master_data:
                raise RuntimeError(
                    f"Mã {product_code} bị lặp trong {MASTER_SHEET}!A."
                )

            master_data[product_code] = {
                "name": str(
                    master.cell(row, 2).value or ""
                ),
                "uom": str(
                    master.cell(row, 3).value or ""
                ),
                "batch": num(
                    master.cell(row, 4).value
                ),
                "per_shift": num(
                    master.cell(row, 5).value
                ),
                "line": str(
                    master.cell(row, 6).value or ""
                ).strip(),
                "group": str(
                    master.cell(row, 7).value or ""
                ).strip(),
                "classification": str(
                    master.cell(row, 8).value or ""
                ).strip(),
                "mold": num(
                    master.cell(row, 9).value
                ),
                "leadtime": num(
                    master.cell(row, 10).value
                ),
                "debt_mode": (
                    debt_mode(
                        master.cell(
                            row,
                            debt_col,
                        ).value
                    )
                    if debt_col
                    else None
                ),
                "profile": (
                    schedule_profile(
                        master.cell(
                            row,
                            profile_col,
                        ).value
                    )
                    if profile_col
                    else None
                ),
            }

        rows: list[WeeklyInputRow] = []
        warnings: list[dict[str, Any]] = []
        spread: set[int] = set()

        for row in range(2, planning.max_row + 1):
            raw = planning.cell(row, 1).value
            if raw in (None, ""):
                continue

            product_code = code(raw)
            meta = master_data.get(product_code)
            if meta is None:
                raise RuntimeError(
                    f"Mã {product_code} không có trong {MASTER_SHEET}."
                )

            debt = num(
                planning.cell(row, 14).value
            )
            mode = meta["debt_mode"]
            if mode is None:
                mode = "SUBTRACT_BOOK_ON_DEBT"
                if debt > EPS:
                    warnings.append({
                        "code": str(product_code),
                        "type": "MISSING_DEBT_MODE",
                        "message": (
                            "Thiếu Danh_muc!Debt mode; proposal tạm "
                            "dùng SUBTRACT_BOOK_ON_DEBT."
                        ),
                    })

            if meta["line"] == "Galon":
                if meta["profile"] == "SPREAD_NON_SUNDAY":
                    spread.add(product_code)
                elif meta["profile"] is None:
                    warnings.append({
                        "code": str(product_code),
                        "type": "MISSING_SCHEDULE_PROFILE",
                        "message": (
                            "Thiếu Danh_muc!Schedule profile cho Galon; "
                            "proposal tạm dùng CONTINUOUS."
                        ),
                    })

            fc = num(
                planning.cell(row, 12).value
            )
            rows.append(
                WeeklyInputRow(
                    source_row=row,
                    ma_sp=product_code,
                    ten_sp=meta["name"],
                    don_vi_tinh=meta["uom"],
                    sl_me=meta["batch"],
                    sl_ca=meta["per_shift"],
                    chuyen=meta["line"],
                    nhom_sp=meta["group"],
                    phan_loai=meta["classification"],
                    quy_cach=meta["mold"],
                    shifts_per_day=num(
                        planning.cell(row, 9).value
                    ),
                    ton_dau_thuc_te=num(
                        planning.cell(row, 10).value
                    ),
                    ton_dau_so_sach=num(
                        planning.cell(row, 11).value
                    ),
                    fc=fc,
                    ton_cuoi_du_kien=num(
                        planning.cell(row, 13).value
                    ),
                    no_kho=debt,
                    avg_daily_sales=(
                        fc / 26.0
                        if abs(fc) > EPS
                        else 0.0
                    ),
                    leadtime=meta["leadtime"],
                    debt_formula_mode=mode,
                )
            )

        return rows, frozenset(spread), warnings
    finally:
        workbook.close()


def compute_planning_inputs_hash(
    workbook_bytes: bytes,
) -> str:
    """Fingerprint user-editable Ke_hoach_SX planning inputs only."""
    workbook = load_workbook(
        BytesIO(workbook_bytes),
        data_only=True,
        read_only=True,
    )
    try:
        if PLANNING_SHEET not in workbook.sheetnames:
            return ""

        worksheet = workbook[PLANNING_SHEET]
        inputs = []
        for row in range(2, worksheet.max_row + 1):
            raw_code = worksheet.cell(row, 1).value
            if raw_code in (None, ""):
                continue
            inputs.append({
                "code": code(raw_code),
                "shifts_per_day": num(
                    worksheet.cell(row, 9).value
                ),
            })

        payload = json.dumps(
            inputs,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        return hashlib.sha256(payload).hexdigest()
    finally:
        workbook.close()
