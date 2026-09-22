"""Weekly schedule workbook patching."""

from __future__ import annotations

import calendar
from collections import defaultdict
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from .inputs import EPS, PLANNING_SHEET, START_COLUMN
from .schedule import (
    WeeklyAnalysis,
    committed_days,
    committed_qty,
    scheduled_by_code,
)


def patch_weekly_workbook(
    workbook_bytes: bytes,
    analysis: WeeklyAnalysis,
) -> bytes:
    workbook = load_workbook(
        BytesIO(workbook_bytes),
        data_only=False,
    )
    try:
        worksheet = workbook[PLANNING_SHEET]
        lookup: dict[tuple[int, date], float] = defaultdict(float)
        for item in analysis.daily_plan:
            lookup[
                (item.ma_sp, item.date)
            ] += float(item.qty)

        days = calendar.monthrange(
            analysis.period_year,
            analysis.period_month,
        )[1]
        committed = scheduled_by_code(
            analysis.daily_plan
        )

        for calc in analysis.calculated:
            row = calc.input.source_row
            committed_value = committed_qty(
                calc,
                committed,
            )
            committed_day_value = committed_days(
                calc,
                committed,
            )

            worksheet.cell(
                row,
                15,
                calc.p_need,
            )
            worksheet.cell(
                row,
                16,
                committed_value,
            )
            worksheet.cell(
                row,
                17,
                committed_day_value,
            )
            worksheet.cell(
                row,
                18,
                calc.start_datetime,
            )

            for day in range(1, days + 1):
                qty = lookup.get(
                    (
                        calc.input.ma_sp,
                        date(
                            analysis.period_year,
                            analysis.period_month,
                            day,
                        ),
                    ),
                    0.0,
                )
                worksheet.cell(
                    row,
                    START_COLUMN + day - 1,
                ).value = (
                    qty
                    if qty > EPS
                    else None
                )

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
