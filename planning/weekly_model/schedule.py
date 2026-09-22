"""Weekly schedule analysis built on the pure weekly_engine."""

from __future__ import annotations

import calendar
import math
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

from planning.weekly_engine import (
    PlannerPolicy,
    WeeklyCalculatedRow,
    build_daily_plan,
    calculate_rows,
)

from .inputs import (
    EPS,
    PLANNING_SHEET,
    START_COLUMN,
    read_inputs,
)


ENGINE_VERSION = "ke_hoach_sx_tuan_v5_khsx_ki_20260908"


@dataclass(frozen=True)
class WeeklyAnalysis:
    calculated: list[WeeklyCalculatedRow]
    daily_plan: list[Any]
    policy_warnings: list[dict[str, Any]]
    changed_cells: int
    period_year: int
    period_month: int


def same_value(current: Any, target: Any) -> bool:
    if target is None:
        return current in (None, "")

    if isinstance(target, datetime):
        if isinstance(current, datetime):
            return (
                abs(
                    (current - target).total_seconds()
                )
                < 1
            )
        if isinstance(current, date):
            return current == target.date()
        if isinstance(current, (int, float)):
            return math.isclose(
                float(current),
                float(to_excel(target)),
                rel_tol=0.0,
                abs_tol=1.0 / 86400.0,
            )
        return False

    if (
        isinstance(current, (int, float))
        and isinstance(target, (int, float))
    ):
        return math.isclose(
            float(current),
            float(target),
            rel_tol=1e-10,
            abs_tol=1e-7,
        )

    return current == target


def scheduled_by_code(
    plan: list[Any],
) -> dict[int, float]:
    result: dict[int, float] = defaultdict(float)
    for item in plan:
        result[int(item.ma_sp)] += float(item.qty)
    return result


def committed_qty(
    calc: WeeklyCalculatedRow,
    scheduled: dict[int, float],
) -> float:
    """Quantity actually committed to the daily plan."""
    value = max(
        0.0,
        float(
            scheduled.get(
                calc.input.ma_sp,
                0.0,
            )
        ),
    )
    return min(
        calc.schedulable_qty,
        value,
    )


def committed_days(
    calc: WeeklyCalculatedRow,
    scheduled: dict[int, float],
) -> float:
    committed = committed_qty(
        calc,
        scheduled,
    )
    return (
        committed
        / calc.input.sl_ca
        / calc.input.shifts_per_day
    )


def analyze_weekly_workbook(
    workbook_bytes: bytes,
    *,
    plan_year: int,
    plan_month: int,
) -> WeeklyAnalysis:
    rows, spread, warnings = read_inputs(
        workbook_bytes,
        plan_year=plan_year,
        plan_month=plan_month,
    )
    calculated = calculate_rows(
        rows,
        period_year=plan_year,
        period_month=plan_month,
    )
    daily = build_daily_plan(
        calculated,
        policy=PlannerPolicy(
            spread_product_codes=spread,
            allow_capacity_trim=True,
        ),
    )

    lookup: dict[tuple[int, date], float] = defaultdict(float)
    for item in daily:
        lookup[
            (item.ma_sp, item.date)
        ] += float(item.qty)

    committed = scheduled_by_code(daily)

    workbook = load_workbook(
        BytesIO(workbook_bytes),
        data_only=True,
        read_only=True,
    )
    try:
        worksheet = workbook[PLANNING_SHEET]
        changed = 0
        days = calendar.monthrange(
            plan_year,
            plan_month,
        )[1]

        for calc in calculated:
            row = calc.input.source_row
            committed_value = committed_qty(
                calc,
                committed,
            )
            committed_day_value = committed_days(
                calc,
                committed,
            )

            for column, target in (
                (15, calc.p_need),
                (16, committed_value),
                (17, committed_day_value),
                (18, calc.start_datetime),
            ):
                if not same_value(
                    worksheet.cell(
                        row,
                        column,
                    ).value,
                    target,
                ):
                    changed += 1

            for day in range(1, days + 1):
                target = (
                    lookup.get(
                        (
                            calc.input.ma_sp,
                            date(
                                plan_year,
                                plan_month,
                                day,
                            ),
                        ),
                        0.0,
                    )
                    or None
                )
                if not same_value(
                    worksheet.cell(
                        row,
                        START_COLUMN + day - 1,
                    ).value,
                    target,
                ):
                    changed += 1

        return WeeklyAnalysis(
            calculated,
            daily,
            warnings,
            changed,
            plan_year,
            plan_month,
        )
    finally:
        workbook.close()
