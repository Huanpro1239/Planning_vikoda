import calendar
import json
import math
import re
import zipfile
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook

from sync_planning_fc import (
    FC_SELECTOR_CELL,
    FC_SHEET,
    PLANNING_SHEET,
    _find_sheet_xml_path,
    _get_or_create_cell,
    _load_shared_strings,
    _read_cell_text,
)
from sync_stock import (
    DEST_PATH,
    SOURCE_ACTUAL_PATH,
    SOURCE_FACTORY_VIKODA_PATH,
    GraphClient,
    clean_number,
    get_access_token,
    normalize_code,
    read_conversion_factors,
    to_number,
)


PLANNING_STATE_FILE = Path("planning_state.json")
PLANNING_STATE_VERSION = 1

# DMSP!J (Leadtime) lấy từ file mẫu
# "File tính kế hoạch - BẢN CẢI TIẾN_V2(2).xlsm".
LEADTIME_BY_CODE = {
    "130100096": 3,
    "130100091": 3,
    "130100169": 3,
    "130200026": 1,
    "130100055": 4,
    "130100017": 3,
    "130100011": 4,
    "130100036": 4,
    "130100116": 3,
    "130100115": 3,
    "130100110": 3,
    "130100172": 3,
    "130100006": 3,
    "130100151": 3,
    "130100152": 3,
    "130100008": 3,
    "130100013": 4,
    "130100149": 3,
    "130300005": 4,
    "130300006": 4,
}

# Cột trong file báo cáo tồn thực tế.
ACTUAL_CODE_COL = 3          # C - MÃ SP VIKODA
ACTUAL_RECEIPT_COL = 7       # G - TỔNG NHẬP THÁNG
ACTUAL_CONSIGNMENT_COL = 17  # Q - SL HÀNG ĐÃ BÁN (GỬI KHO)

# Cột trong NXT_Vikoda.xlsm!Sheet1.
SYSTEM_CODE_COL = 2          # B - Mã vật tư
SYSTEM_RECEIPT_COL = 9       # I - Nhập trong kỳ

# Ke_hoach_SX.
COL_BATCH = 4                # D
COL_PER_SHIFT = 5            # E
COL_CLASSIFICATION = 8       # H
COL_SHIFTS_PER_DAY = 9       # I
COL_FC = 12                  # L

# Ton_kho.
STOCK_ACTUAL_COL = 4         # D
STOCK_BOOK_MIN_COL = 5       # E
STOCK_BOOK_MAX_COL = 8       # H

OUTPUT_COLUMNS = {
    "J": "actual_stock",
    "K": "book_stock",
    "M": "expected_end_stock",
    "N": "warehouse_debt",
    "O": "required_production",
    "P": "rounded_production",
    "Q": "production_days",
    "R": "production_start",
}


def load_planning_state():
    if not PLANNING_STATE_FILE.exists():
        return {
            "version": PLANNING_STATE_VERSION,
            "opening_debt_by_month": {},
            "opening_consignment_by_month": {},
        }

    try:
        data = json.loads(
            PLANNING_STATE_FILE.read_text(encoding="utf-8")
        )
    except Exception as exc:
        raise RuntimeError(
            f"Không đọc được {PLANNING_STATE_FILE}: {exc}"
        ) from exc

    data.setdefault("version", PLANNING_STATE_VERSION)
    data.setdefault("opening_debt_by_month", {})
    data.setdefault("opening_consignment_by_month", {})
    return data


def save_planning_state(state):
    state["version"] = PLANNING_STATE_VERSION
    PLANNING_STATE_FILE.write_text(
        json.dumps(state, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def _month_key(year, month):
    return f"{year:04d}-{month:02d}"


def _next_month_key(year, month):
    if month == 12:
        return _month_key(year + 1, 1)
    return _month_key(year, month + 1)


def _is_month_end(day):
    return day.day == calendar.monthrange(day.year, day.month)[1]


def _find_report_date(worksheet):
    for row in range(1, min(worksheet.max_row, 6) + 1):
        for column in range(1, min(worksheet.max_column, 20) + 1):
            value = worksheet.cell(row=row, column=column).value
            if isinstance(value, datetime):
                return value.date()
            if isinstance(value, date):
                return value

    raise RuntimeError(
        "Không tìm thấy ngày báo cáo trong file tồn thực tế."
    )


def read_actual_planning_source(source_bytes):
    workbook = load_workbook(
        BytesIO(source_bytes),
        data_only=True,
        read_only=True,
    )

    try:
        worksheet = workbook.worksheets[0]
        report_date = _find_report_date(worksheet)
        receipts = {}
        consignments = {}

        for row in range(1, worksheet.max_row + 1):
            code = normalize_code(
                worksheet.cell(row=row, column=ACTUAL_CODE_COL).value
            )
            if not code:
                continue

            if code in receipts:
                raise RuntimeError(
                    f"[Tồn thực tế] Mã {code} bị lặp trong cột C."
                )

            receipts[code] = clean_number(
                to_number(
                    worksheet.cell(
                        row=row,
                        column=ACTUAL_RECEIPT_COL,
                    ).value,
                    f"G{row}",
                )
            )
            consignments[code] = clean_number(
                to_number(
                    worksheet.cell(
                        row=row,
                        column=ACTUAL_CONSIGNMENT_COL,
                    ).value,
                    f"Q{row}",
                )
            )

        if not receipts:
            raise RuntimeError(
                "[Tồn thực tế] Không đọc được dữ liệu sản xuất tháng."
            )

        return report_date, receipts, consignments
    finally:
        workbook.close()


def read_system_receipts(source_bytes, conversion_factors):
    workbook = load_workbook(
        BytesIO(source_bytes),
        data_only=True,
        read_only=True,
        keep_vba=True,
    )

    try:
        if "Sheet1" not in workbook.sheetnames:
            raise RuntimeError(
                "Không tìm thấy Sheet1 trong NXT_Vikoda.xlsm."
            )

        worksheet = workbook["Sheet1"]
        receipts = {}

        for row in range(1, worksheet.max_row + 1):
            code = normalize_code(
                worksheet.cell(row=row, column=SYSTEM_CODE_COL).value
            )
            if not code:
                continue

            # NXT có nhiều mã ngoài 20 mã đang lập kế hoạch. Chỉ đọc
            # những mã đã có trong Danh_muc/conversion_factors.
            if code not in conversion_factors:
                continue

            if code in receipts:
                raise RuntimeError(
                    f"[NXT_Vikoda] Mã {code} bị lặp trong cột B."
                )

            factor = conversion_factors[code]
            if factor <= 0:
                raise RuntimeError(
                    f"Quy cách của mã {code} phải > 0 để quy đổi "
                    "Nhập trong kỳ."
                )

            raw_value = to_number(
                worksheet.cell(
                    row=row,
                    column=SYSTEM_RECEIPT_COL,
                ).value,
                f"I{row}",
            )
            receipts[code] = clean_number(raw_value / factor)

        if not receipts:
            raise RuntimeError(
                "[NXT_Vikoda] Không đọc được Nhập trong kỳ cho các mã kế hoạch."
            )

        return receipts
    finally:
        workbook.close()


def _parse_plan_month(selector):
    if selector is None:
        raise RuntimeError(f"{FC_SHEET}!{FC_SELECTOR_CELL} đang trống.")

    match = re.search(r"(\d{1,2})", str(selector))
    if not match:
        raise RuntimeError(
            f"Không xác định được tháng từ "
            f"{FC_SHEET}!{FC_SELECTOR_CELL}={selector!r}."
        )

    month = int(match.group(1))
    if month < 1 or month > 12:
        raise RuntimeError(f"Tháng kế hoạch không hợp lệ: {month}.")
    return month


def _resolve_plan_year(report_date, plan_month):
    year = report_date.year
    if report_date.month >= 10 and plan_month <= 3:
        year += 1
    elif report_date.month <= 3 and plan_month >= 10:
        year -= 1
    return year


def read_destination_inputs(dest_bytes):
    workbook = load_workbook(
        BytesIO(dest_bytes),
        data_only=True,
        read_only=True,
    )

    try:
        required = {PLANNING_SHEET, "Ton_kho", FC_SHEET}
        missing = required - set(workbook.sheetnames)
        if missing:
            raise RuntimeError(
                "Thiếu sheet trong Sắp kế hoạch.xlsx: "
                + ", ".join(sorted(missing))
            )

        planning = workbook[PLANNING_SHEET]
        stock = workbook["Ton_kho"]
        fc = workbook[FC_SHEET]

        selector = fc[FC_SELECTOR_CELL].value
        plan_month = _parse_plan_month(selector)

        stock_by_code = {}
        for row in range(2, stock.max_row + 1):
            code = normalize_code(stock.cell(row=row, column=1).value)
            if not code:
                continue

            actual_stock = to_number(
                stock.cell(row=row, column=STOCK_ACTUAL_COL).value,
                f"Ton_kho!D{row}",
            )
            book_stock = sum(
                to_number(
                    stock.cell(row=row, column=column).value,
                    f"Ton_kho!{column}{row}",
                )
                for column in range(
                    STOCK_BOOK_MIN_COL,
                    STOCK_BOOK_MAX_COL + 1,
                )
            )
            stock_by_code[code] = {
                "actual_stock": clean_number(actual_stock),
                "book_stock": clean_number(book_stock),
            }

        rows = {}
        for row in range(2, planning.max_row + 1):
            code = normalize_code(planning.cell(row=row, column=1).value)
            if not code:
                continue

            if code in rows:
                raise RuntimeError(
                    f"Mã {code} bị lặp trong {PLANNING_SHEET}!A."
                )
            if code not in stock_by_code:
                raise RuntimeError(
                    f"Mã {code} không tồn tại trong Ton_kho!A."
                )

            rows[code] = {
                "row": row,
                "batch": to_number(
                    planning.cell(row=row, column=COL_BATCH).value,
                    f"{PLANNING_SHEET}!D{row}",
                ),
                "per_shift": to_number(
                    planning.cell(row=row, column=COL_PER_SHIFT).value,
                    f"{PLANNING_SHEET}!E{row}",
                ),
                "classification": str(
                    planning.cell(
                        row=row,
                        column=COL_CLASSIFICATION,
                    ).value
                    or ""
                ).strip(),
                "shifts_per_day": to_number(
                    planning.cell(
                        row=row,
                        column=COL_SHIFTS_PER_DAY,
                    ).value,
                    f"{PLANNING_SHEET}!I{row}",
                ),
                "fc": to_number(
                    planning.cell(row=row, column=COL_FC).value,
                    f"{PLANNING_SHEET}!L{row}",
                ),
                **stock_by_code[code],
            }

        if not rows:
            raise RuntimeError(
                f"Không đọc được dữ liệu {PLANNING_SHEET}!A:L."
            )

        return selector, plan_month, rows
    finally:
        workbook.close()


def excel_roundup_integer(value):
    """ROUNDUP(value, 0): làm tròn ra xa số 0 như Excel."""
    if value > 0:
        return math.ceil(value)
    if value < 0:
        return math.floor(value)
    return 0


def calculate_product_metrics(
    *,
    fc,
    actual_stock,
    book_stock,
    opening_consignment,
    warehouse_debt,
    leadtime,
    batch,
    per_shift,
    shifts_per_day,
    classification,
    plan_year,
    plan_month,
):
    if batch <= 0:
        raise ValueError("Số lượng/mẻ phải > 0.")
    if per_shift <= 0:
        raise ValueError("Số lượng/ca phải > 0.")
    if shifts_per_day <= 0:
        raise ValueError("Số ca theo ngày phải > 0.")
    if leadtime < 0:
        raise ValueError("Leadtime không được âm.")

    # FC thang nay!N = Tổng FC Sales / 26.
    daily_fc = fc / 26 if fc else 0

    # FC thang nay!Q = Xuất TB/ngày * (Leadtime + 2).
    minimum_stock = daily_fc * (leadtime + 2)

    # FC thang nay!R = IF(Gửi kho đầu tháng > Tồn tối thiểu, 0, Tồn tối thiểu).
    expected_end_stock = (
        0 if opening_consignment > minimum_stock else minimum_stock
    )

    # Ke hoach SX tuan!P (công thức chuẩn ở P4):
    # IF(Nợ kho>0, FC+Nợ kho-Tồn sổ sách,
    #              FC+Tồn cuối dự kiến-Tồn sổ sách+Nợ kho)
    if warehouse_debt > 0:
        required_production = fc + warehouse_debt - book_stock
    else:
        required_production = (
            fc + expected_end_stock - book_stock + warehouse_debt
        )

    # Ke hoach SX tuan!Q.
    if required_production == 0:
        rounded_production = 0
    else:
        sugar_product = classification.casefold() == "có đường".casefold()
        rounding_base = batch if sugar_product else per_shift
        rounded_production = (
            excel_roundup_integer(required_production / rounding_base)
            * rounding_base
        )

    # Ke hoach SX tuan!R.
    production_days = rounded_production / per_shift / shifts_per_day

    # Ke hoach SX tuan!S -> Ke_hoach_SX!R.
    if daily_fc <= 0:
        production_start = None
    else:
        first_day = datetime(plan_year, plan_month, 1)
        calculated = first_day + timedelta(
            days=(actual_stock / daily_fc) - leadtime
        )
        production_start = max(first_day, calculated)

    return {
        "actual_stock": clean_number(actual_stock),
        "book_stock": clean_number(book_stock),
        "expected_end_stock": clean_number(expected_end_stock),
        "warehouse_debt": clean_number(warehouse_debt),
        "required_production": clean_number(required_production),
        "rounded_production": clean_number(rounded_production),
        "production_days": clean_number(production_days),
        "production_start": production_start,
    }


def _month_opening_values(state, kind, month_key, codes):
    section = state.get(kind, {})
    month_values = section.get(month_key)
    if month_values is None:
        return None

    return {
        code: float(month_values.get(code, 0) or 0)
        for code in codes
    }


def calculate_planning_metrics(
    *,
    report_date,
    actual_receipts,
    current_consignments,
    system_receipts,
    selector,
    plan_month,
    planning_rows,
    state,
):
    source_key = _month_key(report_date.year, report_date.month)
    plan_year = _resolve_plan_year(report_date, plan_month)
    plan_key = _month_key(plan_year, plan_month)
    codes = sorted(planning_rows)

    opening_debt = _month_opening_values(
        state,
        "opening_debt_by_month",
        source_key,
        codes,
    )
    if opening_debt is None:
        raise RuntimeError(
            f"Chưa có Nợ kho đầu tháng cho {source_key} trong "
            f"{PLANNING_STATE_FILE}."
        )

    current_debt = {}
    for code in codes:
        current_debt[code] = clean_number(
            max(
                opening_debt.get(code, 0)
                - actual_receipts.get(code, 0)
                + system_receipts.get(code, 0),
                0,
            )
        )

    if plan_key == source_key:
        opening_consignment = _month_opening_values(
            state,
            "opening_consignment_by_month",
            source_key,
            codes,
        )
        if opening_consignment is None:
            raise RuntimeError(
                f"Chưa có Gửi kho đầu tháng cho {source_key} trong "
                f"{PLANNING_STATE_FILE}."
            )
    else:
        # Khi lập tháng kế tiếp, số gửi kho hiện tại là ước tính đầu kỳ.
        opening_consignment = {
            code: float(current_consignments.get(code, 0) or 0)
            for code in codes
        }

    metrics = {}
    for code, row in planning_rows.items():
        leadtime = LEADTIME_BY_CODE.get(code)
        if leadtime is None:
            raise RuntimeError(
                f"Chưa cấu hình Leadtime cho mã {code}."
            )

        metrics[code] = calculate_product_metrics(
            fc=row["fc"],
            actual_stock=row["actual_stock"],
            book_stock=row["book_stock"],
            opening_consignment=opening_consignment.get(code, 0),
            warehouse_debt=current_debt.get(code, 0),
            leadtime=leadtime,
            batch=row["batch"],
            per_shift=row["per_shift"],
            shifts_per_day=row["shifts_per_day"],
            classification=row["classification"],
            plan_year=plan_year,
            plan_month=plan_month,
        )

    state_changed = False
    if _is_month_end(report_date):
        next_key = _next_month_key(report_date.year, report_date.month)
        debt_section = state.setdefault("opening_debt_by_month", {})
        consign_section = state.setdefault(
            "opening_consignment_by_month", {}
        )

        next_debt = {
            code: clean_number(current_debt.get(code, 0))
            for code in codes
        }
        next_consignment = {
            code: clean_number(current_consignments.get(code, 0))
            for code in codes
        }

        if debt_section.get(next_key) != next_debt:
            debt_section[next_key] = next_debt
            state_changed = True
        if consign_section.get(next_key) != next_consignment:
            consign_section[next_key] = next_consignment
            state_changed = True

    return {
        "selector": selector,
        "source_month": source_key,
        "plan_month": plan_key,
        "metrics": metrics,
        "state_changed": state_changed,
    }


def _datetime_to_excel_serial(value):
    if value is None:
        return None
    epoch = datetime(1899, 12, 30)
    return (value - epoch).total_seconds() / 86400


def _set_cell_value(row_element, row_number, column_letter, value):
    cell = _get_or_create_cell(row_element, row_number, column_letter)

    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)

    if value is None:
        return

    if isinstance(value, datetime):
        value = _datetime_to_excel_serial(value)

    namespace = etree.QName(cell).namespace
    value_node = etree.SubElement(cell, f"{{{namespace}}}v")

    value = clean_number(value)
    if isinstance(value, int):
        value_node.text = str(value)
    else:
        value_node.text = format(float(value), ".15g")


def patch_planning_metrics_workbook(workbook_bytes, metrics):
    source_buffer = BytesIO(workbook_bytes)
    output_buffer = BytesIO()

    with zipfile.ZipFile(source_buffer, "r") as source_zip:
        sheet_path = _find_sheet_xml_path(source_zip, PLANNING_SHEET)
        shared_strings = _load_shared_strings(source_zip)
        sheet_root = etree.fromstring(source_zip.read(sheet_path))

        seen = set()
        patched = 0
        rows = sheet_root.xpath(
            '//*[local-name()="sheetData"]/*[local-name()="row"]'
        )

        for row_element in rows:
            row_number_text = row_element.get("r")
            if not row_number_text:
                continue
            row_number = int(row_number_text)

            code = None
            for cell in row_element.xpath('./*[local-name()="c"]'):
                if cell.get("r") == f"A{row_number}":
                    code = normalize_code(
                        _read_cell_text(cell, shared_strings)
                    )
                    break

            if not code or code not in metrics:
                continue
            if code in seen:
                raise RuntimeError(
                    f"Mã {code} bị lặp trong XML {PLANNING_SHEET}!A."
                )
            seen.add(code)

            row_metrics = metrics[code]
            for column, key in OUTPUT_COLUMNS.items():
                _set_cell_value(
                    row_element,
                    row_number,
                    column,
                    row_metrics[key],
                )
            patched += 1

        if patched != len(metrics):
            missing = sorted(set(metrics) - seen)
            raise RuntimeError(
                f"Chỉ cập nhật được {patched}/{len(metrics)} mã vào "
                f"{PLANNING_SHEET}!J:K,M:R. Thiếu: "
                + ", ".join(missing[:10])
            )

        new_sheet_xml = etree.tostring(
            sheet_root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        with zipfile.ZipFile(
            output_buffer,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as output_zip:
            for item in source_zip.infolist():
                data = (
                    new_sheet_xml
                    if item.filename == sheet_path
                    else source_zip.read(item.filename)
                )
                output_zip.writestr(item, data)

    return output_buffer.getvalue(), patched


def main():
    token = get_access_token()
    graph = GraphClient(token)

    site_id = graph.get_site_id()
    drive_id = graph.get_default_drive_id(site_id)

    dest_item = graph.get_item_by_path(drive_id, DEST_PATH)
    actual_item = graph.get_item_by_path(drive_id, SOURCE_ACTUAL_PATH)
    system_item = graph.get_item_by_path(
        drive_id,
        SOURCE_FACTORY_VIKODA_PATH,
    )

    dest_bytes = graph.download_file(drive_id, dest_item["id"])
    actual_bytes = graph.download_file(drive_id, actual_item["id"])
    system_bytes = graph.download_file(drive_id, system_item["id"])

    conversion_factors, _ = read_conversion_factors(dest_bytes)
    report_date, actual_receipts, consignments = (
        read_actual_planning_source(actual_bytes)
    )
    system_receipts = read_system_receipts(
        system_bytes,
        conversion_factors,
    )
    selector, plan_month, planning_rows = read_destination_inputs(
        dest_bytes
    )

    state = load_planning_state()
    info = calculate_planning_metrics(
        report_date=report_date,
        actual_receipts=actual_receipts,
        current_consignments=consignments,
        system_receipts=system_receipts,
        selector=selector,
        plan_month=plan_month,
        planning_rows=planning_rows,
        state=state,
    )

    updated_bytes, patched = patch_planning_metrics_workbook(
        dest_bytes,
        info["metrics"],
    )

    result = graph.upload_file(
        drive_id,
        dest_item["id"],
        updated_bytes,
        expected_etag=dest_item["eTag"],
    )

    if info["state_changed"]:
        save_planning_state(state)

    print(
        f"[{PLANNING_SHEET}] Đã tính J, K và M:R cho "
        f"{patched} mã; nguồn tháng {info['source_month']}; "
        f"kế hoạch {info['plan_month']} "
        f"({FC_SHEET}!{FC_SELECTOR_CELL}={selector!r})."
    )
    print(
        "Upload thành công:",
        result.get("name", "Sắp kế hoạch.xlsx"),
    )


if __name__ == "__main__":
    main()
