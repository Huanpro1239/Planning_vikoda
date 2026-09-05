from io import BytesIO

from openpyxl import load_workbook


def remove_sheet_formulas(workbook_bytes, sheet_name):
    """Remove formulas only from one worksheet, preserving other cells/sheets."""
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

    if sheet_name not in workbook.sheetnames:
        raise RuntimeError(f"Không tìm thấy sheet {sheet_name!r} trong file đích.")

    worksheet = workbook[sheet_name]
    removed = 0

    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                cell.value = None
                removed += 1

    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), removed
