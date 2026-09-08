"""Adapter SharePoint -> engine ``Ke hoach SX tuan`` -> proposal/report."""
from __future__ import annotations

import calendar
import hashlib
import json
import math
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from openpyxl.utils.datetime import to_excel
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from weekly_planning_engine import PlannerPolicy, WeeklyCalculatedRow, WeeklyInputRow, build_daily_plan, calculate_rows

PLANNING_SHEET = "Ke_hoach_SX"
MASTER_SHEET = "Danh_muc"
START_COLUMN = 19  # S
EPS = 1e-6
BALANCE_EPS = 1e-5
ENGINE_VERSION = "ke_hoach_sx_tuan_v5_khsx_ki_20260908"


@dataclass(frozen=True)
class WeeklyAnalysis:
    calculated: list[WeeklyCalculatedRow]
    daily_plan: list[Any]
    policy_warnings: list[dict[str, Any]]
    changed_cells: int
    period_year: int
    period_month: int


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _num(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, bool):
        raise RuntimeError("TRUE/FALSE không phải số kế hoạch.")
    result = float(value)
    if not math.isfinite(result):
        raise RuntimeError(f"Giá trị không hữu hạn: {value!r}")
    return result


DEBT_HEADER_NAMES: frozenset[str] = frozenset({
    "debt mode",
    "cách tính nợ",
    "cach tinh no",
})
PROFILE_HEADER_NAMES: frozenset[str] = frozenset({
    "schedule profile",
    "profile lịch",
    "profile lich",
    "lịch sản xuất",
    "lich san xuat",
})


def _code(value: Any) -> int:
    return int(float(value))


def find_header_col(headers: list[Any], names: set[str] | frozenset[str]) -> int | None:
    """Tìm chỉ số cột (0-indexed) trong danh sách tiêu đề khớp với một trong các tên."""
    wanted = {_norm(name) for name in names}
    for idx, h in enumerate(headers):
        if _norm(h) in wanted:
            return idx
    return None


def _header_col(ws, names: set[str] | frozenset[str]) -> int | None:
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    idx = find_header_col(headers, names)
    return idx + 1 if idx is not None else None


def _debt_mode(value: Any) -> str | None:
    text = _norm(value).replace("-", "_").replace(" ", "_")
    if not text:
        return None
    if text in {"subtract_book_on_debt", "tru_ton_so", "trừ_tồn_sổ"}:
        return "SUBTRACT_BOOK_ON_DEBT"
    if text in {"ignore_book_on_debt", "khong_tru_ton_so", "không_trừ_tồn_sổ"}:
        return "IGNORE_BOOK_ON_DEBT"
    raise RuntimeError(f"Debt mode không hợp lệ: {value!r}")


def normalize_debt_mode(value: Any) -> str | None:
    """Chuẩn hóa debt mode an toàn cho cả engine lẫn hash fingerprint."""
    try:
        return _debt_mode(value)
    except Exception:
        return str(value or "").strip() or None


def _profile(value: Any) -> str | None:
    text = _norm(value).replace("-", "_").replace(" ", "_")
    if not text:
        return None
    if text in {"continuous", "lien_tuc", "liên_tục"}:
        return "CONTINUOUS"
    if text in {"spread_non_sunday", "rai_khong_chu_nhat", "rải_không_chủ_nhật"}:
        return "SPREAD_NON_SUNDAY"
    raise RuntimeError(f"Schedule profile không hợp lệ: {value!r}")


def normalize_profile(value: Any) -> str | None:
    """Chuẩn hóa schedule profile an toàn cho cả engine lẫn hash fingerprint."""
    try:
        return _profile(value)
    except Exception:
        return str(value or "").strip() or None


def _same(current: Any, target: Any) -> bool:
    if target is None:
        return current in (None, "")
    if isinstance(target, datetime):
        if isinstance(current, datetime):
            return abs((current - target).total_seconds()) < 1
        if isinstance(current, date):
            return current == target.date()
        if isinstance(current, (int, float)):
            return math.isclose(float(current), float(to_excel(target)), rel_tol=0.0, abs_tol=1.0/86400.0)
        return False
    if isinstance(current, (int, float)) and isinstance(target, (int, float)):
        return math.isclose(float(current), float(target), rel_tol=1e-10, abs_tol=1e-7)
    return current == target



def _scheduled_by_code(plan: list[Any]) -> dict[int, float]:
    result: dict[int, float] = defaultdict(float)
    for item in plan:
        result[int(item.ma_sp)] += float(item.qty)
    return result


def _committed_qty(calc: WeeklyCalculatedRow, scheduled_by_code: dict[int, float]) -> float:
    """Quantity actually committed to the daily plan for this snapshot."""
    scheduled = max(0.0, float(scheduled_by_code.get(calc.input.ma_sp, 0.0)))
    return min(calc.schedulable_qty, scheduled)


def _committed_days(calc: WeeklyCalculatedRow, scheduled_by_code: dict[int, float]) -> float:
    committed = _committed_qty(calc, scheduled_by_code)
    return committed / calc.input.sl_ca / calc.input.shifts_per_day

def _read_inputs(workbook_bytes: bytes, *, plan_year: int, plan_month: int):
    wb = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        for name in (PLANNING_SHEET, MASTER_SHEET):
            if name not in wb.sheetnames:
                raise RuntimeError(f"Không tìm thấy sheet {name!r}.")
        planning, master = wb[PLANNING_SHEET], wb[MASTER_SHEET]
        debt_col = _header_col(master, DEBT_HEADER_NAMES)
        profile_col = _header_col(master, PROFILE_HEADER_NAMES)
        master_data: dict[int, dict[str, Any]] = {}
        for r in range(2, master.max_row + 1):
            raw = master.cell(r, 1).value
            if raw in (None, ""):
                continue
            code = _code(raw)
            if code in master_data:
                raise RuntimeError(f"Mã {code} bị lặp trong {MASTER_SHEET}!A.")
            master_data[code] = {
                "name": str(master.cell(r, 2).value or ""),
                "uom": str(master.cell(r, 3).value or ""),
                "batch": _num(master.cell(r, 4).value),
                "per_shift": _num(master.cell(r, 5).value),
                "line": str(master.cell(r, 6).value or "").strip(),
                "group": str(master.cell(r, 7).value or "").strip(),
                "classification": str(master.cell(r, 8).value or "").strip(),
                "mold": _num(master.cell(r, 9).value),
                "leadtime": _num(master.cell(r, 10).value),
                "debt_mode": _debt_mode(master.cell(r, debt_col).value) if debt_col else None,
                "profile": _profile(master.cell(r, profile_col).value) if profile_col else None,
            }

        rows: list[WeeklyInputRow] = []
        warnings: list[dict[str, Any]] = []
        spread: set[int] = set()
        for r in range(2, planning.max_row + 1):
            raw = planning.cell(r, 1).value
            if raw in (None, ""):
                continue
            code = _code(raw)
            meta = master_data.get(code)
            if meta is None:
                raise RuntimeError(f"Mã {code} không có trong {MASTER_SHEET}.")
            debt = _num(planning.cell(r, 14).value)
            mode = meta["debt_mode"]
            if mode is None:
                mode = "SUBTRACT_BOOK_ON_DEBT"  # proposal-only fallback
                if debt > EPS:
                    warnings.append({"code": str(code), "type": "MISSING_DEBT_MODE", "message": "Thiếu Danh_muc!Debt mode; proposal tạm dùng SUBTRACT_BOOK_ON_DEBT."})
            if meta["line"] == "Galon":
                if meta["profile"] == "SPREAD_NON_SUNDAY":
                    spread.add(code)
                elif meta["profile"] is None:
                    warnings.append({"code": str(code), "type": "MISSING_SCHEDULE_PROFILE", "message": "Thiếu Danh_muc!Schedule profile cho Galon; proposal tạm dùng CONTINUOUS."})
            fc = _num(planning.cell(r, 12).value)
            rows.append(WeeklyInputRow(
                source_row=r, ma_sp=code, ten_sp=meta["name"], don_vi_tinh=meta["uom"],
                sl_me=meta["batch"], sl_ca=meta["per_shift"], chuyen=meta["line"],
                nhom_sp=meta["group"], phan_loai=meta["classification"], quy_cach=meta["mold"],
                shifts_per_day=_num(planning.cell(r, 9).value), ton_dau_thuc_te=_num(planning.cell(r, 10).value),
                ton_dau_so_sach=_num(planning.cell(r, 11).value), fc=fc,
                ton_cuoi_du_kien=_num(planning.cell(r, 13).value), no_kho=debt,
                avg_daily_sales=fc / 26.0 if abs(fc) > EPS else 0.0,
                leadtime=meta["leadtime"], debt_formula_mode=mode,
            ))
        return rows, frozenset(spread), warnings
    finally:
        wb.close()


def analyze_weekly_workbook(workbook_bytes: bytes, *, plan_year: int, plan_month: int) -> WeeklyAnalysis:
    rows, spread, warnings = _read_inputs(workbook_bytes, plan_year=plan_year, plan_month=plan_month)
    calculated = calculate_rows(rows, period_year=plan_year, period_month=plan_month)
    daily = build_daily_plan(
        calculated,
        policy=PlannerPolicy(
            spread_product_codes=spread,
            allow_capacity_trim=True,
        ),
    )
    lookup: dict[tuple[int, date], float] = defaultdict(float)
    for item in daily:
        lookup[(item.ma_sp, item.date)] += float(item.qty)
    scheduled_by_code = _scheduled_by_code(daily)
    wb = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        ws = wb[PLANNING_SHEET]
        changed = 0
        days = calendar.monthrange(plan_year, plan_month)[1]
        for calc in calculated:
            r = calc.input.source_row
            committed = _committed_qty(calc, scheduled_by_code)
            committed_days = _committed_days(calc, scheduled_by_code)
            for c, target in (
                (15, calc.p_need),
                (16, committed),
                (17, committed_days),
                (18, calc.start_datetime),
            ):
                if not _same(ws.cell(r, c).value, target):
                    changed += 1
            for d in range(1, days + 1):
                target = lookup.get((calc.input.ma_sp, date(plan_year, plan_month, d)), 0.0) or None
                if not _same(ws.cell(r, START_COLUMN + d - 1).value, target):
                    changed += 1
        return WeeklyAnalysis(calculated, daily, warnings, changed, plan_year, plan_month)
    finally:
        wb.close()


def patch_weekly_workbook(workbook_bytes: bytes, analysis: WeeklyAnalysis) -> bytes:
    wb = load_workbook(BytesIO(workbook_bytes), data_only=False)
    try:
        ws = wb[PLANNING_SHEET]
        lookup: dict[tuple[int, date], float] = defaultdict(float)
        for item in analysis.daily_plan:
            lookup[(item.ma_sp, item.date)] += float(item.qty)
        days = calendar.monthrange(analysis.period_year, analysis.period_month)[1]
        scheduled_by_code = _scheduled_by_code(analysis.daily_plan)
        for calc in analysis.calculated:
            r = calc.input.source_row
            committed = _committed_qty(calc, scheduled_by_code)
            committed_days = _committed_days(calc, scheduled_by_code)
            # O keeps the desired need including target stock. P/Q are the
            # production commitment that physically fits the computed schedule.
            ws.cell(r, 15, calc.p_need); ws.cell(r, 16, committed)
            ws.cell(r, 17, committed_days); ws.cell(r, 18, calc.start_datetime)
            for d in range(1, days + 1):
                qty = lookup.get((calc.input.ma_sp, date(analysis.period_year, analysis.period_month, d)), 0.0)
                ws.cell(r, START_COLUMN + d - 1).value = qty if qty > EPS else None
        out = BytesIO(); wb.save(out); return out.getvalue()
    finally:
        wb.close()


def _mass_balance(analysis: WeeklyAnalysis):
    """Mass balance with explicit service and safety-stock layers."""

    scheduled: dict[str, float] = defaultdict(float)
    for item in analysis.daily_plan:
        scheduled[str(item.ma_sp)] += float(item.qty)

    mass, service_carryovers, buffer_carryovers = {}, [], []
    for calc in analysis.calculated:
        code = str(calc.input.ma_sp)
        desired = calc.schedulable_qty
        service_target = calc.service_qty
        done = scheduled.get(code, 0.0)

        if done > desired + BALANCE_EPS:
            raise RuntimeError(f"Mã {code} scheduled={done} vượt Q={desired}.")

        service_done = min(done, service_target)
        service_carry = max(0.0, service_target - service_done)
        buffer_target = calc.buffer_qty
        buffer_done = max(0.0, done - service_target)
        buffer_done = min(buffer_done, buffer_target)
        buffer_carry = max(0.0, buffer_target - buffer_done)
        total_carry = max(0.0, desired - done)

        if service_carry <= BALANCE_EPS:
            service_carry = 0.0
        if buffer_carry <= BALANCE_EPS:
            buffer_carry = 0.0
        if total_carry <= BALANCE_EPS:
            total_carry = 0.0

        if service_carry:
            service_carryovers.append(code)
        if buffer_carry:
            buffer_carryovers.append(code)

        mass[code] = {
            "uom": calc.input.don_vi_tinh,
            "planned_qty": desired,
            "scheduled_qty": done,
            "carryover_qty": total_carry,
            "balanced_qty": done + total_carry,
            "service_target_qty": service_target,
            "service_scheduled_qty": service_done,
            "service_carryover_qty": service_carry,
            "buffer_target_qty": buffer_target,
            "buffer_scheduled_qty": buffer_done,
            "buffer_carryover_qty": buffer_carry,
        }

    return (
        mass,
        sorted(service_carryovers),
        sorted(buffer_carryovers),
    )

def _validate_shared_machine(analysis: WeeklyAnalysis, policy: PlannerPolicy | None = None):
    """Kiểm tra capacity vật lý của máy chung KHS + PET 9000.

    Báo cáo cũ đánh dấu resource_validation=True vô điều kiện. Điều đó có thể
    làm proposal xanh dù hai line đã bị xếp song song. Validation này quy đổi
    sản lượng ngày về số ca và kiểm tra tổng usage của *một* máy chung.
    """

    active_policy = policy or PlannerPolicy()
    selected = [
        calc
        for calc in analysis.calculated
        if calc.input.chuyen in active_policy.serialized_lines
        and calc.q_rounded > EPS
    ]
    if not selected:
        return {
            "ok": True,
            "state": "not_used",
            "resource": "KHS/PET 9000 shared machine",
            "capacity_shifts_per_day": 0.0,
            "peak_usage_shifts": 0.0,
            "over_capacity_dates": [],
        }

    capacity = selected[0].input.shifts_per_day
    if capacity <= EPS:
        return {
            "ok": False,
            "state": "invalid_capacity",
            "resource": "KHS/PET 9000 shared machine",
            "capacity_shifts_per_day": capacity,
            "peak_usage_shifts": 0.0,
            "over_capacity_dates": [],
        }

    if any(abs(calc.input.shifts_per_day - capacity) > EPS for calc in selected[1:]):
        return {
            "ok": False,
            "state": "inconsistent_shift_calendar",
            "resource": "KHS/PET 9000 shared machine",
            "capacity_shifts_per_day": capacity,
            "peak_usage_shifts": 0.0,
            "over_capacity_dates": [],
        }

    by_code = {calc.input.ma_sp: calc for calc in selected}
    usage: dict[date, float] = defaultdict(float)
    for item in analysis.daily_plan:
        calc = by_code.get(item.ma_sp)
        if calc is not None:
            usage[item.date] += float(item.qty) / calc.input.sl_ca

    peak = max(usage.values(), default=0.0)
    over = sorted(
        current.isoformat()
        for current, used in usage.items()
        if used > capacity + EPS
    )
    return {
        "ok": not over,
        "state": "shared_machine_passed" if not over else "shared_machine_over_capacity",
        "resource": "KHS/PET 9000 shared machine",
        "capacity_shifts_per_day": capacity,
        "peak_usage_shifts": peak,
        "over_capacity_dates": over,
    }


def build_weekly_schedule_report(workbook_bytes: bytes, analysis: WeeklyAnalysis, *, input_revision=None):
    mass, service_carryovers, buffer_carryovers = _mass_balance(analysis)
    lines = sorted({calc.input.chuyen for calc in analysis.calculated})
    shared_machine = _validate_shared_machine(analysis)
    shared_name = shared_machine["resource"]

    is_capacity_balanced = any(
        getattr(item, "phase", "") == "capacity_balanced"
        for item in analysis.daily_plan
    )

    if analysis.policy_warnings:
        service_state = "policy_metadata_missing"
    elif service_carryovers:
        service_state = "stockout_risk"
    else:
        service_state = "rules_complete"

    if not buffer_carryovers:
        safety_state = "complete"
    elif is_capacity_balanced:
        safety_state = "partially_achieved"
    elif service_carryovers:
        safety_state = "not_achieved_with_service_shortfall"
    else:
        safety_state = "partially_achieved"

    monthly_state = (
        "complete"
        if not service_carryovers and not buffer_carryovers
        else (
            "capacity_constrained_balanced"
            if is_capacity_balanced
            else (
                "service_complete_buffer_shortfall"
                if not service_carryovers
                else "service_carryover"
            )
        )
    )

    status = {
        "monthly_quantity": {
            "ok": not service_carryovers,
            "state": monthly_state,
            "carryover_skus": sorted(
                set(service_carryovers) | set(buffer_carryovers)
            ),
            "service_carryover_skus": service_carryovers,
            "buffer_carryover_skus": buffer_carryovers,
        },
        "resource_validation": {
            "ok": shared_machine["ok"],
            "state": shared_machine["state"],
            "validated_resources": [shared_name],
            "failed_resources": [] if shared_machine["ok"] else [shared_name],
            "shared_machine": shared_machine,
        },
        "service": {
            "ok": not analysis.policy_warnings and not service_carryovers,
            "state": service_state,
            "stockout_skus": service_carryovers,
            "capacity_balanced_skus": service_carryovers if is_capacity_balanced else [],
        },
        "safety_stock": {
            "ok": not buffer_carryovers,
            "state": safety_state,
            "shortfall_skus": buffer_carryovers,
        },
    }

    publish_ready = (
        status["monthly_quantity"]["ok"]
        and status["resource_validation"]["ok"]
        and status["service"]["ok"]
    )

    resources = {
        line: {"meta": {"mode": "ke_hoach_sx_tuan_reference_rules"}}
        for line in lines
    }
    resources[shared_name] = {"meta": shared_machine}

    return {
        "schema_version": 5,
        "algorithm": "ke_hoach_sx_tuan_v2_service_first",
        "plan_month": f"{analysis.period_year:04d}-{analysis.period_month:02d}",
        "input_revision": dict(input_revision or {}),
        "input_sha256": hashlib.sha256(workbook_bytes).hexdigest(),
        "mass_balance": mass,
        "resources": resources,
        "inventory": {},
        "policy_warnings": list(analysis.policy_warnings),
        "status": status,
        "publish_status": (
            "ready_for_publish" if publish_ready else "review_required"
        ),
    }

def prepare_weekly_schedule_update(workbook_bytes: bytes, *, plan_year: int, plan_month: int, input_revision=None):
    analysis = analyze_weekly_workbook(workbook_bytes, plan_year=plan_year, plan_month=plan_month)
    updated = patch_weekly_workbook(workbook_bytes, analysis)
    report = build_weekly_schedule_report(workbook_bytes, analysis, input_revision=input_revision)
    return updated, report, analysis


def verify_weekly_workbook(workbook_bytes: bytes, *, schedule_report: dict[str, Any], plan_year: int, plan_month: int):
    revision = schedule_report.get("input_revision")
    if not isinstance(revision, dict) or not revision:
        raise RuntimeError("Weekly schedule report thiếu input_revision/provenance.")
    if schedule_report.get("plan_month") != f"{plan_year:04d}-{plan_month:02d}":
        raise RuntimeError("Weekly schedule report sai kỳ kế hoạch.")
    analysis = analyze_weekly_workbook(workbook_bytes, plan_year=plan_year, plan_month=plan_month)
    if analysis.changed_cells:
        raise RuntimeError(f"Workbook chưa khớp weekly engine: {analysis.changed_cells} ô sai.")
    expected_mass, service_carryovers, buffer_carryovers = _mass_balance(analysis)
    actual_mass = schedule_report.get("mass_balance") or {}
    for code, expected in expected_mass.items():
        actual = actual_mass.get(code) or {}
        for key in (
            "planned_qty",
            "scheduled_qty",
            "carryover_qty",
            "service_target_qty",
            "service_scheduled_qty",
            "service_carryover_qty",
            "buffer_target_qty",
            "buffer_scheduled_qty",
            "buffer_carryover_qty",
        ):
            if not math.isclose(
                float(actual.get(key, 0) or 0),
                float(expected[key]),
                rel_tol=1e-9,
                abs_tol=BALANCE_EPS,
            ):
                raise RuntimeError(f"Mass balance {code}.{key} không khớp workbook.")
    expected_warnings = sorted((w["code"], w["type"]) for w in analysis.policy_warnings)
    report_warnings = sorted((str(w.get("code")), str(w.get("type"))) for w in schedule_report.get("policy_warnings", []) if isinstance(w, dict))
    if expected_warnings != report_warnings:
        raise RuntimeError("policy_warnings không khớp workbook.")
    expected_resource = _validate_shared_machine(analysis)
    report_resource = (schedule_report.get("status") or {}).get("resource_validation") or {}
    if (
        bool(report_resource.get("ok")) != bool(expected_resource["ok"])
        or report_resource.get("state") != expected_resource["state"]
    ):
        raise RuntimeError("resource_validation không khớp shared-machine schedule.")
    if service_carryovers and schedule_report.get("publish_status") in ("ready_for_publish", "feasible"):
        raise RuntimeError(
            f"Kế hoạch còn thiếu service ({service_carryovers}) nhưng publish_status lại là "
            f"{schedule_report.get('publish_status')}."
        )

    khsx_ki_verify = None
    import sync_planning_khsx_ki
    if sync_planning_khsx_ki.has_khsx_ki_sheet(workbook_bytes):
        khsx_ki_verify = sync_planning_khsx_ki.verify_khsx_ki(
            workbook_bytes,
            plan_year=plan_year,
            plan_month=plan_month,
        )

    return {
        "validated": True,
        "algorithm": ENGINE_VERSION,
        "checked_skus": len(analysis.calculated),
        "service_carryover_skus": service_carryovers,
        "buffer_carryover_skus": buffer_carryovers,
        "policy_warning_count": len(analysis.policy_warnings),
        "publish_status": schedule_report.get("publish_status"),
        "khsx_ki": khsx_ki_verify,
    }


def compute_planning_inputs_hash(workbook_bytes: bytes) -> str:
    """Tính SHA-256 fingerprint của các ô đầu vào do người dùng chỉnh trên sheet Ke_hoach_SX.
    Bao gồm:
    - Danh sách mã SP (cột A)
    - Số ca/ngày (cột I)
    Loại trừ các cột J, K (tính từ Tồn kho), L (tính từ FC), M (tồn cuối dự kiến do pipeline tính),
    N (tính từ Nợ kho), và các cột đầu ra O, P, Q, R, S+ để chống trigger lặp sau khi publish.
    """
    wb = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        if PLANNING_SHEET not in wb.sheetnames:
            return ""
        ws = wb[PLANNING_SHEET]
        inputs = []
        for r in range(2, ws.max_row + 1):
            raw_code = ws.cell(r, 1).value
            if raw_code in (None, ""):
                continue
            code = _code(raw_code)
            shifts = _num(ws.cell(r, 9).value)
            inputs.append({
                "code": code,
                "shifts_per_day": shifts,
            })
        payload = json.dumps(
            inputs,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        return hashlib.sha256(payload).hexdigest()
    finally:
        wb.close()
