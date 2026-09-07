import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from planning_cleanup import PLANNING_SHEET, remove_sheet_formulas


class PlanningFormulaRemovalTests(unittest.TestCase):
    def make_workbook(self):
        workbook = Workbook()
        planning = workbook.active
        planning.title = "Ke_hoach_SX"
        planning["A1"] = "Nhap tay"
        planning["B1"] = "=1+1"
        planning["B1"].number_format = "0.00"
        stock = workbook.create_sheet("Ton_kho")
        stock["A1"] = "=2+2"
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_runner_targets_actual_sheet_name(self):
        self.assertEqual(PLANNING_SHEET, "Ke_hoach_SX")

    def test_removes_formula_only_from_ke_hoach_sx(self):
        updated, removed = remove_sheet_formulas(self.make_workbook(), PLANNING_SHEET)
        workbook = load_workbook(BytesIO(updated), data_only=False)
        self.assertEqual(removed, 1)
        self.assertEqual(workbook[PLANNING_SHEET]["A1"].value, "Nhap tay")
        self.assertIsNone(workbook[PLANNING_SHEET]["B1"].value)
        self.assertEqual(workbook[PLANNING_SHEET]["B1"].number_format, "0.00")
        self.assertEqual(workbook["Ton_kho"]["A1"].value, "=2+2")

    def test_missing_sheet_raises_clear_error(self):
        with self.assertRaisesRegex(RuntimeError, "Khong ton tai"):
            remove_sheet_formulas(self.make_workbook(), "Khong ton tai")


if __name__ == "__main__":
    unittest.main()
