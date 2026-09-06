import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from sync_planning_calendar_all_months import prepare_calendar_update_all_months
from sync_planning_layout import canonicalize_planning_layout, needs_schedule_reset


class PlanningLayoutTests(unittest.TestCase):
    def make_legacy_workbook(self):
        workbook = Workbook()
        fc = workbook.active
        fc.title = "FC"
        fc["R1"] = "Tháng 9"

        planning = workbook.create_sheet("Ke_hoach_SX")
        planning["A1"] = "Mã Sản Phẩm"
        planning["R1"] = "Ngày bắt đầu sản xuất"
        planning["S1"] = "Kỳ kế hoạch"
        planning["T1"] = "Ngày 01"
        planning["AX1"] = "Ngày 31"
        planning["AY1"] = "Tổng SX"
        planning["AZ1"] = "Chênh lệch (SX-P)"
        planning["A2"] = "130100017"
        planning["S2"] = "2026-08"
        planning["T2"] = 12000
        planning["AY2"] = 12000
        planning["AZ2"] = 0
        planning.auto_filter.ref = "A1:AZ2"

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_legacy_layout_is_detected(self):
        self.assertTrue(needs_schedule_reset(self.make_legacy_workbook()))

    def test_calendar_migrates_legacy_layout_to_s_through_av(self):
        updated, info = prepare_calendar_update_all_months(
            self.make_legacy_workbook(),
            plan_year=2026,
        )
        self.assertEqual(info["end"], "AV1")
        self.assertGreater(info["layout_changes"], 0)

        workbook = load_workbook(BytesIO(updated), data_only=True)
        try:
            planning = workbook["Ke_hoach_SX"]
            self.assertEqual(planning["S1"].value, "01/09\nT3")
            self.assertEqual(planning["AV1"].value, "30/09\nT4")
            self.assertIsNone(planning["AW1"].value)
            self.assertIsNone(planning["AX1"].value)
            self.assertIsNone(planning["AY1"].value)
            self.assertIsNone(planning["AZ1"].value)
            self.assertIsNone(planning["S2"].value)
            self.assertIsNone(planning["T2"].value)
            self.assertIsNone(planning["AY2"].value)
            self.assertEqual(planning.auto_filter.ref, "A1:AV2")
        finally:
            workbook.close()

    def test_cleanup_is_idempotent_after_migration(self):
        updated, _ = prepare_calendar_update_all_months(
            self.make_legacy_workbook(),
            plan_year=2026,
        )
        cleaned, changes = canonicalize_planning_layout(
            updated,
            active_days=30,
            reset_schedule=False,
        )
        self.assertEqual(changes, 0)
        self.assertEqual(cleaned, updated)


if __name__ == "__main__":
    unittest.main()
