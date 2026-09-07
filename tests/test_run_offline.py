import json
import tempfile
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

import run_offline
from sync_planning_calendar import build_date_headers


CODE = 130100011          # Vikoda code (starts with 1)
CODE_VKD = 230100011      # VKD code -> normalises to 130100011
FACTOR = 1.0


def _save(wb):
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _target_workbook():
    wb = Workbook()

    ton_kho = wb.active
    ton_kho.title = "Ton_kho"
    ton_kho.append([
        "Mã Sản Phẩm", "Tên", "ĐVT", "Tồn thực tế",
        "Tồn Vikoda (NM)", "Tồn VKD (NM)", "Tồn Vikoda (khác)", "Tồn VKD (khác)",
    ])
    ton_kho.append([CODE, "A", "Thùng", 0, 0, 0, 0, 0])

    master = wb.create_sheet("Danh_muc")
    master.append([
        "Mã Sản Phẩm", "Tên", "ĐVT", "Số lượng /mẻ", "Số lượng/ ca",
        "Chuyền", "Nhóm", "Phân loại SP", "Quy cách", "Leadtime", "Debt mode",
    ])
    master.append([
        CODE, "A", "Thùng", 100, 100, "KHS", "KHS", "Không đường",
        FACTOR, 0, "SUBTRACT_BOOK_ON_DEBT",
    ])

    fc = wb.create_sheet("FC")
    fc.append([
        "STT", "Mã SP", "Tên", "ĐVT",
        "Tháng 1", "Tháng 2", "Tháng 3", "Tháng 4", "Tháng 5", "Tháng 6",
        "Tháng 7", "Tháng 8", "Tháng 9", "Tháng 10", "Tháng 11", "Tháng 12",
        None, "Tháng 9",
    ])
    fc.append([1, CODE, "A", "Thùng", 0, 0, 0, 0, 0, 0, 0, 0, 300, 0, 0, 0])

    debt = wb.create_sheet("No kho")
    debt.append(["Mã Sản Phẩm", "Tên", "ĐVT", "Số lượng nợ"])
    debt.append([CODE, "A", "Thùng", 0])

    planning = wb.create_sheet("Ke_hoach_SX")
    planning.append([
        "Mã Sản Phẩm", "Tên", "ĐVT", "Số lượng /mẻ", "Số lượng/ ca", "Chuyền",
        "Nhóm sản phẩm", "Phân loại SP", "Số ca theo ngày", "Tồn đầu thực tế",
        "Tồn đầu sổ sách", "FC", "Tồn cuối dự kiến", "Nợ kho",
        "Số lượng cần sản xuất", "Số lượng sản xuất theo mẻ/ca",
        "Số ngày cần sản xuất", "Ngày bắt đầu sản xuất",
    ] + build_date_headers(2026, 9))
    planning.append([
        CODE, "A", "Thùng", 100, 100, "KHS", "KHS", "Không đường",
        2, 0, 0, 0, 0, 0, 0, 0, 0, None,
    ] + [None] * 30)

    return _save(wb)


def _actual_source():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = date(2026, 9, 15)  # report date (mid-month)
    header = [None] * 20
    header[2] = "Mã SP Vikoda"       # C
    header[6] = "Tổng nhập tháng"    # G
    header[13] = "Tồn N"             # N
    header[14] = "Tồn O"             # O
    header[16] = "Đã bán (gửi kho)"  # Q
    ws.append(header)
    row = [None] * 20
    row[2] = CODE
    row[6] = 50    # receipt
    row[13] = 20   # N
    row[14] = 30   # O -> Ton_kho!D = 50
    row[16] = 100  # consignment
    ws.append(row)
    return _save(wb)


def _single_value_source(code, *, value_col_index, value, receipt=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ncols = max(value_col_index, 12) + 1
    header = [None] * ncols
    header[1] = "Mã vật tư"  # B
    ws.append(header)
    row = [None] * ncols
    row[1] = code
    row[value_col_index] = value
    if receipt is not None:
        row[8] = receipt  # I - Nhập trong kỳ
    ws.append(row)
    return _save(wb)


class OfflinePipelineTests(unittest.TestCase):
    def test_run_offline_produces_proposal_and_verifies(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            files = {
                "target.xlsx": _target_workbook(),
                "actual.xlsx": _actual_source(),
                # factory: value at L (index 11), receipt at I (index 8)
                "factory_vikoda.xlsx": _single_value_source(CODE, value_col_index=11, value=0, receipt=50),
                "factory_vkd.xlsx": _single_value_source(CODE_VKD, value_col_index=11, value=0),
                # accounting: value at M (index 12)
                "accounting_vikoda.xlsx": _single_value_source(CODE, value_col_index=12, value=0),
                "accounting_vkd.xlsx": _single_value_source(CODE_VKD, value_col_index=12, value=0),
            }
            for name, data in files.items():
                (tmp_path / name).write_bytes(data)

            runtime_state = tmp_path / "runtime.json"
            runtime_state.write_text(
                json.dumps({
                    "version": 1,
                    "opening_debt_by_month": {},
                    "opening_consignment_by_month": {},
                }),
                encoding="utf-8",
            )

            out_dir = tmp_path / "out"
            result = run_offline.main([
                "--target", str(tmp_path / "target.xlsx"),
                "--actual", str(tmp_path / "actual.xlsx"),
                "--factory-vikoda", str(tmp_path / "factory_vikoda.xlsx"),
                "--factory-vkd", str(tmp_path / "factory_vkd.xlsx"),
                "--accounting-vikoda", str(tmp_path / "accounting_vikoda.xlsx"),
                "--accounting-vkd", str(tmp_path / "accounting_vkd.xlsx"),
                "--out", str(out_dir),
                "--runtime-state", str(runtime_state),
                "--verify",
            ])

            report = result["report"]
            self.assertEqual(report["plan_month"], "2026-09")
            self.assertIn(report["publish_status"], {"ready_for_publish", "review_required"})
            self.assertTrue((out_dir / "planning_proposal.xlsx").is_file())
            self.assertTrue((out_dir / "planning_schedule_report.json").is_file())
            # Report on disk must be valid JSON.
            json.loads((out_dir / "planning_schedule_report.json").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
