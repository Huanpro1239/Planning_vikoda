import json
import unittest
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook

import sync_planning_metrics_direct as direct_metrics
import sync_planning_pipeline as pipeline_runner
import sync_stock
import sync_stock_compat
from sync_planning_weekly_model import (
    compute_planning_inputs_hash,
    prepare_weekly_schedule_update,
    verify_weekly_workbook,
)
from weekly_planning_engine import (
    PlannerPolicy,
    WeeklyInputRow,
    build_daily_plan,
    calculate_rows,
)


def make_workbook_bytes(
    *,
    skus=None,
    debt_mode=True,
    galon_profile=True,
    debts=None,
):
    """Create a minimal in-memory Excel workbook for pipeline testing."""
    wb = Workbook()
    planning = wb.active
    planning.title = "Ke_hoach_SX"
    planning.append([
        "Mã Sản Phẩm", "Tên Sản Phẩm", "Đơn vị tính", "Số lượng /mẻ",
        "Số lượng/ ca", "Chuyền", "Nhóm sản phẩm", "Phân loại SP",
        "Số ca theo ngày", "Tồn đầu thực tế", "Tồn đầu sổ sách", "FC",
        "Tồn cuối dự kiến", "Nợ kho", "Số lượng cần sản xuất",
        "Số lượng sản xuất theo mẻ/ca", "Số ngày cần sản xuất",
        "Ngày bắt đầu sản xuất",
    ] + [f"{day:02d}/09" for day in range(1, 31)])

    if skus is None:
        skus = [
            {
                "code": 130100011,
                "name": "SP A",
                "uom": "Thùng",
                "batch": 100,
                "per_shift": 200,
                "line": "KHS",
                "group": "Nhom",
                "classification": "Không đường",
                "shifts_per_day": 1,
                "ton_dau_thuc_te": 0,
                "ton_dau_so_sach": 0,
                "fc": 8000,
                "ton_cuoi_du_kien": 0,
                "no_kho": 0,
                "mold": 10,
                "leadtime": 0,
                "debt_mode": "SUBTRACT_BOOK_ON_DEBT",
                "profile": "",
            }
        ]

    for item in skus:
        planning.append([
            item["code"],
            item.get("name", "SP"),
            item.get("uom", "Thùng"),
            item.get("batch", 100),
            item.get("per_shift", 200),
            item.get("line", "KHS"),
            item.get("group", "Nhom"),
            item.get("classification", "Không đường"),
            item.get("shifts_per_day", 1),
            item.get("ton_dau_thuc_te", 0),
            item.get("ton_dau_so_sach", 0),
            item.get("fc", 0),
            item.get("ton_cuoi_du_kien", 0),
            item.get("no_kho", 0),
            0, 0, 0, None,
        ] + [None] * 30)

    master = wb.create_sheet("Danh_muc")
    headers = [
        "Mã Sản Phẩm", "Tên Sản Phẩm", "Đơn vị tính", "Số lượng /mẻ",
        "Số lượng/ ca", "Chuyền", "Nhóm sản phẩm", "Phân loại SP",
        "Quy cách", "Leadtime",
    ]
    if debt_mode:
        headers.append("Debt mode")
    if galon_profile:
        headers.append("Schedule profile")
    master.append(headers)

    for item in skus:
        row = [
            item["code"],
            item.get("name", "SP"),
            item.get("uom", "Thùng"),
            item.get("batch", 100),
            item.get("per_shift", 200),
            item.get("line", "KHS"),
            item.get("group", "Nhom"),
            item.get("classification", "Không đường"),
            item.get("mold", 10),
            item.get("leadtime", 0),
        ]
        if debt_mode:
            row.append(item.get("debt_mode", "SUBTRACT_BOOK_ON_DEBT"))
        if galon_profile:
            row.append(item.get("profile", ""))
        master.append(row)

    if debts is not None:
        debt_sheet = wb.create_sheet("No kho")
        debt_sheet.append(["Mã Sản Phẩm", "Tên", "ĐVT", "Số lượng nợ"])
        for code, val in debts.items():
            debt_sheet.append([code, "SP", "Thùng", val])

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


class ProductionReadinessRegressionTests(unittest.TestCase):
    def test_capacity_shortfall_blocks_autopublish_requires_approval(self):
        """Defect 1: Capacity shortfall must mark service.ok=False and require approval."""
        # 8000 target demand, but monthly capacity is only 6000 (30 days * 200/day).
        wb_bytes = make_workbook_bytes(
            skus=[{
                "code": 130100011,
                "name": "SP A",
                "uom": "Thùng",
                "batch": 100,
                "per_shift": 200,
                "line": "KHS",
                "group": "Nhom",
                "classification": "Không đường",
                "shifts_per_day": 1,
                "ton_dau_thuc_te": 0,
                "ton_dau_so_sach": 0,
                "fc": 8000,
                "ton_cuoi_du_kien": 0,
                "no_kho": 0,
                "mold": 10,
                "leadtime": 0,
                "debt_mode": "SUBTRACT_BOOK_ON_DEBT",
                "profile": "",
            }],
            debt_mode=True,
        )

        final_bytes, report, analysis = prepare_weekly_schedule_update(
            wb_bytes,
            plan_year=2026,
            plan_month=9,
            input_revision={"target": {"etag": "offline-review"}},
        )

        self.assertEqual(report["publish_status"], "review_required")
        self.assertFalse(report["status"]["service"]["ok"])
        self.assertEqual(report["status"]["service"]["stockout_skus"], ["130100011"])
        self.assertEqual(report["status"]["service"]["capacity_balanced_skus"], ["130100011"])

        # Auto-publish without approval MUST be blocked
        authorized, decision = pipeline_runner._publish_decision(report, None, "scheduler")
        self.assertFalse(authorized)
        self.assertEqual(decision["state"], "blocked")

        # Explicit approval with reason for the exact proposal MUST be accepted
        proposal_id = report.get("proposal_id", "test_prop")
        report["proposal_id"] = proposal_id
        authorized_approved, decision_approved = pipeline_runner._publish_decision(
            report,
            {"proposal_id": proposal_id, "reason": "Chấp nhận thiếu 2000 thùng vì hết công suất"},
            "planner_lead",
        )
        self.assertTrue(authorized_approved)
        self.assertEqual(decision_approved["state"], "authorized")
        self.assertEqual(decision_approved["basis"], "review_approval")

        # verify_weekly_workbook must pass with this report
        verification = verify_weekly_workbook(
            final_bytes,
            schedule_report=report,
            plan_year=2026,
            plan_month=9,
        )
        self.assertTrue(verification["validated"])

    def test_all_input_changes_detected_and_outputs_ignored(self):
        """Defect 2: All master policy, No kho, and Ke_hoach_SX inputs trigger recalculation."""
        base_skus = [{
            "code": 130100011,
            "name": "SP A",
            "uom": "Thùng",
            "batch": 100,
            "per_shift": 200,
            "line": "KHS",
            "group": "Nhom",
            "classification": "Không đường",
            "shifts_per_day": 1,
            "ton_dau_thuc_te": 0,
            "ton_dau_so_sach": 0,
            "fc": 1000,
            "ton_cuoi_du_kien": 0,
            "no_kho": 0,
            "mold": 10,
            "leadtime": 0,
            "debt_mode": "SUBTRACT_BOOK_ON_DEBT",
            "profile": "",
        }]

        base_bytes = make_workbook_bytes(skus=base_skus, debts={130100011: 0})
        _, base_conv_hash = sync_stock_compat.read_conversion_factors_robust(base_bytes)
        _, base_no_kho_hash = direct_metrics.hash_debt_sheet(base_bytes)
        base_planning_hash = compute_planning_inputs_hash(base_bytes)

        old_state = {
            "sources": {k: "etag_1" for k in pipeline_runner.SOURCES},
            "conversion_hash": base_conv_hash,
            "fc_hash": "fc_hash_1",
            "no_kho_hash": base_no_kho_hash,
            "planning_inputs_hash": base_planning_hash,
            "engine_version": "ke_hoach_sx_tuan_v2_service_first_20260908",
        }
        source_items = {k: {"eTag": "etag_1"} for k in pipeline_runner.SOURCES}

        # 1. Same inputs -> No changes
        info_same = {
            "conversion_hash": base_conv_hash,
            "fc_hash": "fc_hash_1",
            "no_kho_hash": base_no_kho_hash,
            "planning_inputs_hash": base_planning_hash,
            "engine_version": "ke_hoach_sx_tuan_v2_service_first_20260908",
        }
        self.assertEqual(
            pipeline_runner.detect_input_changes(old_state, source_items, info_same),
            [],
        )

        # 2. Leadtime changed in Danh_muc -> sheet:Danh_muc detected
        skus_leadtime = [dict(base_skus[0], leadtime=7)]
        bytes_leadtime = make_workbook_bytes(skus=skus_leadtime, debts={130100011: 0})
        _, conv_leadtime_hash = sync_stock_compat.read_conversion_factors_robust(bytes_leadtime)
        self.assertNotEqual(base_conv_hash, conv_leadtime_hash)
        info_leadtime = dict(info_same, conversion_hash=conv_leadtime_hash)
        self.assertEqual(
            pipeline_runner.detect_input_changes(old_state, source_items, info_leadtime),
            ["sheet:Danh_muc"],
        )

        # 3. Debt changed in No kho -> sheet:No_kho detected
        bytes_debt = make_workbook_bytes(skus=base_skus, debts={130100011: 500})
        _, debt_hash = direct_metrics.hash_debt_sheet(bytes_debt)
        self.assertNotEqual(base_no_kho_hash, debt_hash)
        info_debt = dict(info_same, no_kho_hash=debt_hash)
        self.assertEqual(
            pipeline_runner.detect_input_changes(old_state, source_items, info_debt),
            ["sheet:No_kho"],
        )

        # 4. shifts_per_day changed in Ke_hoach_SX -> sheet:Ke_hoach_SX detected
        skus_shifts = [dict(base_skus[0], shifts_per_day=2)]
        bytes_shifts = make_workbook_bytes(skus=skus_shifts, debts={130100011: 0})
        planning_shifts_hash = compute_planning_inputs_hash(bytes_shifts)
        self.assertNotEqual(base_planning_hash, planning_shifts_hash)
        info_shifts = dict(info_same, planning_inputs_hash=planning_shifts_hash)
        self.assertEqual(
            pipeline_runner.detect_input_changes(old_state, source_items, info_shifts),
            ["sheet:Ke_hoach_SX"],
        )

        # 5. Output cells in Ke_hoach_SX do NOT change planning_inputs_hash
        final_bytes, _, _ = prepare_weekly_schedule_update(
            base_bytes,
            plan_year=2026,
            plan_month=9,
            input_revision={"target": {"etag": "offline-review"}},
        )
        output_planning_hash = compute_planning_inputs_hash(final_bytes)
        self.assertEqual(base_planning_hash, output_planning_hash)

        # 6. Legacy old_state without new hashes triggers recalculation safely
        legacy_state = {
            "sources": {k: "etag_1" for k in pipeline_runner.SOURCES},
            "conversion_hash": base_conv_hash,
            "fc_hash": "fc_hash_1",
        }
        detected_legacy = pipeline_runner.detect_input_changes(legacy_state, source_items, info_same)
        self.assertIn("sheet:No_kho", detected_legacy)
        self.assertIn("sheet:Ke_hoach_SX", detected_legacy)

    def test_buffer_does_not_displace_service_timeline(self):
        """Defect 3: Timeline capacity constraints must prioritize service over earlier SKU buffer."""
        from dataclasses import replace
        from tests.test_service_first_safety_stock import row

        inputs = [
            replace(row(1001, 'KHS', 2, fc=100, target=900), ton_dau_thuc_te=2000),
            replace(row(1002, 'PET 9000', 3, fc=900, target=0), ton_dau_thuc_te=2100),
        ]
        calculated = calculate_rows(inputs, period_year=2026, period_month=9)
        plan = build_daily_plan(calculated, policy=PlannerPolicy(setup_shifts=0, allow_capacity_trim=True))

        scheduled_a = sum(p.qty for p in plan if p.ma_sp == 1001)
        scheduled_b = sum(p.qty for p in plan if p.ma_sp == 1002)

        # SKU A gets exactly 100 (service); buffer 900 is dropped because capacity on timeline cannot fit it
        self.assertEqual(scheduled_a, 100.0)
        # SKU B gets full 900 (service)
        self.assertEqual(scheduled_b, 900.0)

    def test_zero_forecast_with_debt_schedules_production(self):
        """Defect 4: When FC=0 and Nợ kho > 0, start_datetime must be day 1 and production must be scheduled."""
        for mode in ("SUBTRACT_BOOK_ON_DEBT", "IGNORE_BOOK_ON_DEBT"):
            wb_bytes = make_workbook_bytes(
                skus=[{
                    "code": 130100011,
                    "name": "SP A",
                    "uom": "Thùng",
                    "batch": 100,
                    "per_shift": 500,
                    "line": "KHS",
                    "group": "Nhom",
                    "classification": "Không đường",
                    "shifts_per_day": 1,
                    "ton_dau_thuc_te": 0,
                    "ton_dau_so_sach": 0,
                    "fc": 0,
                    "ton_cuoi_du_kien": 0,
                    "no_kho": 500,
                    "mold": 10,
                    "leadtime": 0,
                    "debt_mode": mode,
                    "profile": "",
                }],
                debt_mode=True,
            )

            final_bytes, report, analysis = prepare_weekly_schedule_update(
                wb_bytes,
                plan_year=2026,
                plan_month=9,
                input_revision={"target": {"etag": "offline-review"}},
            )

            self.assertEqual(report["publish_status"], "ready_for_publish")
            calc_row = analysis.calculated[0]
            self.assertEqual(calc_row.q_rounded, 500.0)
            self.assertEqual(calc_row.start_datetime, datetime(2026, 9, 1, 0, 0, 0))

            mb = report["mass_balance"]["130100011"]
            self.assertEqual(mb["planned_qty"], 500.0)
            self.assertEqual(mb["scheduled_qty"], 500.0)
            self.assertEqual(mb["carryover_qty"], 0.0)

    def test_vietnamese_debt_and_profile_headers_trigger_change_detection(self):
        """P1: Ensure Vietnamese headers 'Cách tính nợ' and 'Profile lịch' update conversion_hash and detect changes."""
        from openpyxl import load_workbook
        from tests.test_run_offline import _target_workbook

        def _encode(wb):
            out = BytesIO()
            wb.save(out)
            wb.close()
            return out.getvalue()

        # 1. Debt mode under 'Cách tính nợ'
        wb = load_workbook(BytesIO(_target_workbook()))
        wb["Danh_muc"]["K1"] = "Cách tính nợ"
        wb["Danh_muc"]["K2"] = "SUBTRACT_BOOK_ON_DEBT"
        before_bytes = _encode(wb)

        wb2 = load_workbook(BytesIO(before_bytes))
        wb2["Danh_muc"]["K2"] = "IGNORE_BOOK_ON_DEBT"
        after_bytes = _encode(wb2)

        _, h1_compat = sync_stock_compat.read_conversion_factors_robust(before_bytes)
        _, h2_compat = sync_stock_compat.read_conversion_factors_robust(after_bytes)
        self.assertNotEqual(h1_compat, h2_compat, "conversion_hash must change when 'Cách tính nợ' changes (compat)")

        _, h1_sync = sync_stock.read_conversion_factors(before_bytes)
        _, h2_sync = sync_stock.read_conversion_factors(after_bytes)
        self.assertNotEqual(h1_sync, h2_sync, "conversion_hash must change when 'Cách tính nợ' changes (sync_stock)")

        old_state = {
            "sources": {k: "etag_1" for k in pipeline_runner.SOURCES},
            "conversion_hash": h1_compat,
            "fc_hash": "fc_1",
            "no_kho_hash": "debt_1",
            "planning_inputs_hash": "plan_1",
            "engine_version": "ke_hoach_sx_tuan_v2_service_first_20260908",
        }
        source_items = {k: {"eTag": "etag_1"} for k in pipeline_runner.SOURCES}
        info_changed = dict(old_state, conversion_hash=h2_compat)
        self.assertEqual(
            pipeline_runner.detect_input_changes(old_state, source_items, info_changed),
            ["sheet:Danh_muc"],
        )

        # 2. Profile under 'Profile lịch'
        wb3 = load_workbook(BytesIO(_target_workbook()))
        wb3["Danh_muc"]["L1"] = "Profile lịch"
        wb3["Danh_muc"]["L2"] = "SPREAD_WORKING_DAYS"
        prof_before = _encode(wb3)

        wb4 = load_workbook(BytesIO(prof_before))
        wb4["Danh_muc"]["L2"] = ""
        prof_after = _encode(wb4)

        _, hp1 = sync_stock_compat.read_conversion_factors_robust(prof_before)
        _, hp2 = sync_stock_compat.read_conversion_factors_robust(prof_after)
        self.assertNotEqual(hp1, hp2, "conversion_hash must change when 'Profile lịch' changes")

    def test_two_round_pipeline_output_does_not_falsely_detect_column_m_change(self):
        """P2: Column M is calculated by pipeline and excluded from input hash, preventing false changes on round 2."""
        from openpyxl import load_workbook
        from planning_pipeline import prepare_pipeline_output
        from tests.test_run_offline import (
            _actual_source,
            _single_value_source,
            _target_workbook,
            CODE,
            CODE_VKD,
        )

        source_bytes = {
            "actual_stock": _actual_source(),
            "factory_vikoda": _single_value_source(CODE, value_col_index=11, value=0, receipt=50),
            "factory_vkd": _single_value_source(CODE_VKD, value_col_index=11, value=0),
            "accounting_vikoda": _single_value_source(CODE, value_col_index=12, value=0),
            "accounting_vkd": _single_value_source(CODE_VKD, value_col_index=12, value=0),
        }

        wb = load_workbook(BytesIO(_target_workbook()))
        wb["FC"]["M2"] = 2600
        out_buf = BytesIO()
        wb.save(out_buf)
        wb.close()
        original = out_buf.getvalue()

        # Round 1: M starts at 0, pipeline calculates M=200
        out1, report1, state1 = prepare_pipeline_output(
            original,
            source_bytes,
            runtime_state={},
            input_revision={"target": {"etag": "offline-test"}},
        )
        wb1 = load_workbook(BytesIO(out1), data_only=True)
        m_val = wb1["Ke_hoach_SX"]["M2"].value
        wb1.close()
        self.assertGreater(m_val, 0, "Pipeline should have calculated a positive column M")

        # Round 2: Feed round 1's output and state
        out2, report2, state2 = prepare_pipeline_output(
            out1,
            source_bytes,
            runtime_state=state1,
            input_revision={"target": {"etag": "offline-test"}},
        )

        # Change detection with unchanged sources must report NO changes
        state = {
            key: report1["pipeline"][key]
            for key in ("conversion_hash", "fc_hash", "no_kho_hash", "planning_inputs_hash", "engine_version")
        }
        state["sources"] = {key: "same-etag" for key in pipeline_runner.SOURCES}
        source_items = {key: {"eTag": "same-etag"} for key in pipeline_runner.SOURCES}

        changes = pipeline_runner.detect_input_changes(state, source_items, report2["pipeline"])
        self.assertEqual(changes, [], f"Round 2 detected spurious changes: {changes}")
        self.assertEqual(
            report1["pipeline"]["planning_inputs_hash"],
            report2["pipeline"]["planning_inputs_hash"],
        )

    def test_partial_buffer_allocated_on_residual_capacity(self):
        """P2: Residual capacity must be allocated to buffer without violating service priority or contiguous runs."""
        from tests.test_service_first_safety_stock import row

        # Scenario 1: Same mold (no setup)
        # SKU A: service=1000 (10 shifts), buffer=5000 (50 shifts).
        # SKU B: service=1000 (10 shifts), buffer=0 (0 shifts).
        # Total capacity: 30 shifts. Residual capacity: 10 shifts.
        # A should get 2000 (1000 service + 1000 buffer), B should get 1000.
        inputs_same_mold = [
            row(1001, "KHS", 2, fc=1000, target=5000, mold=1.0),
            row(1002, "PET 9000", 3, fc=1000, target=0, mold=1.0),
        ]
        calc_same = calculate_rows(inputs_same_mold, period_year=2026, period_month=9)
        plan_same = build_daily_plan(calc_same, policy=PlannerPolicy(allow_capacity_trim=True))

        a_qty = sum(p.qty for p in plan_same if p.ma_sp == 1001)
        b_qty = sum(p.qty for p in plan_same if p.ma_sp == 1002)
        self.assertEqual(a_qty, 2000.0)
        self.assertEqual(b_qty, 1000.0)
        self.assertEqual(sum(p.qty / 100.0 for p in plan_same), 30.0)

        # Scenario 2: Different mold (setup required = 0.5 shifts)
        # SKU A: service=1000 (10 shifts), buffer=5000 (sl_ca=100, mold=1.0).
        # SKU B: service=1000 (10 shifts), buffer=0 (sl_ca=100, mold=2.0).
        # Total capacity: 30 shifts. Setup between A and B = 0.5 shifts.
        # Service + Setup = 10 + 0.5 + 10 = 20.5 shifts.
        # Residual capacity = 9.5 shifts.
        # Since basis = 100 (1 shift), A can receive floor(9.5) = 9 shifts = 900 buffer.
        inputs_diff_mold = [
            row(1001, "KHS", 2, fc=1000, target=5000, mold=1.0),
            row(1002, "PET 9000", 3, fc=1000, target=0, mold=2.0),
        ]
        calc_diff = calculate_rows(inputs_diff_mold, period_year=2026, period_month=9)
        plan_diff = build_daily_plan(calc_diff, policy=PlannerPolicy(allow_capacity_trim=True))

        a_diff_qty = sum(p.qty for p in plan_diff if p.ma_sp == 1001)
        b_diff_qty = sum(p.qty for p in plan_diff if p.ma_sp == 1002)
        self.assertEqual(a_diff_qty, 1900.0)
        self.assertEqual(b_diff_qty, 1000.0)
        # Total shifts = 19 production + 0.5 setup + 10 production = 29.5 <= 30
        self.assertLessEqual(19.0 + 0.5 + 10.0, 30.0)

        # Scenario 3: Residual capacity smaller than 1 basis (0.5 shift left, basis is 1 shift)
        # SKU A: service=1500 (15 shifts), buffer=1000, mold=1.0.
        # SKU B: service=1400 (14 shifts), buffer=0, mold=2.0.
        # Total capacity = 30 shifts. Setup = 0.5 shift.
        # Service + setup = 15 + 0.5 + 14 = 29.5 shifts.
        # Residual capacity = 0.5 shift < 1 basis (100 qty = 1 shift).
        # Neither receives buffer, 100% service preserved.
        inputs_no_buffer = [
            row(1001, "KHS", 2, fc=1500, target=1000, mold=1.0),
            row(1002, "PET 9000", 3, fc=1400, target=0, mold=2.0),
        ]
        calc_no_buffer = calculate_rows(inputs_no_buffer, period_year=2026, period_month=9)
        plan_no_buffer = build_daily_plan(calc_no_buffer, policy=PlannerPolicy(allow_capacity_trim=True))

        self.assertEqual(sum(p.qty for p in plan_no_buffer if p.ma_sp == 1001), 1500.0)
        self.assertEqual(sum(p.qty for p in plan_no_buffer if p.ma_sp == 1002), 1400.0)
        self.assertLessEqual(sum(p.qty / 100.0 for p in plan_no_buffer) + 0.5, 30.0)


if __name__ == "__main__":
    unittest.main()
