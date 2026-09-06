import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from sync_planning_stock_inputs import prepare_stock_input_update


class PlanningStockInputsTests(unittest.TestCase):
    def make_workbook(self, *, planning_actual=8444, planning_book=13161.8333333333):
        workbook = Workbook()
        stock = workbook.active
        stock.title = "Ton_kho"
        stock.append([
            "Mã Sản Phẩm", "Tên Sản Phẩm", "Đơn vị tính", "Tồn thực tế",
            "Tồn Vikoda (Nhà Máy)", "Tồn VKD (Nhà Máy)",
            "Tồn Vikoda (Kho khác)", "Tồn VKD (Kho khác)",
        ])
        stock.append([130100096, "A", "Thùng", 6534, 5679.16666666667, 100.666666666667, 784.5, 812.5])

        planning = workbook.create_sheet("Ke_hoach_SX")
        planning.append([
            "Mã Sản Phẩm", "Tên Sản Phẩm", "Đơn vị tính", "Số lượng /mẻ",
            "Số lượng/ ca", "Chuyền", "Nhóm sản phẩm", "Phân loại SP",
            "Số ca theo ngày", "Tồn đầu thực tế", "Tồn đầu sổ sách", "FC",
        ])
        planning.append([
            130100096, "A", "Thùng", 4500, 4500, "KHS", "Lon", "Không đường",
            3, planning_actual, planning_book, 9370,
        ])

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_updates_j_and_k_from_ton_kho(self):
        updated, info = prepare_stock_input_update(self.make_workbook())
        self.assertEqual(info["changed_count"], 1)
        self.assertEqual(info["patched_count"], 1)

        workbook = load_workbook(BytesIO(updated), data_only=True)
        try:
            planning = workbook["Ke_hoach_SX"]
            self.assertEqual(planning["J2"].value, 6534)
            self.assertAlmostEqual(planning["K2"].value, 7376.833333333334, places=9)
            self.assertEqual(planning["L2"].value, 9370)
        finally:
            workbook.close()

    def test_is_idempotent_when_j_k_already_match(self):
        updated, info = prepare_stock_input_update(
            self.make_workbook(planning_actual=6534, planning_book=7376.833333333334)
        )
        self.assertEqual(info["changed_count"], 0)

    def test_missing_stock_code_fails(self):
        workbook = Workbook()
        stock = workbook.active
        stock.title = "Ton_kho"
        stock.append(["Mã Sản Phẩm", "Tên", "ĐVT", "Tồn thực tế", "E", "F", "G", "H"])
        stock.append([130100096, "A", "Thùng", 1, 1, 1, 1, 1])
        planning = workbook.create_sheet("Ke_hoach_SX")
        planning.append(["Mã Sản Phẩm"] + [None] * 10)
        planning.append([130100999] + [None] * 10)
        output = BytesIO()
        workbook.save(output)

        with self.assertRaisesRegex(RuntimeError, "không tồn tại trong Ton_kho"):
            prepare_stock_input_update(output.getvalue())


if __name__ == "__main__":
    unittest.main()
