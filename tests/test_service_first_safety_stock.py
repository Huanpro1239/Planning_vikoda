import unittest
from collections import defaultdict
from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook

from sync_planning_weekly_model import (
    analyze_weekly_workbook,
    build_weekly_schedule_report,
    patch_weekly_workbook,
)
from weekly_planning_engine import (
    PlannerPolicy,
    WeeklyInputRow,
    build_daily_plan,
    calculate_rows,
)


def row(
    code,
    line,
    source_row,
    *,
    fc,
    target,
    per_shift=100.0,
    shifts=1.0,
    mold=1.0,
):
    return WeeklyInputRow(
        source_row=source_row,
        ma_sp=code,
        ten_sp=str(code),
        don_vi_tinh="DV",
        sl_me=100.0,
        sl_ca=per_shift,
        chuyen=line,
        nhom_sp=line,
        phan_loai="Không đường",
        quy_cach=mold,
        shifts_per_day=shifts,
        ton_dau_thuc_te=0.0,
        ton_dau_so_sach=0.0,
        fc=fc,
        ton_cuoi_du_kien=target,
        no_kho=0.0,
        avg_daily_sales=100.0,
        leadtime=0.0,
        debt_formula_mode="SUBTRACT_BOOK_ON_DEBT",
    )


class ServiceFirstSafetyStockTests(unittest.TestCase):
    def _analysis(self, *, service_a=1000.0, service_b=1000.0):
        calculated = calculate_rows(
            [
                row(1001, "KHS", 2, fc=service_a, target=5000.0),
                row(1002, "PET 9000", 3, fc=service_b, target=0.0),
            ],
            period_year=2026,
            period_month=9,
        )
        daily = build_daily_plan(calculated, policy=PlannerPolicy(setup_shifts=0.0))
        from sync_planning_weekly_model import WeeklyAnalysis

        return WeeklyAnalysis(
            calculated=calculated,
            daily_plan=daily,
            policy_warnings=[],
            changed_cells=0,
            period_year=2026,
            period_month=9,
        )

    def test_sales_are_complete_when_only_safety_stock_exceeds_capacity(self):
        analysis = self._analysis()
        report = build_weekly_schedule_report(
            b"synthetic",
            analysis,
            input_revision={"target": {"etag": "synthetic"}},
        )

        self.assertEqual(report["publish_status"], "ready_for_publish")
        self.assertTrue(report["status"]["monthly_quantity"]["ok"])
        self.assertEqual(
            report["status"]["monthly_quantity"]["state"],
            "service_complete_buffer_shortfall",
        )
        self.assertTrue(report["status"]["service"]["ok"])
        self.assertEqual(report["status"]["service"]["stockout_skus"], [])
        self.assertFalse(report["status"]["safety_stock"]["ok"])
        self.assertEqual(
            report["status"]["safety_stock"]["state"],
            "partially_achieved",
        )

        a = report["mass_balance"]["1001"]
        b = report["mass_balance"]["1002"]
        self.assertEqual(a["service_target_qty"], 1000.0)
        self.assertEqual(a["service_carryover_qty"], 0.0)
        self.assertGreater(a["buffer_carryover_qty"], 0.0)
        self.assertEqual(b["service_carryover_qty"], 0.0)

    def test_shared_machine_tight_capacity_drops_buffer_and_schedules_contiguously(self):
        analysis = self._analysis()
        days_by_sku = defaultdict(list)
        phases_by_day = defaultdict(set)

        for item in analysis.daily_plan:
            if item.chuyen not in {"KHS", "PET 9000"}:
                continue
            days_by_sku[item.ma_sp].append(item.date)
            phases_by_day[item.date].add(item.phase)

        # When capacity is tight (70 shifts > 30 shifts), buffer is dropped
        # to ensure contiguous runs without split campaigns or tiny fragments.
        self.assertNotIn("buffer", [phase for phases in phases_by_day.values() for phase in phases])

        # Both SKUs run 100% of their service quantity
        self.assertEqual(sum(item.qty for item in analysis.daily_plan if item.ma_sp == 1001), 1000.0)
        self.assertEqual(sum(item.qty for item in analysis.daily_plan if item.ma_sp == 1002), 1000.0)

        # Each SKU runs contiguously in a single block without gaps or interleaving
        for ma_sp, dates in days_by_sku.items():
            for i in range(len(dates) - 1):
                self.assertEqual(
                    (dates[i + 1] - dates[i]).days,
                    1,
                    f"SKU {ma_sp} was interrupted between {dates[i]} and {dates[i + 1]}",
                )

        # 1001 runs first (days 1-10), then 1002 (days 11-20)
        self.assertEqual(days_by_sku[1001][-1], date(2026, 9, 10))
        self.assertEqual(days_by_sku[1002][0], date(2026, 9, 11))

        # No day may exceed the one-shift shared machine.
        per_shift = {c.input.ma_sp: c.input.sl_ca for c in analysis.calculated}
        usage = defaultdict(float)
        for item in analysis.daily_plan:
            if item.chuyen in {"KHS", "PET 9000"}:
                usage[item.date] += item.qty / per_shift[item.ma_sp]
        for used in usage.values():
            self.assertLessEqual(used, 1.0 + 1e-6)

    def test_shared_machine_sufficient_capacity_schedules_full_contiguously(self):
        # Service 1000 (10 shifts) + buffer 500 (5 shifts) + service 1000 (10 shifts) = 25 shifts <= 30 shifts
        analysis = self._analysis(service_a=1000.0, service_b=1000.0)
        # Re-run with target_a=500 so total fits in 30 shifts
        calculated = calculate_rows(
            [
                row(1001, "KHS", 2, fc=1000.0, target=500.0),
                row(1002, "PET 9000", 3, fc=1000.0, target=0.0),
            ],
            period_year=2026,
            period_month=9,
        )
        daily = build_daily_plan(calculated, policy=PlannerPolicy(setup_shifts=0.0))
        days_by_sku = defaultdict(list)
        for item in daily:
            if item.chuyen in {"KHS", "PET 9000"}:
                days_by_sku[item.ma_sp].append(item.date)

        # When capacity is sufficient, schedules full quantity contiguously
        self.assertEqual(sum(item.qty for item in daily if item.ma_sp == 1001), 1500.0)
        self.assertEqual(sum(item.qty for item in daily if item.ma_sp == 1002), 1000.0)
        for ma_sp, dates in days_by_sku.items():
            for i in range(len(dates) - 1):
                self.assertEqual((dates[i + 1] - dates[i]).days, 1)

    def test_service_shortfall_remains_review_required(self):
        analysis = self._analysis(service_a=2000.0, service_b=2000.0)
        report = build_weekly_schedule_report(
            b"synthetic",
            analysis,
            input_revision={"target": {"etag": "synthetic"}},
        )
        self.assertEqual(report["publish_status"], "review_required")
        self.assertFalse(report["status"]["monthly_quantity"]["ok"])
        self.assertEqual(
            report["status"]["monthly_quantity"]["state"],
            "service_carryover",
        )
        self.assertFalse(report["status"]["service"]["ok"])
        self.assertEqual(report["status"]["service"]["state"], "stockout_risk")

    def test_workbook_commits_capacity_feasible_qty_not_unreachable_buffer(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Ke_hoach_SX"
        ws.append([
            "Mã Sản Phẩm", "Tên Sản Phẩm", "Đơn vị tính", "Số lượng /mẻ",
            "Số lượng/ ca", "Chuyền", "Nhóm sản phẩm", "Phân loại SP",
            "Số ca theo ngày", "Tồn đầu thực tế", "Tồn đầu sổ sách", "FC",
            "Tồn cuối dự kiến", "Nợ kho", "Số lượng cần sản xuất",
            "Số lượng sản xuất theo mẻ/ca", "Số ngày cần sản xuất",
            "Ngày bắt đầu sản xuất",
        ] + [f"{day:02d}/09" for day in range(1, 31)])
        ws.append([1001, "A", "DV", 100, 100, "KHS", "KHS", "Không đường",
                   1, 0, 0, 1000, 5000, 0, 0, 0, 0, None] + [None] * 30)
        ws.append([1002, "B", "DV", 100, 100, "PET 9000", "PET 9000", "Không đường",
                   1, 0, 0, 1000, 0, 0, 0, 0, 0, None] + [None] * 30)
        master = wb.create_sheet("Danh_muc")
        master.append([
            "Mã Sản Phẩm", "Tên Sản Phẩm", "Đơn vị tính", "Số lượng /mẻ",
            "Số lượng/ ca", "Chuyền", "Nhóm sản phẩm", "Phân loại SP",
            "Quy cách", "Leadtime", "Debt mode",
        ])
        master.append([1001, "A", "DV", 100, 100, "KHS", "KHS", "Không đường",
                       1, 0, "SUBTRACT_BOOK_ON_DEBT"])
        master.append([1002, "B", "DV", 100, 100, "PET 9000", "PET 9000", "Không đường",
                       1, 0, "SUBTRACT_BOOK_ON_DEBT"])

        raw = BytesIO()
        wb.save(raw)
        analysis = analyze_weekly_workbook(
            raw.getvalue(),
            plan_year=2026,
            plan_month=9,
        )
        updated = patch_weekly_workbook(raw.getvalue(), analysis)

        result = load_workbook(BytesIO(updated), data_only=True, read_only=True)
        try:
            plan = result["Ke_hoach_SX"]
            # O keeps desired need (1000 + 5000); P/Q reflect what fits without buffer when capacity is tight.
            self.assertEqual(plan["O2"].value, 6000)
            self.assertEqual(plan["P2"].value, 1000)
            self.assertEqual(plan["Q2"].value, 10)
            self.assertEqual(plan["P3"].value, 1000)
            self.assertEqual(plan["Q3"].value, 10)
            self.assertEqual(
                sum(float(plan.cell(2, 19 + d).value or 0) for d in range(30)),
                1000.0,
            )
        finally:
            result.close()


if __name__ == "__main__":
    unittest.main()
