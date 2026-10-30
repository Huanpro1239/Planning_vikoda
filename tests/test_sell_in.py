import io
import unittest

from openpyxl import Workbook, load_workbook

from sell_in import build_workbook, doc_danh_muc, tong_hop_nhieu_file

HEADER = ["Vung", "KhuVuc", "TinhThanh", "QuanHuyen", "LoaiHoaDon", "SoHoaDon", "NgayHoaDon",
          "DienGiaiHD", "Kho", "MaKhachHangCu", "MaKhachHangMoi", "TenKhachHang", "TrinhDuocVien",
          "MaSanPhamCu", "MaSanPhamMoi", "TenSanPham", "SoLuong"]


def _sell_in(ten_dn, ky, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([ten_dn])
    for _ in range(3):
        ws.append([None])
    ws.append([ky])
    for _ in range(3):
        ws.append([None])
    ws.append(HEADER)
    for kh, ma, sl in rows:
        r = [None] * len(HEADER)
        r[10], r[14], r[16] = kh, ma, sl
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _danh_muc():
    wb = Workbook()
    ws = wb.active
    ws.title = "DanhMucSanPham"
    ws.append(["MaSanPhamMoi_Vikoda", "MaSanPhamMoi_VKD", "TenSanPham", "Tên viết tắt", "Quy cách",
               "ĐVT", "Brand", "SKUs Group"])
    ws.append([130100110, 230100110, "Kiềm 500ml", "Khoáng kiềm 500ml", 24, "Thùng", "Vikoda", "Pet Vikoda"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class SellInTests(unittest.TestCase):
    def test_filter_convert_and_merge_companies(self):
        dm = doc_danh_muc(_danh_muc())
        vikoda = _sell_in("Tên DN: CÔNG TY CỔ PHẦN NƯỚC KHOÁNG KHÁNH HÒA", "Từ ngày : 01-10-2026", [
            ("VKD3", "130100110", 2400),        # bán cho VKD -> bỏ
            ("B-HCMC-1", "130100110", 480),     # lấy: 20 thùng
            ("B-HCMC-1", "430300037", 50),      # két -> bỏ
        ])
        vkd = _sell_in("Tên DN: CÔNG TY CỔ PHẦN ĐẦU TƯ VÀ THƯƠNG MẠI VKD", "Từ ngày : 01-10-2026", [
            ("B-HCMC-2", "230100110", 240),     # lấy: 10 thùng
            ("B-HCMC-2", "", 0),                # chiết khấu -> bỏ
        ])
        data, log = tong_hop_nhieu_file([("THANG 10 2026.xlsm", vikoda), ("THANG 10 2026 VKD.xlsm", vkd)], dm)
        self.assertEqual(len(log), 2)
        content, n_sp, n_thang = build_workbook(data, log)
        self.assertEqual((n_sp, n_thang), (1, 1))
        rows = list(load_workbook(io.BytesIO(content), read_only=True)["Pivot"].iter_rows(values_only=True))
        self.assertEqual(rows[0][:2], ("MaVikoda", "MaVKD"))
        self.assertEqual(rows[1][0:2], ("130100110", "230100110"))
        self.assertEqual(rows[1][-2:], (30, 30))


if __name__ == "__main__":
    unittest.main()
