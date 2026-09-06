import unittest
from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook

import cleanup_planning
import sync_planning_schedule_priority_v8 as v8
from sync_planning_calendar import build_date_headers
from sync_stock import GraphRequestError
from verify_planning_month import verify_workbook


class FollowupReviewReproductions(unittest.TestCase):
    def _workbook_two_shared_skus(self, *, second_qty=6750, first_qty=6750):
        wb = Workbook()
        fc = wb.active
        fc.title = "FC"
        fc.append([
            "STT", "Mã SP", "Tên SP", "Đơn vị tính",
            "Tháng 1", "Tháng 2", "Tháng 3", "Tháng 4",
            "Tháng 5", "Tháng 6", "Tháng 7", "Tháng 8",
            "Tháng 9", "Tháng 10", "Tháng 11", "Tháng 12", None, "Tháng 9",
        ])
        fc.append([1, 100001, "A", "Thùng", 0, 0, 0, 0, 0, 0, 0, 0, first_qty, 0, 0, 0])
        fc.append([2, 100002, "B", "Thùng", 0, 0, 0, 0, 0, 0, 0, 0, second_qty, 0, 0, 0])

        stock = wb.create_sheet("Ton_kho")
        stock.append([
            "Mã Sản Phẩm", "Tên Sản Phẩm", "Đơn vị tính", "Tồn thực tế",
            "Tồn Vikoda (Nhà Máy)", "Tồn VKD (Nhà Máy)",
            "Tồn Vikoda (Kho khác)", "Tồn VKD (Kho khác)",
        ])
        stock.append([100001, "A", "Thùng", 0, 0, 0, 0, 0])
        stock.append([100002, "B", "Thùng", 0, 0, 0, 0, 0])

        planning = wb.create_sheet("Ke_hoach_SX")
        headers = [
            "Mã Sản Phẩm", "Tên Sản Phẩm", "Đơn vị tính", "Số lượng /mẻ",
            "Số lượng/ ca", "Chuyền", "Nhóm sản phẩm", "Phân loại SP",
            "Số ca theo ngày", "Tồn đầu thực tế", "Tồn đầu sổ sách", "FC",
            "Tồn cuối dự kiến", "Nợ kho", "Số lượng cần sản xuất",
            "Số lượng sản xuất theo mẻ/ca", "Số ngày cần sản xuất", "Ngày bắt đầu sản xuất",
        ] + build_date_headers(2026, 9)
        planning.append(headers)
        planning.append([
            100001, "A", "Thùng", 6750, 4500, "KHS", "G1", "Có đường",
            3, 0, 0, first_qty, 0, 0, first_qty, first_qty, first_qty / 4500 / 3, None,
        ] + [first_qty] + [None] * 29)
        planning.append([
            100002, "B", "Thùng", 6750, 4500, "PET 9000", "G2", "Có đường",
            3, 0, 0, second_qty, 0, 0, second_qty, second_qty, second_qty / 4500 / 3, None,
        ] + [second_qty] + [None] * 29)

        out = BytesIO()
        wb.save(out)
        return out.getvalue()

    def _shared_report(self, *, timeline, first_qty=6750, second_qty=6750):
        return {
            "schema_version": 2,
            "algorithm": "priority_v8",
            "plan_month": "2026-09",
            "input_revision": {"etag": "copy-etag"},
            "mass_balance": {
                "100001": {
                    "planned_qty": first_qty,
                    "scheduled_qty": first_qty,
                    "carryover_qty": 0,
                },
                "100002": {
                    "planned_qty": second_qty,
                    "scheduled_qty": second_qty,
                    "carryover_qty": 0,
                },
            },
            "resources": {
                "KHS + PET 9000": {
                    "capacity_shifts_per_day": 3,
                    "meta": {"timeline": timeline},
                }
            },
        }

    def test_shared_machine_setup_requires_timeline_provenance(self):
        # 3 production shifts exactly fill the machine, but a KHS->PET switch
        # also needs 0.5 shift. Daily qty alone cannot prove this schedule valid.
        with self.assertRaisesRegex(RuntimeError, "timeline|setup|provenance"):
            verify_workbook(self._workbook_two_shared_skus())

    def test_report_timeline_without_setup_is_rejected(self):
        report = self._shared_report(
            timeline=[
                {
                    "kind": "production",
                    "code": "100001",
                    "qty": 6750,
                    "start_shift": 0,
                    "end_shift": 1.5,
                },
                {
                    "kind": "production",
                    "code": "100002",
                    "qty": 6750,
                    "start_shift": 1.5,
                    "end_shift": 3,
                },
            ]
        )
        with self.assertRaisesRegex(RuntimeError, "thiếu setup"):
            verify_workbook(
                self._workbook_two_shared_skus(),
                schedule_report=report,
            )

    def test_valid_carryover_is_mass_balanced_with_report(self):
        data = self._workbook_two_shared_skus(second_qty=0, first_qty=6750)
        wb = load_workbook(BytesIO(data))
        try:
            ws = wb["Ke_hoach_SX"]
            ws["S2"] = 4500
            out = BytesIO()
            wb.save(out)
        finally:
            wb.close()

        report = {
            "schema_version": 2,
            "algorithm": "priority_v8",
            "plan_month": "2026-09",
            "input_revision": {"etag": "copy-etag"},
            "mass_balance": {
                "100001": {
                    "planned_qty": 6750,
                    "scheduled_qty": 4500,
                    "carryover_qty": 2250,
                },
                "100002": {
                    "planned_qty": 0,
                    "scheduled_qty": 0,
                    "carryover_qty": 0,
                },
            },
            "resources": {},
        }
        info = verify_workbook(out.getvalue(), schedule_report=report)
        self.assertEqual(info["schedule_checked"], 2)

    def test_negative_daily_qty_reports_business_error_not_valueerror(self):
        data = self._workbook_two_shared_skus(second_qty=0, first_qty=6750)
        wb = load_workbook(BytesIO(data))
        try:
            ws = wb["Ke_hoach_SX"]
            ws["S2"] = -1
            out = BytesIO()
            wb.save(out)
        finally:
            wb.close()

        with self.assertRaisesRegex(RuntimeError, "âm"):
            verify_workbook(out.getvalue())

    def test_overcapacity_reports_business_error_not_valueerror(self):
        with self.assertRaisesRegex(RuntimeError, r"Resource KHS \+ PET 9000 ngày 01/09"):
            verify_workbook(
                self._workbook_two_shared_skus(
                    first_qty=13500,
                    second_qty=6750,
                )
            )

    def test_cleanup_retries_read_503_and_uses_fresh_snapshot_etag_and_patch(self):
        class FakeGraph:
            def __init__(self):
                self.item_calls = 0
                self.download_calls = 0
                self.uploads = []

            def get_item_by_path(self, drive_id, path):
                self.item_calls += 1
                return {"id": "dest", "eTag": f"etag-{self.item_calls}"}

            def download_file(self, drive_id, item_id):
                self.download_calls += 1
                if self.download_calls == 1:
                    err = GraphRequestError("service unavailable", status_code=503)
                    err.retry_after_seconds = 7
                    raise err
                return b"fresh-snapshot"

            def upload_file(self, drive_id, item_id, content, expected_etag):
                self.uploads.append((content, expected_etag))
                return {"name": "copy.xlsx"}

        graph = FakeGraph()
        cleanup_inputs = []
        sleeps = []

        def fake_cleanup(data, sheet_name):
            cleanup_inputs.append(data)
            return b"patch-from-" + data, 1

        removed = cleanup_planning.cleanup_with_retry(
            graph,
            "drive",
            cleanup_func=fake_cleanup,
            sleep_func=sleeps.append,
            max_attempts=3,
            retry_delay_seconds=2,
        )

        self.assertEqual(removed, 1)
        self.assertEqual(graph.download_calls, 2)
        self.assertEqual(cleanup_inputs, [b"fresh-snapshot"])
        self.assertEqual(graph.uploads, [(b"patch-from-fresh-snapshot", "etag-2")])
        self.assertEqual(sleeps, [7])

    def test_unserved_due_keeps_each_quantum_deadline(self):
        headers = [date(2026, 9, 1), date(2026, 9, 2), date(2026, 9, 3)]
        product = {
            "row": 2,
            "code": "A",
            "uom": "Thùng",
            "batch": 100.0,
            "per_shift": 50.0,
            "line": "KHS",
            "product_group": "G1",
            "classification": "Có đường",
            "is_sugar": True,
            "max_shifts_per_day": 2.0,
            "actual_stock": 0.0,
            "fc": 600.0,
            "target_stock": 0.0,
            "debt": 0.0,
            "planned_qty": 500.0,
            "earliest_date": headers[0],
            "preferred_date": headers[0],
            "demand_by_day": {
                headers[0]: 200.0,
                headers[1]: 200.0,
                headers[2]: 200.0,
            },
            "quantum_qty": 100.0,
            "quantum_shift": 2.0,
            "required_units": 5,
            "existing_daily": [None] * 31,
        }

        schedule, usage, carryover, meta = v8.allocate_shared_deadline_guarded(
            headers, [product], 2.0
        )
        self.assertEqual(sum(schedule["A"].values()), 300)
        self.assertEqual(carryover, {"A": 200})
        self.assertEqual(len(meta["unserved_due"]), 2)
        self.assertEqual(
            [event["demand_due_date"] for event in meta["unserved_due"]],
            [headers[1], headers[2]],
        )
        self.assertTrue(
            all("production_deadline_date" in event for event in meta["unserved_due"])
        )
        self.assertTrue(
            all(
                event["code"] == "A" and event["uom"] == "Thùng"
                for event in meta["unserved_due"]
            )
        )
        self.assertEqual(
            [event["demand_due_date"] for event in meta["shortage_by_day"]],
            [headers[1], headers[2]],
        )


if __name__ == "__main__":
    unittest.main()
