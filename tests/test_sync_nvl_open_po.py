"""Unit tests cho đồng bộ Tồn Đơn hàng Ton_NVL!E."""

from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook

from nvl.open_po import (
    OpenPOConfig,
    patch_target_workbook,
    read_open_po,
    reconcile_target,
    verify_target,
)


def make_config():
    return OpenPOConfig(
        source_name="BCTheodoiDMHANG.xlsm",
        source_sourcedoc="1836226D-FB9D-4F8A-B6D5-3508A975936C",
        source_sheet="REPORT_DONMUAHANG",
        source_start_row=6,
        source_code_col=6,
        source_purchase_col=9,
        source_received_col=11,
        source_status_col=28,
        target_name="Kế hoạch mua hàng.xlsx",
        target_path="Tinh san xuat Mua hang 2027/Kế hoạch mua hàng.xlsx",
        target_sourcedoc="89D1BA7B-006E-4527-B879-ABF120309214",
        target_sheet="Ton_NVL",
        target_start_row=2,
        target_code_col=1,
        target_value_col=5,
    )


def make_source_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "REPORT_DONMUAHANG"
    ws.cell(5, 6, "Mã")
    ws.cell(5, 9, "Số lượng mua")
    ws.cell(5, 11, "Số lượng nhận")
    ws.cell(5, 28, "Trạng thái")

    rows = [
        ("VT1", 100, 20, ""),
        ("VT1", 50, 60, ""),
        ("VT2", 40, 40, ""),
        ("VT3", 10, 0, "Đã hoàn tất"),
    ]
    for row_num, (code, purchase, received, status) in enumerate(rows, 6):
        ws.cell(row_num, 6, code)
        ws.cell(row_num, 9, purchase)
        ws.cell(row_num, 11, received)
        ws.cell(row_num, 28, status)

    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def make_target_bytes(*, formula=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ton_NVL"
    ws.append(["Mã NVL", "Tên NVL", "ĐVT", "Tồn Cuối", "Tồn Đơn hàng"])
    ws.append(["VT1", "Vật tư 1", "Kg", 10, "=1+1" if formula else None])
    ws.append(["VT2", "Vật tư 2", "Kg", 20, None])
    ws.append(["VT4", "Vật tư 4", "Kg", 30, 999])
    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


class OpenPOTests(unittest.TestCase):
    def test_read_open_po_filters_and_clamps_per_line(self):
        totals, info = read_open_po(make_source_bytes(), make_config())
        self.assertEqual(totals, {"VT1": 80.0, "VT2": 0.0})
        self.assertEqual(info["open_rows"], 3)
        self.assertEqual(info["over_received_rows"], 1)
        self.assertEqual(info["positive_rows"], 1)

    def test_reconcile_sets_missing_code_to_zero(self):
        totals, _ = read_open_po(make_source_bytes(), make_config())
        changes, expected = reconcile_target(make_target_bytes(), totals, make_config())
        self.assertEqual(expected["VT1"], 80.0)
        self.assertEqual(expected["VT2"], 0.0)
        self.assertEqual(expected["VT4"], 0.0)
        by_code = {item["code"]: item for item in changes}
        self.assertEqual(by_code["VT4"]["current"], 999)
        self.assertEqual(by_code["VT4"]["target"], 0.0)

    def test_formula_in_target_e_is_rejected(self):
        totals, _ = read_open_po(make_source_bytes(), make_config())
        with self.assertRaises(RuntimeError):
            reconcile_target(make_target_bytes(formula=True), totals, make_config())

    def test_patch_only_e_and_verify(self):
        target = make_target_bytes()
        totals, _ = read_open_po(make_source_bytes(), make_config())
        changes, expected = reconcile_target(target, totals, make_config())
        patched = patch_target_workbook(target, changes, make_config())
        verify_target(patched, expected, make_config())

        wb = load_workbook(BytesIO(patched), data_only=True)
        try:
            ws = wb["Ton_NVL"]
            self.assertEqual(ws["D2"].value, 10)
            self.assertEqual(ws["E2"].value, 80)
            self.assertEqual(ws["E3"].value, 0)
            self.assertEqual(ws["E4"].value, 0)
        finally:
            wb.close()


if __name__ == "__main__":
    unittest.main()
