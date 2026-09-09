"""Tests for KHSX_ki weekly summary aggregation, formatting, and verification."""
import calendar
from io import BytesIO
from pathlib import Path
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

import sync_planning_khsx_ki as ki
from sync_planning_weekly_model import ENGINE_VERSION, compute_planning_inputs_hash

AUDIT_SNAPSHOT = Path("temp_audit/survey_run_34190029619/planning_proposal.xlsx")


def make_mock_khsx_ki_workbook(*, days=None, skus=None, month=9, year=2026, with_merges=True):
    """Tạo workbook mẫu có cả Ke_hoach_SX và KHSX_ki để kiểm thử độc lập."""
    days_in_month = calendar.monthrange(year, month)[1]
    total_days = days if days is not None else days_in_month
    wb = Workbook()
    ws_kh = wb.active
    ws_kh.title = "Ke_hoach_SX"

    headers_kh = [
        "Mã Sản Phẩm", "Tên Sản Phẩm", "Đơn vị tính", "Số lượng /mẻ",
        "Số lượng/ ca", "Chuyền", "Nhóm sản phẩm", "Phân loại SP",
        "Số ca theo ngày", "Tồn đầu thực tế", "Tồn đầu sổ sách", "FC",
        "Tồn cuối dự kiến", "Nợ kho", "Số lượng cần sản xuất",
        "Số lượng sản xuất theo mẻ/ca", "Số ngày cần sản xuất",
        "Ngày bắt đầu sản xuất",
    ] + [f"{d:02d}/{month:02d}" for d in range(1, total_days + 1)]
    ws_kh.append(headers_kh)

    sku_list = skus or [
        {"code": 130100011, "name": "Đảnh Thạnh chanh", "uom": "Thùng", "p": 86400.0, "daily": [12000.0] * 6 + [12000.0, 2400.0] + [0.0] * max(0, total_days - 8)},
        {"code": 130100096, "name": "Vikoda Alkaline", "uom": "Thùng", "p": 0.0, "daily": [0.0] * total_days},
    ]

    for s in sku_list:
        row = [
            s["code"], s["name"], s["uom"], 100, 500, "KHS", "Nhom", "Không đường",
            1, 0, 0, s["p"], 0, 0, s["p"], s["p"], 1, None,
        ] + [s["daily"][d - 1] if d <= len(s["daily"]) and s["daily"][d - 1] > 0 else None for d in range(1, total_days + 1)]
        ws_kh.append(row)

    ws_ki = wb.create_sheet("KHSX_ki")
    ws_ki.append(["CÔNG TY CỔ PHẦN NƯỚC KHOÁNG KHÁNH HÒA"])
    ws_ki.append(["KẾ HOẠCH SẢN XUẤT THÁNG"])
    ws_ki.append([f"Kỳ kế hoạch: Tháng {month}"])
    ws_ki.append([f"Ngày lập: 01/{month:02d}/{year}"])
    ws_ki.append([])

    # Row 6: Tiêu đề
    std_wks = ki.compute_month_weeks(year, month, num_week_cols=5)
    header_ki = ["STT", "Mã SP", "Tên Sản Phẩm", "ĐVT"] + [w["label"] for w in std_wks] + ["Tổng cộng"]
    ws_ki.append(header_ki)

    empty_cols = [None] * (len(std_wks) + 1)
    for idx, s in enumerate(sku_list, start=1):
        ws_ki.append([idx, s["code"], s["name"], s["uom"]] + empty_cols)

    ws_ki.append(["TỔNG CỘNG", None, None, None] + empty_cols)

    # Định dạng theo chuẩn mẫu template SharePoint
    total_col_idx = len(header_ki)
    for c in range(1, total_col_idx + 1):
        ws_ki.cell(6, c).font = Font(bold=True)
    ws_ki.column_dimensions["J"].width = 17.0 if total_col_idx == 10 else 19.44140625
    if total_col_idx == 11:
        ws_ki.column_dimensions["K"].width = 17.0
    ws_ki.column_dimensions["I"].width = 19.44140625
    total_row_idx = 6 + len(sku_list) + 1
    for r in range(7, 7 + len(sku_list)):
        for c in range(5, total_col_idx):
            ws_ki.cell(r, c).font = Font(bold=False)
            ws_ki.cell(r, c).number_format = r'_(* #,##0_);_(* \(#,##0\);_(* "-"_);_(@_)'
        ws_ki.cell(r, total_col_idx).font = Font(bold=True)
        ws_ki.cell(r, total_col_idx).number_format = "#,##0"
    for c in range(5, total_col_idx + 1):
        ws_ki.cell(total_row_idx, c).font = Font(bold=True)
        ws_ki.cell(total_row_idx, c).number_format = "#,##0"

    if with_merges:
        from openpyxl.utils.cell import get_column_letter
        tot_letter = get_column_letter(total_col_idx)
        for r in range(1, 5):
            ws_ki.merge_cells(f"A{r}:{tot_letter}{r}")
        ws_ki.cell(31, 9).value = "TP.KẾ HOẠCH"
        ws_ki.merge_cells(f"I31:{tot_letter}31")

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


class KHSXKiTests(unittest.TestCase):

    def test_has_khsx_ki_sheet(self):
        wb = Workbook()
        buf = BytesIO()
        wb.save(buf)
        self.assertFalse(ki.has_khsx_ki_sheet(buf.getvalue()))

        wb.create_sheet("KHSX_ki")
        buf = BytesIO()
        wb.save(buf)
        self.assertTrue(ki.has_khsx_ki_sheet(buf.getvalue()))

    def test_compute_month_weeks_coverage_all_month_lengths(self):
        """Kiểm tra ranh giới tuần cho các độ dài tháng 28, 29, 30, 31 ngày."""
        test_cases = [
            (2026, 2, 28),   # Tháng 2 năm thường (28 ngày, CN đầu tiên = ngày 1)
            (2028, 2, 29),   # Tháng 2 năm nhuận (29 ngày, CN đầu tiên = ngày 6)
            (2027, 2, 28),   # Tháng 2 bắt đầu vào Thứ Hai (đủ 4 tuần x 7 ngày)
            (2026, 9, 30),   # Tháng 9 (30 ngày)
            (2026, 11, 30),  # Tháng 11 (30 ngày, CN đầu tiên = ngày 1)
            (2026, 1, 31),   # Tháng 1 (31 ngày)
            (2026, 10, 31),  # Tháng 10 (31 ngày)
        ]
        for year, month, expected_days in test_cases:
            weeks = ki.compute_month_weeks(year, month, num_week_cols=5)
            covered = []
            for w in weeks:
                covered.extend(w["days"])
                if w["days"]:
                    self.assertEqual(w["days"], list(range(w["start_day"], w["end_day"] + 1)))
            self.assertEqual(covered, list(range(1, expected_days + 1)), f"Mismatch for {year}-{month:02d}")

    def test_mock_workbook_patch_and_verify(self):
        """Kiểm thử tính toán tuần trên mock workbook."""
        wb_bytes = make_mock_khsx_ki_workbook(days=30, month=9, year=2026)
        patched_bytes, report = ki.patch_khsx_ki_workbook(wb_bytes, plan_year=2026, plan_month=9)
        self.assertTrue(report["ok"])
        self.assertEqual(report["checked_skus"], 2)
        self.assertAlmostEqual(report["grand_total"], 86400.0, places=4)

        verify = ki.verify_khsx_ki(patched_bytes, plan_year=2026, plan_month=9)
        self.assertTrue(verify["ok"])
        self.assertEqual(verify["checked_skus"], 2)

    def test_sku_order_independence(self):
        """Kiểm tra thứ tự SKU khác nhau giữa Ke_hoach_SX và KHSX_ki vẫn khớp chính xác."""
        # Tạo workbook với KHSX_ki có thứ tự đảo ngược
        wb = Workbook()
        ws_kh = wb.active
        ws_kh.title = "Ke_hoach_SX"
        headers_kh = ["Mã SP", "Tên", "ĐVT", "", "", "", "", "", "", "", "", "", "", "", "Nhu cầu", "P", "Q", "Start"] + [f"{d:02d}/09" for d in range(1, 31)]
        ws_kh.append(headers_kh)
        # Ke_hoach_SX: SKU 1 trước, SKU 2 sau
        ws_kh.append([130100001, "SP 1", "Thùng", 0,0,0,0,0,0,0,0,0,0,0,5000, 5000.0, 1, None] + [5000.0] + [None] * 29)
        ws_kh.append([130100002, "SP 2", "Thùng", 0,0,0,0,0,0,0,0,0,0,0,10000, 10000.0, 1, None] + [None] * 7 + [10000.0] + [None] * 22)

        ws_ki = wb.create_sheet("KHSX_ki")
        ws_ki.append(["Title"])
        ws_ki.append(["Subtitle"])
        ws_ki.append(["Kỳ kế hoạch: Tháng 9"])
        ws_ki.append(["Ngày lập: 01/09/2026"])
        ws_ki.append([])
        std_wks = ki.compute_month_weeks(2026, 9)
        ws_ki.append(["STT", "Mã SP", "Tên", "ĐVT"] + [w["label"] for w in std_wks] + ["Tổng cộng"])
        # KHSX_ki: SKU 2 trước, SKU 1 sau
        ws_ki.append([1, 130100002, "SP 2", "Thùng", None, None, None, None, None, None])
        ws_ki.append([2, 130100001, "SP 1", "Thùng", None, None, None, None, None, None])
        ws_ki.append(["TỔNG CỘNG", None, None, None, None, None, None, None, None, None])

        buf = BytesIO()
        wb.save(buf)

        patched_bytes, report = ki.patch_khsx_ki_workbook(buf.getvalue(), plan_year=2026, plan_month=9)
        verify = ki.verify_khsx_ki(patched_bytes, plan_year=2026, plan_month=9)
        self.assertTrue(verify["ok"])

        # Kiểm tra nội dung trực tiếp
        wb_check = load_workbook(BytesIO(patched_bytes), data_only=True)
        ws_check = wb_check["KHSX_ki"]
        # Row 7 is SKU 2 (10000 in Week 2, total 10000)
        self.assertEqual(ws_check.cell(7, 2).value, 130100002)
        self.assertAlmostEqual(float(ws_check.cell(7, 6).value), 10000.0)  # Week 2 is col 6
        self.assertAlmostEqual(float(ws_check.cell(7, 10).value), 10000.0) # Total is col 10

        # Row 8 is SKU 1 (5000 in Week 1, total 5000)
        self.assertEqual(ws_check.cell(8, 2).value, 130100001)
        self.assertAlmostEqual(float(ws_check.cell(8, 5).value), 5000.0)   # Week 1 is col 5
        self.assertAlmostEqual(float(ws_check.cell(8, 10).value), 5000.0)  # Total is col 10

        # Row 9 is TỔNG CỘNG
        self.assertAlmostEqual(float(ws_check.cell(9, 5).value), 5000.0)
        self.assertAlmostEqual(float(ws_check.cell(9, 6).value), 10000.0)
        self.assertAlmostEqual(float(ws_check.cell(9, 10).value), 15000.0)

    def test_real_sharepoint_proposal_snapshot_parity(self):
        """Kiểm tra đối chiếu toàn diện 100% SKU trên file proposal SharePoint thật."""
        if not AUDIT_SNAPSHOT.exists():
            self.skipTest(f"Không tìm thấy file snapshot {AUDIT_SNAPSHOT}")

        wb_bytes = AUDIT_SNAPSHOT.read_bytes()
        patched_bytes, report = ki.patch_khsx_ki_workbook(wb_bytes, plan_year=2026, plan_month=9)
        self.assertTrue(report["ok"])
        self.assertEqual(report["checked_skus"], 20)
        self.assertAlmostEqual(report["grand_total"], 673483.3333333333, places=4)

        verify = ki.verify_khsx_ki(patched_bytes, plan_year=2026, plan_month=9)
        self.assertTrue(verify["ok"])
        self.assertEqual(verify["checked_skus"], 20)
        self.assertAlmostEqual(verify["grand_total"], 673483.3333333333, places=4)

        # Kiểm chứng riêng SKU 130100011
        wb_check = load_workbook(BytesIO(patched_bytes), data_only=True)
        ws_ki = wb_check["KHSX_ki"]
        row_11 = None
        for r in range(7, 27):
            if ws_ki.cell(r, 2).value == 130100011:
                row_11 = r
                break
        self.assertIsNotNone(row_11)
        # Tuần 1: 72000, Tuần 2: 14400, Tuần 3-5: 0, Tổng: 86400
        self.assertAlmostEqual(float(ws_ki.cell(row_11, 5).value), 72000.0)
        self.assertAlmostEqual(float(ws_ki.cell(row_11, 6).value), 14400.0)
        self.assertAlmostEqual(float(ws_ki.cell(row_11, 7).value), 0.0)
        self.assertAlmostEqual(float(ws_ki.cell(row_11, 8).value), 0.0)
        self.assertAlmostEqual(float(ws_ki.cell(row_11, 9).value), 0.0)
        self.assertAlmostEqual(float(ws_ki.cell(row_11, 10).value), 86400.0)

    def test_input_hash_excludes_khsx_ki(self):
        """Khẳng định sửa đổi trên KHSX_ki không làm thay đổi compute_planning_inputs_hash."""
        wb_bytes = make_mock_khsx_ki_workbook(days=30, month=9, year=2026)
        h1 = compute_planning_inputs_hash(wb_bytes)

        patched_bytes, _ = ki.patch_khsx_ki_workbook(wb_bytes, plan_year=2026, plan_month=9)
        h2 = compute_planning_inputs_hash(patched_bytes)

        self.assertEqual(h1, h2, "compute_planning_inputs_hash bị thay đổi bởi KHSX_ki!")

    def test_two_round_pipeline_idempotence(self):
        """Kiểm tra tính lũy đẳng 2 vòng (idempotence)."""
        wb_bytes = make_mock_khsx_ki_workbook(days=30, month=9, year=2026)
        p1, r1 = ki.patch_khsx_ki_workbook(wb_bytes, plan_year=2026, plan_month=9)
        p2, r2 = ki.patch_khsx_ki_workbook(p1, plan_year=2026, plan_month=9)

        self.assertEqual(r1["grand_total"], r2["grand_total"])
        self.assertEqual(r1["checked_skus"], r2["checked_skus"])

        v1 = ki.verify_khsx_ki(p1, plan_year=2026, plan_month=9)
        v2 = ki.verify_khsx_ki(p2, plan_year=2026, plan_month=9)
        self.assertTrue(v1["ok"])
        self.assertTrue(v2["ok"])

    def test_pipeline_integration_with_khsx_ki(self):
        """Kiểm tra toàn bộ luồng prepare_pipeline_output khi workbook có sheet KHSX_ki."""
        from planning_pipeline import prepare_pipeline_output
        from tests.test_run_offline import (
            _actual_source,
            _single_value_source,
            _target_workbook,
            CODE,
            CODE_VKD,
        )

        source_bytes = {
            "actual_stock": _actual_source(),
            "factory_vikoda": _single_value_source(CODE, value_col_index=11, value=0, receipt=50),
            "factory_vkd": _single_value_source(CODE_VKD, value_col_index=11, value=0),
            "accounting_vikoda": _single_value_source(CODE, value_col_index=12, value=0),
            "accounting_vkd": _single_value_source(CODE_VKD, value_col_index=12, value=0),
        }

        wb = load_workbook(BytesIO(_target_workbook()))
        wb["FC"]["M2"] = 2600

        # Thêm sheet KHSX_ki vào target workbook
        ws_ki = wb.create_sheet("KHSX_ki")
        ws_ki.append(["CÔNG TY CỔ PHẦN NƯỚC KHOÁNG KHÁNH HÒA"])
        ws_ki.append(["KẾ HOẠCH SẢN XUẤT THÁNG"])
        ws_ki.append(["Kỳ kế hoạch: Tháng 9"])
        ws_ki.append(["Ngày lập: 01/09/2026"])
        ws_ki.append([])
        std_wks = ki.compute_month_weeks(2026, 9)
        ws_ki.append(["STT", "Mã SP", "Tên", "ĐVT"] + [w["label"] for w in std_wks] + ["Tổng cộng"])
        ws_ki.append([1, CODE, "SP", "Thùng", None, None, None, None, None, None])
        ws_ki.append(["TỔNG CỘNG", None, None, None, None, None, None, None, None, None])

        buf = BytesIO()
        wb.save(buf)
        target_bytes = buf.getvalue()

        # Round 1
        out1, report1, state1 = prepare_pipeline_output(
            target_bytes,
            source_bytes,
            runtime_state={},
            input_revision={"target": {"etag": "test-etag"}},
        )

        pipeline_info1 = report1["pipeline"]
        self.assertIn("khsx_ki", pipeline_info1["steps"])
        self.assertEqual(pipeline_info1["engine_version"], ENGINE_VERSION)
        self.assertTrue(pipeline_info1["verify"]["khsx_ki"]["ok"])

        # Round 2: Idempotence test
        out2, report2, state2 = prepare_pipeline_output(
            out1,
            source_bytes,
            runtime_state=state1,
            input_revision={"target": {"etag": "test-etag"}},
        )
        self.assertTrue(report2["pipeline"]["verify"]["khsx_ki"]["ok"])

    def test_missing_active_sku_blocked(self):
        """[P1 Defect 1] Chặn khi KHSX_ki bị xóa/thiếu mã SKU so với Ke_hoach_SX."""
        base_bytes = make_mock_khsx_ki_workbook()
        wb = load_workbook(BytesIO(base_bytes))
        # Row 7 is SKU 130100011
        wb["KHSX_ki"].delete_rows(7)
        buf = BytesIO()
        wb.save(buf)
        tampered_bytes = buf.getvalue()

        # Writer must block
        with self.assertRaises(RuntimeError) as ctx_patch:
            ki.patch_khsx_ki_workbook(tampered_bytes, plan_year=2026, plan_month=9)
        self.assertIn("thiếu 1 mã SKU", str(ctx_patch.exception))
        self.assertIn("130100011", str(ctx_patch.exception))

        # Verifier must also block
        with self.assertRaises(RuntimeError) as ctx_verify:
            ki.verify_khsx_ki(tampered_bytes, plan_year=2026, plan_month=9)
        self.assertIn("thiếu 1 mã SKU", str(ctx_verify.exception))
        self.assertIn("130100011", str(ctx_verify.exception))

    def test_duplicate_sku_blocked_in_both_sheets(self):
        """[P1 Defect 1] Chặn khi có mã SKU trùng lặp trong KHSX_ki hoặc Ke_hoach_SX."""
        base_bytes = make_mock_khsx_ki_workbook()

        # Ca 1: Trùng trong KHSX_ki
        wb1 = load_workbook(BytesIO(base_bytes))
        ws_ki = wb1["KHSX_ki"]
        ws_ki.cell(8, 2).value = ws_ki.cell(7, 2).value  # Đặt dòng 8 cùng mã dòng 7 (130100011)
        buf1 = BytesIO()
        wb1.save(buf1)
        tampered_ki = buf1.getvalue()

        with self.assertRaises(RuntimeError) as ctx1_patch:
            ki.patch_khsx_ki_workbook(tampered_ki, plan_year=2026, plan_month=9)
        self.assertIn("Trùng mã SKU 130100011 trong sheet KHSX_ki", str(ctx1_patch.exception))

        with self.assertRaises(RuntimeError) as ctx1_ver:
            ki.verify_khsx_ki(tampered_ki, plan_year=2026, plan_month=9)
        self.assertIn("Trùng mã SKU 130100011 trong sheet KHSX_ki", str(ctx1_ver.exception))

        # Ca 2: Trùng trong Ke_hoach_SX
        wb2 = load_workbook(BytesIO(base_bytes))
        ws_kh = wb2["Ke_hoach_SX"]
        ws_kh.cell(3, 1).value = ws_kh.cell(2, 1).value  # Dòng 3 trùng mã dòng 2
        buf2 = BytesIO()
        wb2.save(buf2)
        tampered_kh = buf2.getvalue()

        with self.assertRaises(RuntimeError) as ctx2_patch:
            ki.patch_khsx_ki_workbook(tampered_kh, plan_year=2026, plan_month=9)
        self.assertIn("Trùng mã SKU 130100011 trong sheet Ke_hoach_SX", str(ctx2_patch.exception))

        with self.assertRaises(RuntimeError) as ctx2_ver:
            ki.verify_khsx_ki(tampered_kh, plan_year=2026, plan_month=9)
        self.assertIn("Trùng mã SKU 130100011 trong sheet Ke_hoach_SX", str(ctx2_ver.exception))

    def test_unknown_sku_blocked(self):
        """[P1 Defect 1] Chặn khi KHSX_ki chứa mã SKU lạ không có trong Ke_hoach_SX."""
        base_bytes = make_mock_khsx_ki_workbook()
        wb = load_workbook(BytesIO(base_bytes))
        ws_ki = wb["KHSX_ki"]
        # Chèn thêm 1 dòng SKU lạ trước dòng TỔNG CỘNG
        ws_ki.insert_rows(9)
        ws_ki.cell(9, 1).value = 3
        ws_ki.cell(9, 2).value = 130199999
        ws_ki.cell(9, 3).value = "Mã SKU lạ không có trong Ke_hoach_SX"
        ws_ki.cell(9, 4).value = "Thùng"
        buf = BytesIO()
        wb.save(buf)
        tampered_bytes = buf.getvalue()

        with self.assertRaises(RuntimeError) as ctx_patch:
            ki.patch_khsx_ki_workbook(tampered_bytes, plan_year=2026, plan_month=9)
        self.assertIn("chứa 1 mã SKU không có trong Ke_hoach_SX", str(ctx_patch.exception))
        self.assertIn("130199999", str(ctx_patch.exception))

        with self.assertRaises(RuntimeError) as ctx_ver:
            ki.verify_khsx_ki(tampered_bytes, plan_year=2026, plan_month=9)
        self.assertIn("chứa 1 mã SKU không có trong Ke_hoach_SX", str(ctx_ver.exception))
        self.assertIn("130199999", str(ctx_ver.exception))

    def test_four_week_month_readonly_verification(self):
        """[P1 Defect 2] Tháng có 4 tuần (Tháng 2/2027) xác minh thành công trên read-only workbook."""
        wb_bytes = make_mock_khsx_ki_workbook(month=2, year=2027)
        patched, report = ki.patch_khsx_ki_workbook(wb_bytes, plan_year=2027, plan_month=2)
        self.assertTrue(report["ok"])
        self.assertEqual(report["num_weeks"], 4)

        # Verifier thuần đọc mở read_only=True không được ném AttributeError: Cell is read only
        verify = ki.verify_khsx_ki(patched, plan_year=2027, plan_month=2)
        self.assertTrue(verify["ok"])
        self.assertEqual(verify["num_weeks"], 4)

        # Kiểm tra nội dung chi tiết trên workbook đã patch
        wb_check = load_workbook(BytesIO(patched), data_only=True)
        ws_ki = wb_check["KHSX_ki"]
        # Cột 9 (Tuần 5) là tuần rỗng: giá trị 0.0
        self.assertEqual(ws_ki.cell(6, 9).value, "Tuần 5\n-")
        self.assertAlmostEqual(float(ws_ki.cell(7, 9).value or 0.0), 0.0)
        self.assertAlmostEqual(float(ws_ki.cell(8, 9).value or 0.0), 0.0)
        # Cột 10 (Tổng cộng) có giá trị 86400
        self.assertAlmostEqual(float(ws_ki.cell(7, 10).value), 86400.0)
        # Cột 11 (K) hoàn toàn rỗng
        self.assertIsNone(ws_ki.cell(6, 11).value)
        self.assertIsNone(ws_ki.cell(7, 11).value)
        wb_check.close()

    def test_six_week_month_layout_and_verification(self):
        """[P2 Defect 3] Tháng có 6 tuần (Tháng 11/2026) mở rộng sang Cột K và phân hoạch đúng ranh giới."""
        wb_bytes = make_mock_khsx_ki_workbook(month=11, year=2026)
        patched, report = ki.patch_khsx_ki_workbook(wb_bytes, plan_year=2026, plan_month=11)
        self.assertTrue(report["ok"])
        self.assertEqual(report["num_weeks"], 6)

        verify = ki.verify_khsx_ki(patched, plan_year=2026, plan_month=11)
        self.assertTrue(verify["ok"])
        self.assertEqual(verify["num_weeks"], 6)
        self.assertAlmostEqual(verify["grand_total"], 86400.0)

        wb_check = load_workbook(BytesIO(patched), data_only=True)
        ws_ki = wb_check["KHSX_ki"]

        # Cột 10 (J) là Tuần 6 (30/11-30/11)
        self.assertEqual(ws_ki.cell(6, 10).value, "Tuần 6\n30/11-30/11")
        # Cột 11 (K) là Cột Tổng cộng
        self.assertEqual(ws_ki.cell(6, 11).value, "Tổng cộng")
        self.assertAlmostEqual(float(ws_ki.cell(7, 11).value), 86400.0)

        # Kiểm tra tiêu đề A1..A4 và dòng chữ ký 31 đã mở rộng sang cột K
        merged_coords = [rng.coord for rng in ws_ki.merged_cells.ranges]
        self.assertIn("A1:K1", merged_coords)
        self.assertIn("A2:K2", merged_coords)
        self.assertIn("A3:K3", merged_coords)
        self.assertIn("A4:K4", merged_coords)
        self.assertIn("I31:K31", merged_coords)
        wb_check.close()

        # Kiểm tra thêm Tháng 3/2026 (31 ngày, 6 tuần)
        mar_bytes = make_mock_khsx_ki_workbook(month=3, year=2026)
        mar_patched, mar_rep = ki.patch_khsx_ki_workbook(mar_bytes, plan_year=2026, plan_month=3)
        self.assertEqual(mar_rep["num_weeks"], 6)
        mar_ver = ki.verify_khsx_ki(mar_patched, plan_year=2026, plan_month=3)
        self.assertTrue(mar_ver["ok"])
        self.assertEqual(mar_ver["num_weeks"], 6)

    def test_switching_between_6_5_4_weeks_cleans_managed_cells(self):
        """[P2 Defect 3] Chuyển đổi giữa tháng 6 tuần, 5 tuần và 4 tuần dọn sạch Cột K và phục hồi merge."""
        # 1. Bắt đầu với tháng 6 tuần (Tháng 11/2026) -> Cột 11 có dữ liệu
        wb_bytes = make_mock_khsx_ki_workbook(month=11, year=2026)
        p_6wk, _ = ki.patch_khsx_ki_workbook(wb_bytes, plan_year=2026, plan_month=11)
        self.assertTrue(ki.verify_khsx_ki(p_6wk, plan_year=2026, plan_month=11)["ok"])

        # 2. Lấy output 6 tuần chạy sang Tháng 9/2026 (5 tuần)
        p_5wk, r_5wk = ki.patch_khsx_ki_workbook(p_6wk, plan_year=2026, plan_month=9)
        self.assertTrue(r_5wk["ok"])
        self.assertEqual(r_5wk["num_weeks"], 5)
        v_5wk = ki.verify_khsx_ki(p_5wk, plan_year=2026, plan_month=9)
        self.assertTrue(v_5wk["ok"])

        wb_5 = load_workbook(BytesIO(p_5wk), data_only=True)
        ws_5 = wb_5["KHSX_ki"]
        self.assertEqual(ws_5.cell(6, 10).value, "Tổng cộng")
        # Cột K phải được xóa sạch
        self.assertIsNone(ws_5.cell(6, 11).value)
        self.assertIsNone(ws_5.cell(7, 11).value)
        # Hợp nhất phục hồi về cột J
        merged_5 = [rng.coord for rng in ws_5.merged_cells.ranges]
        self.assertIn("A1:J1", merged_5)
        self.assertIn("I31:J31", merged_5)
        self.assertNotIn("A1:K1", merged_5)
        self.assertNotIn("I31:K31", merged_5)
        wb_5.close()

        # 3. Lấy output 5 tuần chạy sang Tháng 2/2027 (4 tuần)
        p_4wk, r_4wk = ki.patch_khsx_ki_workbook(p_5wk, plan_year=2027, plan_month=2)
        self.assertTrue(r_4wk["ok"])
        self.assertEqual(r_4wk["num_weeks"], 4)
        v_4wk = ki.verify_khsx_ki(p_4wk, plan_year=2027, plan_month=2)
        self.assertTrue(v_4wk["ok"])

        wb_4 = load_workbook(BytesIO(p_4wk), data_only=True)
        ws_4 = wb_4["KHSX_ki"]
        self.assertEqual(ws_4.cell(6, 9).value, "Tuần 5\n-")
        self.assertEqual(ws_4.cell(6, 10).value, "Tổng cộng")
        self.assertIsNone(ws_4.cell(6, 11).value)
        wb_4.close()

    def test_overlapping_or_corrupted_week_header_rejected(self):
        """[P2 Defect 3 & 4] Tiêu đề tuần bị chồng lấn/sai lệch bị verifier từ chối và được patch sửa chuẩn."""
        base_bytes = make_mock_khsx_ki_workbook()
        wb = load_workbook(BytesIO(base_bytes))
        ws = wb["KHSX_ki"]
        # Giả lập người dùng sửa tay tiêu đề tuần 3 và tuần 4 chồng lấn ngày 21/09
        ws["G6"] = "Tuần 3\n14/09-21/09"
        ws["H6"] = "Tuần 4\n21/09-27/09"
        buf = BytesIO()
        wb.save(buf)
        tampered_bytes = buf.getvalue()

        # Verifier chạy trực tiếp trên file bị sửa phải chặn lại
        with self.assertRaises(RuntimeError) as ctx:
            ki.verify_khsx_ki(tampered_bytes, plan_year=2026, plan_month=9)
        self.assertIn("không khớp dải ngày chuẩn", str(ctx.exception))

        # Writer patch lại phải ghi đè tiêu đề chuẩn và kết quả pass verify
        repaired_bytes, report = ki.patch_khsx_ki_workbook(tampered_bytes, plan_year=2026, plan_month=9)
        self.assertTrue(report["ok"])
        verify = ki.verify_khsx_ki(repaired_bytes, plan_year=2026, plan_month=9)
        self.assertTrue(verify["ok"])

    def test_user_edited_headers_do_not_alter_schedule_rules(self):
        """[P2 Defect 4] Code làm chủ tiêu đề tuần; chỉnh sửa ô header không làm thay đổi kết quả kế hoạch."""
        base_bytes = make_mock_khsx_ki_workbook()
        wb = load_workbook(BytesIO(base_bytes))
        ws = wb["KHSX_ki"]
        # Người dùng sửa tiêu đề E6 và F6
        ws["E6"] = "Tuần 1\n01/09-07/09"
        ws["F6"] = "Tuần 2\n08/09-13/09"
        buf = BytesIO()
        wb.save(buf)
        altered_bytes = buf.getvalue()

        orig_out, _ = ki.patch_khsx_ki_workbook(base_bytes, plan_year=2026, plan_month=9)
        altered_out, _ = ki.patch_khsx_ki_workbook(altered_bytes, plan_year=2026, plan_month=9)

        b1 = load_workbook(BytesIO(orig_out), data_only=True)
        b2 = load_workbook(BytesIO(altered_out), data_only=True)
        # Sản lượng tuần 1 phải bằng nhau (72000), không bị ăn lẹm ngày 7
        self.assertEqual(b1["KHSX_ki"]["E7"].value, b2["KHSX_ki"]["E7"].value)
        self.assertEqual(b1["KHSX_ki"]["F7"].value, b2["KHSX_ki"]["F7"].value)
        self.assertEqual(b2["KHSX_ki"]["E6"].value, "Tuần 1\n01/09-06/09")
        b1.close()
        b2.close()

    def test_calendar_partition_invariant_all_24_months(self):
        """[P2 Defect 3] Xác minh phân hoạch tuần Monday-Sunday trên toàn bộ 24 tháng 2026-2027."""
        for year in (2026, 2027):
            for month in range(1, 13):
                days_in_month = calendar.monthrange(year, month)[1]
                weeks = ki.compute_standard_calendar_weeks(year, month)
                self.assertIn(len(weeks), (4, 5, 6), f"Số tuần ngoài dải 4-6: {year}-{month:02d}")

                all_days = []
                for idx, w in enumerate(weeks, start=1):
                    self.assertEqual(w["week_num"], idx)
                    self.assertGreaterEqual(len(w["days"]), 1)
                    self.assertLessEqual(len(w["days"]), 7, f"Tuần dài hơn 7 ngày tại {year}-{month:02d} tuần {idx}")
                    self.assertEqual(w["start_day"], w["days"][0])
                    self.assertEqual(w["end_day"], w["days"][-1])
                    self.assertEqual(w["days"], list(range(w["start_day"], w["end_day"] + 1)))
                    all_days.extend(w["days"])

                self.assertEqual(all_days, list(range(1, days_in_month + 1)))
                self.assertEqual(len(all_days), len(set(all_days)), "Trùng ngày trong phân hoạch tuần!")

    def test_note_outside_table_preserved(self):
        """[Defect 1] Ghi chú của người dùng ngoài bảng (K35) phải được bảo toàn khi patch cho tháng 4, 5, 6 tuần."""
        base_raw = make_mock_khsx_ki_workbook(year=2026, month=9)
        wb = load_workbook(BytesIO(base_raw))
        ws = wb["KHSX_ki"]
        ws["K35"] = "Ghi chú của người lập kế hoạch"
        ws["K36"] = "Công thức phụ"
        buf = BytesIO()
        wb.save(buf)
        wb.close()
        raw_with_note = buf.getvalue()

        # Patch tháng 9 (5 tuần)
        out5, rep5 = ki.patch_khsx_ki_workbook(raw_with_note, plan_year=2026, plan_month=9)
        self.assertTrue(rep5["ok"])
        wb5 = load_workbook(BytesIO(out5))
        self.assertEqual(wb5["KHSX_ki"]["K35"].value, "Ghi chú của người lập kế hoạch")
        self.assertEqual(wb5["KHSX_ki"]["K36"].value, "Công thức phụ")
        wb5.close()

        # Patch tháng 11 (6 tuần)
        out6, rep6 = ki.patch_khsx_ki_workbook(out5, plan_year=2026, plan_month=11)
        self.assertTrue(rep6["ok"])
        wb6 = load_workbook(BytesIO(out6))
        self.assertEqual(wb6["KHSX_ki"]["K35"].value, "Ghi chú của người lập kế hoạch")
        self.assertEqual(wb6["KHSX_ki"]["K36"].value, "Công thức phụ")
        wb6.close()

        # Patch lại tháng 9 (chuyển từ 6 tuần về 5 tuần)
        out_back, rep_back = ki.patch_khsx_ki_workbook(out6, plan_year=2026, plan_month=9)
        self.assertTrue(rep_back["ok"])
        wb_back = load_workbook(BytesIO(out_back))
        self.assertEqual(wb_back["KHSX_ki"]["K35"].value, "Ghi chú của người lập kế hoạch")
        self.assertEqual(wb_back["KHSX_ki"]["K36"].value, "Công thức phụ")
        wb_back.close()

    def test_note_preserved_when_total_row_shifts(self):
        """[Defect 1] Ghi chú dưới bảng được bảo toàn kể cả khi vị trí dòng tổng thay đổi."""
        wb = Workbook()
        ws_kh = wb.active
        ws_kh.title = "Ke_hoach_SX"
        ws_ki = wb.create_sheet("KHSX_ki")

        headers_kh = ["Mã SP", "Tên", "ĐVT"] + [""] * 11 + ["Nhu cầu", "P", "Q", "Start"] + [f"{d:02d}/09" for d in range(1, 31)]
        ws_kh.append(headers_kh)
        headers_ki = ["STT", "Mã SP", "Tên", "ĐVT", "Tuần 1", "Tuần 2", "Tuần 3", "Tuần 4", "Tuần 5", "Tổng cộng"]
        for r in range(1, 6):
            ws_ki.append([f"Header {r}"])
        ws_ki.append(headers_ki)

        # 4 SKU -> dòng 7, 8, 9, 10 là SKU, dòng 11 là TỔNG CỘNG
        for i in range(1, 5):
            code = 130100000 + i
            ws_kh.append([code, f"SP {i}", "Thùng"] + [0] * 11 + [1000.0, 1000.0, 1, None, 1000.0] + [None] * 29)
            ws_ki.append([i, code, f"SP {i}", "Thùng", None, None, None, None, None, None])
        ws_ki.append(["TỔNG CỘNG", None, None, None, None, None, None, None, None, None])
        # Đặt ghi chú tại dòng 15 cột K (dưới dòng tổng 11)
        ws_ki["K15"] = "Ghi chú tại dòng 15"

        buf = BytesIO()
        wb.save(buf)
        wb.close()

        out, rep = ki.patch_khsx_ki_workbook(buf.getvalue(), plan_year=2026, plan_month=9)
        self.assertTrue(rep["ok"])
        wb_out = load_workbook(BytesIO(out))
        self.assertEqual(wb_out["KHSX_ki"]["K15"].value, "Ghi chú tại dòng 15")
        wb_out.close()

    def test_idempotence_six_to_six_weeks(self):
        """[Defect 2] Chạy lại 6 -> 6 tuần giữ nguyên định dạng in đậm, độ rộng và number format của Cột J và K."""
        raw6 = make_mock_khsx_ki_workbook(year=2026, month=11)
        p1, rep1 = ki.patch_khsx_ki_workbook(raw6, plan_year=2026, plan_month=11)
        self.assertTrue(rep1["ok"])

        # Đọc định dạng sau lần chạy 1
        wb1 = load_workbook(BytesIO(p1))
        ws1 = wb1["KHSX_ki"]
        k7_bold_1 = ws1["K7"].font.bold
        j7_bold_1 = ws1["J7"].font.bold
        k7_fmt_1 = ws1["K7"].number_format
        j7_fmt_1 = ws1["J7"].number_format
        j_width_1 = ws1.column_dimensions["J"].width
        k_width_1 = ws1.column_dimensions["K"].width
        wb1.close()

        # Chạy lần 2 trên cùng đầu vào (6 -> 6)
        p2, rep2 = ki.patch_khsx_ki_workbook(p1, plan_year=2026, plan_month=11)
        self.assertTrue(rep2["ok"])

        wb2 = load_workbook(BytesIO(p2))
        ws2 = wb2["KHSX_ki"]
        self.assertEqual(ws2["K7"].font.bold, k7_bold_1)
        self.assertEqual(ws2["J7"].font.bold, j7_bold_1)
        self.assertEqual(ws2["K7"].number_format, k7_fmt_1)
        self.assertEqual(ws2["J7"].number_format, j7_fmt_1)
        self.assertEqual(ws2.column_dimensions["J"].width, j_width_1)
        self.assertEqual(ws2.column_dimensions["K"].width, k_width_1)
        self.assertTrue(ws2["K7"].font.bold, "Cột tổng K7 phải in đậm sau khi chạy lại 6->6")
        self.assertFalse(ws2["J7"].font.bold, "Cột tuần J7 không được in đậm")
        wb2.close()

    def test_switching_five_six_six_five(self):
        """[Defect 2] Chuỗi chuyển đổi 5 -> 6 -> 6 -> 5 tuần khôi phục chuẩn xác định dạng tổng cho Cột 10."""
        raw5 = make_mock_khsx_ki_workbook(year=2026, month=9)
        p1, _ = ki.patch_khsx_ki_workbook(raw5, plan_year=2026, plan_month=9)
        wb1 = load_workbook(BytesIO(p1))
        j7_bold_orig = wb1["KHSX_ki"]["J7"].font.bold
        j7_fmt_orig = wb1["KHSX_ki"]["J7"].number_format
        wb1.close()

        # 5 -> 6
        p2, _ = ki.patch_khsx_ki_workbook(p1, plan_year=2026, plan_month=11)
        # 6 -> 6
        p3, _ = ki.patch_khsx_ki_workbook(p2, plan_year=2026, plan_month=11)
        # 6 -> 5
        p4, _ = ki.patch_khsx_ki_workbook(p3, plan_year=2026, plan_month=9)

        wb4 = load_workbook(BytesIO(p4))
        ws4 = wb4["KHSX_ki"]
        self.assertEqual(ws4["J7"].font.bold, j7_bold_orig)
        self.assertEqual(ws4["J7"].number_format, j7_fmt_orig)
        self.assertEqual(ws4["J6"].value, "Tổng cộng")
        self.assertIsNone(ws4["K7"].value)
        self.assertIsNone(ws4["K6"].value)
        wb4.close()

    def test_switching_four_six_four(self):
        """[Defect 2] Chuỗi chuyển đổi 4 -> 6 -> 4 tuần bảo toàn tính đúng đắn dữ liệu và định dạng."""
        raw4 = make_mock_khsx_ki_workbook(year=2027, month=2)
        p1, _ = ki.patch_khsx_ki_workbook(raw4, plan_year=2027, plan_month=2)
        # 4 -> 6
        p2, _ = ki.patch_khsx_ki_workbook(p1, plan_year=2026, plan_month=11)
        # 6 -> 4
        p3, _ = ki.patch_khsx_ki_workbook(p2, plan_year=2027, plan_month=2)

        v3 = ki.verify_khsx_ki(p3, plan_year=2027, plan_month=2)
        self.assertTrue(v3["ok"])
        self.assertEqual(v3["num_weeks"], 4)

        wb3 = load_workbook(BytesIO(p3))
        ws3 = wb3["KHSX_ki"]
        self.assertEqual(ws3["J6"].value, "Tổng cộng")
        self.assertIn("Tuần 5", str(ws3["I6"].value))
        self.assertEqual(ws3["I7"].value, 0.0)
        self.assertIsNone(ws3["K6"].value)
        self.assertIsNone(ws3["K7"].value)
        wb3.close()

    def test_dynamic_signature_block_merge(self):
        """[Defect 1] Khối chữ ký ở dòng bất kỳ (dòng 35) tự động co/dãn hợp nhất theo số tuần."""
        raw5 = make_mock_khsx_ki_workbook(year=2026, month=9)
        wb = load_workbook(BytesIO(raw5))
        ws = wb["KHSX_ki"]
        # Đặt khối chữ ký ở dòng 35
        ws["I35"] = "TP.KẾ HOẠCH"
        ws.merge_cells("I35:J35")
        buf = BytesIO()
        wb.save(buf)
        wb.close()

        # Patch 6 tuần -> merge phải mở rộng thành I35:K35
        out6, _ = ki.patch_khsx_ki_workbook(buf.getvalue(), plan_year=2026, plan_month=11)
        wb6 = load_workbook(BytesIO(out6))
        ranges6 = [str(r) for r in wb6["KHSX_ki"].merged_cells.ranges]
        self.assertIn("I35:K35", ranges6)
        wb6.close()

        # Patch về 5 tuần -> merge phải thu về I35:J35
        out5, _ = ki.patch_khsx_ki_workbook(out6, plan_year=2026, plan_month=9)
        wb5 = load_workbook(BytesIO(out5))
        ranges5 = [str(r) for r in wb5["KHSX_ki"].merged_cells.ranges]
        self.assertIn("I35:J35", ranges5)
        self.assertNotIn("I35:K35", ranges5)
        wb5.close()

    def test_signature_block_and_adjacent_note_preserved_five_to_six(self):
        """[Defect 1A] Khối chữ ký I35:J35 và ghi chú K35 được bảo toàn nguyên vẹn khi chuyển 5 -> 6 tuần."""
        raw5 = make_mock_khsx_ki_workbook(year=2026, month=9)
        wb = load_workbook(BytesIO(raw5))
        ws = wb["KHSX_ki"]
        ws["I35"] = "TP.KẾ HOẠCH"
        ws.merge_cells("I35:J35")
        ws["K35"] = "Ghi chú cần giữ"
        buf = BytesIO()
        wb.save(buf)
        wb.close()

        # Patch sang 6 tuần (tháng 11/2026)
        out6, rep6 = ki.patch_khsx_ki_workbook(buf.getvalue(), plan_year=2026, plan_month=11)
        self.assertTrue(rep6["ok"])
        ver6 = ki.verify_khsx_ki(out6, plan_year=2026, plan_month=11)
        self.assertTrue(ver6["ok"])
        wb6 = load_workbook(BytesIO(out6))
        ws6 = wb6["KHSX_ki"]
        self.assertEqual(ws6["K35"].value, "Ghi chú cần giữ")
        self.assertEqual(ws6["I35"].value, "TP.KẾ HOẠCH")
        ranges6 = [str(r) for r in ws6.merged_cells.ranges]
        # Không mở rộng đè lên K35, giữ nguyên I35:J35
        self.assertIn("I35:J35", ranges6)
        self.assertNotIn("I35:K35", ranges6)
        wb6.close()

    def test_note_with_formula_and_formatting_at_expansion_preserved(self):
        """[Defect 1A] Ghi chú chứa công thức và định dạng tùy biến tại ô mở rộng K35 được bảo toàn."""
        raw5 = make_mock_khsx_ki_workbook(year=2026, month=9)
        wb = load_workbook(BytesIO(raw5))
        ws = wb["KHSX_ki"]
        ws["I35"] = "TP.KẾ HOẠCH"
        ws.merge_cells("I35:J35")
        ws["K35"].value = '="GHI_CHU_" & "2026"'
        ws["K35"].font = Font(name="Arial", size=12, bold=True, italic=True)
        ws["K35"].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        buf = BytesIO()
        wb.save(buf)
        wb.close()

        out6, rep6 = ki.patch_khsx_ki_workbook(buf.getvalue(), plan_year=2026, plan_month=11)
        self.assertTrue(rep6["ok"])

        wb6 = load_workbook(BytesIO(out6))
        ws6 = wb6["KHSX_ki"]
        self.assertEqual(ws6["K35"].value, '="GHI_CHU_" & "2026"')
        self.assertTrue(ws6["K35"].font.bold)
        self.assertTrue(ws6["K35"].font.italic)
        self.assertEqual(ws6["K35"].fill.fill_type, "solid")
        ranges6 = [str(r) for r in ws6.merged_cells.ranges]
        self.assertIn("I35:J35", ranges6)
        self.assertNotIn("I35:K35", ranges6)
        wb6.close()

    def test_vertical_merged_note_preserved_no_crash(self):
        """[Defect 1B] Vùng gộp dọc K35:K36 không bị sửa và không gây crash ValueError."""
        raw5 = make_mock_khsx_ki_workbook(year=2026, month=9)
        wb = load_workbook(BytesIO(raw5))
        ws = wb["KHSX_ki"]
        ws["K35"] = "Ghi chú 2 dòng"
        ws.merge_cells("K35:K36")
        buf = BytesIO()
        wb.save(buf)
        wb.close()

        # Patch tháng 5 tuần (tháng 9/2026)
        out5, rep5 = ki.patch_khsx_ki_workbook(buf.getvalue(), plan_year=2026, plan_month=9)
        self.assertTrue(rep5["ok"])

        wb5 = load_workbook(BytesIO(out5))
        ws5 = wb5["KHSX_ki"]
        self.assertEqual(ws5["K35"].value, "Ghi chú 2 dòng")
        ranges5 = [str(r) for r in ws5.merged_cells.ranges]
        self.assertIn("K35:K36", ranges5)
        wb5.close()

    def test_unrelated_footer_merges_preserved(self):
        """[Defect 1] Các vùng gộp chữ ký khác và vùng gộp tùy ý dưới bảng được giữ nguyên."""
        raw5 = make_mock_khsx_ki_workbook(year=2026, month=9)
        wb = load_workbook(BytesIO(raw5))
        ws = wb["KHSX_ki"]
        ws["B31"] = "CEO"
        ws.merge_cells("B31:C31")
        ws["D31"] = "TP. SẢN XUẤT"
        ws.merge_cells("D31:F31")
        ws["G31"] = "TP.MUA HÀNG"
        ws.merge_cells("G31:H31")
        ws["D36"] = "Ghi chú kỹ thuật"
        ws.merge_cells("D36:F36")
        buf = BytesIO()
        wb.save(buf)
        wb.close()

        # Patch 6 tuần rồi về 5 tuần
        out6, _ = ki.patch_khsx_ki_workbook(buf.getvalue(), plan_year=2026, plan_month=11)
        wb6 = load_workbook(BytesIO(out6))
        ranges6 = [str(r) for r in wb6["KHSX_ki"].merged_cells.ranges]
        for expected in ("B31:C31", "D31:F31", "G31:H31", "D36:F36"):
            self.assertIn(expected, ranges6)
        wb6.close()

        out5, _ = ki.patch_khsx_ki_workbook(out6, plan_year=2026, plan_month=9)
        wb5 = load_workbook(BytesIO(out5))
        ranges5 = [str(r) for r in wb5["KHSX_ki"].merged_cells.ranges]
        for expected in ("B31:C31", "D31:F31", "G31:H31", "D36:F36"):
            self.assertIn(expected, ranges5)
        wb5.close()

    def test_layout_detection_with_wide_k_and_external_merges(self):
        """[Defect 2] Nới rộng K lên 25 và có merge ngoài bảng không làm nhận diện nhầm cột tổng."""
        raw5 = make_mock_khsx_ki_workbook(year=2026, month=9)
        wb = load_workbook(BytesIO(raw5))
        ws = wb["KHSX_ki"]
        ws.column_dimensions["K"].width = 25.0
        ws["K35"] = "Ghi chú cột K rộng"
        ws.merge_cells("A1:L1")  # Merge tiêu đề kéo dài tới L
        ws.merge_cells("E38:L38")  # Merge ngoài bảng kéo dài tới L

        # Detector phải nhận diện đúng 10, không bị đánh lừa bởi K width hay external merge
        detected = ki._detect_current_layout(ws)
        self.assertEqual(detected, 10)

        buf = BytesIO()
        wb.save(buf)
        wb.close()

        # Patch sang 6 tuần
        out6, _ = ki.patch_khsx_ki_workbook(buf.getvalue(), plan_year=2026, plan_month=11)
        wb6 = load_workbook(BytesIO(out6))
        ws6 = wb6["KHSX_ki"]
        self.assertEqual(ws6["K6"].value, "Tổng cộng")
        self.assertTrue(ws6["K7"].font.bold, "Cột tổng mới K7 phải in đậm")
        self.assertFalse(ws6["J7"].font.bold, "Cột tuần J7 không được in đậm")
        self.assertEqual(ws6.column_dimensions["K"].width, 17.0)
        self.assertEqual(ws6.column_dimensions["J"].width, 19.44140625)
        wb6.close()

    def test_conflicting_or_corrupt_headers_resolved_safely(self):
        """[Defect 2] Tiêu đề mâu thuẫn hoặc hỏng được phân giải chính xác qua đối soát SKU rows."""
        # Trường hợp 1: Cả J6 và K6 đều ghi "Tổng cộng", nhưng col 11 có =SUM formula
        raw6_bytes, _ = ki.patch_khsx_ki_workbook(make_mock_khsx_ki_workbook(year=2026, month=11), plan_year=2026, plan_month=11)
        wb6 = load_workbook(BytesIO(raw6_bytes))
        ws6 = wb6["KHSX_ki"]
        ws6["J6"].value = "Tổng cộng"  # Gây mâu thuẫn tiêu đề
        ws6["K6"].value = "Tổng cộng"
        # Detector vẫn phải nhận diện đúng 11 vì col 11 chứa công thức tổng SKU
        self.assertEqual(ki._detect_current_layout(ws6), 11)
        wb6.close()

        # Trường hợp 2: Cả J6 và K6 đều ghi "Tổng cộng", nhưng col 11 trống rỗng
        raw5 = make_mock_khsx_ki_workbook(year=2026, month=9)
        wb5 = load_workbook(BytesIO(raw5))
        ws5 = wb5["KHSX_ki"]
        ws5["J6"].value = "Tổng cộng"
        ws5["K6"].value = "Tổng cộng"
        # Detector phải nhận diện đúng 10 vì col 11 không có dữ liệu
        self.assertEqual(ki._detect_current_layout(ws5), 10)
        wb5.close()

        # Trường hợp 3: Tiêu đề dòng 6 bị xóa sạch (None) và không có công thức =SUM trong các dòng SKU
        # Thay vì đoán mò 10 như phiên bản cũ, hệ thống phải từ chối an toàn bằng ValueError
        # nêu rõ sheet và tọa độ J6:K6 để bảo vệ dữ liệu người dùng.
        raw5_blank = make_mock_khsx_ki_workbook(year=2026, month=9)
        wb_blank = load_workbook(BytesIO(raw5_blank))
        ws_blank = wb_blank["KHSX_ki"]
        ws_blank["J6"].value = None
        ws_blank["K6"].value = None
        with self.assertRaises(ValueError) as ctx_blank:
            ki._detect_current_layout(ws_blank)
        self.assertIn("J6:K6", str(ctx_blank.exception))
        wb_blank.close()

        # Trường hợp 4: Tiêu đề dòng 6 bị xóa sạch (None), nhưng cột 10 có công thức =SUM
        wb_sum10 = load_workbook(BytesIO(raw5_blank))
        ws_sum10 = wb_sum10["KHSX_ki"]
        ws_sum10["J6"].value = None
        ws_sum10["K6"].value = None
        ws_sum10["J7"].value = "=SUM(E7:I7)"
        self.assertEqual(ki._detect_current_layout(ws_sum10), 10)
        wb_sum10.close()

        # Trường hợp 5: Tiêu đề dòng 6 bị xóa sạch (None), nhưng cột 11 có công thức =SUM
        wb_sum11 = load_workbook(BytesIO(raw5_blank))
        ws_sum11 = wb_sum11["KHSX_ki"]
        ws_sum11["J6"].value = None
        ws_sum11["K6"].value = None
        ws_sum11["K7"].value = "=SUM(E7:J7)"
        self.assertEqual(ki._detect_current_layout(ws_sum11), 11)
        wb_sum11.close()

    def test_conflicting_merge_inside_table_raises_clear_error(self):
        """[An toàn dữ liệu] Vùng merge xung đột bên trong bảng SKU báo lỗi rõ sheet và tọa độ."""
        raw5 = make_mock_khsx_ki_workbook(year=2026, month=9)
        wb = load_workbook(BytesIO(raw5))
        ws = wb["KHSX_ki"]
        # Đặt một ô gộp bất thường bên trong dòng SKU (dòng 7, cột J..K)
        ws.merge_cells("J7:K7")
        buf = BytesIO()
        wb.save(buf)
        wb.close()

        with self.assertRaises(ValueError) as ctx:
            ki.patch_khsx_ki_workbook(buf.getvalue(), plan_year=2026, plan_month=11)
        err_msg = str(ctx.exception)
        self.assertIn("KHSX_ki", err_msg)
        self.assertIn("J7:K7", err_msg)

    def test_defect_p2_footer_note_does_not_affect_layout_or_styles(self):
        """[Defect P2 Regression] Ghi chú K35 ngoài bảng không được làm sai lệch bố cục hay suy giảm định dạng."""
        results = []
        for note in (None, "Ghi chú cần giữ"):
            wb = load_workbook(BytesIO(make_mock_khsx_ki_workbook(year=2026, month=9)))
            ws = wb["KHSX_ki"]
            ws["J6"] = "Tổng cộng"
            ws["K6"] = "Tổng cộng"
            if note is not None:
                ws["K35"] = note
            detected = ki._detect_current_layout(ws)
            buf = BytesIO()
            wb.save(buf)
            wb.close()

            out, _ = ki.patch_khsx_ki_workbook(buf.getvalue(), plan_year=2026, plan_month=11)
            ver = ki.verify_khsx_ki(out, plan_year=2026, plan_month=11)
            self.assertTrue(ver["ok"], f"Verification phải pass cho note={note!r}")

            wb_out = load_workbook(BytesIO(out))
            ws_out = wb_out["KHSX_ki"]
            results.append({
                "note": note,
                "detected": detected,
                "k7_bold": ws_out["K7"].font.bold,
                "k7_nf": ws_out["K7"].number_format,
                "k_width": ws_out.column_dimensions["K"].width,
                "k35_val": ws_out["K35"].value if note is not None else None,
            })
            wb_out.close()

        # Cả hai trường hợp phải cho kết quả nhận diện và định dạng đồng nhất 100%
        self.assertEqual(results[0]["detected"], 10, "Khi không có note, detected phải là 10")
        self.assertEqual(results[1]["detected"], 10, "Khi có K35='Ghi chú cần giữ', detected vẫn phải là 10 (không bị nhảy sang 11)")
        self.assertTrue(results[0]["k7_bold"])
        self.assertTrue(results[1]["k7_bold"], "K7 phải giữ font.bold=True")
        self.assertEqual(results[0]["k7_nf"], "#,##0")
        self.assertEqual(results[1]["k7_nf"], "#,##0", "K7 phải giữ number_format=#,##0")
        self.assertEqual(results[0]["k_width"], 17.0)
        self.assertEqual(results[1]["k_width"], 17.0, "Cột K phải có width=17.0")
        self.assertEqual(results[1]["k35_val"], "Ghi chú cần giữ", "Ghi chú K35 phải được bảo toàn nguyên vẹn")

    def test_footer_sum_formula_does_not_affect_layout_detection(self):
        """[Defect P2] Công thức =SUM đặt ngoài bảng (K35) không được đánh lừa detector."""
        wb = load_workbook(BytesIO(make_mock_khsx_ki_workbook(year=2026, month=9)))
        ws = wb["KHSX_ki"]
        ws["J6"] = "Tổng cộng"
        ws["K6"] = "Tổng cộng"
        ws["K35"] = "=SUM(K7:K34)"  # Công thức ngoài bảng
        self.assertEqual(ki._detect_current_layout(ws), 10)
        wb.close()

    def test_distinguish_sku_zero_qty_from_empty_cell(self):
        """[Defect P2] Phân biệt rõ dữ liệu SKU = 0 với ô trống."""
        # 1. Ô trống (None hoặc chuỗi rỗng) không phải dữ liệu SKU
        self.assertTrue(ki._is_empty_cell(None))
        self.assertTrue(ki._is_empty_cell(""))
        self.assertTrue(ki._is_empty_cell("   "))
        self.assertFalse(ki._has_sku_data(None))
        self.assertFalse(ki._has_sku_data(""))

        # 2. Giá trị 0, 0.0, "0" là dữ liệu SKU hợp lệ, không phải ô trống
        self.assertFalse(ki._is_empty_cell(0))
        self.assertFalse(ki._is_empty_cell(0.0))
        self.assertTrue(ki._has_sku_data(0))
        self.assertTrue(ki._has_sku_data(0.0))

        # 3. Khi tiêu đề mâu thuẫn (cả hai là Tổng cộng), SKU có sản lượng 0 ở cột 11 chứng tỏ cột 11 có dữ liệu
        wb = load_workbook(BytesIO(make_mock_khsx_ki_workbook(year=2026, month=9)))
        ws = wb["KHSX_ki"]
        ws["J6"] = "Tổng cộng"
        ws["K6"] = "Tổng cộng"
        ws["K7"] = 0  # Sản lượng 0 ở cột K dòng SKU 7
        # Cột K có sản lượng 0 -> nhận diện đúng 11
        self.assertEqual(ki._detect_current_layout(ws), 11)
        wb.close()

    def test_identical_table_contents_with_varied_external_notes_yield_identical_layout(self):
        """[Defect P2] Hai workbook có bảng giống hệt nhau nhưng ghi chú ngoài bảng khác nhau phải cho kết quả nhận diện giống hệt."""
        notes = [
            None,
            "Ghi chú người lập",
            123456,
            "=SUM(A1:B1)",
            "Ghi chú rất dài kèm các ký tự đặc biệt !@#$%^&*()",
        ]
        results = []
        for n in notes:
            wb = load_workbook(BytesIO(make_mock_khsx_ki_workbook(year=2026, month=9)))
            ws = wb["KHSX_ki"]
            ws["J6"] = "Tổng cộng"
            ws["K6"] = "Tổng cộng"
            if n is not None:
                ws["K35"] = n
                ws["L40"] = "Ghi chú cột L"
            results.append(ki._detect_current_layout(ws))
            wb.close()

        self.assertTrue(all(r == 10 for r in results), f"Tất cả kết quả phải bằng 10, thực tế: {results}")

    def test_large_sku_table_beyond_30_rows_evaluated_completely(self):
        """[Defect P2] Bảng có số lượng SKU lớn (>30 SKU) được quét toàn vẹn, không bị giới hạn cứng ở dòng 36."""
        wb = load_workbook(BytesIO(make_mock_khsx_ki_workbook(year=2026, month=9)))
        ws = wb["KHSX_ki"]
        ws["J6"] = "Tổng cộng"
        ws["K6"] = "Tổng cộng"

        # Thêm 35 dòng SKU (từ dòng 7 đến dòng 41) và dòng tổng tại dòng 42
        for i in range(1, 36):
            r = 6 + i
            ws.cell(r, 1, value=i)
            ws.cell(r, 2, value=130100000 + i)
            ws.cell(r, 3, value=f"SKU {i}")
            ws.cell(r, 10, value=None)
            ws.cell(r, 11, value=None)
        total_r = 42
        ws.cell(total_r, 1, value="Tổng cộng")
        ws.cell(total_r, 2, value=None)

        # Đặt công thức =SUM tại dòng SKU 40 cột 11 (vượt quá giới hạn cũ row 36)
        ws.cell(40, 11, value="=SUM(E40:J40)")

        # Detector mới phải phát hiện =SUM tại dòng 40 và trả về 11
        self.assertEqual(ki._detect_current_layout(ws), 11)

        # Xóa =SUM tại dòng 40 và đặt =SUM tại dòng 45 (ngoài bảng): detector phải trả về 10
        ws.cell(40, 11).value = None
        ws.cell(45, 11).value = "=SUM(E45:J45)"
        self.assertEqual(ki._detect_current_layout(ws), 10)
        wb.close()

    def test_corrupted_headers_or_both_sum_formulas_raises_value_error(self):
        """[Defect P2] Khi tiêu đề bị hỏng/thiếu không có =SUM hoặc cả hai cột đều có =SUM, phải báo lỗi rõ ràng."""
        # 1. Cả hai cột đều có công thức =SUM mâu thuẫn
        wb = load_workbook(BytesIO(make_mock_khsx_ki_workbook(year=2026, month=9)))
        ws = wb["KHSX_ki"]
        ws["J6"] = "Tổng cộng"
        ws["K6"] = "Tổng cộng"
        ws["J7"] = "=SUM(E7:I7)"
        ws["K7"] = "=SUM(E7:J7)"
        with self.assertRaises(ValueError) as ctx1:
            ki._detect_current_layout(ws)
        err_msg1 = str(ctx1.exception)
        self.assertIn("KHSX_ki", err_msg1)
        self.assertIn("J6:K6", err_msg1)
        self.assertIn("=SUM", err_msg1)

        # 2. Tiêu đề không có 'Tổng cộng' và không có =SUM (tiêu đề bị ghi đè / hỏng)
        ws["J6"] = "Ghi chú cột J"
        ws["K6"] = "Ghi chú cột K"
        ws["J7"] = 1000
        ws["K7"] = 2000
        with self.assertRaises(ValueError) as ctx2:
            ki._detect_current_layout(ws)
        err_msg2 = str(ctx2.exception)
        self.assertIn("KHSX_ki", err_msg2)
        self.assertIn("J6:K6", err_msg2)
        wb.close()


if __name__ == "__main__":
    unittest.main()
