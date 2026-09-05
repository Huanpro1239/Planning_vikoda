import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from sync_planning_fc import prepare_planning_fc_update


class SyncPlanningFcTests(unittest.TestCase):
    def make_workbook(self, selector="Tháng 9", include_missing=False):
        workbook = Workbook()

        fc = workbook.active
        fc.title = "FC"
        fc["A1"] = "STT"
        fc["B1"] = "Mã SP"
        fc["C1"] = "Tên SP"
        fc["D1"] = "Đơn vị tính"

        for month in range(1, 13):
            fc.cell(row=1, column=4 + month).value = f"Tháng {month}"

        fc["R1"] = selector

        fc["B2"] = 130100096
        fc["M2"] = 9370
        fc["B3"] = 130100091
        fc["M3"] = 1421

        planning = workbook.create_sheet("Ke_hoach_SX")
        planning["A1"] = "Mã Sản Phẩm"
        planning["L1"] = "FC"

        # Đảo thứ tự để kiểm tra ghép theo mã, không theo số dòng.
        planning["A2"] = 130100091
        planning["L2"] = 0
        planning["A3"] = 130100096
        planning["L3"] = "=1+1"

        if include_missing:
            planning["A4"] = 130199999

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_matches_r1_header_and_updates_l_by_product_code(self):
        updated, info = prepare_planning_fc_update(
            self.make_workbook()
        )

        self.assertEqual(info["selector"], "Tháng 9")
        self.assertEqual(info["source_column_letter"], "M")
        self.assertEqual(info["changed_count"], 2)
        self.assertEqual(info["patched_count"], 2)

        workbook = load_workbook(BytesIO(updated), data_only=False)
        planning = workbook["Ke_hoach_SX"]

        self.assertEqual(planning["L2"].value, 1421)
        self.assertEqual(planning["L3"].value, 9370)
        self.assertNotEqual(planning["L3"].data_type, "f")

    def test_second_run_is_idempotent(self):
        updated, _ = prepare_planning_fc_update(
            self.make_workbook()
        )
        updated_again, info = prepare_planning_fc_update(updated)

        self.assertEqual(info["changed_count"], 0)
        self.assertEqual(updated_again, updated)

    def test_selector_must_match_e_to_p_header(self):
        with self.assertRaisesRegex(
            RuntimeError,
            "không khớp tiêu đề nào",
        ):
            prepare_planning_fc_update(
                self.make_workbook(selector="Tháng 13")
            )

    def test_missing_planning_code_raises_clear_error(self):
        with self.assertRaisesRegex(
            RuntimeError,
            "không tồn tại trong FC!B",
        ):
            prepare_planning_fc_update(
                self.make_workbook(include_missing=True)
            )


if __name__ == "__main__":
    unittest.main()
