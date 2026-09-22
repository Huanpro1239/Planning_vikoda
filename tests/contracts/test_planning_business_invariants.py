"""Business-invariant contracts for Vikoda production planning.

These tests intentionally exercise public planning behavior across engine,
weekly-model workbook adapter, verification and publish policy. They describe
business contracts, not implementation details.
"""

import copy
import unittest
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from planning.publish.policy import publish_decision
from planning.publish.proposal import with_proposal_identity
from planning.weekly_model import (
    prepare_weekly_schedule_update,
    verify_weekly_workbook,
)


PLAN_YEAR = 2026
PLAN_MONTH = 9
MONTH_DAYS = 30


def build_workbook(skus):
    workbook = Workbook()
    planning = workbook.active
    planning.title = "Ke_hoach_SX"
    planning.append([
        "Mã Sản Phẩm",
        "Tên Sản Phẩm",
        "Đơn vị tính",
        "Số lượng /mẻ",
        "Số lượng/ ca",
        "Chuyền",
        "Nhóm sản phẩm",
        "Phân loại SP",
        "Số ca theo ngày",
        "Tồn đầu thực tế",
        "Tồn đầu sổ sách",
        "FC",
        "Tồn cuối dự kiến",
        "Nợ kho",
        "Số lượng cần sản xuất",
        "Số lượng sản xuất theo mẻ/ca",
        "Số ngày cần sản xuất",
        "Ngày bắt đầu sản xuất",
    ] + [f"{day:02d}/09" for day in range(1, MONTH_DAYS + 1)])

    master = workbook.create_sheet("Danh_muc")
    master.append([
        "Mã Sản Phẩm",
        "Tên Sản Phẩm",
        "Đơn vị tính",
        "Số lượng /mẻ",
        "Số lượng/ ca",
        "Chuyền",
        "Nhóm sản phẩm",
        "Phân loại SP",
        "Quy cách",
        "Leadtime",
        "Debt mode",
        "Schedule profile",
    ])

    for index, sku in enumerate(skus, start=2):
        code = sku["code"]
        name = sku.get("name", f"SKU {code}")
        uom = sku.get("uom", "Thùng")
        batch = sku.get("batch", 100.0)
        per_shift = sku.get("per_shift", 100.0)
        line = sku.get("line", "KHS")
        group = sku.get("group", line)
        classification = sku.get("classification", "Không đường")
        shifts = sku.get("shifts_per_day", 1.0)
        actual = sku.get("actual_stock", 0.0)
        book = sku.get("book_stock", 0.0)
        fc = sku.get("fc", 0.0)
        target = sku.get("target_stock", 0.0)
        debt = sku.get("debt", 0.0)
        mold = sku.get("mold", 1.0)
        leadtime = sku.get("leadtime", 0.0)
        debt_mode = sku.get("debt_mode", "SUBTRACT_BOOK_ON_DEBT")
        profile = sku.get("profile", "")

        planning.append([
            code,
            name,
            uom,
            batch,
            per_shift,
            line,
            group,
            classification,
            shifts,
            actual,
            book,
            fc,
            target,
            debt,
            0,
            0,
            0,
            None,
        ] + [None] * MONTH_DAYS)

        master.append([
            code,
            name,
            uom,
            batch,
            per_shift,
            line,
            group,
            classification,
            mold,
            leadtime,
            debt_mode,
            profile,
        ])

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def run_contract(skus, *, revision="contract"):
    raw = build_workbook(skus)
    final_bytes, report, analysis = prepare_weekly_schedule_update(
        raw,
        plan_year=PLAN_YEAR,
        plan_month=PLAN_MONTH,
        input_revision={"target": {"etag": revision}},
    )
    verification = verify_weekly_workbook(
        final_bytes,
        schedule_report=report,
        plan_year=PLAN_YEAR,
        plan_month=PLAN_MONTH,
    )
    if not verification["validated"]:
        raise AssertionError("Weekly contract workbook did not verify")
    return final_bytes, report, analysis


def scheduled_qty(analysis, code):
    return sum(
        float(item.qty)
        for item in analysis.daily_plan
        if int(item.ma_sp) == int(code)
    )


def schedule_dates(analysis, code):
    return [
        item.date
        for item in analysis.daily_plan
        if int(item.ma_sp) == int(code)
    ]


class PlanningBusinessInvariantContracts(unittest.TestCase):
    def test_contract_service_first_and_safety_stock_use_only_residual_capacity(self):
        """CONTRACT-01/03: service is complete before any safety-stock buffer."""
        final_bytes, report, analysis = run_contract([
            {
                "code": 1101,
                "line": "KHS",
                "fc": 1000,
                "target_stock": 5000,
                "mold": 1,
            },
            {
                "code": 1102,
                "line": "PET 9000",
                "fc": 1000,
                "target_stock": 0,
                "mold": 1,
            },
        ])

        self.assertEqual(report["publish_status"], "ready_for_publish")
        self.assertTrue(report["status"]["service"]["ok"])
        self.assertTrue(report["status"]["monthly_quantity"]["ok"])
        self.assertFalse(report["status"]["safety_stock"]["ok"])
        self.assertEqual(
            report["status"]["monthly_quantity"]["state"],
            "service_complete_buffer_shortfall",
        )

        first = report["mass_balance"]["1101"]
        second = report["mass_balance"]["1102"]
        self.assertEqual(first["service_target_qty"], 1000.0)
        self.assertEqual(first["service_scheduled_qty"], 1000.0)
        self.assertEqual(first["service_carryover_qty"], 0.0)
        self.assertEqual(second["service_target_qty"], 1000.0)
        self.assertEqual(second["service_scheduled_qty"], 1000.0)
        self.assertEqual(second["service_carryover_qty"], 0.0)

        # September has 30 shifts at one shift/day. After 20 service shifts,
        # only 10 shifts remain for safety stock.
        self.assertEqual(scheduled_qty(analysis, 1101), 2000.0)
        self.assertEqual(scheduled_qty(analysis, 1102), 1000.0)
        self.assertEqual(first["buffer_scheduled_qty"], 1000.0)
        self.assertEqual(first["buffer_carryover_qty"], 4000.0)

        workbook = load_workbook(
            BytesIO(final_bytes),
            data_only=True,
            read_only=True,
        )
        try:
            sheet = workbook["Ke_hoach_SX"]
            self.assertEqual(sheet["O2"].value, 6000)
            self.assertEqual(sheet["P2"].value, 2000)
            self.assertEqual(sheet["P3"].value, 1000)
            self.assertEqual(
                sum(
                    float(sheet.cell(2, 19 + offset).value or 0)
                    for offset in range(MONTH_DAYS)
                ),
                2000.0,
            )
        finally:
            workbook.close()

    def test_contract_debt_has_priority_and_zero_fc_debt_starts_day_one(self):
        """CONTRACT-02: debt is mandatory service and wins equal-time priority."""
        _, report, analysis = run_contract([
            {
                "code": 1201,
                "line": "KHS",
                "fc": 0,
                "debt": 1000,
                "target_stock": 5000,
                "mold": 1,
            },
            {
                "code": 1202,
                "line": "PET 9000",
                "fc": 1000,
                "debt": 0,
                "target_stock": 0,
                "mold": 1,
            },
        ])

        debt_calc = next(
            row for row in analysis.calculated
            if row.input.ma_sp == 1201
        )
        self.assertEqual(
            debt_calc.start_datetime,
            datetime(PLAN_YEAR, PLAN_MONTH, 1),
        )
        # Debt branch excludes target safety stock: only debt service is planned.
        self.assertEqual(debt_calc.p_need, 1000.0)
        self.assertEqual(debt_calc.service_qty, 1000.0)
        self.assertEqual(debt_calc.buffer_qty, 0.0)

        debt_dates = schedule_dates(analysis, 1201)
        normal_dates = schedule_dates(analysis, 1202)
        self.assertEqual(min(debt_dates), date(PLAN_YEAR, PLAN_MONTH, 1))
        self.assertLess(max(debt_dates), min(normal_dates))
        self.assertEqual(report["mass_balance"]["1201"]["scheduled_qty"], 1000.0)
        self.assertEqual(report["mass_balance"]["1201"]["carryover_qty"], 0.0)

    def test_contract_debt_formula_modes_remain_distinct(self):
        """CONTRACT-02B: debt formula mode controls whether book stock is deducted."""
        _, _, subtract_analysis = run_contract([
            {
                "code": 1211,
                "line": "KHS",
                "fc": 0,
                "debt": 500,
                "book_stock": 200,
                "debt_mode": "SUBTRACT_BOOK_ON_DEBT",
            },
        ])
        _, _, ignore_analysis = run_contract([
            {
                "code": 1212,
                "line": "KHS",
                "fc": 0,
                "debt": 500,
                "book_stock": 200,
                "debt_mode": "IGNORE_BOOK_ON_DEBT",
            },
        ])

        self.assertEqual(subtract_analysis.calculated[0].p_need, 300.0)
        self.assertEqual(subtract_analysis.calculated[0].q_rounded, 300.0)
        self.assertEqual(ignore_analysis.calculated[0].p_need, 500.0)
        self.assertEqual(ignore_analysis.calculated[0].q_rounded, 500.0)

    def test_contract_shared_machine_is_one_serial_capacity_pool(self):
        """CONTRACT-04: KHS and PET9000 cannot consume the same one-shift slot."""
        _, report, analysis = run_contract([
            {
                "code": 1301,
                "line": "KHS",
                "fc": 100,
                "mold": 1,
                "shifts_per_day": 1,
            },
            {
                "code": 1302,
                "line": "PET 9000",
                "fc": 100,
                "mold": 1,
                "shifts_per_day": 1,
            },
        ])

        first_dates = schedule_dates(analysis, 1301)
        second_dates = schedule_dates(analysis, 1302)
        self.assertEqual(first_dates, [date(PLAN_YEAR, PLAN_MONTH, 1)])
        self.assertEqual(second_dates, [date(PLAN_YEAR, PLAN_MONTH, 2)])

        per_shift = {
            calc.input.ma_sp: calc.input.sl_ca
            for calc in analysis.calculated
        }
        usage = defaultdict(float)
        active_codes = defaultdict(set)
        for item in analysis.daily_plan:
            if item.chuyen in {"KHS", "PET 9000"}:
                usage[item.date] += (
                    float(item.qty)
                    / per_shift[item.ma_sp]
                )
                active_codes[item.date].add(item.ma_sp)

        for current, used in usage.items():
            self.assertLessEqual(used, 1.0 + 1e-6)
            self.assertLessEqual(
                len(active_codes[current]),
                1,
                f"Shared machine ran parallel SKUs on {current}",
            )

        resource = report["status"]["resource_validation"]
        self.assertTrue(resource["ok"])
        self.assertEqual(resource["state"], "shared_machine_passed")

    def test_contract_shared_machine_changeover_consumes_capacity(self):
        """CONTRACT-04B: mold changeover reduces same-day productive capacity."""
        _, _, analysis = run_contract([
            {
                "code": 1311,
                "line": "KHS",
                "fc": 100,
                "mold": 1,
                "shifts_per_day": 2,
            },
            {
                "code": 1312,
                "line": "PET 9000",
                "fc": 100,
                "mold": 2,
                "shifts_per_day": 2,
            },
        ])

        pet_day_1 = sum(
            float(item.qty)
            for item in analysis.daily_plan
            if item.ma_sp == 1312
            and item.date == date(PLAN_YEAR, PLAN_MONTH, 1)
        )
        pet_day_2 = sum(
            float(item.qty)
            for item in analysis.daily_plan
            if item.ma_sp == 1312
            and item.date == date(PLAN_YEAR, PLAN_MONTH, 2)
        )
        self.assertEqual(pet_day_1, 50.0)
        self.assertEqual(pet_day_2, 50.0)

    def test_contract_publish_boundary_ready_and_stockout_review(self):
        """CONTRACT-05: ready auto-authorizes; stockout needs exact approval."""
        final_ready, ready_report, _ = run_contract([
            {
                "code": 1401,
                "line": "KHS",
                "fc": 500,
                "target_stock": 0,
            },
        ])
        ready_report = with_proposal_identity(
            ready_report,
            final_ready,
        )
        authorized, decision = publish_decision(
            ready_report,
            None,
            "contract-runner",
        )
        self.assertTrue(authorized)
        self.assertEqual(decision["basis"], "publish_status")

        final_short, short_report, _ = run_contract([
            {
                "code": 1411,
                "line": "KHS",
                "fc": 2000,
                "target_stock": 0,
                "shifts_per_day": 1,
            },
            {
                "code": 1412,
                "line": "PET 9000",
                "fc": 2000,
                "target_stock": 0,
                "shifts_per_day": 1,
            },
        ])
        short_report = with_proposal_identity(
            short_report,
            final_short,
        )
        self.assertEqual(short_report["publish_status"], "review_required")
        self.assertFalse(short_report["status"]["service"]["ok"])
        self.assertEqual(
            short_report["status"]["service"]["state"],
            "stockout_risk",
        )

        authorized, decision = publish_decision(
            short_report,
            None,
            "contract-runner",
        )
        self.assertFalse(authorized)
        self.assertEqual(decision["state"], "blocked")

        stale_approval = {
            "proposal_id": "stale-proposal",
            "reason": "Reviewed capacity shortage.",
            "approved_by": "planner",
        }
        authorized, _ = publish_decision(
            short_report,
            stale_approval,
            "contract-runner",
        )
        self.assertFalse(authorized)

        exact_approval = {
            "proposal_id": short_report["proposal_id"],
            "reason": "Reviewed unavoidable service shortage under fixed capacity.",
            "approved_by": "planner",
        }
        authorized, decision = publish_decision(
            short_report,
            exact_approval,
            "contract-runner",
        )
        self.assertTrue(authorized)
        self.assertEqual(decision["basis"], "review_approval")

    def test_contract_resource_failure_is_never_waivable(self):
        """CONTRACT-05B: review approval cannot bypass resource validation."""
        final_bytes, report, _ = run_contract([
            {
                "code": 1501,
                "line": "KHS",
                "fc": 2000,
                "shifts_per_day": 1,
            },
            {
                "code": 1502,
                "line": "PET 9000",
                "fc": 2000,
                "shifts_per_day": 1,
            },
        ])
        report = with_proposal_identity(
            report,
            final_bytes,
        )
        blocked = copy.deepcopy(report)
        blocked["status"]["resource_validation"]["ok"] = False
        blocked["status"]["resource_validation"]["state"] = (
            "shared_machine_over_capacity"
        )
        blocked["publish_status"] = "review_required"

        approval = {
            "proposal_id": blocked["proposal_id"],
            "reason": "Attempted resource override.",
            "approved_by": "planner",
        }
        authorized, decision = publish_decision(
            blocked,
            approval,
            "contract-runner",
        )
        self.assertFalse(authorized)
        self.assertEqual(
            decision["basis"],
            "non_waivable_validation",
        )


if __name__ == "__main__":
    unittest.main()
