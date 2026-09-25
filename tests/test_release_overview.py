"""Contracts for the unified Planning + NVL release overview."""

from __future__ import annotations

from io import BytesIO
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook

from nvl.release import (
    build_release_manifest as build_nvl_manifest,
    write_release_manifest as write_nvl_manifest,
)
from ops.release_overview import build_release_overview
from planning.publish.proposal import with_proposal_identity
from planning.publish.release import (
    build_release_manifest as build_planning_manifest,
    write_release_manifest as write_planning_manifest,
)


ROOT = Path(__file__).resolve().parents[1]


def workbook_bytes(value: float) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ton_NVL"
    ws["A1"] = "Mã NVL"
    ws["D1"] = "Tồn cuối"
    ws["E1"] = "Tồn đơn hàng"
    ws["A2"] = "VT001"
    ws["D2"] = value
    ws["E2"] = value
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


class ReleaseOverviewTests(unittest.TestCase):
    def _readiness(self, root: Path, name: str, gate: str):
        path = root / name
        path.write_text(
            json.dumps(
                {
                    "status": "passed",
                    "gate_version": gate,
                    "git_sha": "deadbeef",
                },
                sort_keys=True,
            ),
            encoding="utf-8",
        )
        return path

    def _planning_release(
        self,
        root: Path,
        *,
        timestamp: str,
        run_id: str,
    ) -> Path:
        releases = root / "releases"
        releases.mkdir(parents=True, exist_ok=True)
        readiness = self._readiness(
            root,
            "production_readiness_report.json",
            "production_readiness_v1",
        )
        proposal = workbook_bytes(10)
        report = with_proposal_identity(
            {
                "plan_month": "2026-09",
                "algorithm": "weekly",
                "input_revision": {"source": f"rev-{run_id}"},
                "pipeline": {"engine_version": "engine-v1"},
            },
            proposal,
        )
        decision = {
            "state": "published",
            "proposal_id": report["proposal_id"],
            "reason": "test",
        }
        manifest = build_planning_manifest(
            report,
            decision,
            proposal,
            published_at=timestamp,
            environ={
                "PLANNING_REQUIRE_READINESS": "1",
                "GITHUB_SHA": "deadbeef",
                "GITHUB_RUN_ID": run_id,
                "GITHUB_RUN_ATTEMPT": "1",
            },
            readiness_path=readiness,
        )
        path = releases / f"{manifest['release_id']}.json"
        write_planning_manifest(manifest, path=path)
        shutil.copyfile(path, root / "latest_release.json")
        return path

    def _nvl_release(
        self,
        root: Path,
        *,
        timestamp: str,
        run_id: str,
        previous: Path | None = None,
    ) -> Path:
        nvl_root = root / "nvl"
        releases = nvl_root / "releases"
        releases.mkdir(parents=True, exist_ok=True)
        readiness = self._readiness(
            root,
            "nvl_production_readiness_report.json",
            "nvl_production_readiness_v1",
        )
        stock_bytes = workbook_bytes(float(run_id))
        final_bytes = workbook_bytes(float(run_id) + 1)
        final_sha = hashlib.sha256(final_bytes).hexdigest()

        stock = {
            "mode": "publish",
            "status": "published",
            "source": {"revision": f'"stock-source-{run_id}"'},
            "target": {"revision": f'"target-before-stock-{run_id}"'},
            "metrics": {"changed_count": 1},
            "post_upload_verified": True,
            "upload_result": {
                "eTag": f'"target-after-stock-{run_id}"',
            },
        }
        open_po = {
            "publish_requested": True,
            "published": True,
            "source": {"revision": f'"open-source-{run_id}"'},
            "target": {
                "revision_before": f'"target-before-open-{run_id}"',
                "revision_after": f'"target-after-open-{run_id}"',
                "changed_cells": 1,
            },
            "publish_evidence": {
                "expected_etag": f'"target-before-open-{run_id}"',
                "upload_skipped": False,
                "upload_etag": f'"target-after-open-{run_id}"',
                "target_revision_after": f'"target-after-open-{run_id}"',
                "post_upload_verified": True,
                "server_sha256": final_sha,
            },
        }

        previous_manifest = None
        previous_sha = None
        if previous:
            previous_raw = previous.read_bytes()
            previous_manifest = json.loads(previous_raw.decode("utf-8"))
            previous_sha = hashlib.sha256(previous_raw).hexdigest()

        manifest = build_nvl_manifest(
            stock,
            open_po,
            stock_bytes,
            final_bytes,
            published_at=timestamp,
            environ={
                "NVL_REQUIRE_READINESS": "1",
                "GITHUB_SHA": "deadbeef",
                "GITHUB_RUN_ID": run_id,
                "GITHUB_RUN_ATTEMPT": "1",
            },
            readiness_path=readiness,
            previous_manifest=previous_manifest,
            previous_manifest_sha256=previous_sha,
        )
        path = releases / f"{manifest['release_id']}.json"
        write_nvl_manifest(manifest, path=path)
        shutil.copyfile(path, nvl_root / "latest_release.json")
        return path

    def test_combines_planning_and_nvl_into_one_timeline(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            planning = self._planning_release(
                root,
                timestamp="2026-09-25T01:00:00+00:00",
                run_id="100",
            )
            nvl = self._nvl_release(
                root,
                timestamp="2026-09-25T02:00:00+00:00",
                run_id="101",
            )
            overview = build_release_overview(
                root,
                verify_commit=False,
            )

        self.assertTrue(planning.name.endswith(".json"))
        self.assertTrue(nvl.name.endswith(".json"))
        self.assertEqual(overview["status"], "passed")
        self.assertEqual(
            overview["summary"]["planning_release_count"],
            1,
        )
        self.assertEqual(
            overview["summary"]["nvl_release_count"],
            1,
        )
        self.assertEqual(
            [row["domain"] for row in overview["timeline"]],
            ["planning", "nvl"],
        )
        self.assertEqual(
            overview["latest"]["planning"]["domain"],
            "planning",
        )
        self.assertEqual(
            overview["latest"]["nvl"]["domain"],
            "nvl",
        )

    def test_stale_planning_latest_pointer_fails_overview(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            first = self._planning_release(
                root,
                timestamp="2026-09-25T01:00:00+00:00",
                run_id="100",
            )
            self._planning_release(
                root,
                timestamp="2026-09-25T02:00:00+00:00",
                run_id="101",
            )
            shutil.copyfile(first, root / "latest_release.json")

            overview = build_release_overview(
                root,
                verify_commit=False,
            )

        self.assertEqual(overview["status"], "failed")
        self.assertEqual(
            overview["planning"]["latest"]["status"],
            "failed",
        )

    def test_broken_nvl_chain_propagates_failure(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            first = self._nvl_release(
                root,
                timestamp="2026-09-25T01:00:00+00:00",
                run_id="100",
            )
            self._nvl_release(
                root,
                timestamp="2026-09-25T02:00:00+00:00",
                run_id="101",
                previous=first,
            )
            payload = json.loads(first.read_text(encoding="utf-8"))
            payload["commit"]["actor"] = "tampered"
            first.write_text(
                json.dumps(
                    payload,
                    ensure_ascii=False,
                    indent=2,
                    sort_keys=True,
                )
                + "\n",
                encoding="utf-8",
            )

            overview = build_release_overview(
                root,
                verify_commit=False,
            )

        self.assertEqual(overview["status"], "failed")
        self.assertEqual(overview["nvl"]["status"], "failed")

    def test_empty_runtime_state_is_supported(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            overview = build_release_overview(
                tmpdir,
                verify_commit=False,
            )

        self.assertEqual(overview["status"], "empty")
        self.assertEqual(
            overview["summary"]["total_release_count"],
            0,
        )

    def test_cli_writes_json_csv_and_markdown(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            self._planning_release(
                root,
                timestamp="2026-09-25T01:00:00+00:00",
                run_id="100",
            )
            self._nvl_release(
                root,
                timestamp="2026-09-25T02:00:00+00:00",
                run_id="101",
            )
            out_json = root / "overview.json"
            out_csv = root / "overview.csv"
            out_md = root / "overview.md"

            result = subprocess.run(
                [
                    sys.executable,
                    "-X",
                    "utf8",
                    str(ROOT / "scripts" / "release_overview.py"),
                    str(root),
                    "--no-git",
                    "--out-json",
                    str(out_json),
                    "--out-csv",
                    str(out_csv),
                    "--out-md",
                    str(out_md),
                ],
                cwd=ROOT,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                encoding="utf-8",
                errors="replace",
                check=False,
            )

            self.assertEqual(
                result.returncode,
                0,
                msg=result.stderr or result.stdout,
            )
            self.assertTrue(out_json.is_file())
            self.assertTrue(out_csv.is_file())
            self.assertTrue(out_md.is_file())
            self.assertIn(
                "Planning + NVL Release Overview",
                out_md.read_text(encoding="utf-8"),
            )


if __name__ == "__main__":
    unittest.main()
