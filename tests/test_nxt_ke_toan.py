import io
import unittest

from openpyxl import Workbook, load_workbook

from nxt_ke_toan import build_workbook, doc_luong_ky_gui, ton_ky_gui, tong_hop_nhieu_file
from sell_in import doc_danh_muc
from tests.test_xuat_kho import _xuat_kho


def _nxt(ten_dn, ky, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append([ten_dn])
    for _ in range(5):
        ws.append([None])
    ws.append([ky])
    ws.append([None])
    ws.append(["Stt", "Mã vật tư", "Tên vật tư", "Loại sản phẩm", "Tài khoản kế toán ", "Đvt",
               "Tồn đầu kỳ", None, "Nhập trong kỳ", None, "Xuất trong kỳ", None, "Tồn cuối kỳ", None])
    ws.append([None] * 6 + ["Số lượng", "Thành tiền"] * 4)
    for i, (ma, loai, dau, nhap, xuat) in enumerate(rows, start=1):
        ws.append([str(i), ma, "SP " + ma, loai, "155", "CI", dau, 0, nhap, 0, xuat, 0, dau + nhap - xuat, 0])
    wb.create_sheet("Sheet2")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _danh_muc():
    wb = Workbook()
    ws = wb.active
    ws.title = "DanhMucSanPham"
    ws.append(["MaSanPhamMoi_Vikoda", "MaSanPhamMoi_VKD", "TenSanPham", "Tên viết tắt", "Quy cách",
               "ĐVT", "Brand", "SKUs Group"])
    ws.append([130100110, 230100110, "Kiềm 500ml", "Khoáng kiềm 500ml", 24, "Thùng", "Vikoda", "Pet Vikoda"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


VIK = "Tên DN: CÔNG TY CỔ PHẦN NƯỚC KHOÁNG KHÁNH HÒA"
VKD = "Tên DN: CÔNG TY CỔ PHẦN ĐẦU TƯ VÀ THƯƠNG MẠI VKD"


class NxtKeToanTests(unittest.TestCase):
    def test_nxt_merge_and_filters(self):
        dm = doc_danh_muc(_danh_muc())
        ky = "Từ ngày 01-09-2026 đến ngày 30-09-2026"
        data, log = tong_hop_nhieu_file([
            ("Vikoda/Thang 9.xlsm", _nxt(VIK, ky, [("130100110", "Thành phẩm", 2400, 0, 1200),
                                                  ("230100110", "Thành phẩm", 999, 0, 0),   # sai đầu mã
                                                  ("310100001", "Nguyên liệu", 5, 5, 5)])),
            ("VKD/Thang 9.xlsm", _nxt(VKD, ky, [("230100110", "Hàng hóa", 240, 1200, 720)])),
        ], dm)
        self.assertEqual(len(data), 4)                       # Data giữ mọi vật tư
        qd = {(r[0], r[5]): r[23] for r in data}
        self.assertEqual(qd[("Vikoda", "130100110")], 50)    # 1200 / 24
        self.assertIsNone(qd[("Vikoda", "230100110")])       # mã 2 ở Vikoda không quy đổi
        self.assertIsNone(qd[("Vikoda", "310100001")])
        self.assertEqual(qd[("VKD", "230100110")], 30)
        self.assertEqual([r[4] for r in data[:1]], [30])     # số ngày của kỳ

        content, n_sp, n_thang, _ = build_workbook(data, log, "fp")
        self.assertEqual((n_sp, n_thang), (1, 1))
        wb = load_workbook(io.BytesIO(content))
        r = list(wb["Ton thong minh"].iter_rows(values_only=True))
        h = r[0]
        row = dict(zip(h, r[1]))
        self.assertEqual(row["TongTonKhaDung"], 80)          # 50 + 30 gộp 1 sản phẩm
        self.assertEqual(row["NhuCauNgayHeThong"], 1)        # xuất VKD 720/24 = 30 thùng / 30 ngày

    def test_ton_ky_gui(self):
        dm = doc_danh_muc(_danh_muc())
        ckg = "Chuyển qua kho ký gửi"
        xk = _xuat_kho("CÔNG TY CỔ PHẦN ĐẦU TƯ VÀ THƯƠNG MẠI VKD", [
            (ckg, "B", "Đơn hàng bán", "230100110", 480),     # chuyển vào ký gửi 20 thùng
            (None, "B", "Ký gửi", "230100110", 720),          # xuất từ ký gửi 30 thùng
            (None, "B", "Đơn hàng bán", "230100110", 24),     # bán thường, không liên quan
        ])
        # _xuat_kho không đặt KHOXUAT -> đặt cho dòng "Ký gửi"
        wb = load_workbook(xk)
        ws = wb.active
        for row in ws.iter_rows(min_row=6):
            if row[11].value == "Ký gửi":
                row[13].value = "KHO KÝ GỬI"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        flow = doc_luong_ky_gui(buf, "VKD/Xuat kho thang 9.xlsm")
        self.assertEqual(flow[:3], ("VKD", 2026, 9))
        self.assertEqual(flow[3], {"230100110": [480, 720]})
        kg, nguon = ton_ky_gui([flow])
        self.assertEqual(kg[("VKD", (2026, 9), "230100110")], [240, 480, 720, 0])  # đầu kỳ tối thiểu 240
        self.assertEqual(nguon, "Ước tính tối thiểu")

        ky = "Từ ngày 01-09-2026 đến ngày 30-09-2026"
        data, log = tong_hop_nhieu_file([("VKD/Thang 9.xlsm", _nxt(VKD, ky, [("230100110", "Hàng hóa", 0, 2400, 960)]))],
                                        dm)
        content, *_ = build_workbook(data, log, "fp", kg, nguon)
        r = list(load_workbook(io.BytesIO(content))["Ton thong minh"].iter_rows(values_only=True))
        row = dict(zip(r[0], r[1]))
        self.assertEqual(row["TongTonKhaDung"], 60)
        self.assertEqual(row["TonKyGuiChoGiao"], 0)
        self.assertEqual(row["TongTonVatLy"], 60)
        # nhu cầu = xuất kế toán (bán ra) 960 chai = 40 thùng / 30 ngày
        self.assertAlmostEqual(row["NhuCauNgayHeThong"], 1.33, places=2)
        self.assertEqual(row["SoNgayTonKhaDung"], 45)

    def test_ton_nha_may(self):
        import datetime as dt
        wb = Workbook()
        ws = wb.active
        ws.title = "THÁNG 10"
        ws.append(["BÁO CÁO KHO THÀNH PHẨM"])
        ws.append([None] * 7 + ["Ngày:", dt.datetime(2026, 10, 28)])
        ws.append(["STT", "BRAND", "MÃ SP VIKODA", "MÃ SP VKD", "TÊN SẢN PHẨM", "ĐVT", "SỐ LƯỢNG SẢN XUẤT", None,
                   "SỐ LƯỢNG XUẤT KHO", None, None, None, "TỒN KHO THỰC TẾ", None, "SL CHƯA NHẬP KHO (CHỜ KQ QA/QC)",
                   "KẾ HOẠCH XUẤT KHO 29/10/2026", "SL HÀNG ĐÃ BÁN (GỬI KHO)", "TỒN KHO DỰ KIẾN CUỐI NGÀY 29/10/2026"])
        ws.append([None] * 6 + ["TỔNG NHẬP T10/26", "NHẬP TRONG NGÀY", "KẾ HOẠCH", "THỰC TẾ", "% KẾ HOẠCH",
                                "TỔNG XUẤT T10/26", "ĐẦU NGÀY", "CUỐI NGÀY"])
        ws.append([1, "Vikoda", 130100110, 230100110, "Kiềm 500ml", "Thùng", 1000, 50, 30, 20, 0.6, 800, 170, 200,
                   10, 40, 60, 100])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        npp = io.BytesIO()
        w2 = Workbook()
        w2.active.append(["BÁO CÁO TỒN KHO THÁNG 10/2026"])
        w2.save(npp)
        npp.seek(0)
        from nxt_ke_toan import tong_hop_nha_may
        nm, bo_qua = tong_hop_nha_may([("Thang 10.xlsx", buf), ("Thang 1.xlsx", npp)])
        self.assertEqual(bo_qua, ["Thang 1.xlsx"])
        bc = nm[(2026, 10)]
        self.assertEqual(bc["ngay"], dt.date(2026, 10, 28))
        d = bc["rows"][0]
        self.assertEqual((d["ma_vik"], d["cuoi_ngay"], d["cho_qc"], d["kh_xuat"], d["da_ban"], d["du_kien"],
                          d["nhap_thang"], d["xuat_thang"]), ("130100110", 200, 10, 40, 60, 100, 1000, 800))

        dm = doc_danh_muc(_danh_muc())
        ky = "Từ ngày 01-10-2026 đến ngày 29-10-2026"
        data, log = tong_hop_nhieu_file([("VKD/Thang 10.xlsm", _nxt(VKD, ky, [("230100110", "Hàng hóa", 0, 2400, 960)]))],
                                        dm)
        content, *_ = build_workbook(data, log, "fp", nha_may=nm, nm_bo_qua=bo_qua)
        wbo = load_workbook(io.BytesIO(content))
        r = list(wbo["Ton thong minh"].iter_rows(values_only=True))
        row = dict(zip(r[0], r[1]))
        self.assertEqual((row["NM_TonThucTe"], row["NM_DaBanGuiKho"], row["NM_TonDuKien"]), (200, 60, 100))
        self.assertIn("Ton nha may", wbo.sheetnames)
        self.assertEqual(list(wbo["Pivot Ton nha may"].iter_rows(values_only=True))[1][4], 200)
        kt = [x for x in wbo["Kiem tra"].iter_rows(values_only=True) if x[3] == "Thang 1.xlsx"]
        self.assertEqual(len(kt), 1)


if __name__ == "__main__":
    unittest.main()
