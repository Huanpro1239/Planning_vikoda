import unittest
from io import BytesIO

from openpyxl import Workbook, load_workbook

from sync_planning_calendar import (
    build_date_headers,
    parse_plan_month,
    prepare_calendar_update,
)


class PlanningCalendarTests(unittest.TestCase):
    def make_workbook(self, selector="Tháng 9"):
        workbook = Workbook()
        fc = workbook.active
        fc.title = "FC"
        fc["R1"] = selector

        planning = workbook.create_sheet("Ke_hoach_SX")
        planning["R1"] = "Ngày bắt đầu sản xuất"
        planning["S1"] = "OLD"
        planning["AW1"] = "OLD31"

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_parse_plan_month(self):
        self.assertEqual(parse_plan_month("Tháng 9"), 9)
        self.assertEqual(parse_plan_month("9"), 9)

    def test_september_2026_headers_match_sample_pattern(self):
        headers = build_date_headers(2026, 9)
        self.assertEqual(len(headers), 30)
        self.assertEqual(headers[0], "01/09\nT3")
        self.assertEqual(headers[-1], "30/09\nT4")

    def test_prepare_update_writes_s1_to_av1_and_clears_aw1(self):
        updated, info = prepare_calendar_update(
            self.make_workbook(),
            plan_year=2026,
        )

        self.assertEqual(info["start"], "S1")
        self.assertEqual(info["end"], "AV1")
        self.assertEqual(info["days"], 30)

        workbook = load_workbook(BytesIO(updated), data_only=True)
        try:
            planning = workbook["Ke_hoach_SX"]
            self.assertEqual(planning["S1"].value, "01/09\nT3")
            self.assertEqual(planning["AV1"].value, "30/09\nT4")
            self.assertIsNone(planning["AW1"].value)
        finally:
            workbook.close()

    def test_february_leap_year_uses_29_days(self):
        headers = build_date_headers(2028, 2)
        self.assertEqual(len(headers), 29)
        self.assertEqual(headers[-1].split("\n")[0], "29/02")


if __name__ == "__main__":
    unittest.main()
