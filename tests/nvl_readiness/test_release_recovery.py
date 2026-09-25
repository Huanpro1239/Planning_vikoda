"""Safety contracts for controlled NVL release recovery."""

from __future__ import annotations

from io import BytesIO
import hashlib
import json
from pathlib import Path
import shutil
import tempfile
import unittest

from openpyxl import load_workbook

from nvl.recovery import (
    build_recovery_proposal,
    run_recovery,
    verify_recovery_workbook,
)
from nvl.release import build_release_manifest, write_release_manifest
from sharepoint.client import GraphRequestError
from tests.test_sync_nvl_stock import make_mock_target_bytes


class RecoveryGraph:
    def __init__(self, *, drift_before_upload=False, tamper_after_upload=False):
        self.site_id = "site"
        self.drive_id = "drive"
        self.target_path = "path/to/Kế hoạch mua hàng.xlsx"
        self.item_id = "target-1"
        self.etag = '"{11111111-2222-3333-4444-555555555555},10"'
        self.content = b""
        self.uploads = []
        self.fetch_count = 0
        self.drift_before_upload = drift_before_upload
        self.tamper_after_upload = tamper_after_upload

    def get_site_id(self):
        return self.site_id

    def get_default_drive_id(self, site_id):
        return self.drive_id

    def get_item_by_path(self, drive_id, path):
        self.fetch_count += 1
        etag = self.etag
        if self.drift_before_upload and self.fetch_count >= 2:
            etag = '"{11111111-2222-3333-4444-555555555555},11"'
        return {
            "id": self.item_id,
            "name": "Kế hoạch mua hàng.xlsx",
            "eTag": etag,
            "sharepointIds": {
                "listItemUniqueId": "11111111-2222-3333-4444-555555555555",
            },
            "parentReference": {"driveId": self.drive_id},
        }

    def download_file(self, drive_id, item_id):
        return self.content

    def upload_file(self, drive_id, item_id, content, expected_etag):
        self.uploads.append(
            {
                "drive_id": drive_id,
                "item_id": item_id,
                "content": content,
                "expected_etag": expected_etag,
            }
        )
        self.content = content
        self.etag = '"{11111111-2222-3333-4444-555555555555},11"'
        if self.tamper_after_upload:
            self.content = _set_de(
                self.content,
                {"VT001": (9999, 8888)},
            )
        return {
            "id": item_id,
            "name": "Kế hoạch mua hàng.xlsx",
            "eTag": self.etag,
        }


def _set_de(data: bytes, values: dict[str, tuple[object, object]]) -> bytes:
    wb = load_workbook(BytesIO(data))
    try:
        ws = wb["Ton_NVL"]
        for row in range(2, ws.max_row + 1):
            code = str(ws.cell(row=row, column=1).value or "").strip()
            if code in values:
                d, e = values[code]
                ws.cell(row=row, column=4).value = d
                ws.cell(row=row, column=5).value = e
        out = BytesIO()
        wb.save(out)
        return out.getvalue()
    finally:
        wb.close()


class NVLReleaseRecoveryTests(unittest.TestCase):
    def _bundle(self, root: Path):
        releases = root / "runtime-state" / "nvl" / "releases"
        releases.mkdir(parents=True)

        historical = make_mock_target_bytes(
            [("VT001", 100), ("VT002", 200)]
        )
        historical = _set_de(
            historical,
            {
                "VT001": (100, 11),
                "VT002": (200, 22),
            },
        )
        current = _set_de(
            historical,
            {
                "VT001": (500, 55),
                "VT002": (600, 66),
            },
        )

        readiness = root / "nvl_production_readiness_report.json"
        readiness.write_text(
            json.dumps(
                {
                    "status": "passed",
                    "gate_version": "nvl_production_readiness_v1",
                    "git_sha": "deadbeef",
                },
                sort_keys=True,
            ),
            encoding="utf-8",
        )

        stock_report = {
            "mode": "publish",
            "status": "published",
            "source": {"revision": "stock-src"},
            "target": {"revision": "target-before"},
            "metrics": {"changed_count": 2},
            "post_upload_verified": True,
            "upload_result": {"eTag": "target-after-stock"},
        }
        open_po_report = {
            "publish_requested": True,
            "published": True,
            "source": {"revision": "po-src"},
            "target": {
                "revision_before": "target-after-stock",
                "revision_after": "target-after-po",
                "changed_cells": 2,
            },
            "publish_evidence": {
                "expected_etag": "target-after-stock",
                "upload_skipped": False,
                "upload_etag": "target-after-po",
                "target_revision_after": "target-after-po",
                "post_upload_verified": True,
                "server_sha256": hashlib.sha256(historical).hexdigest(),
            },
        }

        manifest = build_release_manifest(
            stock_report,
            open_po_report,
            b"stock-proposal",
            historical,
            published_at="2026-09-25T01:00:00+00:00",
            environ={
                "NVL_REQUIRE_READINESS": "1",
                "GITHUB_SHA": "deadbeef",
                "GITHUB_RUN_ID": "100",
                "GITHUB_RUN_ATTEMPT": "1",
            },
            readiness_path=readiness,
        )
        manifest_path = releases / f"{manifest['release_id']}.json"
        write_release_manifest(manifest, path=manifest_path)
        latest = releases.parent / "latest_release.json"
        shutil.copyfile(manifest_path, latest)

        historical_path = root / "historical.xlsx"
        historical_path.write_bytes(historical)

        config = root / "recovery_config.json"
        config.write_text(
            json.dumps(
                {
                    "source": {
                        "name": "unused.xlsm",
                        "sharepoint_path": "unused",
                        "sheet_name": "Sheet1",
                        "code_column": 2,
                        "code_column_letter": "B",
                        "value_column": 13,
                        "value_column_letter": "M",
                        "start_row": 2,
                    },
                    "target": {
                        "name": "Kế hoạch mua hàng.xlsx",
                        "sharepoint_path": "path/to/Kế hoạch mua hàng.xlsx",
                        "sourcedoc": "11111111-2222-3333-4444-555555555555",
                        "sheet_name": "Ton_NVL",
                        "code_column": 1,
                        "code_column_letter": "A",
                        "value_column": 4,
                        "value_column_letter": "D",
                        "start_row": 2,
                    },
                }
            ),
            encoding="utf-8",
        )
        return {
            "releases": releases,
            "latest": latest,
            "manifest": manifest_path,
            "release_id": manifest["release_id"],
            "historical": historical_path,
            "historical_bytes": historical,
            "current_bytes": current,
            "config": config,
        }

    def test_proposal_restores_only_d_and_e(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            bundle = self._bundle(Path(tmpdir))
            from nvl.config import load_nvl_config

            cfg = load_nvl_config(bundle["config"])
            proposal, changes = build_recovery_proposal(
                bundle["current_bytes"],
                bundle["historical_bytes"],
                cfg,
            )
            verification = verify_recovery_workbook(
                bundle["current_bytes"],
                proposal,
                bundle["historical_bytes"],
                cfg,
            )

        self.assertTrue(verification["ok"])
        self.assertEqual(len(changes), 4)
        self.assertEqual(
            {(item["code"], item["column"]) for item in changes},
            {
                ("VT001", "D"),
                ("VT001", "E"),
                ("VT002", "D"),
                ("VT002", "E"),
            },
        )

    def test_layout_drift_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            bundle = self._bundle(Path(tmpdir))
            from nvl.config import load_nvl_config

            cfg = load_nvl_config(bundle["config"])
            moved = load_workbook(BytesIO(bundle["current_bytes"]))
            try:
                ws = moved["Ton_NVL"]
                ws["A2"], ws["A3"] = ws["A3"].value, ws["A2"].value
                out = BytesIO()
                moved.save(out)
                drifted = out.getvalue()
            finally:
                moved.close()

            with self.assertRaisesRegex(RuntimeError, "layout"):
                build_recovery_proposal(
                    drifted,
                    bundle["historical_bytes"],
                    cfg,
                )

    def test_dry_run_creates_proposal_report_and_zero_upload(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            bundle = self._bundle(root)
            graph = RecoveryGraph()
            graph.content = bundle["current_bytes"]
            out = root / "out"

            report = run_recovery(
                bundle["release_id"],
                historical_workbook=bundle["historical"],
                config_path=bundle["config"],
                releases_dir=bundle["releases"],
                latest_path=bundle["latest"],
                out_dir=out,
                publish=False,
                graph=graph,
            )

            self.assertEqual(report["status"], "proposal_ready")
            self.assertEqual(report["proposal"]["changed_cells"], 4)
            self.assertEqual(len(graph.uploads), 0)
            self.assertTrue((out / "nvl_recovery_proposal.xlsx").is_file())
            self.assertTrue((out / "nvl_recovery_report.json").is_file())
            self.assertTrue((out / "nvl_recovery_current_backup.xlsx").is_file())

    def test_tampered_historical_workbook_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            bundle = self._bundle(root)
            bundle["historical"].write_bytes(b"tampered")
            graph = RecoveryGraph()
            graph.content = bundle["current_bytes"]

            with self.assertRaises(Exception):
                run_recovery(
                    bundle["release_id"],
                    historical_workbook=bundle["historical"],
                    config_path=bundle["config"],
                    releases_dir=bundle["releases"],
                    latest_path=bundle["latest"],
                    out_dir=root / "out",
                    graph=graph,
                )
            self.assertEqual(len(graph.uploads), 0)

    def test_broken_ledger_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            bundle = self._bundle(root)
            bundle["latest"].write_text(
                "{}\n",
                encoding="utf-8",
            )
            graph = RecoveryGraph()
            graph.content = bundle["current_bytes"]

            with self.assertRaisesRegex(RuntimeError, "ledger"):
                run_recovery(
                    bundle["release_id"],
                    historical_workbook=bundle["historical"],
                    config_path=bundle["config"],
                    releases_dir=bundle["releases"],
                    latest_path=bundle["latest"],
                    out_dir=root / "out",
                    graph=graph,
                )
            self.assertEqual(len(graph.uploads), 0)

    def test_publish_requires_exact_release_approval(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            bundle = self._bundle(root)
            graph = RecoveryGraph()
            graph.content = bundle["current_bytes"]

            with self.assertRaisesRegex(RuntimeError, "approve-release-id"):
                run_recovery(
                    bundle["release_id"],
                    historical_workbook=bundle["historical"],
                    config_path=bundle["config"],
                    releases_dir=bundle["releases"],
                    latest_path=bundle["latest"],
                    out_dir=root / "out",
                    publish=True,
                    approve_release_id="wrong-release",
                    graph=graph,
                )
            self.assertEqual(len(graph.uploads), 0)

    def test_etag_drift_before_upload_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            bundle = self._bundle(root)
            graph = RecoveryGraph(drift_before_upload=True)
            graph.content = bundle["current_bytes"]

            with self.assertRaises(GraphRequestError):
                run_recovery(
                    bundle["release_id"],
                    historical_workbook=bundle["historical"],
                    config_path=bundle["config"],
                    releases_dir=bundle["releases"],
                    latest_path=bundle["latest"],
                    out_dir=root / "out",
                    publish=True,
                    approve_release_id=bundle["release_id"],
                    graph=graph,
                )
            self.assertEqual(len(graph.uploads), 0)

    def test_valid_publish_uses_if_match_and_post_verifies(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            bundle = self._bundle(root)
            graph = RecoveryGraph()
            graph.content = bundle["current_bytes"]

            report = run_recovery(
                bundle["release_id"],
                historical_workbook=bundle["historical"],
                config_path=bundle["config"],
                releases_dir=bundle["releases"],
                latest_path=bundle["latest"],
                out_dir=root / "out",
                publish=True,
                approve_release_id=bundle["release_id"],
                graph=graph,
            )

        self.assertEqual(report["status"], "published")
        self.assertEqual(len(graph.uploads), 1)
        self.assertEqual(
            graph.uploads[0]["expected_etag"],
            '"{11111111-2222-3333-4444-555555555555},10"',
        )
        self.assertTrue(
            report["publish_evidence"]["post_upload_verified"]
        )

    def test_post_upload_tamper_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            bundle = self._bundle(root)
            graph = RecoveryGraph(tamper_after_upload=True)
            graph.content = bundle["current_bytes"]

            with self.assertRaisesRegex(RuntimeError, "verify"):
                run_recovery(
                    bundle["release_id"],
                    historical_workbook=bundle["historical"],
                    config_path=bundle["config"],
                    releases_dir=bundle["releases"],
                    latest_path=bundle["latest"],
                    out_dir=root / "out",
                    publish=True,
                    approve_release_id=bundle["release_id"],
                    graph=graph,
                )
            self.assertEqual(len(graph.uploads), 1)


if __name__ == "__main__":
    unittest.main()
