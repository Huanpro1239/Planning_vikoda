import unittest
from dataclasses import replace
from io import BytesIO

from openpyxl import Workbook, load_workbook

from planning.weekly_engine import DailyPlanRow
from planning.weekly_model import patch_weekly_workbook
from tests import test_service_first_safety_stock as fixtures


class PlanningShortfallNotesTests(unittest.TestCase):
    def _source(self):
        workbook = Workbook()
        workbook.active.title = "Ke_hoach_SX"
        workbook.create_sheet("Manual_notes")["A1"] = "Giữ ghi chú của người dùng"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        return output.getvalue()

    def test_notes_match_service_and_buffer_shortfall_without_changing_schedule(self):
        analysis = fixtures.ServiceFirstSafetyStockTests()._analysis(service_a=2000, service_b=2000)
        updated = patch_weekly_workbook(self._source(), analysis)
        workbook = load_workbook(BytesIO(updated))
        try:
            notes = list(workbook["Ghi_chu_Planning"].values)[1:]
            self.assertTrue(any(values[6] > 0 for values in notes))
            self.assertTrue(any(values[7] > 0 for values in notes))
            for values in notes:
                code = int(values[1])
                calc = next(c for c in analysis.calculated if c.input.ma_sp == code)
                actual = sum(p.qty for p in analysis.daily_plan if p.ma_sp == code)
                self.assertAlmostEqual(values[5], actual)
                self.assertAlmostEqual(values[4] - values[5], values[6] + values[7])
                self.assertIn("Vẫn chạy theo lịch đã xếp", values[8])
                self.assertEqual(workbook["Ke_hoach_SX"].cell(calc.input.source_row, 16).value, actual)
            self.assertEqual(workbook["Manual_notes"]["A1"].value, "Giữ ghi chú của người dùng")
        finally:
            workbook.close()

    def test_resolved_shortfall_removes_old_notes_and_repeated_patch_does_not_duplicate(self):
        analysis = fixtures.ServiceFirstSafetyStockTests()._analysis()
        first = patch_weekly_workbook(self._source(), analysis)
        second = patch_weekly_workbook(first, analysis)
        workbook = load_workbook(BytesIO(second))
        self.assertEqual(workbook.sheetnames.count("Ghi_chu_Planning"), 1)
        self.assertEqual(workbook["Ghi_chu_Planning"].max_row, 2)
        workbook.close()
        complete = replace(analysis, daily_plan=[
            DailyPlanRow(c.input.ma_sp, c.input.source_row, c.input.chuyen,
                         c.start_date, c.schedulable_qty)
            for c in analysis.calculated
        ])
        workbook = load_workbook(BytesIO(patch_weekly_workbook(second, complete)))
        self.assertEqual(workbook["Ghi_chu_Planning"].max_row, 1)
        workbook.close()
