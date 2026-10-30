import io
import unittest

from openpyxl import Workbook, load_workbook

from sell_in import doc_danh_muc
from xuat_kho import build_workbook, tong_hop_nhieu_file

HEADER = ["MACHINHANH", "TENCHINHANH", "SOPHIEUXUATKHO", "NGAYXUATKHO", "CHUYENKHOKYGUI", "MAKHACHHANG",
          "KHACHHANG", "DONVIVANCHUYEN", "SOXETAU", "NGUOIVANCHUYEN", "GHICHU", "LOAIDONHANG", "SODONHANG",
          "KHOXUAT", "MADOITUONG", "DOITUONG", "LOAISANPHAM", "MASANPHAM", "TENSANPHAM", "DONVITINH",
          "SOLUONG", "TRONGLUONG", "THETICH", "SCHEME"]


def _xuat_kho(ten_cn, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "REPORT_XUATBANHANG"
    ws.append(["BÁO CÁO XUẤT KHO BÁN HÀNG"])
    ws.append([None])
    ws.append(["Từ ngày : 01-09-2026   Đến ngày: 30-09-2026"])
    ws.append([None])
    ws.append(HEADER)
    for ckg, kh, ldh, ma, sl in rows:
        r = [None] * len(HEADER)
        r[1], r[4], r[5], r[11], r[17], r[20] = ten_cn, ckg, kh, ldh, ma, sl
        ws.append(r)
    wb.create_sheet("Sheet1")["A1"] = "Evaluation Warning"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _danh_muc():
    wb = Workbook()
    ws = wb.active
    ws.title = "DanhMucSanPham"
    ws.append(["MaSanPhamMoi_Vikoda", "MaSanPhamMoi_VKD", "TenSanPham", "Tên viết tắt", "Quy cách",
               "ĐVT", "Brand", "SKUs Group"])
    ws.append([130100110, 230100110, "Kiềm 500ml", "Khoáng kiềm 500ml", 24, "Thùng", "Vikoda", "Pet Vikoda"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class XuatKhoTests(unittest.TestCase):
    def test_actual_shipments_only(self):
        dm = doc_danh_muc(_danh_muc())
        ckg = "Chuyển qua kho ký gửi"
        vikoda = _xuat_kho("CÔNG TY CỔ PHẦN NƯỚC KHOÁNG KHÁNH HÒA", [
            (None, "VKD3", "Đơn hàng bán", "130100110", 2400),     # bán nội bộ cho VKD -> bỏ
            (None, "B-1", "Đơn hàng bán", "130100110", 480),       # 20 thùng
            (None, "B-1", "Đơn hàng bán", "430300037", 50),        # két -> bỏ
            (ckg, "B-1", "Đơn hàng bán", "130100110", 240),        # chuyển ký gửi -> bỏ
            (None, "B-1", "Ký gửi", "130100110", 120),             # xuất từ kho ký gửi: 5 thùng
        ])
        vkd = _xuat_kho("CÔNG TY CỔ PHẦN ĐẦU TƯ VÀ THƯƠNG MẠI VKD", [
            (None, "B-2", "Đơn hàng bán", "230100110", 240),       # 10 thùng
            (None, "B-2", "Đơn hàng khuyến mại", "230100110", 24),  # 1 thùng
            (ckg, "B-2", "Đơn hàng bán", "230100110", 9999),       # chuyển ký gửi -> bỏ
            (None, "B-2", "Ký gửi", "230100110", 48),              # 2 thùng
            (None, "B-2", "Đơn hàng bán", "230100014", 999),       # bán thành phẩm -> bỏ
            (None, "B-2", "Đơn hàng bán", "130100110", 999),       # mã Vikoda ở VKD -> bỏ
        ])
        data, log = tong_hop_nhieu_file([("Vikoda/Xuat kho thang 9.xlsm", vikoda),
                                         ("VKD/Xuat kho thang 9.xlsm", vkd)], dm)
        got = {(r[0], r[4], r[5]): r[13] for r in data}
        self.assertEqual(got, {
            ("Vikoda", "Đơn hàng bán", "130100110"): 20,
            ("Vikoda", "Ký gửi", "130100110"): 5,
            ("VKD", "Đơn hàng bán", "230100110"): 10,
            ("VKD", "Đơn hàng khuyến mại", "230100110"): 1,
            ("VKD", "Ký gửi", "230100110"): 2,
        })
        self.assertEqual([(r[1], r[3], r[6]) for r in log], [("Vikoda", 9, 1), ("VKD", 9, 1)])

        content, n_sp, n_thang = build_workbook(data, log, "fp")
        self.assertEqual((n_sp, n_thang), (1, 1))
        ws = load_workbook(io.BytesIO(content))["Pivot"]
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][-2:], ("T09/2026", "Tong"))
        self.assertEqual(rows[1][:2], ("130100110", "230100110"))
        self.assertEqual(rows[1][-2:], (38, 38))

    def test_rerun_month_replaces(self):
        dm = doc_danh_muc(_danh_muc())
        a = _xuat_kho("VKD", [(None, "B", "Đơn hàng bán", "230100110", 240)])
        b = _xuat_kho("VKD", [(None, "B", "Đơn hàng bán", "230100110", 480)])
        data, log = tong_hop_nhieu_file([("VKD/a.xlsm", a), ("VKD/b.xlsm", b)], dm)
        self.assertEqual([r[13] for r in data], [20])
        self.assertEqual(len(log), 1)


if __name__ == "__main__":
    unittest.main()
