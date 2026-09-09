"""Unit tests for sync_nvl_stock module."""

from io import BytesIO
from typing import Any
import unittest
import zipfile
from lxml import etree
from openpyxl import Workbook, load_workbook

from sync_nvl_stock import (
    NVLConfig,
    load_nvl_config,
    normalize_nvl_code,
    parse_nvl_quantity,
    patch_nvl_destination_workbook,
    read_nvl_source_stock,
    reconcile_nvl_target,
    verify_nvl_patched_workbook,
)


def make_mock_config() -> NVLConfig:
    return NVLConfig(
        source_name="XNT_ketoan_Vikoda.xlsm",
        source_path="path/to/XNT_ketoan_Vikoda.xlsm",
        source_sheet="Sheet1",
        source_code_col=2,
        source_code_col_letter="B",
        source_value_col=13,
        source_value_col_letter="M",
        source_start_row=2,
        source_sourcedoc="SOURCE-GUID",
        target_name="Kế hoạch mua hàng.xlsx",
        target_path="path/to/Kế hoạch mua hàng.xlsx",
        target_sheet="Ton_NVL",
        target_code_col=1,
        target_code_col_letter="A",
        target_value_col=4,
        target_value_col_letter="D",
        target_start_row=2,
        target_sourcedoc="TARGET-GUID",
    )


def make_mock_source_bytes(rows: list[tuple[Any, Any]], *, with_header=True) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    start_r = 1
    if with_header:
        ws.cell(1, 2, value="Mã VT")
        ws.cell(1, 13, value="Số lượng tồn")
        start_r = 2

    for idx, (code, qty) in enumerate(rows):
        r = start_r + idx
        ws.cell(r, 2, value=code)
        ws.cell(r, 13, value=qty)

    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def make_mock_target_bytes(rows: list[tuple[Any, Any]], *, with_header=True, extra_sheets=True) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ton_NVL"
    start_r = 1
    if with_header:
        ws.cell(1, 1, value="Mã vật tư")
        ws.cell(1, 2, value="Tên vật tư")
        ws.cell(1, 3, value="ĐVT")
        ws.cell(1, 4, value="Tồn kho")
        start_r = 2

    for idx, (code, qty) in enumerate(rows):
        r = start_r + idx
        ws.cell(r, 1, value=code)
        ws.cell(r, 2, value=f"Tên {code}")
        ws.cell(r, 3, value="Cái")
        ws.cell(r, 4, value=qty)

    if extra_sheets:
        ws_extra = wb.create_sheet("DanhMuc")
        ws_extra["A1"] = "Danh mục vật tư"
        ws_extra["B1"] = 12345

    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


class SyncNVLStockTests(unittest.TestCase):
    def test_normalize_nvl_code(self):
        """Kiểm tra quy tắc chuẩn hóa mã vật tư NVL."""
        # 1. Số nguyên và float nguyên
        self.assertEqual(normalize_nvl_code(12345), "12345")
        self.assertEqual(normalize_nvl_code(12345.0), "12345")
        self.assertEqual(normalize_nvl_code("12345.0"), "12345")
        self.assertEqual(normalize_nvl_code("12345.000"), "12345")

        # 2. Chuỗi có số 0 ở đầu được bảo toàn
        self.assertEqual(normalize_nvl_code("012345"), "012345")
        self.assertEqual(normalize_nvl_code("00987"), "00987")

        # 3. Trim khoảng trắng đầu và cuối, giữ khoảng trắng ở giữa
        self.assertEqual(normalize_nvl_code("  NVL-001  "), "NVL-001")
        self.assertEqual(normalize_nvl_code("  VT 123 ABC  "), "VT 123 ABC")

        # 4. None, boolean, chuỗi rỗng
        self.assertIsNone(normalize_nvl_code(None))
        self.assertIsNone(normalize_nvl_code(""))
        self.assertIsNone(normalize_nvl_code("   "))
        self.assertIsNone(normalize_nvl_code(True))
        self.assertIsNone(normalize_nvl_code(False))

    def test_parse_nvl_quantity(self):
        """Kiểm tra phân tích số lượng tồn kho NVL."""
        # 1. Số 0, số âm, số lẻ
        self.assertEqual(parse_nvl_quantity(0), 0.0)
        self.assertEqual(parse_nvl_quantity(0.0), 0.0)
        self.assertEqual(parse_nvl_quantity(150), 150.0)
        self.assertEqual(parse_nvl_quantity(-25.5), -25.5)
        self.assertAlmostEqual(parse_nvl_quantity(123.456), 123.456)

        # 2. Chuỗi số có định dạng
        self.assertEqual(parse_nvl_quantity("0"), 0.0)
        self.assertEqual(parse_nvl_quantity("1,250"), 1250.0)
        self.assertEqual(parse_nvl_quantity("-50.25"), -50.25)
        self.assertEqual(parse_nvl_quantity("  3,456.789  "), 3456.789)

        # 3. Lỗi chặn: rỗng, boolean, NaN, văn bản không phải số
        with self.assertRaises(ValueError):
            parse_nvl_quantity(None)
        with self.assertRaises(ValueError):
            parse_nvl_quantity("")
        with self.assertRaises(ValueError):
            parse_nvl_quantity("   ")
        with self.assertRaises(ValueError):
            parse_nvl_quantity(True)
        with self.assertRaises(ValueError):
            parse_nvl_quantity(False)
        with self.assertRaises(ValueError):
            parse_nvl_quantity("abc")
        with self.assertRaises(ValueError):
            parse_nvl_quantity(float("nan"))
        with self.assertRaises(ValueError):
            parse_nvl_quantity(float("inf"))

    def test_read_nvl_source_stock_success(self):
        """Đọc nguồn thành công bao gồm tồn 0, âm, lẻ, bỏ qua dòng tổng cộng."""
        cfg = make_mock_config()
        rows = [
            ("VT001", 100),
            ("VT002", 0),
            ("VT003", -15.5),
            ("VT004", 250.75),
            ("Tổng cộng", 335.25),  # Dòng tổng cộng
        ]
        src_bytes = make_mock_source_bytes(rows)
        stock, meta = read_nvl_source_stock(src_bytes, cfg)

        self.assertEqual(len(stock), 4)
        self.assertEqual(stock["VT001"], 100.0)
        self.assertEqual(stock["VT002"], 0.0)
        self.assertEqual(stock["VT003"], -15.5)
        self.assertEqual(stock["VT004"], 250.75)
        self.assertNotIn("Tổng cộng", stock)

    def test_read_source_stock_rejects_duplicates(self):
        """Bắt lỗi khi mã vật tư bị lặp trong nguồn."""
        cfg = make_mock_config()
        rows = [
            ("VT001", 100),
            ("VT002", 50),
            ("VT001", 80),  # Lặp mã VT001
        ]
        src_bytes = make_mock_source_bytes(rows)
        with self.assertRaises(RuntimeError) as ctx:
            read_nvl_source_stock(src_bytes, cfg)
        self.assertIn("VT001", str(ctx.exception))
        self.assertIn("lặp", str(ctx.exception).lower())

    def test_read_source_stock_formula_without_cache_rejected(self):
        """Công thức trong cột M không có cached value bị từ chối."""
        cfg = make_mock_config()
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["B2"] = "VT001"
        ws["M2"] = "=SUM(C2:L2)"  # Formula without cached value
        buf = BytesIO()
        wb.save(buf)
        wb.close()

        with self.assertRaises(RuntimeError) as ctx:
            read_nvl_source_stock(buf.getvalue(), cfg)
        self.assertIn("M2", str(ctx.exception))
        self.assertIn("cached", str(ctx.exception).lower())

    def test_reconcile_and_patch_nvl_workbook_full_flow(self):
        """Đối soát và patch hoàn chỉnh: ghép theo mã, thứ tự dòng khác nhau, giữ nguyên mã thiếu."""
        cfg = make_mock_config()

        # Nguồn: VT001=150, VT002=0, VT003=75.5, VT_EXTRA=999
        source_rows = [
            ("VT001", 150),
            ("VT002", 0),
            ("VT003", 75.5),
            ("VT_EXTRA", 999),
        ]
        source_bytes = make_mock_source_bytes(source_rows)
        source_stock, _ = read_nvl_source_stock(source_bytes, cfg)

        # Đích: thứ tự khác, có VT004 (thiếu ở nguồn), VT001 đã có 150 (không đổi)
        target_rows = [
            ("VT003", 50),     # Cần đổi 50 -> 75.5
            ("VT001", 150),    # Không đổi (đã là 150)
            ("VT004", 888),    # Thiếu ở nguồn -> giữ nguyên 888
            ("VT002", 10),     # Cần đổi 10 -> 0
        ]
        target_bytes = make_mock_target_bytes(target_rows)

        rec = reconcile_nvl_target(target_bytes, source_stock, cfg)
        self.assertEqual(rec.status, "completed_with_warnings")
        self.assertEqual(len(rec.changes), 2)  # VT003 và VT002
        self.assertEqual(len(rec.unchanged), 1)  # VT001
        self.assertEqual(len(rec.missing_in_source), 1)  # VT004
        self.assertEqual(rec.source_only, ["VT_EXTRA"])

        # Patch XML
        patched_bytes = patch_nvl_destination_workbook(target_bytes, rec, cfg)
        verify_res = verify_nvl_patched_workbook(target_bytes, patched_bytes, rec, cfg)
        self.assertTrue(verify_res["ok"])

        # Kiểm tra workbook sau patch
        wb = load_workbook(BytesIO(patched_bytes), data_only=True)
        ws = wb["Ton_NVL"]
        self.assertEqual(ws["D2"].value, 75.5)  # VT003 đổi sang 75.5
        self.assertEqual(ws["D3"].value, 150)   # VT001 giữ nguyên 150
        self.assertEqual(ws["D4"].value, 888)   # VT004 thiếu ở nguồn -> giữ nguyên 888
        self.assertEqual(ws["D5"].value, 0)     # VT002 đổi sang 0

        # Kiểm tra sheet phụ DanhMuc hoàn toàn nguyên vẹn
        self.assertIn("DanhMuc", wb.sheetnames)
        self.assertEqual(wb["DanhMuc"]["B1"].value, 12345)
        wb.close()

    def test_reject_formula_in_target_cell(self):
        """Bắt lỗi khi ô đích cột D chứa công thức người dùng."""
        cfg = make_mock_config()
        source_stock = {"VT001": 100}

        wb = Workbook()
        ws = wb.active
        ws.title = "Ton_NVL"
        ws["A2"] = "VT001"
        ws["D2"] = "=VLOOKUP(A2, Kho!A:B, 2, FALSE)"  # Formula in target cell
        buf = BytesIO()
        wb.save(buf)
        wb.close()

        with self.assertRaises(RuntimeError) as ctx:
            reconcile_nvl_target(buf.getvalue(), source_stock, cfg)
        self.assertIn("D2", str(ctx.exception))
        self.assertIn("công thức", str(ctx.exception).lower())

    def test_reject_duplicate_codes_in_target(self):
        """Bắt lỗi khi mã vật tư bị lặp trong đích."""
        cfg = make_mock_config()
        source_stock = {"VT001": 100, "VT002": 200}
        target_rows = [
            ("VT001", 10),
            ("VT002", 20),
            ("VT001", 30),  # Lặp VT001
        ]
        target_bytes = make_mock_target_bytes(target_rows)

        with self.assertRaises(RuntimeError) as ctx:
            reconcile_nvl_target(target_bytes, source_stock, cfg)
        self.assertIn("VT001", str(ctx.exception))
        self.assertIn("lặp", str(ctx.exception).lower())

    def test_no_matching_codes_raises_error(self):
        """Dừng và báo lỗi khi không có mã vật tư nào khớp giữa nguồn và đích."""
        cfg = make_mock_config()
        source_stock = {"VT_SRC_1": 100}
        target_rows = [("VT_TGT_1", 50), ("VT_TGT_2", 60)]
        target_bytes = make_mock_target_bytes(target_rows)

        with self.assertRaises(RuntimeError) as ctx:
            reconcile_nvl_target(target_bytes, source_stock, cfg)
        self.assertIn("không có mã vật tư nào", str(ctx.exception).lower())

    def test_idempotence_second_run_has_zero_changes(self):
        """Chạy lần thứ hai trên file đã patch cho kết quả unchanged, 0 thay đổi."""
        cfg = make_mock_config()
        source_stock = {"VT001": 100.0, "VT002": 200.0}
        target_rows = [("VT001", 50.0), ("VT002", 70.0)]
        target_bytes = make_mock_target_bytes(target_rows)

        # Lần 1: 2 thay đổi
        rec1 = reconcile_nvl_target(target_bytes, source_stock, cfg)
        self.assertEqual(len(rec1.changes), 2)
        patched_bytes = patch_nvl_destination_workbook(target_bytes, rec1, cfg)

        # Lần 2: 0 thay đổi, status = 'unchanged'
        rec2 = reconcile_nvl_target(patched_bytes, source_stock, cfg)
        self.assertEqual(len(rec2.changes), 0)
        self.assertEqual(len(rec2.unchanged), 2)
        self.assertEqual(rec2.status, "unchanged")

        # Patch lần 2 trả về chính xác bytes cũ
        second_patched_bytes = patch_nvl_destination_workbook(patched_bytes, rec2, cfg)
        self.assertEqual(patched_bytes, second_patched_bytes)


if __name__ == "__main__":
    unittest.main()
