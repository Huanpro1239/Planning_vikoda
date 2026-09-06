import math
import unittest
from datetime import date, timedelta
from io import BytesIO
from unittest.mock import patch

from openpyxl import Workbook

from planning_schedule_report import build_schedule_report
import sync_planning_pipeline as pipeline_runner
import sync_planning_schedule as base
import sync_planning_schedule_priority as priority
import sync_planning_schedule_priority_debt as debt_priority
import sync_planning_schedule_production as production
import sync_planning_schedule_galon_v9 as galon_v9
import sync_stock
from sync_planning_calendar import build_date_headers
from sync_stock import GraphRequestError
from verify_planning_month import verify_workbook


class DebtProductionInstallerTests(unittest.TestCase):
    def _galon_19l(self):
        headers = [date(2026, 9, 1) + timedelta(days=i) for i in range(30)]
        demand = {
            day: (0.0 if day.weekday() == 6 else 6000.0)
            for day in headers
        }
        product = {
            "row": 2,
            "code": "130100006",
            "uom": "Bình",
            "batch": 5400.0,
            "per_shift": 5400.0,
            "line": "Galon",
            "product_group": "Galon",
            "classification": "Không đường",
            "is_sugar": False,
            "max_shifts_per_day": 2.0,
            "actual_stock": 24843.0,
            "fc": 156000.0,
            "target_stock": 0.0,
            "debt": 26135.0,
            "planned_qty": 162000.0,
            "earliest_date": headers[0],
            "preferred_date": headers[0],
            "demand_by_day": demand,
            "existing_daily": [None] * 31,
        }
        base._validate_quantum(product)
        return headers, product

    def test_production_installer_keeps_debt_risk_and_galon_v9_after_v8(self):
        headers, product = self._galon_19l()

        production.install_production_output_cleanup()
        self.assertIs(priority._risk_snapshot, debt_priority.debt_aware_risk_snapshot)
        self.assertIs(base._allocate_galon_line, galon_v9.allocate_galon_quantized_service)

        schedule, _, carryover, meta = base._allocate_galon_line(
            headers,
            [product],
            2.0,
        )
        self.assertEqual(schedule[product["code"]][headers[0]], 10800)
        day1_net = (
            product["actual_stock"]
            + schedule[product["code"]][headers[0]]
            - product["debt"]
            - product["demand_by_day"][headers[0]]
        )
        self.assertEqual(day1_net, 3508)
        self.assertEqual(carryover, {})
        self.assertEqual(meta["mode"], "galon_quantized_service_v9")
        self.assertTrue(meta["timeline"])

    def test_repeated_production_installer_does_not_lose_debt_or_v9_hooks(self):
        production.install_production_output_cleanup()
        production.install_production_output_cleanup()

        self.assertIs(priority._risk_snapshot, debt_priority.debt_aware_risk_snapshot)
        self.assertIs(
            priority.shortage_aware_weekly_unit_targets,
            debt_priority.debt_aware_weekly_unit_targets,
        )
        self.assertIs(base._allocate_galon_line, galon_v9.allocate_galon_quantized_service)


class CarryoverMassBalanceGuardTests(unittest.TestCase):
    def _workbook(self, *, daily_qty):
        wb = Workbook()
        fc = wb.active
        fc.title = "FC"
        fc.append([
            "STT", "Mã SP", "Tên SP", "Đơn vị tính",
            "Tháng 1", "Tháng 2", "Tháng 3", "Tháng 4",
            "Tháng 5", "Tháng 6", "Tháng 7", "Tháng 8",
            "Tháng 9", "Tháng 10", "Tháng 11", "Tháng 12", None, "Tháng 9",
        ])
        fc.append([1, 100001, "A", "Thùng", 0, 0, 0, 0, 0, 0, 0, 0, 4500, 0, 0, 0])

        stock = wb.create_sheet("Ton_kho")
        stock.append([
            "Mã Sản Phẩm", "Tên Sản Phẩm", "Đơn vị tính", "Tồn thực tế",
            "Tồn Vikoda (Nhà Máy)", "Tồn VKD (Nhà Máy)",
            "Tồn Vikoda (Kho khác)", "Tồn VKD (Kho khác)",
        ])
        stock.append([100001, "A", "Thùng", 0, 0, 0, 0, 0])

        planning = wb.create_sheet("Ke_hoach_SX")
        planning.append([
            "Mã Sản Phẩm", "Tên Sản Phẩm", "Đơn vị tính", "Số lượng /mẻ",
            "Số lượng/ ca", "Chuyền", "Nhóm sản phẩm", "Phân loại SP",
            "Số ca theo ngày", "Tồn đầu thực tế", "Tồn đầu sổ sách", "FC",
            "Tồn cuối dự kiến", "Nợ kho", "Số lượng cần sản xuất",
            "Số lượng sản xuất theo mẻ/ca", "Số ngày cần sản xuất", "Ngày bắt đầu sản xuất",
        ] + build_date_headers(2026, 9))
        planning.append([
            100001, "A", "Thùng", 4500, 4500, "Galon", "Galon", "Không đường",
            2, 0, 0, 4500, 0, 0, 4500, 4500, 0.5, None,
        ] + [daily_qty] + [None] * 29)

        out = BytesIO()
        wb.save(out)
        return out.getvalue()

    @staticmethod
    def _report(*, scheduled, carryover):
        return {
            "schema_version": 2,
            "algorithm": "priority_v8",
            "plan_month": "2026-09",
            "input_revision": {"etag": "copy-etag"},
            "mass_balance": {
                "100001": {
                    "planned_qty": 4500,
                    "scheduled_qty": scheduled,
                    "carryover_qty": carryover,
                }
            },
            "resources": {},
            "publish_status": "valid_but_infeasible" if carryover else "feasible",
        }

    def test_verifier_rejects_overproduction_hidden_by_negative_carryover(self):
        with self.assertRaisesRegex(RuntimeError, "âm|vượt P"):
            verify_workbook(
                self._workbook(daily_qty=9000),
                schedule_report=self._report(scheduled=9000, carryover=-4500),
            )

    def test_verifier_rejects_nan_and_infinite_carryover(self):
        for bad in (float("nan"), float("inf")):
            with self.subTest(bad=bad):
                with self.assertRaisesRegex(RuntimeError, "hữu hạn"):
                    verify_workbook(
                        self._workbook(daily_qty=0),
                        schedule_report=self._report(scheduled=0, carryover=bad),
                    )

    def test_verifier_accepts_valid_positive_carryover(self):
        info = verify_workbook(
            self._workbook(daily_qty=0),
            schedule_report=self._report(scheduled=0, carryover=4500),
        )
        self.assertEqual(info["publish_status"], "valid_but_infeasible")

    def test_report_builder_rejects_negative_carryover(self):
        workbook = Workbook()
        workbook.active.title = "Ke_hoach_SX"
        out = BytesIO()
        workbook.save(out)
        day = date(2026, 9, 1)
        info = {
            "headers": [day],
            "products": [{"code": "100001", "planned_qty": 4500, "uom": "Thùng"}],
            "schedule": {"100001": {day: 9000}},
            "carryover": {"100001": -4500},
            "line_capacity": {},
            "optimizer_meta": {},
            "inventory": {},
            "utilization": {},
        }
        with self.assertRaisesRegex(RuntimeError, "âm|vượt P"):
            build_schedule_report(out.getvalue(), info)

    def test_report_builder_accepts_valid_positive_carryover(self):
        workbook = Workbook()
        workbook.active.title = "Ke_hoach_SX"
        out = BytesIO()
        workbook.save(out)
        day = date(2026, 9, 1)
        info = {
            "headers": [day],
            "products": [{"code": "100001", "planned_qty": 4500, "uom": "Thùng"}],
            "schedule": {"100001": {day: 0}},
            "carryover": {"100001": 4500},
            "line_capacity": {},
            "optimizer_meta": {},
            "inventory": {},
            "utilization": {},
        }
        report = build_schedule_report(out.getvalue(), info)
        self.assertEqual(report["mass_balance"]["100001"]["carryover_qty"], 4500)


class PipelineTransactionIntegrationTests(unittest.TestCase):
    class FakeGraph:
        def __init__(self, *, first_upload_412=False):
            self.attempt = 0
            self.first_upload_412 = first_upload_412
            self.uploads = []
            self.downloads = []
            self.events = []

        def get_item_by_path(self, drive_id, path):
            if path == sync_stock.DEST_PATH:
                self.attempt += 1
                return {
                    "id": f"target-r{self.attempt}",
                    "eTag": f"target-etag-r{self.attempt}",
                    "lastModifiedDateTime": f"target-time-r{self.attempt}",
                }
            key = next(key for key, value in pipeline_runner.SOURCES.items() if value == path)
            return {
                "id": f"{key}-r{self.attempt}",
                "eTag": f"{key}-etag-r{self.attempt}",
                "lastModifiedDateTime": f"{key}-time-r{self.attempt}",
            }

        def download_file(self, drive_id, item_id):
            self.downloads.append(item_id)
            return f"bytes:{item_id}".encode()

        def upload_file(self, drive_id, item_id, content, expected_etag):
            self.uploads.append((item_id, content, expected_etag))
            self.events.append(f"upload-attempt-{len(self.uploads)}")
            if self.first_upload_412 and len(self.uploads) == 1:
                raise GraphRequestError(
                    "precondition failed",
                    status_code=412,
                    error_code="preconditionFailed",
                )
            self.events.append(f"upload-success-{len(self.uploads)}")
            return {"name": "copy.xlsx"}

    @staticmethod
    def _fake_report(revision):
        return {
            "schema_version": 2,
            "algorithm": "priority_v8",
            "plan_month": "2026-09",
            "input_revision": revision,
            "mass_balance": {},
            "resources": {},
            "publish_status": "feasible",
            "pipeline": {"conversion_hash": "hash"},
        }

    def test_validation_failure_uploads_zero_and_saves_no_state(self):
        graph = self.FakeGraph()
        saved = []

        def fail_prepare(*args, **kwargs):
            raise RuntimeError("validation failed")

        with (
            patch.object(pipeline_runner, "prepare_pipeline_output", fail_prepare),
            patch.object(pipeline_runner.metrics, "load_runtime_state", lambda: {"state": "old"}),
            patch.object(pipeline_runner, "_save_states_after_success", lambda *args: saved.append(args)),
            patch.object(pipeline_runner, "print_operational_report", lambda report: None),
        ):
            with self.assertRaisesRegex(RuntimeError, "validation failed"):
                pipeline_runner.run_pipeline_with_retry(
                    graph,
                    "drive",
                    sleep_func=lambda _: None,
                    max_attempts=2,
                )

        self.assertEqual(graph.uploads, [])
        self.assertEqual(saved, [])

    def test_success_uploads_exactly_once_then_saves_state(self):
        graph = self.FakeGraph()
        saved = []

        def prepare(target_bytes, source_bytes, *, runtime_state, input_revision):
            return (
                b"final:" + target_bytes,
                self._fake_report(input_revision),
                {"state": "new"},
            )

        def save(*args):
            graph.events.append("state-save")
            saved.append(args)

        with (
            patch.object(pipeline_runner, "prepare_pipeline_output", prepare),
            patch.object(pipeline_runner.metrics, "load_runtime_state", lambda: {"state": "old"}),
            patch.object(pipeline_runner, "_save_states_after_success", save),
            patch.object(pipeline_runner, "print_operational_report", lambda report: None),
        ):
            result = pipeline_runner.run_pipeline_with_retry(
                graph,
                "drive",
                sleep_func=lambda _: None,
                max_attempts=2,
            )

        self.assertTrue(result["uploaded"])
        self.assertEqual(len(graph.uploads), 1)
        self.assertEqual(graph.uploads[0][2], "target-etag-r1")
        self.assertEqual(len(saved), 1)
        self.assertEqual(graph.events[-2:], ["upload-success-1", "state-save"])

    def test_412_reloads_all_snapshots_recomputes_and_saves_state_only_after_success(self):
        graph = self.FakeGraph(first_upload_412=True)
        prepare_calls = []
        saved = []
        sleeps = []

        def prepare(target_bytes, source_bytes, *, runtime_state, input_revision):
            prepare_calls.append((target_bytes, dict(source_bytes), input_revision))
            return (
                b"final:" + target_bytes,
                self._fake_report(input_revision),
                {"state": f"new-r{len(prepare_calls)}"},
            )

        def save(*args):
            graph.events.append("state-save")
            saved.append(args)

        with (
            patch.object(pipeline_runner, "prepare_pipeline_output", prepare),
            patch.object(pipeline_runner.metrics, "load_runtime_state", lambda: {"state": "old"}),
            patch.object(pipeline_runner, "_save_states_after_success", save),
            patch.object(pipeline_runner, "print_operational_report", lambda report: None),
        ):
            result = pipeline_runner.run_pipeline_with_retry(
                graph,
                "drive",
                sleep_func=sleeps.append,
                max_attempts=2,
                retry_delay_seconds=3,
            )

        self.assertTrue(result["uploaded"])
        self.assertEqual(len(prepare_calls), 2)
        self.assertEqual(prepare_calls[0][0], b"bytes:target-r1")
        self.assertEqual(prepare_calls[1][0], b"bytes:target-r2")
        self.assertTrue(all(value.endswith(b"-r1") for value in prepare_calls[0][1].values()))
        self.assertTrue(all(value.endswith(b"-r2") for value in prepare_calls[1][1].values()))
        self.assertEqual(
            [upload[2] for upload in graph.uploads],
            ["target-etag-r1", "target-etag-r2"],
        )
        self.assertEqual(len(saved), 1)
        self.assertEqual(graph.events[-2:], ["upload-success-2", "state-save"])
        self.assertEqual(sleeps, [3])


if __name__ == "__main__":
    unittest.main()
