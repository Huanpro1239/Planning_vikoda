"""Report date from the actual warehouse Excel header, not delivery-plan date."""
import unittest
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook

from planning.metrics.readers import find_report_date, read_actual_inputs


def make_source(report_date="23-24/10/2026"):
    book = Workbook()
    ws = book.active
    ws.title = "THÁNG 10"
    ws["A1"] = "BÁO CÁO KHO THÀNH PHẨM"
    ws["I2"] = report_date
    ws["P3"] = "KẾ HOẠCH XUẤT KHO 26/10/2026"
    ws["C5"] = 130100006
    ws["G5"] = 100
    ws["Q5"] = 7
    ws["C6"] = 130100071
    ws["G6"] = 0
    ws["Q6"] = 0
    output = BytesIO()
    book.save(output)
    return ws, output.getvalue()


class WarehouseReportingDateTests(unittest.TestCase):
    def test_real_october_text_period_uses_end_day(self):
        ws, data = make_source()
        self.assertEqual(find_report_date(ws), date(2026, 10, 24))
        actual_date, receipts, consignments = read_actual_inputs(data)
        self.assertEqual(actual_date, date(2026, 10, 24))
        self.assertEqual(receipts, {"130100006": 100, "130100071": 0})
        self.assertEqual(consignments, {"130100006": 7, "130100071": 0})

    def test_single_text_dates_supported(self):
        for text in ("24/10/2026", "24-10-2026", "24.10.2026"):
            with self.subTest(text=text):
                ws, _ = make_source(text)
                self.assertEqual(find_report_date(ws), date(2026, 10, 24))

    def test_native_excel_dates_still_supported(self):
        for typed_date in (date(2026, 10, 24), datetime(2026, 10, 24, 15)):
            with self.subTest(typed_date=typed_date):
                ws, _ = make_source(typed_date)
                self.assertEqual(find_report_date(ws), date(2026, 10, 24))

    def test_legacy_header_date_outside_I2(self):
        ws, _ = make_source(None)
        ws["A1"] = date(2026, 10, 24)
        self.assertEqual(find_report_date(ws), date(2026, 10, 24))

    def test_do_not_use_future_shipping_plan_date(self):
        ws, _ = make_source(None)
        with self.assertRaisesRegex(RuntimeError, "Không tìm thấy ngày báo cáo"):
            find_report_date(ws)

    def test_reject_reversed_period(self):
        ws, _ = make_source("25-23/10/2026")
        with self.assertRaisesRegex(RuntimeError, "ngược thứ tự"):
            find_report_date(ws)

    def test_reject_impossible_calendar_date(self):
        ws, _ = make_source("30-31/02/2026")
        with self.assertRaisesRegex(RuntimeError, "không hợp lệ"):
            find_report_date(ws)


if __name__ == "__main__":
    unittest.main()
