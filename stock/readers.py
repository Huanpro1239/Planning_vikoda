"""Source and master-data readers for finished-goods stock."""

import hashlib
import json
import re
from io import BytesIO

from openpyxl import load_workbook

from excel.openpyxl_io import safe_close_workbook

from .constants import MASTER_SHEET
from .values import clean_number, normalize_code, to_number


DEBT_HEADER_NAMES = frozenset({
    "debt mode",
    "cách tính nợ",
    "cach tinh no",
})
PROFILE_HEADER_NAMES = frozenset({
    "schedule profile",
    "profile lịch",
    "profile lich",
    "lịch sản xuất",
    "lich san xuat",
})


def _norm(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _find_header_col(headers, names):
    wanted = {_norm(name) for name in names}
    for index, header in enumerate(headers):
        if _norm(header) in wanted:
            return index
    return None


def _normalize_debt_mode(value):
    text = _norm(value).replace("-", "_").replace(" ", "_")
    if not text:
        return None
    mapping = {
        "subtract_book_on_debt": "SUBTRACT_BOOK_ON_DEBT",
        "tru_ton_so": "SUBTRACT_BOOK_ON_DEBT",
        "trừ_tồn_sổ": "SUBTRACT_BOOK_ON_DEBT",
        "ignore_book_on_debt": "IGNORE_BOOK_ON_DEBT",
        "khong_tru_ton_so": "IGNORE_BOOK_ON_DEBT",
        "không_trừ_tồn_sổ": "IGNORE_BOOK_ON_DEBT",
    }
    return mapping.get(text, str(value or "").strip() or None)


def _normalize_profile(value):
    text = _norm(value).replace("-", "_").replace(" ", "_")
    if not text:
        return None
    mapping = {
        "continuous": "CONTINUOUS",
        "lien_tuc": "CONTINUOUS",
        "liên_tục": "CONTINUOUS",
        "spread_non_sunday": "SPREAD_NON_SUNDAY",
        "rai_khong_chu_nhat": "SPREAD_NON_SUNDAY",
        "rải_không_chủ_nhật": "SPREAD_NON_SUNDAY",
    }
    return mapping.get(text, str(value or "").strip() or None)


def read_actual_stock(source_bytes):
    workbook = load_workbook(
        BytesIO(source_bytes),
        data_only=True,
        read_only=True,
    )
    try:
        worksheet = workbook.worksheets[0]
        print(f"[Tồn thực tế] Sheet nguồn: {worksheet.title}")

        result = {}
        for row in range(1, worksheet.max_row + 1):
            code = normalize_code(
                worksheet.cell(row=row, column=3).value
            )
            if not code:
                continue
            if code in result:
                raise RuntimeError(
                    f"[Tồn thực tế] Mã {code} bị lặp trong cột C."
                )

            value_n = to_number(
                worksheet.cell(row=row, column=14).value,
                f"N{row}",
            )
            value_o = to_number(
                worksheet.cell(row=row, column=15).value,
                f"O{row}",
            )
            result[code] = clean_number(value_n + value_o)

        if not result:
            raise RuntimeError(
                "[Tồn thực tế] Không đọc được mã từ cột C."
            )

        print(
            f"[Tồn thực tế] Đọc {len(result)} mã; "
            "Ton_kho!D = N + O."
        )
        return result
    finally:
        safe_close_workbook(workbook)


def read_single_value_source(
    source_bytes,
    *,
    label,
    source_name,
    sheet_name,
    code_column,
    value_column,
    value_column_letter,
    vkd_to_vikoda=False,
):
    workbook = load_workbook(
        BytesIO(source_bytes),
        data_only=True,
        read_only=True,
        keep_vba=True,
    )
    try:
        if sheet_name not in workbook.sheetnames:
            raise RuntimeError(
                f"[{label}] Không tìm thấy {sheet_name} "
                f"trong {source_name}."
            )

        worksheet = workbook[sheet_name]
        result = {}
        for row in range(1, worksheet.max_row + 1):
            raw_code = worksheet.cell(
                row=row,
                column=code_column,
            ).value
            code = normalize_code(
                raw_code,
                vkd_to_vikoda=vkd_to_vikoda,
            )
            if not code:
                continue
            if code in result:
                raise RuntimeError(
                    f"[{label}] Mã {code} bị lặp sau chuẩn hóa."
                )

            value = to_number(
                worksheet.cell(
                    row=row,
                    column=value_column,
                ).value,
                f"{value_column_letter}{row}",
            )
            result[code] = clean_number(value)

        if not result:
            raise RuntimeError(
                f"[{label}] Không đọc được mã sản phẩm."
            )

        mode = (
            " (chuẩn hóa 2xxxxxxxx → 1xxxxxxxx)"
            if vkd_to_vikoda
            else ""
        )
        print(
            f"[{label}] Đọc {len(result)} mã từ "
            f"{source_name}!{sheet_name}{mode}."
        )
        return result
    finally:
        safe_close_workbook(workbook)


def read_conversion_factors(dest_bytes):
    """Read Danh_muc robustly even when worksheet dimension metadata is stale."""
    workbook = load_workbook(
        BytesIO(dest_bytes),
        data_only=True,
        read_only=True,
    )
    try:
        if MASTER_SHEET not in workbook.sheetnames:
            raise RuntimeError(
                f"Không tìm thấy sheet {MASTER_SHEET!r} trong file đích."
            )

        worksheet = workbook[MASTER_SHEET]
        headers = [
            str(value or "").strip()
            for value in next(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                ),
                (),
            )
        ]
        debt_col = _find_header_col(headers, DEBT_HEADER_NAMES)
        profile_col = _find_header_col(headers, PROFILE_HEADER_NAMES)

        factors = {}
        master_meta = {}

        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if not values:
                continue

            code = normalize_code(values[0])
            if not code:
                continue
            if code in factors:
                raise RuntimeError(
                    f"[{MASTER_SHEET}] Mã {code} bị lặp trong cột A."
                )

            factor = to_number(
                values[8] if len(values) >= 9 else None,
                f"{MASTER_SHEET}!I{row_number}",
            )
            if factor <= 0:
                raise RuntimeError(
                    f"[{MASTER_SHEET}] Quy cách của mã {code} "
                    f"phải > 0, hiện là {factor!r}."
                )

            factors[code] = factor
            raw_debt_mode = (
                values[debt_col]
                if debt_col is not None and debt_col < len(values)
                else None
            )
            raw_profile = (
                values[profile_col]
                if profile_col is not None and profile_col < len(values)
                else None
            )
            master_meta[code] = {
                "batch": to_number(
                    values[3] if len(values) >= 4 else None,
                    default=0.0,
                ),
                "per_shift": to_number(
                    values[4] if len(values) >= 5 else None,
                    default=0.0,
                ),
                "line": (
                    str(values[5] or "").strip()
                    if len(values) >= 6
                    else ""
                ),
                "group": (
                    str(values[6] or "").strip()
                    if len(values) >= 7
                    else ""
                ),
                "classification": (
                    str(values[7] or "").strip()
                    if len(values) >= 8
                    else ""
                ),
                "mold": factor,
                "leadtime": to_number(
                    values[9] if len(values) >= 10 else None,
                    default=0.0,
                ),
                "debt_mode": _normalize_debt_mode(raw_debt_mode) or "",
                "profile": _normalize_profile(raw_profile) or "",
            }

        if not factors:
            raise RuntimeError(
                f"[{MASTER_SHEET}] Không đọc được mã/quy cách từ A:I."
            )

        payload = json.dumps(
            {
                key: master_meta[key]
                for key in sorted(master_meta)
            },
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        conversion_hash = hashlib.sha256(payload).hexdigest()

        print(
            f"[{MASTER_SHEET}] Đọc {len(factors)} quy cách "
            "bằng chế độ tương thích sheet không có dimension."
        )
        return factors, conversion_hash
    finally:
        safe_close_workbook(workbook)
