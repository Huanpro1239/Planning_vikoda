import math
import zipfile
from datetime import datetime
from io import BytesIO

from lxml import etree
from openpyxl import load_workbook

import sync_planning_metrics as metrics_core
import sync_planning_schedule_priority as priority
import sync_planning_schedule_priority_v3 as v3
import sync_planning_schedule_priority_v4 as v4
from sync_planning_calendar_all_months import resolve_plan_year
from sync_planning_fc import _find_sheet_xml_path, _load_shared_strings, _read_cell_text
from sync_stock import normalize_code


_ORIGINAL_READ_INPUTS = priority.base.read_schedule_inputs
_ORIGINAL_PATCH_SCHEDULE = priority.base.patch_schedule_workbook


def _resolve_schedule_year(workbook_bytes):
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        worksheet = workbook[priority.base.PLANNING_SHEET]
        raw = str(
            worksheet.cell(row=1, column=priority.base.START_COLUMN_NUMBER).value or ""
        ).strip()
        if not raw:
            return datetime.now(priority.base.TIMEZONE).year
        first_line = raw.splitlines()[0].strip()
        _, month_text = first_line.split("/", 1)
        return resolve_plan_year(int(month_text))
    finally:
        workbook.close()


def _preferred_index(headers, preferred_date):
    if preferred_date is None or preferred_date <= headers[0]:
        return 0
    for index, current_day in enumerate(headers):
        if current_day >= preferred_date:
            return index
    return len(headers)


def _effective_release_index(headers, product, preferred_date):
    """
    R là ngày cần/start-by danh nghĩa, không phải rào cấm chạy trước.
    Pull-forward chỉ khi tồn/rủi ro cho thấy workload cần được mở sớm hơn.
    """
    preferred = _preferred_index(headers, preferred_date)
    if product["planned_qty"] <= priority.base.EPSILON:
        return preferred

    empty_schedule = priority.base._empty_schedule(headers, [product])
    first_stockout, first_safety = v3._dynamic_risk(headers, product, empty_schedule)

    if float(product.get("debt", 0) or 0) > 0:
        return 0

    risk_index = first_stockout if first_stockout is not None else first_safety
    if risk_index is None:
        return preferred

    required_shifts = product["planned_qty"] / product["per_shift"]
    daily_capacity = max(
        float(product.get("max_shifts_per_day", 0) or 0),
        priority.base.EPSILON,
    )
    workload_days = required_shifts / daily_capacity
    setup_days = priority.base.SETUP_SHIFTS / daily_capacity

    release = int(math.floor(risk_index - workload_days - setup_days))
    return min(preferred, max(0, release))


def read_schedule_inputs_v5(workbook_bytes, *, plan_year=None):
    if plan_year is None:
        plan_year = _resolve_schedule_year(workbook_bytes)

    headers, products = _ORIGINAL_READ_INPUTS(workbook_bytes, plan_year=plan_year)
    for product in products:
        preferred = product["earliest_date"]
        product["preferred_date"] = preferred
        release_index = _effective_release_index(headers, product, preferred)
        product["release_index"] = release_index
        if release_index < len(headers):
            product["earliest_date"] = headers[release_index]
    return headers, products


def _actual_first_date(headers, product, schedule):
    for current_day in headers:
        if schedule[product["code"]][current_day] > priority.base.EPSILON:
            return current_day
    return None


def _same_day(left, right):
    if left is None and right is None:
        return True
    if left is None or right is None:
        return False
    if isinstance(left, datetime):
        left = left.date()
    if isinstance(right, datetime):
        right = right.date()
    return left == right


def _patch_actual_start_dates(workbook_bytes, headers, products, schedule):
    targets = {}
    for product in products:
        if product["planned_qty"] <= priority.base.EPSILON:
            continue
        actual = _actual_first_date(headers, product, schedule)
        if actual is None:
            continue
        if not _same_day(product.get("preferred_date"), actual):
            targets[product["code"]] = datetime(actual.year, actual.month, actual.day)

    if not targets:
        return workbook_bytes, 0

    source_buffer = BytesIO(workbook_bytes)
    output_buffer = BytesIO()
    with zipfile.ZipFile(source_buffer, "r") as source_zip:
        sheet_path = _find_sheet_xml_path(source_zip, priority.base.PLANNING_SHEET)
        shared_strings = _load_shared_strings(source_zip)
        root = etree.fromstring(source_zip.read(sheet_path))
        seen = set()

        rows = root.xpath('//*[local-name()="sheetData"]/*[local-name()="row"]')
        for row in rows:
            row_number_text = row.get("r")
            if not row_number_text:
                continue
            row_number = int(row_number_text)
            code = None
            for cell in row.xpath('./*[local-name()="c"]'):
                if cell.get("r") == f"A{row_number}":
                    code = normalize_code(_read_cell_text(cell, shared_strings))
                    break
            if code not in targets:
                continue
            metrics_core._set_cell_value(row, row_number, "R", targets[code])
            seen.add(code)

        missing = sorted(set(targets) - seen)
        if missing:
            raise RuntimeError(
                "Không ghi được R thực tế cho: " + ", ".join(missing[:10])
            )

        new_xml = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )
        with zipfile.ZipFile(output_buffer, "w", compression=zipfile.ZIP_DEFLATED) as output_zip:
            for item in source_zip.infolist():
                data = new_xml if item.filename == sheet_path else source_zip.read(item.filename)
                output_zip.writestr(item, data)

    return output_buffer.getvalue(), len(targets)


def patch_schedule_workbook_v5(workbook_bytes, headers, products, schedule):
    updated, daily_changed = _ORIGINAL_PATCH_SCHEDULE(
        workbook_bytes,
        headers,
        products,
        schedule,
    )
    updated, r_changed = _patch_actual_start_dates(
        updated,
        headers,
        products,
        schedule,
    )
    return updated, daily_changed + r_changed


def install_priority_scheduler_v5():
    v4.install_priority_scheduler_v4()
    priority.base.read_schedule_inputs = read_schedule_inputs_v5
    priority.base.patch_schedule_workbook = patch_schedule_workbook_v5


if __name__ == "__main__":
    install_priority_scheduler_v5()
    priority.base.main_with_retry()
