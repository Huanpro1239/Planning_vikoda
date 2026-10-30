# -*- coding: utf-8 -*-
"""Tổng hợp Sell In theo tháng x sản phẩm (logic thuần, không phụ thuộc SharePoint).

Quy tắc:
  - Công ty đọc từ ô A1 (Tên DN): có "VKD" -> VKD, có "KHÁNH HÒA" -> Vikoda.
  - Kỳ đọc từ dòng "Từ ngày : dd-mm-yyyy".
  - VKD   : chỉ lấy MaSanPhamMoi bắt đầu bằng "2".
  - Vikoda: chỉ lấy MaSanPhamMoi bắt đầu bằng "1" và MaKhachHangMoi khác "VKD3".
  - SanLuongQuyDoi = tổng SoLuong / Quy cách (Danh mục của đúng công ty).
  - Pivot gộp Vikoda + VKD theo phần đuôi mã (1xxxxxxxx = 2xxxxxxxx).
"""
import datetime as dt
import io
import re
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

DATA_COLS = ["CongTy", "Nam", "Thang", "KyBaoCao", "MaSanPham", "TenSanPham", "TenVietTat",
             "Brand", "SKUsGroup", "DVTQuyDoi", "QuyCach", "SoLuong", "SanLuongQuyDoi"]
LOG_COLS = ["SourceFile", "CongTy", "Nam", "Thang", "TongDongNguon", "SoDongLay",
            "SoSanPham", "MaThieuQuyCach", "CapNhatLuc"]
# Pivot gộp 2 công ty: mã Vikoda (1xxxxxxxx) và mã VKD (2xxxxxxxx) cùng phần đuôi là 1 sản phẩm
PIVOT_FIXED = ["MaVikoda", "MaVKD", "TenSanPham", "TenVietTat", "Brand", "SKUsGroup", "DVTQuyDoi", "QuyCach"]


def gio_vn():
    return (dt.datetime.now(dt.timezone.utc) + dt.timedelta(hours=7)).replace(tzinfo=None, microsecond=0)


def norm(s):
    s = unicodedata.normalize("NFD", "" if s is None else str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").replace("đ", "d").replace("Đ", "D")
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def ma_sp(x):
    if x is None:
        return ""
    if isinstance(x, float) and x.is_integer():
        x = int(x)
    return str(x).strip()


def to_num(x):
    if isinstance(x, (int, float)):
        return x
    try:
        return float(str(x).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------- Danh mục
def doc_danh_muc(src):
    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb["DanhMucSanPham"] if "DanhMucSanPham" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    head = [norm(h) for h in next(rows)]

    def c(name):
        if name not in head:
            raise ValueError("Danh mục thiếu cột " + name)
        return head.index(name)

    i_vik, i_vkd, i_ten, i_tvt = c("MASANPHAMMOIVIKODA"), c("MASANPHAMMOIVKD"), c("TENSANPHAM"), c("TENVIETTAT")
    i_qc, i_dvt, i_brand, i_sku = c("QUYCACH"), c("DVT"), c("BRAND"), c("SKUSGROUP")
    dm = {"Vikoda": {}, "VKD": {}}
    for r in rows:
        qc = to_num(r[i_qc])
        info = {
            "TenSanPham": (r[i_ten] or ""), "TenVietTat": (r[i_tvt] or ""), "Brand": (r[i_brand] or ""),
            "SKUsGroup": (r[i_sku] or ""), "DVTQuyDoi": (r[i_dvt] or ""),
            "QuyCach": qc if qc and qc > 0 else None,
        }
        ma_vik, ma_vkd = ma_sp(r[i_vik]), ma_sp(r[i_vkd])
        # Cột mã Vikoda có thể là công thức ="1"&đuôi mã VKD chưa có giá trị cache -> tự suy ra
        if not ma_vik and ma_vkd.startswith("2"):
            ma_vik = "1" + ma_vkd[1:]
        if ma_vik:
            dm["Vikoda"][ma_vik] = info
        if ma_vkd:
            dm["VKD"][ma_vkd] = info
    wb.close()
    return dm


# ---------------------------------------------------------------- Đọc 1 file Sell In
def tong_hop_file(src, dm, source_name):
    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    top, hdr = [], None
    for _ in range(30):
        try:
            row = next(it)
        except StopIteration:
            break
        top.append(row)
        if any(norm(v) == "MASANPHAMMOI" for v in row):
            hdr = row
            break
    if hdr is None:
        wb.close()
        raise ValueError("Không tìm thấy dòng tiêu đề MaSanPhamMoi trong " + source_name)
    h = [norm(v) for v in hdr]
    try:
        i_kh, i_ma, i_sl = h.index("MAKHACHHANGMOI"), h.index("MASANPHAMMOI"), h.index("SOLUONG")
    except ValueError:
        wb.close()
        raise ValueError("Thiếu cột MaKhachHangMoi/MaSanPhamMoi/SoLuong trong " + source_name)

    ten_dn = norm(top[0][0] if top and top[0] else "")
    if "VKD" in ten_dn:
        cty = "VKD"
    elif "KHANHHOA" in ten_dn:
        cty = "Vikoda"
    else:
        cty = "VKD" if re.search(r"VKD\s*\.xls[mxb]?$", source_name, re.I) else "Vikoda"
    dau = "2" if cty == "VKD" else "1"

    nam = thang = None
    for row in top:
        s = " ".join("" if v is None else str(v) for v in row)
        m = re.search(r"T[ừu]\s*ng[àa]y\s*:\s*(\d{1,2})[-/](\d{1,2})[-/](\d{4})", s, re.I)
        if m:
            thang, nam = int(m.group(2)), int(m.group(3))
            break
    if nam is None:
        m = re.search(r"(\d{1,2})[ .\-_]+(\d{4})", source_name)
        if not m:
            wb.close()
            raise ValueError("Không xác định được tháng/năm của " + source_name)
        thang, nam = int(m.group(1)), int(m.group(2))

    tong, n_nguon, n_lay = {}, 0, 0
    need = max(i_kh, i_ma, i_sl)
    for row in it:
        if row is None or len(row) <= need:
            continue
        ma, kh = ma_sp(row[i_ma]), row[i_kh]
        if ma == "" and (kh in (None, "")):
            continue
        n_nguon += 1
        if not ma.startswith(dau):
            continue
        if cty == "Vikoda" and norm(kh) == "VKD3":
            continue
        sl = to_num(row[i_sl])
        if sl is None:
            continue
        tong[ma] = tong.get(ma, 0) + sl
        n_lay += 1
    wb.close()

    ky = dt.datetime(nam, thang, 1)
    rows, thieu = [], []
    for ma in sorted(tong):
        info = dm[cty].get(ma)
        qc = info["QuyCach"] if info else None
        if not qc:
            thieu.append(ma)
        rows.append([cty, nam, thang, ky, ma,
                     info["TenSanPham"] if info else "", info["TenVietTat"] if info else "",
                     info["Brand"] if info else "", info["SKUsGroup"] if info else "",
                     info["DVTQuyDoi"] if info else "", qc, tong[ma], tong[ma] / qc if qc else None])
    log = [source_name, cty, nam, thang, n_nguon, n_lay, len(rows), ", ".join(thieu),
           gio_vn()]
    return rows, log


# ---------------------------------------------------------------- File đầu ra
def _cung_ky(r, cty, nam, thang):
    return r[0] == cty and int(r[1] or 0) == nam and int(r[2] or 0) == thang


def build_workbook(data, log, fingerprint=""):
    """Trả về (bytes .xlsx, số sản phẩm Pivot, số tháng). fingerprint lưu ở sheet ẩn Meta."""
    thu_tu = {"Vikoda": 0, "VKD": 1}
    data = sorted(data, key=lambda r: (int(r[1]), int(r[2]), thu_tu.get(r[0], 9), str(r[4])))
    log = sorted(log, key=lambda r: (int(r[2]), int(r[3]), thu_tu.get(r[1], 9)))

    # Pivot: sản phẩm (gộp Vikoda + VKD theo phần đuôi mã) x tháng
    ky_list = sorted({(int(r[1]), int(r[2])) for r in data})
    sp, nhan = {}, {}
    for r in data:
        ma = str(r[4])
        key = ma[1:]
        if key not in nhan or (not nhan[key][2] and r[5]):
            nhan[key] = ["1" + key, "2" + key, r[5], r[6], r[7], r[8], r[9], r[10]]
        d = sp.setdefault(key, {})
        if r[12] is not None:
            k = (int(r[1]), int(r[2]))
            d[k] = d.get(k, 0) + r[12]
    keys = sorted(nhan, key=lambda k: (str(nhan[k][4]), str(nhan[k][5]), str(nhan[k][3]), k))
    p_head = PIVOT_FIXED + ["T%02d/%d" % (t, n) for n, t in ky_list] + ["Tong"]
    p_rows = []
    for k in keys:
        vals = [sp[k].get(x) for x in ky_list]
        co = [v for v in vals if v is not None]
        p_rows.append(nhan[k] + [round(v, 2) if v is not None else None for v in vals] +
                      [round(sum(co), 2) if co else None])

    wb = Workbook()
    ws = wb.active
    ws.title = "Pivot"
    _ghi_bang(ws, "tblPivot", p_head, p_rows, "TableStyleMedium2")
    for j in range(len(PIVOT_FIXED) + 1, len(p_head) + 1):
        for i in range(2, len(p_rows) + 2):
            ws.cell(i, j).number_format = "#,##0.00"
        ws.column_dimensions[get_column_letter(j)].width = 11
    for j, w in enumerate([11, 11, 44, 28, 11, 20, 9, 8], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "E2"

    wd = wb.create_sheet("Data")
    _ghi_bang(wd, "tblData", DATA_COLS, data, "TableStyleLight9")
    for i in range(2, len(data) + 2):
        wd.cell(i, 4).number_format = "mm/yyyy"
        wd.cell(i, 12).number_format = "#,##0"
        wd.cell(i, 13).number_format = "#,##0.00"
    for j, w in enumerate([8, 6, 6, 9, 11, 44, 28, 11, 20, 9, 8, 12, 14], start=1):
        wd.column_dimensions[get_column_letter(j)].width = w
    wd.freeze_panes = "A2"

    wl = wb.create_sheet("Log")
    _ghi_bang(wl, "tblLog", LOG_COLS, log, "TableStyleLight9")
    for i in range(2, len(log) + 2):
        wl.cell(i, 9).number_format = "dd/mm/yyyy hh:mm"
    for j, w in enumerate([30, 8, 6, 6, 14, 11, 10, 22, 17], start=1):
        wl.column_dimensions[get_column_letter(j)].width = w

    wm = wb.create_sheet("Meta")
    wm["A1"], wm["B1"] = "fingerprint", fingerprint
    wm.sheet_state = "hidden"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(p_rows), len(ky_list)


def doc_fingerprint(src):
    """Đọc fingerprint đã lưu trong sheet Meta của file tổng ("" nếu chưa có)."""
    wb = load_workbook(src, read_only=True, data_only=True)
    try:
        if "Meta" not in wb.sheetnames:
            return ""
        for row in wb["Meta"].iter_rows(values_only=True):
            if row and row[0] == "fingerprint":
                return str(row[1] or "")
        return ""
    finally:
        wb.close()


def _ghi_bang(ws, name, head, rows, style):
    ws.append(head)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ref = "A1:%s%d" % (get_column_letter(len(head)), max(len(rows), 1) + 1)
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)


def tong_hop_nhieu_file(sources, dm):
    """sources: list (tên file, file-like). Trả về (data, log); tháng trùng thì file sau thay file trước."""
    data, log = [], []
    for name, src in sources:
        rows, log_row = tong_hop_file(src, dm, name)
        cty, nam, thang = log_row[1], log_row[2], log_row[3]
        data = [r for r in data if not _cung_ky(r, cty, nam, thang)] + rows
        log = [r for r in log if not (r[1] == cty and r[2] == nam and r[3] == thang)] + [log_row]
    return data, log
