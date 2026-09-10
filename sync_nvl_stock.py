"""Module đồng bộ số lượng tồn kho nguyên vật liệu (NVL).

Nguồn: XNT_ketoan_Vikoda.xlsm, Sheet1
  - Cột B: Mã vật tư
  - Cột M: Số tồn kho

Đích: Kế hoạch mua hàng.xlsx, sheet Ton_NVL
  - Cột A: Mã vật tư
  - Cột D: Số tồn kho

Quy tắc:
  - Ghép mã vật tư Ton_NVL!A với Sheet1!B (chuẩn hóa trim khoảng trắng, tương đương float/int, giữ số 0 ở đầu).
  - Gán Ton_NVL!D = Sheet1!M: chép trực tiếp số tồn, không chia quy cách, không trừ tồn nhà máy, không đổi tiền tố mã 2 thành 1.
  - Patch trực tiếp XML của sheet Ton_NVL trong file ZIP, bảo toàn 100% macro, định dạng, merge, drawing và các sheet khác.
  - Kiểm tra an toàn trước khi patch: từ chối ô merge ở cột D, từ chối sheet bị khóa (sheetProtection), từ chối ô có công thức.
  - Kiểm soát tương tranh và toàn vẹn:
      + Kiểm tra revision nguồn trước download, sau download và trước upload để tránh dữ liệu nguồn cũ.
      + Sử dụng If-Match ETag file đích để chống ghi đè tương tranh trên SharePoint.
      + Tải lại file đích sau upload để xác minh kết quả ghi nhận trên server.
      + Phân tách rõ ràng trạng thái đối soát (proposal) và trạng thái publish, ghi nhận báo cáo lỗi chi tiết khi thất bại.
  - Lưu ý kiến trúc: Do hai file nằm độc lập trong Microsoft Graph, không tồn tại giao dịch nguyên tử phân tán (2-phase commit).
    Module áp dụng cơ chế snapshot validation, retry có phân loại, và verification sau upload để đạt độ nhất quán cao nhất.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass, field
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import re
import socket
import sys
import time
from typing import Any
import zipfile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from graph_retry import install_retry_after_support, retry_delay_seconds
from sync_stock import (
    GraphClient,
    GraphRequestError,
    column_number,
    find_sheet_xml_path,
    get_access_token,
    is_retryable_graph_error,
    set_numeric_cell,
)

DEFAULT_CONFIG_FILE = "nvl_stock_config.json"
EPS = 1e-6


@dataclass
class NVLConfig:
    source_name: str
    source_path: str
    source_sheet: str
    source_code_col: int
    source_code_col_letter: str
    source_value_col: int
    source_value_col_letter: str
    source_start_row: int
    source_sourcedoc: str

    target_name: str
    target_path: str
    target_sheet: str
    target_code_col: int
    target_code_col_letter: str
    target_value_col: int
    target_value_col_letter: str
    target_start_row: int
    target_sourcedoc: str

    direct_copy: bool = True
    allow_zero: bool = True
    allow_negative: bool = True
    preserve_missing_in_source: bool = True
    reject_duplicates: bool = True
    reject_formula_in_target_cell: bool = True
    number_convention: str = "strict"


@dataclass
class NVLReconcileResult:
    changes: list[dict[str, Any]] = field(default_factory=list)
    unchanged: list[dict[str, Any]] = field(default_factory=list)
    missing_in_source: list[dict[str, Any]] = field(default_factory=list)
    source_only: list[str] = field(default_factory=list)
    target_codes: dict[str, int] = field(default_factory=dict)
    source_codes: dict[str, int] = field(default_factory=dict)
    status: str = "success"
    message: str = ""


def load_nvl_config(config_path: str | Path | None = None) -> NVLConfig:
    """Đọc cấu hình đồng bộ NVL từ file JSON."""
    if config_path is None:
        config_path = Path(__file__).parent / DEFAULT_CONFIG_FILE
    else:
        config_path = Path(config_path)

    if not config_path.exists():
        raise FileNotFoundError(f"Không tìm thấy file cấu hình: {config_path}")

    with open(config_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    src = data.get("source", {})
    tgt = data.get("target", {})
    pol = data.get("policies", {})

    return NVLConfig(
        source_name=src.get("name", "XNT_ketoan_Vikoda.xlsm"),
        source_path=src.get("sharepoint_path", ""),
        source_sheet=src.get("sheet_name", "Sheet1"),
        source_code_col=int(src.get("code_column", 2)),
        source_code_col_letter=src.get("code_column_letter", "B"),
        source_value_col=int(src.get("value_column", 13)),
        source_value_col_letter=src.get("value_column_letter", "M"),
        source_start_row=int(src.get("start_row", 2)),
        source_sourcedoc=src.get("sourcedoc", ""),
        target_name=tgt.get("name", "Kế hoạch mua hàng.xlsx"),
        target_path=tgt.get("sharepoint_path", ""),
        target_sheet=tgt.get("sheet_name", "Ton_NVL"),
        target_code_col=int(tgt.get("code_column", 1)),
        target_code_col_letter=tgt.get("code_column_letter", "A"),
        target_value_col=int(tgt.get("value_column", 4)),
        target_value_col_letter=tgt.get("value_column_letter", "D"),
        target_start_row=int(tgt.get("start_row", 2)),
        target_sourcedoc=tgt.get("sourcedoc", ""),
        direct_copy=bool(pol.get("direct_copy", True)),
        allow_zero=bool(pol.get("allow_zero", True)),
        allow_negative=bool(pol.get("allow_negative", True)),
        preserve_missing_in_source=bool(pol.get("preserve_missing_in_source", True)),
        reject_duplicates=bool(pol.get("reject_duplicates", True)),
        reject_formula_in_target_cell=bool(pol.get("reject_formula_in_target_cell", True)),
        number_convention=str(pol.get("number_convention", "strict")).strip().lower(),
    )


def normalize_nvl_code(value: Any) -> str | None:
    """Chuẩn hóa mã vật tư NVL:

    - None, chuỗi rỗng hoặc boolean -> None
    - float với phần thập phân nguyên (ví dụ 12345.0) -> '12345'
    - int (ví dụ 12345) -> '12345'
    - str -> trim khoảng trắng đầu/cuối, giữ nguyên số 0 ở đầu (ví dụ '012345'),
      không chuyển đổi hoa/thường để bảo toàn đúng mã kế toán.
    """
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        if math.isnan(value) or math.isinf(value):
            return None
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        if isinstance(value, int):
            return str(value)
        return str(value).strip()

    text = str(value).strip()
    if not text:
        return None

    # Nếu chuỗi số float nguyên (ví dụ '12345.0') -> chuẩn hóa thành '12345'
    if re.match(r"^\d+\.0+$", text):
        return text.split(".")[0]

    return text


def parse_nvl_quantity(
    value: Any,
    convention: str = "strict",
    cell_name: str = "",
) -> float:
    """Phân tích và kiểm tra giá trị số lượng tồn kho NVL.

    Quy tắc:
    - Giá trị số (int, float) nguyên bản từ Excel được chép trực tiếp.
    - Cho phép số 0, số âm, số thập phân lẻ.
    - None, chuỗi rỗng, boolean, lỗi Excel, NaN, Infinity -> raise ValueError.
    - Chuỗi văn bản phải tuân thủ phân nhóm số nghiêm ngặt theo quy ước:
        + 'vi': Dấu chấm phân nhóm hàng nghìn (1.234.567), dấu phẩy thập phân (1,5; 1.234,56).
        + 'en': Dấu phẩy phân nhóm hàng nghìn (1,234,567), dấu chấm thập phân (1.5; 1,234.56).
        + 'strict': Nhận diện định dạng không mơ hồ (cả hai dấu phân cách hoặc số chữ số lẻ != 3).
                    Từ chối các chuỗi mơ hồ ('1,234' hoặc '1.234') và từ chối phân nhóm sai ('1,2,3').
    """
    prefix = f"[{cell_name}] " if cell_name else ""

    if value is None or (isinstance(value, str) and value.strip() == ""):
        raise ValueError(f"{prefix}Ô số lượng tồn kho rỗng.")

    if isinstance(value, bool):
        raise ValueError(f"{prefix}Chứa giá trị boolean ({value}), không phải số.")

    # 1. Số dạng numeric trong Excel được nhận trực tiếp
    if isinstance(value, (int, float)):
        fval = float(value)
        if math.isnan(fval) or math.isinf(fval):
            raise ValueError(f"{prefix}Chứa giá trị NaN hoặc vô cực.")
        return fval

    # 2. Xử lý chuỗi văn bản
    if not isinstance(value, str):
        raise ValueError(f"{prefix}Kiểu dữ liệu không được hỗ trợ: {type(value)}")

    cleaned = value.strip()
    if not cleaned:
        raise ValueError(f"{prefix}Chuỗi rỗng không phải số hợp lệ.")

    # Kiểm tra dấu âm/dương
    sign = 1.0
    if cleaned.startswith("-"):
        sign = -1.0
        cleaned = cleaned[1:].strip()
    elif cleaned.startswith("+"):
        cleaned = cleaned[1:].strip()

    if not cleaned:
        raise ValueError(f"{prefix}Chuỗi số không hợp lệ: {value!r}")

    # Chuỗi số nguyên thuần túy
    if cleaned.isdigit():
        return sign * float(cleaned)

    # Đếm dấu phân cách
    dot_count = cleaned.count(".")
    comma_count = cleaned.count(",")

    if dot_count == 0 and comma_count == 0:
        raise ValueError(f"{prefix}Giá trị {value!r} không thể chuyển đổi thành số tồn hợp lệ.")

    valid_chars = set("0123456789.,")
    if not set(cleaned).issubset(valid_chars):
        raise ValueError(f"{prefix}Chuỗi số chứa ký tự không hợp lệ: {value!r}")

    # Chặn phân cách nhóm liên tiếp: '..', ',,', '.,', ',.'
    if re.search(r"[.,]{2,}", cleaned):
        raise ValueError(f"{prefix}Phân cách nhóm số không hợp lệ: {value!r}")

    conv = (convention or "strict").strip().lower()

    # Trường hợp 1: Chứa cả dấu chấm và dấu phẩy
    if dot_count > 0 and comma_count > 0:
        is_en_pattern = bool(re.match(r"^\d{1,3}(,\d{3})+(\.\d+)$", cleaned))
        is_vi_pattern = bool(re.match(r"^\d{1,3}(\.\d{3})+(,\d+)$", cleaned))

        if is_en_pattern:
            if conv == "vi":
                raise ValueError(f"{prefix}Định dạng số kiểu Anh {value!r} không khớp với quy ước Việt Nam đã cấu hình.")
            return sign * float(cleaned.replace(",", ""))
        elif is_vi_pattern:
            if conv == "en":
                raise ValueError(f"{prefix}Định dạng số kiểu Việt Nam {value!r} không khớp với quy ước Anh đã cấu hình.")
            return sign * float(cleaned.replace(".", "").replace(",", "."))
        else:
            raise ValueError(f"{prefix}Phân nhóm dấu phân cách số không đúng quy cách: {value!r}")

    # Trường hợp 2: Chỉ chứa dấu phẩy
    if comma_count > 0:
        if conv == "en":
            # Trong chuẩn EN, dấu phẩy chỉ có thể là phân nhóm hàng nghìn (mỗi nhóm đúng 3 chữ số)
            if not re.match(r"^\d{1,3}(,\d{3})+$", cleaned):
                raise ValueError(f"{prefix}Phân nhóm dấu phẩy không đúng quy cách số tiếng Anh: {value!r}")
            return sign * float(cleaned.replace(",", ""))
        elif conv == "vi":
            if comma_count > 1:
                raise ValueError(f"{prefix}Dấu phẩy phân nhóm không hợp lệ trong quy ước Việt Nam: {value!r}")
            if not re.match(r"^\d+,\d+$", cleaned):
                raise ValueError(f"{prefix}Dấu phẩy thập phân không đúng định dạng: {value!r}")
            return sign * float(cleaned.replace(",", "."))
        else:  # strict / auto
            if comma_count > 1:
                raise ValueError(f"{prefix}Nhiều dấu phẩy không hợp lệ trong chế độ nghiêm ngặt: {value!r}")
            m = re.match(r"^(\d+),(\d+)$", cleaned)
            if not m:
                raise ValueError(f"{prefix}Định dạng số chứa dấu phẩy không hợp lệ: {value!r}")
            if len(m.group(2)) == 3:
                raise ValueError(
                    f"{prefix}Chuỗi số mơ hồ giữa phân nhóm hàng nghìn và số thập phân: {value!r}. "
                    "Vui lòng cấu hình number_convention ('vi' hoặc 'en')."
                )
            # Số chữ số sau dấu phẩy != 3 -> chắc chắn là dấu phẩy thập phân
            return sign * float(cleaned.replace(",", "."))

    # Trường hợp 3: Chỉ chứa dấu chấm
    if dot_count > 0:
        if conv == "vi":
            # Trong chuẩn VI, dấu chấm chỉ có thể là phân nhóm hàng nghìn
            if not re.match(r"^\d{1,3}(\.\d{3})+$", cleaned):
                raise ValueError(f"{prefix}Phân nhóm dấu chấm không đúng quy cách Việt Nam: {value!r}")
            return sign * float(cleaned.replace(".", ""))
        elif conv == "en":
            if dot_count > 1:
                raise ValueError(f"{prefix}Nhiều dấu chấm không hợp lệ trong quy ước tiếng Anh: {value!r}")
            if not re.match(r"^\d+\.\d+$", cleaned):
                raise ValueError(f"{prefix}Dấu chấm thập phân không đúng định dạng: {value!r}")
            return sign * float(cleaned)
        else:  # strict / auto
            if dot_count > 1:
                raise ValueError(f"{prefix}Nhiều dấu chấm không hợp lệ trong chế độ nghiêm ngặt: {value!r}")
            m = re.match(r"^(\d+)\.(\d+)$", cleaned)
            if not m:
                raise ValueError(f"{prefix}Định dạng số chứa dấu chấm không hợp lệ: {value!r}")
            if len(m.group(2)) == 3:
                raise ValueError(
                    f"{prefix}Chuỗi số mơ hồ giữa phân nhóm hàng nghìn và số thập phân: {value!r}. "
                    "Vui lòng cấu hình number_convention ('vi' hoặc 'en')."
                )
            # Số chữ số sau dấu chấm != 3 -> chắc chắn là dấu chấm thập phân
            return sign * float(cleaned)

    raise ValueError(f"{prefix}Không thể xử lý giá trị số: {value!r}")


def _is_finite_number(val: Any) -> bool:
    """Kiểm tra một giá trị có phải là số hữu hạn (không phải bool, nan, inf)."""
    if val is None or isinstance(val, bool):
        return False
    if isinstance(val, (int, float)):
        return not (math.isnan(val) or math.isinf(val))
    if isinstance(val, str):
        try:
            f = float(val)
            return not (math.isnan(f) or math.isinf(f))
        except ValueError:
            return False
    return False


def _cells_match(actual: Any, expected: Any, eps: float = EPS) -> bool:
    """So sánh hai giá trị ô mà không bắt buộc ô phi số phải chuyển thành float."""
    # Cả hai là None hoặc chuỗi rỗng
    is_act_empty = actual is None or (isinstance(actual, str) and actual.strip() == "")
    is_exp_empty = expected is None or (isinstance(expected, str) and expected.strip() == "")
    if is_act_empty and is_exp_empty:
        return True
    if is_act_empty != is_exp_empty:
        return False

    # Cả hai là số hữu hạn: so sánh có dung sai EPS
    if _is_finite_number(actual) and _is_finite_number(expected):
        return abs(float(actual) - float(expected)) <= eps

    # So sánh chuỗi hoặc giá trị trực tiếp (ví dụ '-', 'N/A', ghi chú text)
    return str(actual).strip() == str(expected).strip()


def check_target_sheet_safety(
    target_bytes: bytes,
    sheet_name: str,
    target_col_letter: str,
    target_start_row: int,
) -> None:
    """Kiểm tra an toàn cấu trúc sheet đích trước khi patch:

    - Kiểm tra bảo vệ sheet (<sheetProtection>).
    - Kiểm tra gộp ô (<mergeCell>): không cho phép cột đích nằm trong vùng gộp ô.
    """
    target_col = column_number(target_col_letter)
    buf = BytesIO(target_bytes)
    z = zipfile.ZipFile(buf, "r")
    try:
        sheet_path = find_sheet_xml_path(z, sheet_name)
        sheet_root = etree.fromstring(z.read(sheet_path))

        # 1. Kiểm tra sheetProtection
        for prot in sheet_root.xpath('//*[local-name()="sheetProtection"]'):
            is_prot = any(
                prot.get(attr) in ("1", "true")
                for attr in ("sheet", "objects", "scenarios")
            )
            if is_prot:
                raise RuntimeError(
                    f"[Đích] Sheet '{sheet_name}' đang được bật bảo vệ (sheetProtection). "
                    "Không thể ghi đè dữ liệu tồn kho."
                )

        # 2. Kiểm tra mergeCell phủ lên cột đích
        for m in sheet_root.xpath('//*[local-name()="mergeCell"]'):
            ref = m.get("ref", "")
            if not ref:
                continue
            try:
                min_c, min_r, max_c, max_r = range_boundaries(ref)
                if min_c <= target_col <= max_c and max_r >= target_start_row:
                    raise RuntimeError(
                        f"[Đích] Ô trong vùng {sheet_name}!{target_col_letter}{max(min_r, target_start_row)} "
                        f"nằm trong dải ô gộp '{ref}'. Cột {target_col_letter} phải là cột đơn lẻ, "
                        "không được gộp ô."
                    )
            except Exception as exc:
                if "nằm trong dải ô gộp" in str(exc):
                    raise
    finally:
        z.close()
        z.fp = None


def read_nvl_source_stock(
    source_bytes: bytes,
    config: NVLConfig,
) -> tuple[dict[str, float], dict[str, Any]]:
    """Đọc dữ liệu số tồn kho từ file nguồn XNT_ketoan_Vikoda.xlsm!Sheet1.

    Sử dụng streaming reader qua iter_rows() sau khi reset_dimensions() để
    hoàn toàn miễn nhiễm với dimension bị thiếu hoặc khai báo hụt.
    """
    wb_val = load_workbook(BytesIO(source_bytes), data_only=True, read_only=True, keep_vba=True)
    wb_raw = load_workbook(BytesIO(source_bytes), data_only=False, read_only=True, keep_vba=True)

    try:
        if config.source_sheet not in wb_val.sheetnames:
            raise RuntimeError(
                f"[Nguồn: {config.source_name}] Không tìm thấy sheet '{config.source_sheet}'."
            )

        ws_val = wb_val[config.source_sheet]
        ws_raw = wb_raw[config.source_sheet]

        # Xóa dimension cache để đọc trọn vẹn dữ liệu từ sheet XML
        ws_val.reset_dimensions()
        ws_raw.reset_dimensions()

        reporting_period = None
        for r_hdr in range(1, min(config.source_start_row, 15)):
            c_val = ws_val.cell(r_hdr, 1).value
            if isinstance(c_val, str):
                c_str = c_val.strip()
                if any(k in c_str.lower() for k in ("từ ngày", "đến ngày", "kỳ", "tháng")):
                    reporting_period = c_str
                    break

        stock: dict[str, float] = {}
        seen_rows: dict[str, int] = {}
        total_scanned = 0

        val_stream = ws_val.iter_rows(min_row=config.source_start_row, values_only=True)
        raw_stream = ws_raw.iter_rows(min_row=config.source_start_row, values_only=True)

        code_idx = config.source_code_col - 1
        val_idx = config.source_value_col - 1

        for r, (val_row, raw_row) in enumerate(zip(val_stream, raw_stream), start=config.source_start_row):
            total_scanned += 1
            if not val_row or len(val_row) <= code_idx:
                continue

            raw_code = val_row[code_idx]
            if isinstance(raw_code, str):
                lower_code = raw_code.strip().lower()
                if any(k in lower_code for k in ("tổng cộng", "tong cong", "total", "cộng")):
                    continue

            code = normalize_nvl_code(raw_code)
            if not code:
                continue

            if config.reject_duplicates and code in stock:
                raise RuntimeError(
                    f"[Nguồn: {config.source_name}] Mã vật tư '{code}' bị lặp tại dòng {r} "
                    f"(đã xuất hiện tại dòng {seen_rows[code]})."
                )

            seen_rows[code] = r
            cell_ref = f"{config.source_sheet}!{config.source_value_col_letter}{r}"

            c_val = val_row[val_idx] if len(val_row) > val_idx else None
            c_raw = raw_row[val_idx] if len(raw_row) > val_idx else None

            if isinstance(c_raw, str) and c_raw.startswith("="):
                if c_val is None:
                    raise RuntimeError(
                        f"[Nguồn: {config.source_name}] Ô {cell_ref} chứa công thức '{c_raw}' "
                        f"nhưng không có giá trị cached (chưa được Excel tính toán và lưu)."
                    )

            qty = parse_nvl_quantity(
                c_val,
                convention=config.number_convention,
                cell_name=cell_ref,
            )

            if not config.allow_negative and qty < 0:
                raise ValueError(f"[Nguồn: {config.source_name}] Ô {cell_ref} có số lượng âm ({qty}) bị từ chối.")
            if not config.allow_zero and abs(qty) < EPS:
                raise ValueError(f"[Nguồn: {config.source_name}] Ô {cell_ref} có số lượng 0 bị từ chối.")

            stock[code] = qty

        if not stock:
            raise RuntimeError(
                f"[Nguồn: {config.source_name}] Không đọc được bất kỳ mã vật tư hợp lệ nào "
                f"trong {config.source_sheet}!{config.source_code_col_letter}."
            )

        metadata = {
            "source_name": config.source_name,
            "sheet_name": config.source_sheet,
            "total_scanned_rows": total_scanned,
            "valid_codes_count": len(stock),
            "seen_rows": seen_rows,
            "reporting_period": reporting_period,
        }
        return stock, metadata
    finally:
        try:
            val_stream.close()
        except Exception:
            pass
        try:
            raw_stream.close()
        except Exception:
            pass
        _safe_close_workbook(wb_val)
        _safe_close_workbook(wb_raw)


def _safe_close_workbook(wb: Any) -> None:
    """Đóng openpyxl workbook an toàn và dọn dẹp ZipFile archive / vba_archive trên Python 3.12."""
    if wb is None:
        return
    try:
        vba_archive = getattr(wb, "vba_archive", None)
        if vba_archive is not None:
            try:
                vba_archive.close()
            except Exception:
                pass
            vba_archive.fp = None
    except Exception:
        pass
    try:
        archive = getattr(wb, "_archive", None)
        if archive is not None:
            try:
                archive.close()
            except Exception:
                pass
            archive.fp = None
    except Exception:
        pass
    try:
        wb.close()
    except Exception:
        pass


def reconcile_nvl_target(
    target_bytes: bytes,
    source_stock: dict[str, float],
    config: NVLConfig,
) -> NVLReconcileResult:
    """Đối chiếu mã vật tư và lập danh sách thay đổi cho Kế hoạch mua hàng.xlsx!Ton_NVL.

    Sử dụng streaming reader với reset_dimensions() và kiểm tra cấu trúc an toàn
    (bảo vệ sheet, gộp ô, công thức người dùng).
    """
    # 1. Kiểm tra an toàn cấu trúc sheet đích
    check_target_sheet_safety(
        target_bytes,
        config.target_sheet,
        config.target_value_col_letter,
        config.target_start_row,
    )

    # 2. Đọc dữ liệu đích
    wb_val = load_workbook(BytesIO(target_bytes), data_only=True, read_only=True)
    wb_raw = load_workbook(BytesIO(target_bytes), data_only=False, read_only=True)

    try:
        if config.target_sheet not in wb_val.sheetnames:
            raise RuntimeError(
                f"[Đích: {config.target_name}] Không tìm thấy sheet '{config.target_sheet}'."
            )

        ws_val = wb_val[config.target_sheet]
        ws_raw = wb_raw[config.target_sheet]

        ws_val.reset_dimensions()
        ws_raw.reset_dimensions()

        seen_target_codes: dict[str, int] = {}
        changes: list[dict[str, Any]] = []
        unchanged: list[dict[str, Any]] = []
        missing_in_source: list[dict[str, Any]] = []

        val_stream = ws_val.iter_rows(min_row=config.target_start_row, values_only=True)
        raw_stream = ws_raw.iter_rows(min_row=config.target_start_row, values_only=True)

        code_idx = config.target_code_col - 1
        val_idx = config.target_value_col - 1

        for r, (val_row, raw_row) in enumerate(zip(val_stream, raw_stream), start=config.target_start_row):
            if not val_row or len(val_row) <= code_idx:
                continue

            raw_code = val_row[code_idx]
            if isinstance(raw_code, str):
                lower_code = raw_code.strip().lower()
                if any(k in lower_code for k in ("tổng cộng", "tong cong", "total", "cộng")):
                    continue

            code = normalize_nvl_code(raw_code)
            if not code:
                continue

            if config.reject_duplicates and code in seen_target_codes:
                raise RuntimeError(
                    f"[Đích: {config.target_name}] Mã vật tư '{code}' bị lặp trong "
                    f"{config.target_sheet}!{config.target_code_col_letter} tại dòng {r} "
                    f"(đã xuất hiện tại dòng {seen_target_codes[code]})."
                )

            seen_target_codes[code] = r
            cell_ref = f"{config.target_sheet}!{config.target_value_col_letter}{r}"

            target_val = val_row[val_idx] if len(val_row) > val_idx else None
            target_raw = raw_row[val_idx] if len(raw_row) > val_idx else None

            if config.reject_formula_in_target_cell and isinstance(target_raw, str) and target_raw.startswith("="):
                raise RuntimeError(
                    f"[Đích: {config.target_name}] Ô {cell_ref} chứa công thức '{target_raw}'. "
                    f"Cột {config.target_value_col_letter} phải là cột nhập tồn, không được ghi đè lên công thức người dùng."
                )

            if code not in source_stock:
                if not config.preserve_missing_in_source:
                    raise RuntimeError(
                        f"[Đích: {config.target_name}] Mã vật tư '{code}' không có trong file nguồn và "
                        f"cấu hình 'preserve_missing_in_source' đang tắt."
                    )
                missing_in_source.append({
                    "row": r,
                    "code": code,
                    "current_value": target_val,
                })
                continue

            new_qty = source_stock[code]
            needs_change = False

            if target_val is None or (isinstance(target_val, str) and target_val.strip() == ""):
                needs_change = True
            else:
                try:
                    curr_float = parse_nvl_quantity(
                        target_val,
                        convention=config.number_convention,
                        cell_name=cell_ref,
                    )
                    if abs(curr_float - new_qty) > EPS:
                        needs_change = True
                except Exception:
                    # Nếu ô đích hiện tại chứa text không parse được thành số hợp lệ -> cần cập nhật
                    needs_change = True

            if needs_change:
                changes.append({
                    "row": r,
                    "code": code,
                    "before": target_val,
                    "after": new_qty,
                })
            else:
                unchanged.append({
                    "row": r,
                    "code": code,
                    "value": target_val,
                })

        if not seen_target_codes:
            raise RuntimeError(
                f"[Đích: {config.target_name}] Không đọc được bất kỳ mã vật tư hợp lệ nào trong "
                f"{config.target_sheet}!{config.target_code_col_letter}."
            )

        matched_count = len(changes) + len(unchanged)
        if matched_count == 0:
            raise RuntimeError(
                f"[Đích: {config.target_name}] Không có mã vật tư nào trong "
                f"{config.target_sheet}!{config.target_code_col_letter} khớp với nguồn {config.source_name}. Bị chặn cập nhật."
            )

        source_only = [c for c in source_stock if c not in seen_target_codes]

        if missing_in_source:
            status = "completed_with_warnings"
            msg = f"Đối soát hoàn tất: {len(changes)} ô cần cập nhật, còn {len(missing_in_source)} mã đích không có trong nguồn."
        elif changes:
            status = "success"
            msg = f"Đối soát thành công: {len(changes)} ô cần cập nhật, {len(unchanged)} ô không đổi."
        else:
            status = "unchanged"
            msg = "Dữ liệu tồn kho khớp hoàn toàn, không có ô nào cần thay đổi."

        return NVLReconcileResult(
            changes=changes,
            unchanged=unchanged,
            missing_in_source=missing_in_source,
            source_only=source_only,
            target_codes=seen_target_codes,
            status=status,
            message=msg,
        )
    finally:
        try:
            val_stream.close()
        except Exception:
            pass
        try:
            raw_stream.close()
        except Exception:
            pass
        _safe_close_workbook(wb_val)
        _safe_close_workbook(wb_raw)


def patch_nvl_destination_workbook(
    target_bytes: bytes,
    reconcile_result: NVLReconcileResult,
    config: NVLConfig,
) -> bytes:
    """Cập nhật các ô số tồn kho vào file Excel đích ở cấp độ ZIP XML:

    - Chỉ giải nén và sửa file XML của sheet Ton_NVL.
    - Chèn thẻ <row> và <c> đúng vị trí tuần tự theo XML schema của OpenXML.
    - Cập nhật <dimension ref="..."> để bao phủ các ô mới tạo nếu vượt quá phạm vi cũ.
    - Bảo toàn 100% byte-for-byte các file ZIP parts còn lại.
    """
    if not reconcile_result.changes:
        return target_bytes

    source_buffer = BytesIO(target_bytes)
    output_buffer = BytesIO()

    source_zip = zipfile.ZipFile(source_buffer, "r")
    output_zip = zipfile.ZipFile(output_buffer, "w", compression=zipfile.ZIP_DEFLATED)
    try:
        sheet_path = find_sheet_xml_path(source_zip, config.target_sheet)
        sheet_root = etree.fromstring(source_zip.read(sheet_path))

        sheet_data_nodes = sheet_root.xpath('//*[local-name()="sheetData"]')
        if not sheet_data_nodes:
            raise RuntimeError(f"Không tìm thấy thẻ sheetData trong {sheet_path}.")
        sheet_data = sheet_data_nodes[0]

        row_map: dict[int, Any] = {}
        for r_elem in sheet_data.xpath('./*[local-name()="row"]'):
            r_num_text = r_elem.get("r")
            if r_num_text and r_num_text.isdigit():
                row_map[int(r_num_text)] = r_elem

        changes_by_row = {item["row"]: item["after"] for item in reconcile_result.changes}

        for row_num, new_val in changes_by_row.items():
            row_element = row_map.get(row_num)
            if row_element is None:
                namespace = etree.QName(sheet_data).namespace
                row_element = etree.Element(f"{{{namespace}}}row", r=str(row_num))
                inserted = False
                for existing_row in sheet_data.xpath('./*[local-name()="row"]'):
                    er = existing_row.get("r")
                    if er and int(er) > row_num:
                        existing_row.addprevious(row_element)
                        inserted = True
                        break
                if not inserted:
                    sheet_data.append(row_element)
                row_map[row_num] = row_element

            set_numeric_cell(
                row_element,
                row_num,
                config.target_value_col_letter,
                new_val,
            )

            spans_attr = row_element.get("spans")
            if spans_attr and ":" in spans_attr:
                try:
                    s_min, s_max = map(int, spans_attr.split(":"))
                    val_col_num = column_number(config.target_value_col_letter)
                    if val_col_num > s_max:
                        row_element.set("spans", f"{s_min}:{val_col_num}")
                    elif val_col_num < s_min:
                        row_element.set("spans", f"{val_col_num}:{s_max}")
                except Exception:
                    pass

        # Cập nhật <dimension ref="..."> nếu cần mở rộng phạm vi
        dim_nodes = sheet_root.xpath('//*[local-name()="dimension"]')
        if dim_nodes:
            dim_elem = dim_nodes[0]
            ref = dim_elem.get("ref", "")
            if ref and ":" in ref:
                try:
                    min_c, min_r, max_c, max_r = range_boundaries(ref)
                    val_col_num = column_number(config.target_value_col_letter)
                    max_c_new = max(max_c, val_col_num)
                    max_r_new = max([max_r] + list(changes_by_row.keys()))
                    if max_c_new != max_c or max_r_new != max_r:
                        dim_elem.set("ref", f"{get_column_letter(min_c)}{min_r}:{get_column_letter(max_c_new)}{max_r_new}")
                except Exception:
                    pass

        new_sheet_xml = etree.tostring(
            sheet_root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        for name in source_zip.namelist():
            data = new_sheet_xml if name == sheet_path else source_zip.read(name)
            output_zip.writestr(name, data)
    finally:
        output_zip.close()
        output_zip.fp = None
        source_zip.close()
        source_zip.fp = None

    return output_buffer.getvalue()


SHAREPOINT_METADATA_NAMESPACES = (
    "http://schemas.microsoft.com/office/2006/metadata/contentType",
    "http://schemas.microsoft.com/office/2006/metadata/properties",
    "http://schemas.microsoft.com/office/2006/metadata/properties/metaAttributes",
    "http://schemas.microsoft.com/sharepoint/v3/contenttype/forms",
    "http://schemas.microsoft.com/sharepoint/",
)


def identify_sharepoint_metadata_exemption(
    name: str,
    orig_bytes: bytes,
    patch_bytes: bytes,
    z_orig: zipfile.ZipFile,
    z_patch: zipfile.ZipFile,
) -> dict[str, str] | None:
    """Xác thực một phần tử trong file ZIP có phải là metadata SharePoint được máy chủ
    SharePoint tự động cập nhật khi upload hay không.

    Yêu cầu nhận diện chính xác theo:
    1. Cấu trúc XML và root element.
    2. Namespace chính thức của SharePoint metadata.
    3. Quan hệ liên kết trong file quan hệ (_rels/.rels, workbook.xml.rels hoặc customXml/_rels/*.rels).

    TUYỆT ĐỐI không bỏ qua toàn bộ customXml hay docProps:
    - Nếu phần tử không chứa cấu trúc/namespace SharePoint chuẩn -> trả về None (bị chặn).
    - Nếu phần tử nằm ngoài customXml và docProps (như xl/worksheets, xl/styles...) -> trả về None.
    """
    name_clean = name.strip("/").lower()

    # 1. Kiểm tra nếu là customXml/_rels/*.rels
    if name_clean.startswith("customxml/_rels/") and name_clean.endswith(".rels"):
        try:
            root_orig = etree.fromstring(orig_bytes)
            root_patch = etree.fromstring(patch_bytes)
            rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
            if root_orig.tag != f"{{{rel_ns}}}Relationships" or root_patch.tag != f"{{{rel_ns}}}Relationships":
                return None
            prop_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps"
            rels = root_orig.findall(f"{{{rel_ns}}}Relationship")
            if not rels or not all(r.get("Type") == prop_type for r in rels):
                return None
            return {
                "part": name,
                "type": "sharepoint_customxml_rels",
                "namespace": prop_type,
                "reason": "File quan hệ customXmlProps của SharePoint metadata",
            }
        except Exception:
            return None

    # 2. Kiểm tra nếu là customXml/item*.xml hoặc customXml/itemProps*.xml
    if name_clean.startswith("customxml/") and name_clean.endswith(".xml"):
        try:
            root_orig = etree.fromstring(orig_bytes)
            root_patch = etree.fromstring(patch_bytes)
        except Exception:
            return None

        # 2a. Trường hợp itemProps*.xml (ds:datastoreItem)
        ds_ns = "http://schemas.openxmlformats.org/officeDocument/2006/customXml"
        if root_orig.tag == f"{{{ds_ns}}}datastoreItem" and root_patch.tag == f"{{{ds_ns}}}datastoreItem":
            schema_refs = root_orig.findall(f".//{{{ds_ns}}}schemaRef")
            matched_ns = None
            for sref in schema_refs:
                uri = sref.get(f"{{{ds_ns}}}uri") or sref.get("uri") or ""
                if any(sp_ns in uri for sp_ns in SHAREPOINT_METADATA_NAMESPACES):
                    matched_ns = uri
                    break
            if matched_ns:
                return {
                    "part": name,
                    "type": "sharepoint_datastore_item",
                    "namespace": matched_ns,
                    "reason": f"SharePoint datastore item properties tham chiếu schema '{matched_ns}'",
                }
            return None

        # 2b. Trường hợp item*.xml (ct:contentTypeSchema, p:properties, FormTemplates...)
        orig_nsmap = root_orig.nsmap.values()
        patch_nsmap = root_patch.nsmap.values()
        root_tag_orig = root_orig.tag
        root_tag_patch = root_patch.tag

        matched_ns = None
        for sp_ns in SHAREPOINT_METADATA_NAMESPACES:
            if (
                sp_ns in root_tag_orig
                or sp_ns in root_tag_patch
                or any(sp_ns in str(v) for v in orig_nsmap)
                or any(sp_ns in str(v) for v in patch_nsmap)
            ):
                matched_ns = sp_ns
                break

        if matched_ns:
            # Kiểm tra quan hệ từ xl/_rels/workbook.xml.rels hoặc _rels/.rels
            has_valid_rel = False
            try:
                rels_checked = 0
                for rels_name in ("xl/_rels/workbook.xml.rels", "_rels/.rels"):
                    if rels_name in z_orig.namelist():
                        rels_checked += 1
                        rels_root = etree.fromstring(z_orig.read(rels_name))
                        rel_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"
                        for rel in rels_root.findall(".//{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"):
                            if rel.get("Type") == rel_type:
                                target = (rel.get("Target") or "").replace("../", "").strip("/").lower()
                                if target == name_clean or target.endswith(name_clean):
                                    has_valid_rel = True
                                    break
                    if has_valid_rel:
                        break
                if rels_checked == 0:
                    has_valid_rel = True
            except Exception:
                has_valid_rel = False

            if has_valid_rel:
                tag_short = root_tag_orig.split("}")[-1] if "}" in root_tag_orig else root_tag_orig
                return {
                    "part": name,
                    "type": "sharepoint_customxml_metadata",
                    "namespace": matched_ns,
                    "reason": f"SharePoint document contentType/DIP metadata ({tag_short})",
                }
        return None

    # 3. Kiểm tra nếu là docProps/core.xml
    if name_clean == "docprops/core.xml":
        try:
            root_orig = etree.fromstring(orig_bytes)
            root_patch = etree.fromstring(patch_bytes)
            core_ns = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
            if root_orig.tag == f"{{{core_ns}}}coreProperties" and root_patch.tag == f"{{{core_ns}}}coreProperties":
                return {
                    "part": name,
                    "type": "server_core_properties",
                    "namespace": core_ns,
                    "reason": "Dublin Core properties cập nhật bởi SharePoint/Office",
                }
        except Exception:
            return None

    return None


def verify_nvl_patched_workbook(
    original_bytes: bytes,
    patched_bytes: bytes,
    reconcile_result: NVLReconcileResult,
    config: NVLConfig,
    *,
    is_server_comparison: bool = False,
) -> dict[str, Any]:
    """Kiểm tra tính toàn vẹn của workbook sau khi patch XML:

    1. Tất cả các ô changed_cells đều nhận đúng giá trị số mới.
    2. Các ô unchanged và missing_in_source giữ nguyên giá trị ban đầu (hỗ trợ cả văn bản/blank/lỗi).
    3. Tất cả các ô không thuộc danh sách thay đổi trong chính sheet Ton_NVL giữ nguyên định dạng và nội dung XML.
    4. Tất cả các phần tử trong file ZIP ngoài sheet đích đều giống nhau từng byte (SHA-256 đối chiếu 1-1).
       Khi đối chiếu với file tải lại từ server SharePoint (is_server_comparison=True), chỉ miễn trừ các metadata
       được xác thực chính xác là do máy chủ SharePoint tự động đóng dấu (theo cấu trúc/namespace và quan hệ liên kết).
       Ghi nhận danh sách các phần tử được miễn trừ trong kết quả trả về.
    """
    if not reconcile_result.changes:
        return {"ok": True, "message": "Không có thay đổi cần xác minh.", "exempted_parts": []}

    exempted_parts: list[dict[str, str]] = []

    # 1. So sánh các ZIP parts không liên quan (byte-level SHA-256)
    orig_buf = BytesIO(original_bytes)
    patch_buf = BytesIO(patched_bytes)
    z_orig = zipfile.ZipFile(orig_buf, "r")
    z_patch = zipfile.ZipFile(patch_buf, "r")
    try:
        target_sheet_path = find_sheet_xml_path(z_orig, config.target_sheet)
        for name in z_orig.namelist():
            if name != target_sheet_path:
                if name not in z_patch.namelist():
                    raise RuntimeError(f"Phần tử '{name}' bị thiếu trong file sau khi ghi/tải lại!")
                orig_raw = z_orig.read(name)
                patch_raw = z_patch.read(name)
                orig_hash = hashlib.sha256(orig_raw).hexdigest()
                patch_hash = hashlib.sha256(patch_raw).hexdigest()
                if orig_hash != patch_hash:
                    if is_server_comparison:
                        exemption = identify_sharepoint_metadata_exemption(
                            name, orig_raw, patch_raw, z_orig, z_patch
                        )
                        if exemption is not None:
                            exempted_parts.append(exemption)
                            continue
                    raise RuntimeError(
                        f"Phần tử không liên quan '{name}' trong file ZIP bị thay đổi ngoài ý muốn!"
                    )

        # Kiểm tra không có phần tử lạ xuất hiện trong z_patch ngoài các metadata hợp lệ
        for name in z_patch.namelist():
            if name not in z_orig.namelist() and name != target_sheet_path:
                if is_server_comparison:
                    exemption = identify_sharepoint_metadata_exemption(
                        name, b"", z_patch.read(name), z_orig, z_patch
                    )
                    if exemption is not None:
                        exempted_parts.append(exemption)
                        continue
                raise RuntimeError(
                    f"Phần tử lạ ngoài ý muốn '{name}' xuất hiện trong file ZIP sau khi ghi!"
                )

        # Kiểm tra tính toàn vẹn của các ô không thay đổi trong chính sheet đích
        orig_sheet_root = etree.fromstring(z_orig.read(target_sheet_path))
        patch_sheet_root = etree.fromstring(z_patch.read(target_sheet_path))

        changed_refs = {f"{config.target_value_col_letter}{item['row']}" for item in reconcile_result.changes}
        orig_cells = {c.get("r"): etree.tostring(c) for c in orig_sheet_root.xpath('//*[local-name()="c"]') if c.get("r")}
        patch_cells = {c.get("r"): etree.tostring(c) for c in patch_sheet_root.xpath('//*[local-name()="c"]') if c.get("r")}

        for r_coord, orig_c_xml in orig_cells.items():
            if r_coord not in changed_refs:
                if r_coord not in patch_cells:
                    raise RuntimeError(f"Ô không liên quan '{r_coord}' trong sheet '{config.target_sheet}' bị mất sau khi patch!")
                if patch_cells[r_coord] != orig_c_xml:
                    raise RuntimeError(f"Ô không liên quan '{r_coord}' trong sheet '{config.target_sheet}' bị biến đổi cấu trúc XML!")

        for r_coord in patch_cells:
            if r_coord not in orig_cells and r_coord not in changed_refs:
                raise RuntimeError(
                    f"Ô lạ ngoài ý muốn '{r_coord}' xuất hiện trong sheet '{config.target_sheet}' sau khi patch!"
                )
    finally:
        z_orig.close()
        z_orig.fp = None
        z_patch.close()
        z_patch.fp = None

    # 2. Đọc lại workbook đích bằng openpyxl để xác minh dữ liệu giá trị ô
    wb = load_workbook(BytesIO(patched_bytes), data_only=True)
    try:
        ws = wb[config.target_sheet]

        # Kiểm tra các ô thay đổi
        for item in reconcile_result.changes:
            row = item["row"]
            expected = item["after"]
            actual = ws.cell(row=row, column=config.target_value_col).value
            if actual is None or abs(float(actual) - float(expected)) > EPS:
                raise RuntimeError(
                    f"Xác minh thất bại tại dòng {row}: kỳ vọng {expected}, thực tế {actual}."
                )

        # Kiểm tra các ô không đổi (unchanged)
        for item in reconcile_result.unchanged:
            row = item["row"]
            expected = item["value"]
            actual = ws.cell(row=row, column=config.target_value_col).value
            if not _cells_match(actual, expected):
                raise RuntimeError(
                    f"Ô không đổi tại dòng {row} bị thay đổi: cũ {expected!r}, mới {actual!r}."
                )

        # Kiểm tra các ô thiếu nguồn (missing_in_source)
        for item in reconcile_result.missing_in_source:
            row = item["row"]
            expected = item["current_value"]
            actual = ws.cell(row=row, column=config.target_value_col).value
            if not _cells_match(actual, expected):
                raise RuntimeError(
                    f"Ô giữ nguyên (mã thiếu nguồn) tại dòng {row} bị thay đổi: cũ {expected!r}, mới {actual!r}."
                )
    finally:
        _safe_close_workbook(wb)

    return {
        "ok": True,
        "message": "Xác minh toàn vẹn thành công 100%.",
        "exempted_parts": exempted_parts,
    }


def generate_nvl_report(
    reconcile_result: NVLReconcileResult,
    config: NVLConfig,
    *,
    mode: str,
    source_revision: str | None = None,
    target_revision: str | None = None,
    reporting_period: str | None = None,
    exempted_parts: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """Tạo báo cáo JSON đối soát chi tiết đồng bộ tồn NVL."""
    return {
        "schema_version": 1,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "mode": mode,
        "status": reconcile_result.status,
        "message": reconcile_result.message,
        "source": {
            "name": config.source_name,
            "sharepoint_path": config.source_path,
            "sourcedoc": config.source_sourcedoc,
            "sheet_name": config.source_sheet,
            "code_column": config.source_code_col_letter,
            "value_column": config.source_value_col_letter,
            "revision": source_revision,
            "reporting_period": reporting_period,
        },
        "target": {
            "name": config.target_name,
            "sharepoint_path": config.target_path,
            "sourcedoc": config.target_sourcedoc,
            "sheet_name": config.target_sheet,
            "code_column": config.target_code_col_letter,
            "value_column": config.target_value_col_letter,
            "revision": target_revision,
        },
        "metrics": {
            "matched_count": len(reconcile_result.changes) + len(reconcile_result.unchanged),
            "changed_count": len(reconcile_result.changes),
            "unchanged_count": len(reconcile_result.unchanged),
            "missing_in_source_count": len(reconcile_result.missing_in_source),
            "source_only_count": len(reconcile_result.source_only),
        },
        "changes": reconcile_result.changes,
        "warnings": [
            {
                "type": "missing_in_source",
                "row": item["row"],
                "code": item["code"],
                "message": f"Mã vật tư '{item['code']}' ở đích không có trong file nguồn; giữ nguyên số tồn cũ ({item['current_value']}).",
            }
            for item in reconcile_result.missing_in_source
        ],
        "source_only_codes": reconcile_result.source_only,
        "exempted_server_metadata_parts": exempted_parts or [],
    }


def generate_nvl_error_report(
    config: NVLConfig | None = None,
    *,
    mode: str = "unknown",
    phase: str,
    attempt: int,
    error: Exception,
    source_revision: str | None = None,
    target_revision: str | None = None,
    extra: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Tạo báo cáo lỗi JSON khi đồng bộ hoặc publish thất bại."""
    src_name = getattr(config, "source_name", "XNT_ketoan_Vikoda.xlsm") if config else "XNT_ketoan_Vikoda.xlsm"
    src_path = getattr(config, "source_path", "") if config else ""
    src_doc = getattr(config, "source_sourcedoc", "") if config else ""

    tgt_name = getattr(config, "target_name", "Kế hoạch mua hàng.xlsx") if config else "Kế hoạch mua hàng.xlsx"
    tgt_path = getattr(config, "target_path", "") if config else ""
    tgt_doc = getattr(config, "target_sourcedoc", "") if config else ""

    rep: dict[str, Any] = {
        "schema_version": 1,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "mode": mode,
        "status": "failed",
        "phase": phase,
        "attempt": attempt,
        "error_type": type(error).__name__,
        "message": str(error),
        "error_message": str(error),
        "source": {
            "name": src_name,
            "sharepoint_path": src_path,
            "sourcedoc": src_doc,
            "revision": source_revision,
        },
        "target": {
            "name": tgt_name,
            "sharepoint_path": tgt_path,
            "sourcedoc": tgt_doc,
            "revision": target_revision,
        },
    }
    if extra:
        rep.update(extra)
    return rep


def _is_network_timeout_or_reset(exc: Exception) -> bool:
    """Xác định các lỗi timeout hoặc ngắt kết nối mạng tạm thời."""
    if isinstance(exc, (TimeoutError, socket.timeout, ConnectionResetError, ConnectionRefusedError)):
        return True
    name = type(exc).__name__.lower()
    msg = str(exc).lower()
    if any(k in name for k in ("timeout", "connectionerror", "urlerror", "readtimeouterror")):
        return True
    if any(k in msg for k in ("timed out", "timeout", "connection reset", "connection refused", "remotely closed")):
        return True
    return False


def run_nvl_sync(
    config: NVLConfig,
    *,
    source_file: str | Path | None = None,
    target_file: str | Path | None = None,
    out_dir: str | Path | None = None,
    publish: bool = False,
    graph: GraphClient | None = None,
    max_publish_attempts: int = 3,
) -> dict[str, Any]:
    """Chạy quy trình đồng bộ tồn NVL ở chế độ offline hoặc online."""
    install_retry_after_support()

    if out_dir is None:
        out_dir = Path("offline_out/nvl") if (source_file or target_file) else Path(".")
    else:
        out_dir = Path(out_dir)

    try:
        out_dir.mkdir(parents=True, exist_ok=True)
    except OSError:
        pass

    proposal_path = out_dir / "nvl_stock_proposal.xlsx"
    report_path = out_dir / "nvl_stock_report.json"

    # Dọn dẹp artifact cũ để không gây hiểu nhầm nếu lần chạy hiện tại lỗi sớm
    if proposal_path.exists():
        try:
            proposal_path.unlink()
        except OSError:
            pass
    if report_path.exists():
        try:
            report_path.unlink()
        except OSError:
            pass

    mode = "publish" if publish else ("offline" if (source_file or target_file) else "dry_run")
    source_rev_final = None
    target_rev_final = None
    current_phase = "init"
    last_error: Exception | None = None

    def _emit_error_and_raise(
        phase: str,
        err: Exception,
        attempt: int = 1,
        src_rev: str | None = None,
        tgt_rev: str | None = None,
        extra: dict[str, Any] | None = None,
    ) -> None:
        if proposal_path.exists():
            try:
                proposal_path.unlink()
            except OSError:
                pass
        err_rep = generate_nvl_error_report(
            config,
            mode=mode,
            phase=phase,
            attempt=attempt,
            error=err,
            source_revision=src_rev or source_rev_final,
            target_revision=tgt_rev or target_rev_final,
            extra=extra,
        )
        try:
            report_path.write_text(json.dumps(err_rep, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        except OSError:
            pass
        raise err

    # 1. Kiểm tra cấu hình và tham số đầu vào
    if (source_file or target_file) and publish:
        _emit_error_and_raise(
            "validate_input",
            ValueError("Chế độ offline (--source-file/--target-file) không thể kết hợp với cờ --publish."),
        )

    # 2. Chế độ OFFLINE (chạy từ file cục bộ)
    if source_file or target_file:
        if not (source_file and target_file):
            _emit_error_and_raise(
                "validate_input",
                ValueError("Phải cung cấp đồng thời cả --source-file và --target-file."),
            )
        src_path = Path(source_file)
        tgt_path = Path(target_file)
        if not src_path.exists():
            _emit_error_and_raise("offline_input", FileNotFoundError(f"Không tìm thấy file nguồn: {src_path}"))
        if not tgt_path.exists():
            _emit_error_and_raise("offline_input", FileNotFoundError(f"Không tìm thấy file đích: {tgt_path}"))

        current_phase = "offline_read"
        try:
            print(f"[OFFLINE] Đọc nguồn: {src_path}")
            source_bytes = src_path.read_bytes()
            print(f"[OFFLINE] Đọc đích: {tgt_path}")
            target_bytes = tgt_path.read_bytes()

            current_phase = "offline_reconcile"
            source_stock, source_metadata = read_nvl_source_stock(source_bytes, config)
            print(f"[OFFLINE] Đọc được {len(source_stock)} mã vật tư từ nguồn.")

            reconcile_res = reconcile_nvl_target(target_bytes, source_stock, config)
            print(
                f"[OFFLINE] Đối soát đích: {len(reconcile_res.changes)} ô đổi, "
                f"{len(reconcile_res.unchanged)} ô không đổi, "
                f"{len(reconcile_res.missing_in_source)} mã thiếu ở nguồn."
            )

            current_phase = "offline_patch"
            patched_bytes = patch_nvl_destination_workbook(target_bytes, reconcile_res, config)

            current_phase = "offline_verify"
            verify_nvl_patched_workbook(target_bytes, patched_bytes, reconcile_res, config)

            proposal_path.write_bytes(patched_bytes)
            print(f"[OFFLINE] Đã ghi proposal: {proposal_path}")

            report = generate_nvl_report(
                reconcile_res,
                config,
                mode="offline",
                reporting_period=source_metadata.get("reporting_period"),
            )
            report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            print(f"[OFFLINE] Đã ghi báo cáo: {report_path}")
            return report
        except Exception as exc:
            _emit_error_and_raise(current_phase, exc, attempt=1)

    # 3. Chế độ ONLINE (kết nối Microsoft Graph)
    if not getattr(config, "source_path", None):
        _emit_error_and_raise("validate_config", ValueError("Chưa cấu hình 'source.sharepoint_path' trong file cấu hình."))
    if not getattr(config, "target_path", None):
        _emit_error_and_raise(
            "validate_config",
            ValueError(
                f"Chưa cấu hình đường dẫn 'target.sharepoint_path' cho '{config.target_name}' trên SharePoint.\n"
                f"(Sourcedoc ID đã biết: {config.target_sourcedoc}). Vui lòng điền đường dẫn thư mục chính xác."
            ),
        )

    current_phase = "init_graph"
    try:
        if graph is None:
            token = get_access_token()
            graph = GraphClient(token)
        site_id = graph.get_site_id()
        drive_id = graph.get_default_drive_id(site_id)
    except Exception as exc:
        _emit_error_and_raise("init_graph", exc, attempt=1)

    for attempt in range(1, max_publish_attempts + 1):
        try:
            # 3.1. Đọc metadata nguồn trước khi download
            current_phase = "fetch_source_metadata"
            print(f"[ONLINE] Lượt {attempt}/{max_publish_attempts}: Đọc metadata nguồn...")
            source_item_before = graph.get_item_by_path(drive_id, config.source_path)
            source_etag_before = source_item_before.get("eTag")

            # 3.2. Tải snapshot nguồn
            current_phase = "download_source"
            source_bytes = graph.download_file(drive_id, source_item_before["id"])

            # 3.3. Kiểm tra tính tươi mới của nguồn ngay sau download
            current_phase = "verify_source_freshness"
            source_item_after = graph.get_item_by_path(drive_id, config.source_path)
            if source_item_after.get("eTag") != source_etag_before:
                print(f"[ONLINE] Nguồn đã thay đổi trong lúc download snapshot (lượt {attempt}/{max_publish_attempts}); tải lại...")
                last_error = RuntimeError(
                    f"Nguồn đã thay đổi trong lúc download snapshot (eTag trước: {source_etag_before}, sau: {source_item_after.get('eTag')})."
                )
                source_rev_final = source_item_after.get("eTag")
                if proposal_path.exists():
                    try:
                        proposal_path.unlink()
                    except OSError:
                        pass
                if attempt < max_publish_attempts:
                    time.sleep(1.0)
                    continue
                else:
                    _emit_error_and_raise(
                        "verify_source_freshness",
                        last_error,
                        attempt=attempt,
                        src_rev=source_rev_final,
                        tgt_rev=target_rev_final,
                    )
            source_rev_final = source_item_after.get("eTag")

            # 3.4. Tải snapshot đích
            current_phase = "download_target"
            target_item = graph.get_item_by_path(drive_id, config.target_path)
            target_rev_final = target_item.get("eTag")
            target_bytes = graph.download_file(drive_id, target_item["id"])
            target_sha256 = hashlib.sha256(target_bytes).hexdigest()

            # 3.4.1. Lưu và xác minh bản backup thực tế đích trước khi patch/upload
            current_phase = "pre_upload_backup"
            backup_raw_path = out_dir / "official_backup_target_raw.xlsx"
            backup_info_path = out_dir / "official_target_backup_info.json"
            try:
                backup_raw_path.write_bytes(target_bytes)
                saved_raw_bytes = backup_raw_path.read_bytes()
                saved_raw_sha = hashlib.sha256(saved_raw_bytes).hexdigest()
                if saved_raw_sha != target_sha256:
                    raise RuntimeError(
                        f"Xác minh SHA-256 backup đích thất bại: kỳ vọng {target_sha256}, thực tế {saved_raw_sha}"
                    )
                backup_meta = {
                    "name": target_item.get("name", config.target_name),
                    "sharepoint_path": config.target_path,
                    "item_id": target_item.get("id"),
                    "eTag": target_rev_final,
                    "sourcedoc": config.target_sourcedoc,
                    "sha256": target_sha256,
                    "size_bytes": len(target_bytes),
                    "timestamp": datetime.now(timezone.utc).isoformat(),
                }
                backup_info_path.write_text(
                    json.dumps(backup_meta, indent=2, ensure_ascii=False), encoding="utf-8"
                )
                saved_meta = json.loads(backup_info_path.read_text(encoding="utf-8"))
                if saved_meta.get("sha256") != target_sha256:
                    raise RuntimeError(
                        "Xác minh metadata backup thất bại: SHA-256 trong file JSON không khớp với target_sha256."
                    )
                print(f"[ONLINE] Đã lưu và xác minh bản backup thực tế đích trước khi patch: {backup_raw_path} (SHA-256: {target_sha256})")
            except Exception as backup_exc:
                err = RuntimeError(f"Lưu hoặc xác minh backup thực tế đích thất bại trước khi upload: {backup_exc}")
                _emit_error_and_raise(
                    "pre_upload_backup",
                    err,
                    attempt=attempt,
                    src_rev=source_rev_final,
                    tgt_rev=target_rev_final,
                )

            # 3.5. Đối soát và lập kế hoạch patch
            current_phase = "reconcile"
            source_stock, source_metadata = read_nvl_source_stock(source_bytes, config)
            reconcile_res = reconcile_nvl_target(target_bytes, source_stock, config)

            # 3.6. Patch và xác minh proposal cục bộ
            current_phase = "patch"
            patched_bytes = patch_nvl_destination_workbook(target_bytes, reconcile_res, config)

            current_phase = "verify_proposal"
            verify_nvl_patched_workbook(target_bytes, patched_bytes, reconcile_res, config)
            proposal_path.write_bytes(patched_bytes)

            # Trường hợp dry-run: dừng tại đây và ghi nhận báo cáo đề xuất
            if not publish:
                current_phase = "finalize_dry_run"
                report = generate_nvl_report(
                    reconcile_res,
                    config,
                    mode="dry_run",
                    source_revision=source_rev_final,
                    target_revision=target_rev_final,
                    reporting_period=source_metadata.get("reporting_period"),
                )
                report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
                print(f"[ONLINE] Dry-run hoàn tất. Proposal: {proposal_path}, Report: {report_path}")
                return report

            # Trường hợp publish nhưng không có thay đổi (dữ liệu đã khớp)
            if not reconcile_res.changes:
                current_phase = "finalize_unchanged"
                report = generate_nvl_report(
                    reconcile_res,
                    config,
                    mode="publish",
                    source_revision=source_rev_final,
                    target_revision=target_rev_final,
                    reporting_period=source_metadata.get("reporting_period"),
                )
                report["status"] = "unchanged"
                report["message"] = "Dữ liệu tồn kho khớp hoàn toàn, không có ô nào cần upload."
                report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
                print("[ONLINE] Dữ liệu khớp 100%, không có ô nào cần upload.")
                return report

            # 3.7. Chuẩn bị publish: Kiểm tra lại nguồn trước khi upload đích
            current_phase = "pre_upload_source_check"
            source_item_pre = graph.get_item_by_path(drive_id, config.source_path)
            if source_item_pre.get("eTag") != source_rev_final:
                print(
                    f"[ONLINE] Nguồn đã bị thay đổi trước khi upload (lượt {attempt}/{max_publish_attempts}); "
                    "hủy lượt tải lên và làm mới snapshot..."
                )
                last_error = RuntimeError(
                    f"Nguồn đã bị thay đổi trước khi upload (eTag snapshot: {source_rev_final}, hiện tại: {source_item_pre.get('eTag')})."
                )
                source_rev_final = source_item_pre.get("eTag")
                if proposal_path.exists():
                    try:
                        proposal_path.unlink()
                    except OSError:
                        pass
                if attempt < max_publish_attempts:
                    time.sleep(1.0)
                    continue
                else:
                    _emit_error_and_raise(
                        "pre_upload_source_check",
                        last_error,
                        attempt=attempt,
                        src_rev=source_rev_final,
                        tgt_rev=target_rev_final,
                    )

            # 3.8. Upload file đích với If-Match ETag
            current_phase = "upload_target"
            print(f"[ONLINE] Đang upload file đích với ETag {target_rev_final}...")
            upload_response = None
            is_timeout_upload = False
            try:
                upload_response = graph.upload_file(
                    drive_id,
                    target_item["id"],
                    patched_bytes,
                    expected_etag=target_rev_final,
                )
                print(f"[ONLINE] Upload hoàn tất: {upload_response.get('name')} (Lượt {attempt})")
            except Exception as up_exc:
                if _is_network_timeout_or_reset(up_exc):
                    print(f"[ONLINE] Upload bị ngắt kết nối/timeout ({up_exc}). Chuyển sang xác minh trạng thái server...")
                    is_timeout_upload = True
                    upload_response = {
                        "id": target_item["id"],
                        "name": target_item.get("name"),
                        "status": "timeout_unconfirmed_ack",
                        "raw_timeout_error": str(up_exc),
                    }
                else:
                    # Lỗi không phải timeout (ví dụ 412 ETag conflict, 403 Forbidden...):
                    # Chắc chắn server chưa ghi nhận nội dung này. Ném lỗi ra ngoài xử lý.
                    raise up_exc

            # 3.9. Tải lại file đích sau upload để xác minh toàn vẹn trên server
            # LƯU Ý BẢO TOÀN VÀ ĐỐI SOÁT (Áp dụng cho CẢ trường hợp có ACK và trường hợp timeout):
            # - Khi server đã nhận file hoặc gặp timeout chờ phản hồi: tuyệt đối KHÔNG tự ý upload lại toàn bộ
            #   với ETag mới hoặc rollback ghi đè file của người khác nếu phát hiện xung đột.
            # - Nếu GET tải lại gặp lỗi mạng tạm thời (429, 503, timeout): retry đọc tối đa 3 lần.
            # - Nếu dữ liệu trên server không khớp với patch hoặc bị chỉnh sửa đồng thời ngoài cột D: dừng ngay,
            #   báo lỗi, ghi report failed với phase='post_upload_verify' và giữ upload_acknowledged=True.
            # - Nếu GET thất bại sau 3 lần thử: dừng ngay với trạng thái unverified.
            current_phase = "post_upload_verify"
            max_verify_download_attempts = 3
            post_upload_err = None
            verify_success = False
            verify_res: dict[str, Any] | None = None

            for v_attempt in range(1, max_verify_download_attempts + 1):
                try:
                    server_bytes = graph.download_file(drive_id, target_item["id"])
                    verify_res = verify_nvl_patched_workbook(
                        target_bytes, server_bytes, reconcile_res, config, is_server_comparison=True
                    )
                    verify_success = True
                    if is_timeout_upload:
                        print("[ONLINE] Server đã nhận đủ dữ liệu trước khi timeout; xác nhận thành công 100%.")
                        upload_response["status"] = "verified_post_timeout"
                    else:
                        print("[ONLINE] Đã tải lại file đích từ SharePoint và đối soát thành công 100%.")
                    break
                except Exception as verr:
                    post_upload_err = verr
                    # Kiểm tra nếu là lỗi đối soát dữ liệu (mismatch/conflict): dừng ngay, không retry GET!
                    is_data_mismatch = isinstance(verr, RuntimeError) and (
                        any(k in str(verr).lower() for k in ("xác minh thất bại", "bị thay đổi", "khác biệt", "không khớp", "biến đổi", "bị mất"))
                        or (not _is_network_timeout_or_reset(verr) and not isinstance(verr, GraphRequestError))
                    )
                    if is_data_mismatch:
                        prefix = "LỖI XÁC MINH SAU TIMEOUT" if is_timeout_upload else "LỖI XÁC MINH SAU UPLOAD"
                        print(f"[ONLINE] {prefix} (Xung đột dữ liệu trên server): {verr}")
                        break

                    # Nếu là lỗi Graph hoặc timeout tạm thời khi download: retry GET
                    is_temp = (
                        isinstance(verr, GraphRequestError) and is_retryable_graph_error(verr)
                    ) or _is_network_timeout_or_reset(verr)
                    if is_temp and v_attempt < max_verify_download_attempts:
                        delay = retry_delay_seconds(verr, 1.0) if isinstance(verr, GraphRequestError) else 1.0
                        print(
                            f"[ONLINE] Lỗi tạm thời khi tải lại để xác minh ({verr}). "
                            f"Thử lại lượt {v_attempt + 1}/{max_verify_download_attempts} sau {delay}s..."
                        )
                        time.sleep(delay)
                        continue
                    break

            if not verify_success:
                is_data_mismatch = isinstance(post_upload_err, RuntimeError) and (
                    any(k in str(post_upload_err).lower() for k in ("xác minh thất bại", "bị thay đổi", "khác biệt", "không khớp", "biến đổi", "bị mất"))
                    or (not _is_network_timeout_or_reset(post_upload_err) and not isinstance(post_upload_err, GraphRequestError))
                )
                v_status = "conflict_or_mismatch" if is_data_mismatch else "unverified"
                if is_timeout_upload:
                    err_msg = (
                        f"Upload bị timeout/mất kết nối và xác minh sau đó thất bại do dữ liệu trên server bị chỉnh sửa "
                        f"đồng thời hoặc không khớp với bản patch: {post_upload_err}. Dữ liệu trên server được giữ nguyên, "
                        f"tuyệt đối không tự ý upload lại."
                        if is_data_mismatch
                        else (
                            f"Upload bị timeout/mất kết nối và không thể tải lại file để xác minh sau "
                            f"{max_verify_download_attempts} lần thử: {post_upload_err}. Upload có thể đã được server ghi nhận."
                        )
                    )
                else:
                    err_msg = (
                        f"Upload đã được SharePoint ghi nhận nhưng xác minh sau upload thất bại do dữ liệu bị sửa đổi "
                        f"đồng thời hoặc không khớp: {post_upload_err}"
                        if is_data_mismatch
                        else (
                            f"Upload đã gửi thành công nhưng không thể tải lại file để xác minh sau "
                            f"{max_verify_download_attempts} lần thử: {post_upload_err}"
                        )
                    )
                final_err = RuntimeError(err_msg)
                _emit_error_and_raise(
                    "post_upload_verify",
                    final_err,
                    attempt=attempt,
                    src_rev=source_rev_final,
                    tgt_rev=target_rev_final,
                    extra={
                        "upload_acknowledged": not is_timeout_upload,
                        "upload_may_have_committed": True,
                        "upload_result": upload_response,
                        "verification_status": v_status,
                        "is_timeout_upload": is_timeout_upload,
                        "raw_verification_error": str(post_upload_err),
                    },
                )

            # 3.10. Ghi nhận báo cáo thành công (Chỉ khi xác minh thành công!)
            current_phase = "finalize_published"
            exempted_list = verify_res.get("exempted_parts", []) if verify_res else []
            report = generate_nvl_report(
                reconcile_res,
                config,
                mode="publish",
                source_revision=source_rev_final,
                target_revision=target_rev_final,
                reporting_period=source_metadata.get("reporting_period"),
                exempted_parts=exempted_list,
            )
            report["status"] = "published" if not reconcile_res.missing_in_source else "published_with_warnings"
            report["message"] = f"Đồng bộ và publish thành công: {len(reconcile_res.changes)} ô đã cập nhật lên SharePoint."
            report["upload_result"] = upload_response
            report["post_upload_verified"] = True
            report["exempted_server_metadata_parts"] = exempted_list
            report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            return report

        except Exception as exc:
            if current_phase in ("pre_upload_backup", "post_upload_verify"):
                raise

            last_error = exc
            is_retryable = False
            delay = 1.0

            if isinstance(exc, GraphRequestError):
                if exc.status_code == 412:
                    print(
                        f"[ONLINE] HTTP 412: File đích đã thay đổi đồng thời trên SharePoint "
                        f"(Lượt {attempt}/{max_publish_attempts})."
                    )
                    is_retryable = True
                    delay = 1.0
                elif is_retryable_graph_error(exc):
                    is_retryable = True
                    delay = retry_delay_seconds(exc, 3.0)
                    print(f"[ONLINE] Lỗi tạm thời Graph HTTP {exc.status_code} ({exc}). Chờ {delay}s...")
            elif _is_network_timeout_or_reset(exc):
                is_retryable = True
                delay = 2.0
                print(f"[ONLINE] Lỗi mạng tạm thời ({exc}). Chờ {delay}s...")

            if is_retryable and attempt < max_publish_attempts:
                if proposal_path.exists():
                    try:
                        proposal_path.unlink()
                    except OSError:
                        pass
                time.sleep(delay)
                continue

            # Fail fast cho lỗi không retryable hoặc khi đã hết lượt thử
            _emit_error_and_raise(
                current_phase,
                exc,
                attempt=attempt,
                src_rev=source_rev_final,
                tgt_rev=target_rev_final,
            )

    # Thoát vòng lặp do vượt quá số lần thử
    _emit_error_and_raise(
        current_phase,
        last_error or RuntimeError("Vượt quá số lần thử tải lên SharePoint do xung đột hoặc lỗi tạm thời liên tục."),
        attempt=max_publish_attempts,
        src_rev=source_rev_final,
        tgt_rev=target_rev_final,
    )


def main():
    parser = argparse.ArgumentParser(
        description="Đồng bộ tồn kho nguyên vật liệu (NVL) từ XNT_ketoan_Vikoda sang Kế hoạch mua hàng."
    )
    parser.add_argument("--config", default=DEFAULT_CONFIG_FILE, help="Đường dẫn file cấu hình JSON.")
    parser.add_argument("--source-file", help="File nguồn cục bộ (chế độ offline).")
    parser.add_argument("--target-file", help="File đích cục bộ (chế độ offline).")
    parser.add_argument("--target-path", help="Ghi đè đường dẫn file đích trên SharePoint.")
    parser.add_argument("--out", help="Thư mục xuất proposal và report.")
    parser.add_argument("--publish", action="store_true", help="Publish lên SharePoint (mặc định là dry-run).")

    args = parser.parse_args()

    out_dir = Path(args.out) if args.out else (Path("offline_out/nvl") if (args.source_file or args.target_file) else Path("."))

    sync_started = False
    try:
        cfg = load_nvl_config(args.config)
        if args.target_path:
            cfg.target_path = args.target_path

        sync_started = True
        rep = run_nvl_sync(
            cfg,
            source_file=args.source_file,
            target_file=args.target_file,
            out_dir=args.out,
            publish=args.publish,
        )
        print(f"Hoàn thành ({rep.get('status')}): {rep.get('message')}")
    except Exception as exc:
        report_file = out_dir / "nvl_stock_report.json"
        proposal_file = out_dir / "nvl_stock_proposal.xlsx"

        if not sync_started:
            # Lỗi xảy ra TRƯỚC KHI sync khởi động (ví dụ load config lỗi, thiếu file config, JSON hỏng):
            # 1. Xóa proposal cũ từ lần chạy trước nếu có để tránh proposal không hợp lệ tồn tại
            try:
                if proposal_file.exists():
                    proposal_file.unlink(missing_ok=True)
            except OSError:
                pass

            # 2. Luôn ghi đè báo cáo lỗi hiện tại (phase="init_cli", status="failed") thay thế report cũ
            try:
                out_dir.mkdir(parents=True, exist_ok=True)
                err_rep = generate_nvl_error_report(
                    None,
                    mode="publish" if args.publish else ("offline" if (args.source_file or args.target_file) else "dry_run"),
                    phase="init_cli",
                    attempt=1,
                    error=exc,
                )
                report_file.write_text(json.dumps(err_rep, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            except OSError:
                pass
        else:
            # Nếu sync đã khởi động, run_nvl_sync đã chủ động ghi báo cáo lỗi chi tiết chuyên sâu (như post_upload_verify)
            # Chỉ ghi fallback nếu report_file vì lý do nào đó chưa tồn tại
            if not report_file.exists():
                try:
                    out_dir.mkdir(parents=True, exist_ok=True)
                    err_rep = generate_nvl_error_report(
                        None,
                        mode="publish" if args.publish else ("offline" if (args.source_file or args.target_file) else "dry_run"),
                        phase="cli_unhandled",
                        attempt=1,
                        error=exc,
                    )
                    report_file.write_text(json.dumps(err_rep, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
                except OSError:
                    pass

        print(f"LỖI: {exc}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
