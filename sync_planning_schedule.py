import math
import time
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

from lxml import etree
from openpyxl import load_workbook

from sync_planning_fc import (
    _find_sheet_xml_path,
    _get_or_create_cell,
    _load_shared_strings,
    _read_cell_text,
)
from sync_stock import DEST_PATH, GraphClient, get_access_token, normalize_code


PLANNING_SHEET = "Ke_hoach_SX"
START_COLUMN_NUMBER = 19  # S
MAX_DAYS = 31
TIMEZONE = ZoneInfo("Asia/Ho_Chi_Minh")
EPSILON = 1e-9

# File mẫu đang dùng 0,5 ca khi đổi mã/khuôn trên KHS/PET.
SETUP_SHIFTS = 0.5
CONTINUOUS_LINES = {"KHS", "PET 9000"}
WEEKLY_LINES = {"RGB"}
GALON_LINE = "Galon"
GALON_SPREAD_CODE = "130100006"


def _column_letter(number):
    letters = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _to_number(value, label):
    if value in (None, ""):
        return 0.0
    if isinstance(value, bool):
        raise RuntimeError(f"{label} chứa TRUE/FALSE, không phải số.")
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise RuntimeError(f"{label} không phải số: {value!r}") from exc


def _clean_number(value):
    value = float(value)
    if abs(value) < EPSILON:
        return 0
    if math.isclose(value, round(value), rel_tol=1e-12, abs_tol=1e-8):
        return int(round(value))
    return value


def _as_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise RuntimeError(f"Ngày bắt đầu sản xuất không hợp lệ: {value!r}")


def _parse_header_date(value, plan_year):
    text = str(value or "").strip()
    if not text:
        return None
    first_line = text.splitlines()[0].strip()
    try:
        day_text, month_text = first_line.split("/", 1)
        return date(plan_year, int(month_text), int(day_text))
    except Exception as exc:
        raise RuntimeError(f"Tiêu đề ngày không hợp lệ: {value!r}") from exc


def _is_sugar(classification):
    return str(classification or "").strip().casefold() == "có đường".casefold()


def _demand_profile(headers, fc):
    """Giữ convention FC/26: nhu cầu T2-T7, CN vẫn có thể sản xuất."""
    demand_days = sum(1 for current_day in headers if current_day.weekday() != 6)
    if fc <= 0 or demand_days <= 0:
        return {current_day: 0.0 for current_day in headers}
    daily = fc / demand_days
    return {
        current_day: (0.0 if current_day.weekday() == 6 else daily)
        for current_day in headers
    }


def _validate_quantum(product):
    if product["planned_qty"] <= 0:
        product["quantum_qty"] = 0.0
        product["quantum_shift"] = 0.0
        product["required_units"] = 0
        return

    if not product["line"]:
        raise RuntimeError(f"Mã {product['code']} chưa có Chuyền tại cột F.")
    if product["per_shift"] <= 0:
        raise RuntimeError(f"Mã {product['code']} có Số lượng/ca <= 0.")
    if product["max_shifts_per_day"] <= 0:
        raise RuntimeError(f"Mã {product['code']} có Số ca theo ngày <= 0.")

    quantum_qty = product["batch"] if product["is_sugar"] else product["per_shift"]
    if quantum_qty <= 0:
        source = "Số lượng/mẻ" if product["is_sugar"] else "Số lượng/ca"
        raise RuntimeError(f"Mã {product['code']} có {source} <= 0.")

    unit_count = product["planned_qty"] / quantum_qty
    rounded_units = round(unit_count)
    if not math.isclose(unit_count, rounded_units, rel_tol=1e-9, abs_tol=1e-7):
        raise RuntimeError(
            f"Mã {product['code']} có P={product['planned_qty']} không chia hết "
            f"cho quantum={quantum_qty}."
        )

    quantum_shift = quantum_qty / product["per_shift"]
    if quantum_shift - product["max_shifts_per_day"] > EPSILON:
        raise RuntimeError(
            f"Mã {product['code']} cần {quantum_shift:.3f} ca cho một quantum "
            f"nhưng giới hạn ngày chỉ {product['max_shifts_per_day']:.3f} ca."
        )

    product["quantum_qty"] = quantum_qty
    product["quantum_shift"] = quantum_shift
    product["required_units"] = int(rounded_units)


def read_schedule_inputs(workbook_bytes, *, plan_year=None):
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    try:
        if PLANNING_SHEET not in workbook.sheetnames:
            raise RuntimeError(f"Không tìm thấy sheet {PLANNING_SHEET!r}.")
        worksheet = workbook[PLANNING_SHEET]

        if plan_year is None:
            plan_year = datetime.now(TIMEZONE).year

        headers = []
        for offset in range(MAX_DAYS):
            value = worksheet.cell(
                row=1,
                column=START_COLUMN_NUMBER + offset,
            ).value
            parsed = _parse_header_date(value, plan_year)
            if parsed is None:
                if headers:
                    break
                continue
            headers.append(parsed)

        if not headers:
            raise RuntimeError(
                f"{PLANNING_SHEET}!S1 trở đi chưa có dải ngày. "
                "Hãy chạy sync_planning_calendar.py trước."
            )

        first_day = headers[0]
        expected = [first_day + timedelta(days=i) for i in range(len(headers))]
        if headers != expected:
            raise RuntimeError("Dải ngày S1:... không liên tục theo ngày.")

        products = []
        seen_codes = set()
        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=2,
                max_col=START_COLUMN_NUMBER + MAX_DAYS - 1,
                values_only=True,
            ),
            start=2,
        ):
            code = normalize_code(values[0] if values else None)
            if not code:
                continue
            if code in seen_codes:
                raise RuntimeError(f"Mã {code} bị lặp trong {PLANNING_SHEET}!A.")
            seen_codes.add(code)

            classification = str(values[7] or "").strip()
            product = {
                "row": row_number,
                "code": code,
                "batch": _to_number(values[3], f"{PLANNING_SHEET}!D{row_number}"),
                "per_shift": _to_number(values[4], f"{PLANNING_SHEET}!E{row_number}"),
                "line": str(values[5] or "").strip(),
                "product_group": str(values[6] or "").strip(),
                "classification": classification,
                "is_sugar": _is_sugar(classification),
                "max_shifts_per_day": _to_number(values[8], f"{PLANNING_SHEET}!I{row_number}"),
                "actual_stock": _to_number(values[9], f"{PLANNING_SHEET}!J{row_number}"),
                "fc": _to_number(values[11], f"{PLANNING_SHEET}!L{row_number}"),
                "target_stock": _to_number(values[12], f"{PLANNING_SHEET}!M{row_number}"),
                "debt": max(_to_number(values[13], f"{PLANNING_SHEET}!N{row_number}"), 0.0),
                "planned_qty": max(_to_number(values[15], f"{PLANNING_SHEET}!P{row_number}"), 0.0),
                "earliest_date": _as_date(values[17]),
                "existing_daily": list(values[18:18 + MAX_DAYS]),
            }
            product["demand_by_day"] = _demand_profile(headers, product["fc"])
            _validate_quantum(product)
            products.append(product)

        if not products:
            raise RuntimeError(f"Không đọc được SKU trong {PLANNING_SHEET}.")

        return headers, products
    finally:
        workbook.close()


def _earliest_index(headers, product):
    earliest = product["earliest_date"]
    if earliest is None or earliest <= headers[0]:
        return 0
    for index, current_day in enumerate(headers):
        if current_day >= earliest:
            return index
    return len(headers)


def _empty_schedule(headers, products):
    return {
        product["code"]: {current_day: 0.0 for current_day in headers}
        for product in products
    }


def _add_interval_usage(usage, headers, start_shift, duration, capacity):
    end_shift = start_shift + duration
    for index, current_day in enumerate(headers):
        day_start = index * capacity
        day_end = day_start + capacity
        overlap = max(0.0, min(end_shift, day_end) - max(start_shift, day_start))
        if overlap > EPSILON:
            usage[current_day] += overlap


def _add_production_interval(schedule, usage, headers, product, start_shift, duration, capacity):
    end_shift = start_shift + duration
    produced = 0.0
    for index, current_day in enumerate(headers):
        day_start = index * capacity
        day_end = day_start + capacity
        overlap = max(0.0, min(end_shift, day_end) - max(start_shift, day_start))
        if overlap <= EPSILON:
            continue
        qty = overlap * product["per_shift"]
        schedule[product["code"]][current_day] += qty
        usage[current_day] += overlap
        produced += qty
    return produced


def _campaign_candidate(remaining, headers, cursor_shift, capacity, last_group):
    if not remaining:
        return None, cursor_shift

    current_index = min(len(headers), int(math.floor(cursor_shift / capacity + EPSILON)))
    eligible = []
    for product in remaining:
        earliest_index = _earliest_index(headers, product)
        if earliest_index <= current_index:
            eligible.append(product)

    if not eligible:
        next_index = min(_earliest_index(headers, product) for product in remaining)
        cursor_shift = max(cursor_shift, next_index * capacity)
        current_index = next_index
        eligible = [
            product
            for product in remaining
            if _earliest_index(headers, product) <= current_index
        ]

    # File mẫu Helper_SX_2Line sắp theo Chuyền -> R -> dòng gốc.
    # Không kéo một SKU R muộn lên trước chỉ để giữ cùng nhóm; tính liên tục
    # được đảm bảo bằng việc mỗi SKU chỉ chạy đúng một campaign.
    eligible.sort(
        key=lambda product: (
            _earliest_index(headers, product),
            0 if product["debt"] > 0 else 1,
            product["row"],
        )
    )
    return eligible[0], cursor_shift


def _allocate_continuous_campaign_line(headers, products, capacity):
    """Mirror file mẫu: mỗi SKU là một block liên tục, đổi SKU mất 0,5 ca."""
    schedule = _empty_schedule(headers, products)
    usage = defaultdict(float)
    carryover = {}
    campaigns = []
    remaining = list(products)
    cursor_shift = 0.0
    last_code = None
    last_group = None
    month_end_shift = len(headers) * capacity
    setup_total = 0.0

    while remaining:
        product, cursor_shift = _campaign_candidate(
            remaining, headers, cursor_shift, capacity, last_group
        )
        remaining.remove(product)

        earliest_shift = _earliest_index(headers, product) * capacity
        cursor_shift = max(cursor_shift, earliest_shift)

        if last_code is not None and product["code"] != last_code:
            setup_start = cursor_shift
            setup_duration = min(SETUP_SHIFTS, max(0.0, month_end_shift - setup_start))
            if setup_duration > EPSILON:
                _add_interval_usage(usage, headers, setup_start, setup_duration, capacity)
                setup_total += setup_duration
            cursor_shift += SETUP_SHIFTS

        production_start = max(cursor_shift, earliest_shift)
        required_duration = product["planned_qty"] / product["per_shift"]

        if production_start >= month_end_shift - EPSILON:
            carryover[product["code"]] = _clean_number(product["planned_qty"])
            campaigns.append(
                {
                    "code": product["code"],
                    "group": product["product_group"],
                    "start_shift": production_start,
                    "end_shift": production_start,
                    "scheduled_qty": 0,
                }
            )
            last_code = product["code"]
            last_group = product["product_group"]
            continue

        actual_duration = min(required_duration, month_end_shift - production_start)
        produced = _add_production_interval(
            schedule,
            usage,
            headers,
            product,
            production_start,
            actual_duration,
            capacity,
        )
        cursor_shift = production_start + actual_duration

        missing = max(0.0, product["planned_qty"] - produced)
        if missing > 1e-6:
            carryover[product["code"]] = _clean_number(missing)

        campaigns.append(
            {
                "code": product["code"],
                "group": product["product_group"],
                "start_shift": production_start,
                "end_shift": cursor_shift,
                "scheduled_qty": _clean_number(produced),
            }
        )
        last_code = product["code"]
        last_group = product["product_group"]

    return schedule, usage, carryover, {
        "mode": "continuous_campaign",
        "setup_shifts": setup_total,
        "campaigns": campaigns,
    }


def _new_remaining(headers, capacity):
    return {current_day: float(capacity) for current_day in headers}


def _consume_setup(remaining, usage, headers, start_index, *, allow_sunday, max_index=None):
    need = SETUP_SHIFTS
    end_index = len(headers) - 1 if max_index is None else min(max_index, len(headers) - 1)
    for index in range(max(0, start_index), end_index + 1):
        current_day = headers[index]
        if not allow_sunday and current_day.weekday() == 6:
            continue
        take = min(remaining[current_day], need)
        if take > EPSILON:
            remaining[current_day] -= take
            usage[current_day] += take
            need -= take
        if need <= EPSILON:
            return index, True
    return end_index + 1, False


def _place_units(
    schedule,
    usage,
    remaining,
    headers,
    product,
    units,
    start_index,
    *,
    allow_sunday,
    max_index=None,
):
    if units <= 0:
        return 0, start_index

    quantum_shift = product["quantum_shift"]
    placed = 0
    cursor = max(0, start_index)
    end_index = len(headers) - 1 if max_index is None else min(max_index, len(headers) - 1)

    while placed < units and cursor <= end_index:
        current_day = headers[cursor]
        if (
            (allow_sunday or current_day.weekday() != 6)
            and current_day >= (product["earliest_date"] or headers[0])
            and remaining[current_day] + EPSILON >= quantum_shift
        ):
            max_fit = int(math.floor((remaining[current_day] + EPSILON) / quantum_shift))
            max_fit = min(max_fit, units - placed)
            if max_fit > 0:
                shift_used = max_fit * quantum_shift
                qty = max_fit * product["quantum_qty"]
                schedule[product["code"]][current_day] += qty
                usage[current_day] += shift_used
                remaining[current_day] -= shift_used
                placed += max_fit
                if placed >= units:
                    return placed, cursor
        cursor += 1

    return placed, cursor


def _month_week_buckets(headers):
    buckets = []
    for start in range(0, len(headers), 7):
        buckets.append((start, min(start + 6, len(headers) - 1)))
    return buckets


def _weekly_unit_targets(headers, product):
    buckets = _month_week_buckets(headers)
    active = [
        bucket for bucket in buckets
        if headers[bucket[1]] >= (product["earliest_date"] or headers[0])
    ]
    if not active:
        return {}

    total = product["required_units"]
    targets = {}
    done_target = 0
    for ordinal, bucket in enumerate(active, start=1):
        cumulative = int(math.ceil(total * ordinal / len(active) - EPSILON))
        targets[bucket] = cumulative - done_target
        done_target = cumulative
    return targets


def _allocate_weekly_rgb_line(headers, products, capacity):
    """Mirror file mẫu RGB: chia sản lượng theo tuần rồi gom từng mã thành mini-campaign."""
    schedule = _empty_schedule(headers, products)
    usage = defaultdict(float)
    remaining_capacity = _new_remaining(headers, capacity)
    weekly_targets = {product["code"]: _weekly_unit_targets(headers, product) for product in products}
    pending = defaultdict(int)
    last_code = None
    last_group = None
    setup_total = 0.0
    campaigns = []

    for bucket in _month_week_buckets(headers):
        start_index, end_index = bucket
        weekly_units = {}
        for product in products:
            pending[product["code"]] += weekly_targets[product["code"]].get(bucket, 0)
            if pending[product["code"]] > 0:
                weekly_units[product["code"]] = pending[product["code"]]

        remaining_products = [
            product for product in products
            if weekly_units.get(product["code"], 0) > 0
        ]

        while remaining_products:
            same_group = [
                product for product in remaining_products
                if last_group and product["product_group"] == last_group
            ]
            candidates = same_group or remaining_products
            candidates.sort(
                key=lambda product: (
                    _earliest_index(headers, product),
                    0 if product["debt"] > 0 else 1,
                    product["product_group"],
                    product["row"],
                )
            )
            product = candidates[0]
            remaining_products.remove(product)

            campaign_start = start_index
            if last_code is not None and last_code != product["code"]:
                setup_index, ok = _consume_setup(
                    remaining_capacity,
                    usage,
                    headers,
                    campaign_start,
                    allow_sunday=False,
                    max_index=end_index,
                )
                if ok:
                    setup_total += SETUP_SHIFTS
                    campaign_start = setup_index
                else:
                    # Không phí công cố đổi mã trong tuần nếu hết capacity; dồn qua tuần sau.
                    continue

            requested = weekly_units[product["code"]]
            placed, last_index = _place_units(
                schedule,
                usage,
                remaining_capacity,
                headers,
                product,
                requested,
                campaign_start,
                allow_sunday=False,
                max_index=end_index,
            )
            pending[product["code"]] -= placed
            if placed > 0:
                campaigns.append(
                    {
                        "code": product["code"],
                        "week": f"{headers[start_index]:%d/%m}-{headers[end_index]:%d/%m}",
                        "units": placed,
                    }
                )
                last_code = product["code"]
                last_group = product["product_group"]

    # Nếu capacity T2-T7 không đủ mới dùng Chủ nhật, giống file mẫu.
    for product in sorted(products, key=lambda item: (_earliest_index(headers, item), item["row"])):
        missing = pending[product["code"]]
        if missing <= 0:
            continue
        start_index = _earliest_index(headers, product)
        if last_code is not None and last_code != product["code"]:
            setup_index, ok = _consume_setup(
                remaining_capacity,
                usage,
                headers,
                start_index,
                allow_sunday=True,
            )
            if ok:
                setup_total += SETUP_SHIFTS
                start_index = setup_index

        placed, _ = _place_units(
            schedule,
            usage,
            remaining_capacity,
            headers,
            product,
            missing,
            start_index,
            allow_sunday=True,
        )
        pending[product["code"]] -= placed
        if placed > 0:
            last_code = product["code"]
            last_group = product["product_group"]

    carryover = {}
    for product in products:
        missing = pending[product["code"]]
        if missing > 0:
            carryover[product["code"]] = _clean_number(missing * product["quantum_qty"])

    return schedule, usage, carryover, {
        "mode": "weekly_rgb",
        "setup_shifts": setup_total,
        "campaigns": campaigns,
    }


def _place_fractional_shift(
    schedule,
    usage,
    remaining,
    headers,
    product,
    day_index,
    shift_qty,
):
    if shift_qty <= EPSILON:
        return 0.0
    current_day = headers[day_index]
    shift_need = shift_qty / product["per_shift"]
    shift_take = min(remaining[current_day], shift_need)
    if shift_take <= EPSILON:
        return 0.0
    qty = shift_take * product["per_shift"]
    schedule[product["code"]][current_day] += qty
    usage[current_day] += shift_take
    remaining[current_day] -= shift_take
    return qty


def _allocate_galon_line(headers, products, capacity):
    """19L rải đều theo ngày làm việc; các mã Galon khác chạy block liên tục vào capacity còn lại."""
    schedule = _empty_schedule(headers, products)
    usage = defaultdict(float)
    remaining_capacity = _new_remaining(headers, capacity)
    carryover = {}
    campaigns = []
    setup_total = 0.0
    last_code = None

    spread = next((product for product in products if product["code"] == GALON_SPREAD_CODE), None)
    others = [product for product in products if product["code"] != GALON_SPREAD_CODE]

    if spread is not None:
        eligible = [
            index for index, current_day in enumerate(headers)
            if index >= _earliest_index(headers, spread) and current_day.weekday() != 6
        ]
        remaining_qty = spread["planned_qty"]

        # Rải đều giống công thức file mẫu: mỗi ngày nhận số ca nền như nhau,
        # phần ca dư được phân theo vị trí tỷ lệ trên toàn dải ngày.
        if eligible and remaining_qty > EPSILON:
            full_units = int(math.floor(remaining_qty / spread["per_shift"] + EPSILON))
            partial_qty = remaining_qty - full_units * spread["per_shift"]
            base_units, extra_units = divmod(full_units, len(eligible))

            planned_units = {index: base_units for index in eligible}
            for extra_no in range(1, extra_units + 1):
                rank = int(math.ceil(extra_no * len(eligible) / extra_units)) if extra_units else 0
                rank = min(max(rank, 1), len(eligible))
                planned_units[eligible[rank - 1]] += 1

            for index in eligible:
                units_today = planned_units[index]
                for _ in range(units_today):
                    made = _place_fractional_shift(
                        schedule,
                        usage,
                        remaining_capacity,
                        headers,
                        spread,
                        index,
                        spread["per_shift"],
                    )
                    if made <= EPSILON:
                        break
                    remaining_qty -= made

            if partial_qty > EPSILON and remaining_qty > EPSILON:
                # Phần lẻ đặt ở ngày làm việc cuối còn capacity để không tạo nhiều lần đổi mã.
                for index in reversed(eligible):
                    made = _place_fractional_shift(
                        schedule,
                        usage,
                        remaining_capacity,
                        headers,
                        spread,
                        index,
                        min(partial_qty, remaining_qty),
                    )
                    if made > EPSILON:
                        remaining_qty -= made
                        break

        # Chỉ dùng Chủ nhật nếu phần còn lại không thể nhét vào T2-T7.
        if remaining_qty > EPSILON:
            for index, current_day in enumerate(headers):
                if index < _earliest_index(headers, spread) or current_day.weekday() != 6:
                    continue
                while remaining_qty > EPSILON and remaining_capacity[current_day] > EPSILON:
                    qty = min(spread["per_shift"], remaining_qty)
                    made = _place_fractional_shift(
                        schedule,
                        usage,
                        remaining_capacity,
                        headers,
                        spread,
                        index,
                        qty,
                    )
                    if made <= EPSILON:
                        break
                    remaining_qty -= made

        if remaining_qty > 1e-6:
            carryover[spread["code"]] = _clean_number(remaining_qty)

        last_code = spread["code"]
        campaigns.append({"code": spread["code"], "mode": "spread_workdays"})

    for product in sorted(
        others,
        key=lambda item: (
            _earliest_index(headers, item),
            0 if item["debt"] > 0 else 1,
            item["row"],
        ),
    ):
        start_index = _earliest_index(headers, product)

        if last_code is not None and last_code != product["code"]:
            setup_index, ok = _consume_setup(
                remaining_capacity,
                usage,
                headers,
                start_index,
                allow_sunday=True,
            )
            if ok:
                setup_total += SETUP_SHIFTS
                start_index = setup_index

        placed, _ = _place_units(
            schedule,
            usage,
            remaining_capacity,
            headers,
            product,
            product["required_units"],
            start_index,
            allow_sunday=True,
        )
        missing_units = product["required_units"] - placed
        if missing_units > 0:
            carryover[product["code"]] = _clean_number(missing_units * product["quantum_qty"])
        if placed > 0:
            last_code = product["code"]
            campaigns.append({"code": product["code"], "mode": "continuous_remaining_capacity"})

    return schedule, usage, carryover, {
        "mode": "galon_hybrid",
        "setup_shifts": setup_total,
        "campaigns": campaigns,
    }


def _allocate_generic_line(headers, products, capacity):
    # Unknown line: safer default is one continuous block per SKU, not fragmented optimization.
    return _allocate_continuous_campaign_line(headers, products, capacity)


def build_schedule(headers, products):
    schedule = _empty_schedule(headers, products)
    active_products = [product for product in products if product["planned_qty"] > 0]

    by_line = defaultdict(list)
    for product in active_products:
        by_line[product["line"]].append(product)

    line_capacity = {}
    line_usage = defaultdict(lambda: defaultdict(float))
    carryover = {}
    optimizer_meta = {}

    for line, line_products in sorted(by_line.items()):
        if line in CONTINUOUS_LINES:
            # File mẫu KHS/PET sử dụng một capacity chung (hiện tại đều 3 ca/ngày).
            capacity = min(product["max_shifts_per_day"] for product in line_products)
            line_schedule, usage, line_carryover, meta = _allocate_continuous_campaign_line(
                headers, line_products, capacity
            )
        elif line in WEEKLY_LINES:
            capacity = max(product["max_shifts_per_day"] for product in line_products)
            line_schedule, usage, line_carryover, meta = _allocate_weekly_rgb_line(
                headers, line_products, capacity
            )
        elif line == GALON_LINE:
            capacity = max(product["max_shifts_per_day"] for product in line_products)
            line_schedule, usage, line_carryover, meta = _allocate_galon_line(
                headers, line_products, capacity
            )
        else:
            capacity = min(product["max_shifts_per_day"] for product in line_products)
            line_schedule, usage, line_carryover, meta = _allocate_generic_line(
                headers, line_products, capacity
            )

        line_capacity[line] = capacity
        optimizer_meta[line] = meta
        carryover.update(line_carryover)

        for current_day, used_shift in usage.items():
            line_usage[line][current_day] = used_shift
        for product in line_products:
            code = product["code"]
            for current_day in headers:
                schedule[code][current_day] = line_schedule[code][current_day]

    validate_schedule(headers, products, schedule, line_capacity, line_usage)
    inventory = simulate_inventory(headers, products, schedule)
    return schedule, line_capacity, line_usage, carryover, inventory, optimizer_meta


def simulate_inventory(headers, products, schedule):
    result = {}
    for product in products:
        stock = product["actual_stock"]
        min_stock = stock
        first_stockout = None
        first_below_safety = None
        for current_day in headers:
            stock += (
                schedule[product["code"]][current_day]
                - product["demand_by_day"][current_day]
            )
            min_stock = min(min_stock, stock)
            if stock < -EPSILON and first_stockout is None:
                first_stockout = current_day
            if (
                stock < product["target_stock"] - EPSILON
                and first_below_safety is None
            ):
                first_below_safety = current_day

        result[product["code"]] = {
            "ending_stock": _clean_number(stock),
            "min_stock": _clean_number(min_stock),
            "first_stockout": first_stockout,
            "first_below_safety": first_below_safety,
        }
    return result


def validate_schedule(headers, products, schedule, line_capacity, line_usage):
    for line, days in line_usage.items():
        capacity = line_capacity[line]
        for current_day, used in days.items():
            if used - capacity > 1e-6:
                raise RuntimeError(
                    f"Chuyền {line} ngày {current_day:%d/%m} dùng {used:.3f} ca > "
                    f"capacity {capacity:.3f}."
                )

    for product in products:
        code = product["code"]
        total = sum(schedule[code].values())
        if total - product["planned_qty"] > 1e-5:
            raise RuntimeError(
                f"Mã {code} được xếp {total} > P={product['planned_qty']}."
            )
        earliest = product["earliest_date"]
        if earliest:
            for current_day, qty in schedule[code].items():
                if qty > EPSILON and current_day < earliest:
                    raise RuntimeError(
                        f"Mã {code} bị xếp ngày {current_day:%d/%m} trước R={earliest:%d/%m}."
                    )


def _number_equal(current, target):
    current = 0.0 if current in (None, "") else float(current)
    target = 0.0 if target in (None, "") else float(target)
    return math.isclose(current, target, rel_tol=1e-12, abs_tol=1e-8)


def _set_numeric_or_blank(cell, value):
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)
    if value in (None, "") or abs(float(value)) < EPSILON:
        return

    namespace = etree.QName(cell).namespace
    node = etree.SubElement(cell, f"{{{namespace}}}v")
    cleaned = _clean_number(value)
    node.text = str(cleaned) if isinstance(cleaned, int) else format(float(cleaned), ".15g")


def patch_schedule_workbook(workbook_bytes, headers, products, schedule):
    target_values = {}
    day_count = len(headers)
    changed_count = 0

    for product in products:
        values = []
        for index in range(MAX_DAYS):
            target = schedule[product["code"]][headers[index]] if index < day_count else 0
            target = _clean_number(target)
            values.append(target)
            current = product["existing_daily"][index] if index < len(product["existing_daily"]) else None
            if not _number_equal(current, target):
                changed_count += 1
        target_values[product["code"]] = values

    if changed_count == 0:
        return workbook_bytes, 0

    source_buffer = BytesIO(workbook_bytes)
    output_buffer = BytesIO()
    with zipfile.ZipFile(source_buffer, "r") as source_zip:
        sheet_path = _find_sheet_xml_path(source_zip, PLANNING_SHEET)
        shared_strings = _load_shared_strings(source_zip)
        root = etree.fromstring(source_zip.read(sheet_path))

        seen = set()
        rows = root.xpath('//*[local-name()="sheetData"]/*[local-name()="row"]')
        for row in rows:
            row_number_text = row.get("r")
            if not row_number_text or int(row_number_text) == 1:
                continue
            row_number = int(row_number_text)
            code = None
            for cell in row.xpath('./*[local-name()="c"]'):
                if cell.get("r") == f"A{row_number}":
                    code = normalize_code(_read_cell_text(cell, shared_strings))
                    break
            if not code or code not in target_values:
                continue
            seen.add(code)

            p_nodes = row.xpath(f'./*[local-name()="c"][@r="P{row_number}"]')
            body_style = p_nodes[0].get("s") if p_nodes else None
            for index, value in enumerate(target_values[code]):
                column_letter = _column_letter(START_COLUMN_NUMBER + index)
                cell = _get_or_create_cell(row, row_number, column_letter)
                if body_style is not None:
                    cell.set("s", body_style)
                _set_numeric_or_blank(cell, value)

        missing = sorted(set(target_values) - seen)
        if missing:
            raise RuntimeError(
                f"Không tìm thấy {len(missing)} mã trong XML {PLANNING_SHEET}: "
                + ", ".join(missing[:10])
            )

        new_xml = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )
        with zipfile.ZipFile(output_buffer, "w", compression=zipfile.ZIP_DEFLATED) as output_zip:
            for item in source_zip.infolist():
                data = new_xml if item.filename == sheet_path else source_zip.read(item.filename)
                output_zip.writestr(item, data)

    return output_buffer.getvalue(), changed_count


def prepare_schedule_update(workbook_bytes, *, plan_year=None):
    headers, products = read_schedule_inputs(workbook_bytes, plan_year=plan_year)
    (
        schedule,
        line_capacity,
        line_usage,
        carryover,
        inventory,
        optimizer_meta,
    ) = build_schedule(headers, products)
    updated_bytes, changed_count = patch_schedule_workbook(
        workbook_bytes,
        headers,
        products,
        schedule,
    )

    utilization = {}
    for line, capacity in line_capacity.items():
        total_capacity = capacity * len(headers)
        total_used = sum(line_usage[line].values())
        utilization[line] = 0 if total_capacity <= 0 else total_used / total_capacity

    return updated_bytes, {
        "headers": headers,
        "products": products,
        "schedule": schedule,
        "line_capacity": line_capacity,
        "line_usage": line_usage,
        "utilization": utilization,
        "carryover": carryover,
        "inventory": inventory,
        "optimizer_meta": optimizer_meta,
        "changed_count": changed_count,
    }


def main_with_retry(*, sleep_func=time.sleep, max_attempts=6, retry_delay_seconds=10):
    token = get_access_token()
    graph = GraphClient(token)
    site_id = graph.get_site_id()
    drive_id = graph.get_default_drive_id(site_id)

    for attempt in range(1, max_attempts + 1):
        dest_item = graph.get_item_by_path(drive_id, DEST_PATH)
        dest_bytes = graph.download_file(drive_id, dest_item["id"])
        updated_bytes, info = prepare_schedule_update(dest_bytes)

        for line, utilization in sorted(info["utilization"].items()):
            capacity = info["line_capacity"][line]
            mode = info["optimizer_meta"].get(line, {}).get("mode", "")
            setup = info["optimizer_meta"].get(line, {}).get("setup_shifts", 0)
            print(
                f"[Scheduler] {line}: mode={mode}; {capacity:g} ca/ngày; "
                f"setup {setup:g} ca; utilization tháng {utilization * 100:.1f}%."
            )

        if info["carryover"]:
            print("[Scheduler][CARRYOVER] Chưa xếp hết trong tháng:")
            for code, qty in sorted(info["carryover"].items()):
                print(f"  - {code}: {qty}")
        else:
            print("[Scheduler] Tất cả P > 0 đã được xếp hết trong tháng.")

        stockout_codes = []
        safety_codes = []
        for product in info["products"]:
            result = info["inventory"][product["code"]]
            if result["first_stockout"]:
                stockout_codes.append((product["code"], result["first_stockout"]))
            elif result["first_below_safety"]:
                safety_codes.append((product["code"], result["first_below_safety"]))

        if stockout_codes:
            print("[Scheduler][INVENTORY_RISK] Các mã vẫn có nguy cơ tồn âm:")
            for code, risk_day in stockout_codes:
                print(f"  - {code}: từ {risk_day:%d/%m}")
        if safety_codes:
            print("[Scheduler][SAFETY_RISK] Các mã xuống dưới tồn mục tiêu:")
            for code, risk_day in safety_codes:
                print(f"  - {code}: từ {risk_day:%d/%m}")
        if not stockout_codes and not safety_codes:
            print("[Scheduler] Inventory simulation không phát hiện stockout/safety breach.")

        if info["changed_count"] == 0:
            print(f"[{PLANNING_SHEET}] Lịch S:... đã đúng; không cần upload lại.")
            return

        try:
            graph.upload_file(
                drive_id,
                dest_item["id"],
                updated_bytes,
                expected_etag=dest_item["eTag"],
            )
            print(
                f"[{PLANNING_SHEET}] Đã cập nhật {info['changed_count']} ô lịch sản xuất S:..."
            )
            return
        except RuntimeError as exc:
            message = str(exc)
            retryable = "423" in message or "resourceLocked" in message or "412" in message
            if not retryable or attempt == max_attempts:
                raise
            print(
                f"[{PLANNING_SHEET}] File đang khóa/thay đổi; thử scheduler lại "
                f"({attempt}/{max_attempts})."
            )
            sleep_func(retry_delay_seconds)


if __name__ == "__main__":
    main_with_retry()
