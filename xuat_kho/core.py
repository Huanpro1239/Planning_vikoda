# -*- coding: utf-8 -*-
"""Tổng hợp Xuất kho thực tế theo tháng x sản phẩm (logic thuần, không phụ thuộc SharePoint).

Nguồn: "Bao cao xuat kho/<Vikoda|VKD>/Xuat kho thang N.xlsm", sheet REPORT_XUATBANHANG.

Quy tắc (giống Sell In về mã):
  - Công ty đọc từ TENCHINHANH (có "VKD" -> VKD, "KHÁNH HÒA" -> Vikoda), dự phòng theo thư mục.
  - Kỳ đọc từ dòng "Từ ngày : dd-mm-yyyy".
  - VKD   : chỉ lấy MASANPHAM bắt đầu bằng "2".
  - Vikoda: chỉ lấy MASANPHAM bắt đầu bằng "1" và MAKHACHHANG khác "VKD3".
  - Bỏ mã 130100014 / 230100014 (Nước khoáng bán thành phẩm).
  - Chỉ lấy xuất kho THỰC TẾ: bỏ các dòng CHUYENKHOKYGUI = "Chuyển qua kho ký gửi"
    (chỉ là chuyển hàng sang kho ký gửi); hàng ký gửi được tính khi xuất từ KHO KÝ GỬI
    (LOAIDONHANG = "Ký gửi") ở kỳ thực xuất.
  - SanLuongQuyDoi = tổng SOLUONG / Quy cách (Danh mục của đúng công ty).
  - Pivot gộp Vikoda + VKD theo phần đuôi mã, cộng mọi loại đơn; sheet Data tách theo LoaiDonHang.
"""
import datetime as dt
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from sell_in.core import MA_LOAI_TRU, PIVOT_FIXED, _ghi_bang, gio_vn, ma_sp, norm, to_num

SHEET = "REPORT_XUATBANHANG"
DATA_COLS = ["CongTy", "Nam", "Thang", "KyBaoCao", "LoaiDonHang", "MaSanPham", "TenSanPham", "TenVietTat",
             "Brand", "SKUsGroup", "DVTQuyDoi", "QuyCach", "SoLuong", "SanLuongQuyDoi"]
LOG_COLS = ["SourceFile", "CongTy", "Nam", "Thang", "TongDongNguon", "SoDongLay", "SoDongChuyenKyGui",
            "SoSanPham", "MaThieuQuyCach", "CapNhatLuc"]


def tong_hop_file(src, dm, source_name):
    """Đọc 1 file Xuất kho. Trả về (rows theo DATA_COLS, log_row theo LOG_COLS)."""
    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    top, hdr = [], None
    for _ in range(30):
        try:
            row = next(it)
        except StopIteration:
            break
        top.append(row)
        if row and any(norm(v) == "MASANPHAM" for v in row):
            hdr = row
            break
    if hdr is None:
        wb.close()
        raise ValueError("Không tìm thấy dòng tiêu đề MASANPHAM trong " + source_name)
    h = [norm(v) for v in hdr]

    def c(name, bat_buoc=True):
        if name in h:
            return h.index(name)
        if bat_buoc:
            wb.close()
            raise ValueError(f"Thiếu cột {name} trong {source_name}")
        return None

    i_cn, i_kh, i_ma, i_sl = c("TENCHINHANH"), c("MAKHACHHANG"), c("MASANPHAM"), c("SOLUONG")
    i_ldh, i_ckg = c("LOAIDONHANG"), c("CHUYENKHOKYGUI")

    nam = thang = None
    for row in top:
        s = " ".join("" if v is None else str(v) for v in row)
        m = re.search(r"T[ừu]\s*ng[àa]y\s*:\s*(\d{1,2})[-/](\d{1,2})[-/](\d{4})", s, re.I)
        if m:
            thang, nam = int(m.group(2)), int(m.group(3))
            break
    if nam is None:
        wb.close()
        raise ValueError("Không xác định được tháng/năm (dòng 'Từ ngày') của " + source_name)

    cty = None
    tong, n_nguon, n_lay, n_ckg = {}, 0, 0, 0
    need = max(i_cn, i_kh, i_ma, i_sl, i_ldh, i_ckg)
    for row in it:
        if row is None or len(row) <= need:
            continue
        ma = ma_sp(row[i_ma])
        if ma == "":
            continue
        n_nguon += 1
        if cty is None:
            cn = norm(row[i_cn])
            cty = "VKD" if "VKD" in cn else ("Vikoda" if "KHANHHOA" in cn else None)
        if cty is None:
            cty = "VKD" if re.search(r"(^|/)VKD/", source_name, re.I) else "Vikoda"
        if not ma.startswith("2" if cty == "VKD" else "1") or ma in MA_LOAI_TRU:
            continue
        if cty == "Vikoda" and norm(row[i_kh]) == "VKD3":
            continue
        if norm(row[i_ckg]):  # "Chuyển qua kho ký gửi": chưa phải xuất thực tế
            n_ckg += 1
            continue
        sl = to_num(row[i_sl])
        if sl is None:
            continue
        ldh = str(row[i_ldh] or "").strip() or "(trống)"
        tong[(ldh, ma)] = tong.get((ldh, ma), 0) + sl
        n_lay += 1
    wb.close()
    if cty is None:
        cty = "VKD" if re.search(r"(^|/)VKD/", source_name, re.I) else "Vikoda"

    ky = dt.datetime(nam, thang, 1)
    rows, thieu = [], set()
    for ldh, ma in sorted(tong, key=lambda k: (k[1], k[0])):
        info = dm[cty].get(ma)
        qc = info["QuyCach"] if info else None
        if not qc:
            thieu.add(ma)
        sl = tong[(ldh, ma)]
        rows.append([cty, nam, thang, ky, ldh, ma,
                     info["TenSanPham"] if info else "", info["TenVietTat"] if info else "",
                     info["Brand"] if info else "", info["SKUsGroup"] if info else "",
                     info["DVTQuyDoi"] if info else "", qc, sl, sl / qc if qc else None])
    n_sp = len({r[5] for r in rows})
    log = [source_name, cty, nam, thang, n_nguon, n_lay, n_ckg, n_sp, ", ".join(sorted(thieu)), gio_vn()]
    return rows, log


def _cung_ky(r, cty, nam, thang):
    return r[0] == cty and int(r[1] or 0) == nam and int(r[2] or 0) == thang


def tong_hop_nhieu_file(sources, dm):
    """sources: list (tên file, file-like). Tháng trùng (cùng công ty) thì file sau thay file trước."""
    data, log = [], []
    for name, src in sources:
        rows, log_row = tong_hop_file(src, dm, name)
        cty, nam, thang = log_row[1], log_row[2], log_row[3]
        data = [r for r in data if not _cung_ky(r, cty, nam, thang)] + rows
        log = [r for r in log if not (r[1] == cty and r[2] == nam and r[3] == thang)] + [log_row]
    return data, log


def build_workbook(data, log, fingerprint=""):
    """Trả về (bytes .xlsx, số sản phẩm Pivot, số tháng)."""
    thu_tu = {"Vikoda": 0, "VKD": 1}
    data = sorted(data, key=lambda r: (int(r[1]), int(r[2]), thu_tu.get(r[0], 9), str(r[5]), str(r[4])))
    log = sorted(log, key=lambda r: (int(r[2]), int(r[3]), thu_tu.get(r[1], 9)))

    ky_list = sorted({(int(r[1]), int(r[2])) for r in data})
    sp, nhan = {}, {}
    for r in data:
        key = str(r[5])[1:]
        if key not in nhan or (not nhan[key][2] and r[6]):
            nhan[key] = ["1" + key, "2" + key, r[6], r[7], r[8], r[9], r[10], r[11]]
        d = sp.setdefault(key, {})
        if r[13] is not None:
            k = (int(r[1]), int(r[2]))
            d[k] = d.get(k, 0) + r[13]
    keys = sorted(nhan, key=lambda k: (str(nhan[k][4]), str(nhan[k][5]), str(nhan[k][3]), k))
    p_head = PIVOT_FIXED + ["T%02d/%d" % (t, n) for n, t in ky_list] + ["Tong"]
    p_rows = []
    for k in keys:
        vals = [sp[k].get(x) for x in ky_list]
        co = [v for v in vals if v is not None]
        p_rows.append(nhan[k] + [round(v, 2) if v is not None else None for v in vals] +
                      [round(sum(co), 2) if co else None])

    wb = Workbook()
    ws = wb.active
    ws.title = "Pivot"
    _ghi_bang(ws, "tblPivotXK", p_head, p_rows, "TableStyleMedium2")
    for j in range(len(PIVOT_FIXED) + 1, len(p_head) + 1):
        for i in range(2, len(p_rows) + 2):
            ws.cell(i, j).number_format = "#,##0.00"
        ws.column_dimensions[get_column_letter(j)].width = 11
    for j, w in enumerate([11, 11, 44, 28, 11, 20, 9, 8], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "E2"

    wd = wb.create_sheet("Data")
    _ghi_bang(wd, "tblDataXK", DATA_COLS, data, "TableStyleLight9")
    for i in range(2, len(data) + 2):
        wd.cell(i, 4).number_format = "mm/yyyy"
        wd.cell(i, 13).number_format = "#,##0"
        wd.cell(i, 14).number_format = "#,##0.00"
    for j, w in enumerate([8, 6, 6, 9, 18, 11, 44, 28, 11, 20, 9, 8, 12, 14], start=1):
        wd.column_dimensions[get_column_letter(j)].width = w
    wd.freeze_panes = "A2"

    wl = wb.create_sheet("Log")
    _ghi_bang(wl, "tblLogXK", LOG_COLS, log, "TableStyleLight9")
    for i in range(2, len(log) + 2):
        wl.cell(i, 10).number_format = "dd/mm/yyyy hh:mm"
    for j, w in enumerate([34, 8, 6, 6, 14, 11, 18, 10, 22, 17], start=1):
        wl.column_dimensions[get_column_letter(j)].width = w

    wm = wb.create_sheet("Meta")
    wm["A1"], wm["B1"] = "fingerprint", fingerprint
    wm.sheet_state = "hidden"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(p_rows), len(ky_list)
