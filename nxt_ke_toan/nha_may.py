# -*- coding: utf-8 -*-
"""Tồn kho thực tế nhà máy ("Bao cao ton kho/Ton thuc te/Thang N.xlsx").

Mẫu "BÁO CÁO KHO THÀNH PHẨM": ảnh chụp tại 1 ngày (ô "Ngày:"), mỗi dòng 1 sản phẩm có
MÃ SP VIKODA / MÃ SP VKD, đơn vị đã là Thùng/Két/Bình:
  Tổng nhập tháng, Tổng xuất tháng, Tồn thực tế đầu/cuối ngày, SL chờ QA/QC,
  Kế hoạch xuất kho, SL hàng đã bán (gửi kho), Tồn kho dự kiến.
File không đúng mẫu (vd. báo cáo tồn NPP đặt nhầm thư mục) -> trả về None và bị bỏ qua.
"""
import datetime as dt

from openpyxl import load_workbook

from sell_in.core import ma_sp, norm, to_num

TIEU_DE = "BAOCAOKHOTHANHPHAM"
COT = {  # tên chuẩn hoá ở dòng tiêu đề chính -> khoá
    "MASPVIKODA": "ma_vik", "MASPVKD": "ma_vkd", "TENSANPHAM": "ten", "DVT": "dvt",
    "SLCHUANHAPKHOCHOKQQAQC": "cho_qc", "SLHANGDABANGUIKHO": "da_ban",
}
COT_DAU = {  # bắt đầu bằng ... ở dòng tiêu đề chính
    "KEHOACHXUATKHO": "kh_xuat", "TONKHODUKIEN": "du_kien",
}
COT_PHU = {  # dòng tiêu đề phụ, bắt đầu bằng ...
    "TONGNHAP": "nhap_thang", "TONGXUAT": "xuat_thang", "DAUNGAY": "dau_ngay", "CUOINGAY": "cuoi_ngay",
}


def doc_ton_nha_may(src, source_name):
    """-> {"ngay": date, "source": name, "rows": [dict]} hoặc None nếu không đúng mẫu."""
    wb = load_workbook(src, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        top = []
        for row in ws.iter_rows(max_row=8, values_only=True):
            top.append(row)
        if not top or not any(norm(v) == TIEU_DE for v in top[0] if v is not None):
            return None
        ngay = None
        for row in top:
            for i, v in enumerate(row):
                if norm(v) == "NGAY" and i + 1 < len(row) and isinstance(row[i + 1], (dt.datetime, dt.date)):
                    ngay = row[i + 1]
                    break
            if ngay:
                break
        i_hdr = next((i for i, r in enumerate(top) if any(norm(v) == "MASPVIKODA" for v in r)), None)
        if i_hdr is None or ngay is None or i_hdr + 1 >= len(top):
            return None
        h, h2 = [norm(v) for v in top[i_hdr]], [norm(v) for v in top[i_hdr + 1]]
        col = {}
        for j, v in enumerate(h):
            if v in COT:
                col[COT[v]] = j
            for k, ten in COT_DAU.items():
                if v.startswith(k):
                    col[ten] = j
        for j, v in enumerate(h2):
            for k, ten in COT_PHU.items():
                if v.startswith(k):
                    col[ten] = j
        thieu = {"ma_vik", "ma_vkd", "cuoi_ngay", "da_ban"} - set(col)
        if thieu:
            raise ValueError(f"{source_name}: thiếu cột {', '.join(sorted(thieu))}")
        rows = []
        for r in ws.iter_rows(min_row=i_hdr + 3, values_only=True):
            if r is None or len(r) <= max(col.values()):
                continue
            mv, mk = ma_sp(r[col["ma_vik"]]), ma_sp(r[col["ma_vkd"]])
            if not (mv.isdigit() or mk.isdigit()):
                continue
            d = {"ma_vik": mv, "ma_vkd": mk}
            for k in ("ten", "dvt"):
                d[k] = str(r[col[k]] or "").strip() if k in col else ""
            for k in ("nhap_thang", "xuat_thang", "dau_ngay", "cuoi_ngay", "cho_qc", "kh_xuat", "da_ban", "du_kien"):
                d[k] = (to_num(r[col[k]]) or 0) if k in col else None
            rows.append(d)
        ngay = ngay.date() if isinstance(ngay, dt.datetime) else ngay
        return {"ngay": ngay, "source": source_name, "rows": rows}
    finally:
        wb.close()


def tong_hop_nha_may(sources):
    """sources: list (tên, file-like). -> (bao_cao theo (nam, thang) lấy ngày mới nhất trong tháng, file_bo_qua)."""
    theo_thang, bo_qua = {}, []
    for name, src in sources:
        bc = doc_ton_nha_may(src, name)
        if bc is None:
            bo_qua.append(name)
            continue
        k = (bc["ngay"].year, bc["ngay"].month)
        if k not in theo_thang or bc["ngay"] >= theo_thang[k]["ngay"]:
            theo_thang[k] = bc
    return theo_thang, bo_qua


def theo_duoi_ma(bc):
    """Báo cáo 1 ngày -> {đuôi mã: dòng} (đuôi = mã bỏ ký tự đầu, dùng để gộp với mã 1xxx/2xxx)."""
    out = {}
    for d in bc["rows"]:
        ma = d["ma_vik"] or d["ma_vkd"]
        out[ma[1:]] = d
    return out
