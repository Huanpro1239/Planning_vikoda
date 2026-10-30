# -*- coding: utf-8 -*-
"""Tồn ở KHO KÝ GỬI, dựng từ báo cáo Xuất kho bán hàng.

Kế toán ghi "Chuyển qua kho ký gửi" là XUẤT (đã trừ tồn), nhưng hàng chỉ thực bán khi xuất từ
KHO KÝ GỬI. Vì vậy:
    Tồn ký gửi cuối kỳ = Tồn ký gửi đầu kỳ + Chuyển vào ký gửi − Xuất từ ký gửi
    Tồn thực           = Tồn kế toán + Tồn ký gửi
    Xuất thực (bán)    = Xuất kế toán − Chuyển vào ký gửi + Xuất từ ký gửi

Tồn ký gửi đầu kỳ (mốc):
  - Có báo cáo tồn kho ký gửi (thư mục "Bao cao ton kho/Ton ky gui") -> dùng số đó.
  - Chưa có -> ƯỚC TÍNH TỐI THIỂU: mức đầu kỳ nhỏ nhất để tồn ký gửi không âm ở bất kỳ tháng nào.
"""
import re

from openpyxl import load_workbook

from sell_in.core import MA_LOAI_TRU, ma_sp, norm, to_num

SHEET_XK = "REPORT_XUATBANHANG"
NGUON_UOC_TINH = "Ước tính tối thiểu"
NGUON_HE_THONG = "Báo cáo tồn ký gửi"


def doc_luong_ky_gui(src, source_name):
    """Đọc 1 file Xuất kho -> (cty, nam, thang, {ma: [chuyen_vao, xuat_ra]}) theo số lượng gốc."""
    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb[SHEET_XK] if SHEET_XK in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    top, hdr = [], None
    for _ in range(30):
        row = next(it, None)
        if row is None:
            break
        top.append(row)
        if any(norm(v) == "MASANPHAM" for v in row):
            hdr = row
            break
    if hdr is None:
        wb.close()
        raise ValueError("Không tìm thấy tiêu đề MASANPHAM trong " + source_name)
    h = [norm(v) for v in hdr]
    i_cn, i_ma, i_sl = h.index("TENCHINHANH"), h.index("MASANPHAM"), h.index("SOLUONG")
    i_ckg, i_kho = h.index("CHUYENKHOKYGUI"), h.index("KHOXUAT")
    nam = thang = None
    for row in top:
        m = re.search(r"T[ừu]\s*ng[àa]y\s*:\s*(\d{1,2})[-/](\d{1,2})[-/](\d{4})",
                      " ".join("" if v is None else str(v) for v in row), re.I)
        if m:
            thang, nam = int(m.group(2)), int(m.group(3))
            break
    if nam is None:
        wb.close()
        raise ValueError("Không xác định được kỳ của " + source_name)
    cty, luong = None, {}
    need = max(i_cn, i_ma, i_sl, i_ckg, i_kho)
    for row in it:
        if row is None or len(row) <= need:
            continue
        ma = ma_sp(row[i_ma])
        if not ma:
            continue
        if cty is None:
            cn = norm(row[i_cn])
            cty = "VKD" if "VKD" in cn else ("Vikoda" if "KHANHHOA" in cn else None)
        c = cty or ("VKD" if re.search(r"(^|/)VKD/", source_name, re.I) else "Vikoda")
        if not ma.startswith("2" if c == "VKD" else "1") or ma in MA_LOAI_TRU:
            continue
        sl = to_num(row[i_sl]) or 0
        if norm(row[i_ckg]):
            luong.setdefault(ma, [0, 0])[0] += sl
        elif norm(row[i_kho]) == "KHOKYGUI":
            luong.setdefault(ma, [0, 0])[1] += sl
    wb.close()
    cty = cty or ("VKD" if re.search(r"(^|/)VKD/", source_name, re.I) else "Vikoda")
    return cty, nam, thang, luong


def ton_ky_gui(flows, moc=None):
    """flows: list (cty, nam, thang, {ma: [vao, ra]}).
    moc: {(cty, ma): số lượng tồn ký gửi ĐẦU kỳ của tháng đầu tiên} hoặc None (ước tính).
    Trả về ({(cty, (nam, thang), ma): [dau, vao, ra, cuoi]}, nguon_moc)."""
    luong = {}
    for cty, nam, thang, d in flows:
        for ma, (vao, ra) in d.items():
            luong[(cty, (nam, thang), ma)] = [vao, ra]  # file sau thay file trước cùng kỳ
    ky_list = sorted({k[1] for k in luong})
    mas = sorted({(k[0], k[2]) for k in luong})
    nguon = NGUON_HE_THONG if moc else NGUON_UOC_TINH
    out = {}
    for cty, ma in mas:
        if moc:
            dau = moc.get((cty, ma), 0)
        else:
            cum, mn = 0, 0
            for k in ky_list:
                vao, ra = luong.get((cty, k, ma), (0, 0))
                cum += vao - ra
                mn = min(mn, cum)
            dau = -mn
        for k in ky_list:
            vao, ra = luong.get((cty, k, ma), (0, 0))
            cuoi = dau + vao - ra
            out[(cty, k, ma)] = [dau, vao, ra, cuoi]
            dau = cuoi
    return out, nguon
