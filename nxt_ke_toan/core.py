# -*- coding: utf-8 -*-
"""Tổng hợp Nhập - Xuất - Tồn kế toán (Vikoda + VKD) theo tháng và phân tích tồn kho thông minh.

Nguồn: "Bao cao ton kho/Ton ke toan/<Vikoda|VKD>/Thang N.xlsm", Sheet1:
  - Dòng "Từ ngày dd-mm-yyyy đến ngày dd-mm-yyyy" -> kỳ và số ngày của kỳ.
  - Tiêu đề 2 dòng: Mã vật tư, Tên vật tư, Loại sản phẩm, Tài khoản, Đvt,
    Tồn đầu / Nhập / Xuất / Tồn cuối (mỗi nhóm có Số lượng + Thành tiền).

Đầu ra:
  - Data: mọi vật tư mọi tháng (số lượng gốc + thành tiền); thành phẩm/hàng hóa có thêm cột quy đổi Thùng/Két/Bình.
  - Ton thong minh: ảnh chụp tháng mới nhất cho thành phẩm/hàng hóa (gộp mã 1xxx/2xxx):
      tồn cuối từng công ty, xuất bình quân/ngày 3 tháng gần nhất, số ngày tồn, trạng thái, gợi ý.
  - Pivot Ton kha dung / Ton ky gui / So ngay ton / Ton nha may: sản phẩm x tháng (gộp 2 công ty).
  - Ton nha may: báo cáo kho thành phẩm nhà máy (Ton thuc te) theo ngày báo cáo.
  - Kiem tra: cân đối Đầu + Nhập - Xuất = Cuối, Cuối tháng trước = Đầu tháng sau, mã thiếu quy cách.
"""
import datetime as dt
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from sell_in.core import MA_LOAI_TRU, _ghi_bang, gio_vn, ma_sp, norm, to_num

LOAI_TP = {"THANHPHAM", "HANGHOA"}
SO_THANG_BQ = 3
# Ngưỡng số ngày tồn (có thể chỉnh)
NGUONG = [(7, "Thiếu hàng"), (15, "Thấp"), (30, "An toàn"), (60, "Cao")]
TT_DU, TT_HET, TT_CHAM, TT_KHONG = "Dư tồn", "Hết hàng", "Không luân chuyển", "Không phát sinh"
UU_TIEN = {TT_HET: 0, "Thiếu hàng": 1, "Thấp": 2, TT_CHAM: 3, TT_DU: 4, "Cao": 5, "An toàn": 6, TT_KHONG: 7}
GOI_Y = {
    TT_HET: "Cần nhập/sản xuất ngay",
    "Thiếu hàng": "Ưu tiên nhập/sản xuất",
    "Thấp": "Lên kế hoạch bổ sung",
    "An toàn": "Duy trì",
    "Cao": "Giảm nhập kỳ tới",
    TT_DU: "Dừng nhập, đẩy bán",
    TT_CHAM: "Rà soát hàng tồn đọng",
    TT_KHONG: "",
}
MAU = {TT_HET: "F8696B", "Thiếu hàng": "FBB4A5", "Thấp": "FFE699", "An toàn": "C6EFCE",
       "Cao": "BDD7EE", TT_DU: "9BC2E6", TT_CHAM: "D9D9D9"}

DATA_COLS = ["CongTy", "Nam", "Thang", "KyBaoCao", "SoNgayKy", "MaVatTu", "TenVatTu", "LoaiSanPham", "TaiKhoan",
             "Dvt", "TonDauSL", "TonDauTien", "NhapSL", "NhapTien", "XuatSL", "XuatTien", "TonCuoiSL", "TonCuoiTien",
             "QuyCach", "DVTQuyDoi", "TonDauQD", "NhapQD", "XuatQD", "TonCuoiQD"]
LOG_COLS = ["SourceFile", "CongTy", "Nam", "Thang", "TuNgay", "DenNgay", "SoDong", "SoDongLechCanDoi",
            "SoMaTP_HH", "MaThieuQuyCach", "CapNhatLuc"]


def _ngay(s):
    d, m, y = (int(x) for x in re.split(r"[-/]", s))
    return dt.date(y, m, d)


def tong_hop_file(src, dm, source_name):
    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    top, hdr = [], None
    for _ in range(30):
        try:
            row = next(it)
        except StopIteration:
            break
        top.append(row)
        if row and any(norm(v) == "MAVATTU" for v in row):
            hdr = row
            break
    if hdr is None:
        wb.close()
        raise ValueError("Không tìm thấy dòng tiêu đề 'Mã vật tư' trong " + source_name)
    h = [norm(v) for v in hdr]
    i_ma, i_ten, i_loai, i_dvt = h.index("MAVATTU"), h.index("TENVATTU"), h.index("LOAISANPHAM"), h.index("DVT")
    i_tk = h.index("TAIKHOANKETOAN") if "TAIKHOANKETOAN" in h else None
    i_dau = h.index("TONDAUKY")  # 8 cột liên tiếp: SL, Tiền x (Đầu, Nhập, Xuất, Cuối)
    next(it, None)  # dòng tiêu đề phụ "Số lượng / Thành tiền"

    ten_dn = norm(top[0][0] if top and top[0] else "")
    if "VKD" in ten_dn:
        cty = "VKD"
    elif "KHANHHOA" in ten_dn:
        cty = "Vikoda"
    else:
        cty = "VKD" if re.search(r"(^|/)VKD/", source_name, re.I) else "Vikoda"

    tu = den = None
    for row in top:
        s = " ".join("" if v is None else str(v) for v in row)
        m = re.search(r"T[ừu]\s*ng[àa]y\s*:?\s*(\d{1,2}[-/]\d{1,2}[-/]\d{4}).*?[đd][ếe]n\s*ng[àa]y\s*:?\s*(\d{1,2}[-/]\d{1,2}[-/]\d{4})",
                      s, re.I)
        if m:
            tu, den = _ngay(m.group(1)), _ngay(m.group(2))
            break
    if tu is None:
        wb.close()
        raise ValueError("Không xác định được kỳ (dòng 'Từ ngày ... đến ngày') của " + source_name)
    so_ngay = (den - tu).days + 1
    ky = dt.datetime(tu.year, tu.month, 1)

    rows, lech, thieu, n_tp = [], 0, set(), 0
    for row in it:
        if row is None or len(row) < i_dau + 8:
            continue
        ma = ma_sp(row[i_ma])
        if not ma:
            continue
        v = [to_num(row[i_dau + k]) or 0 for k in range(8)]
        if abs(v[0] + v[2] - v[4] - v[6]) > 0.001:
            lech += 1
        loai = str(row[i_loai] or "").strip()
        qc = dvt_qd = None
        qd = [None] * 4
        # Giống Sell In / Xuất kho: Vikoda chỉ mã đầu "1", VKD chỉ mã đầu "2"; bỏ bán thành phẩm
        if norm(loai) in LOAI_TP and ma.startswith("2" if cty == "VKD" else "1") and ma not in MA_LOAI_TRU:
            n_tp += 1
            info = dm[cty].get(ma)
            qc = info["QuyCach"] if info else None
            if qc:
                dvt_qd = info["DVTQuyDoi"]
                qd = [v[0] / qc, v[2] / qc, v[4] / qc, v[6] / qc]
            else:
                thieu.add(ma)
        rows.append([cty, tu.year, tu.month, ky, so_ngay, ma, str(row[i_ten] or "").strip(), loai,
                     str(row[i_tk] or "").strip() if i_tk is not None else "", str(row[i_dvt] or "").strip(),
                     *v, qc, dvt_qd, *qd])
    wb.close()
    log = [source_name, cty, tu.year, tu.month, dt.datetime.combine(tu, dt.time()), dt.datetime.combine(den, dt.time()),
           len(rows), lech, n_tp, ", ".join(sorted(thieu)), gio_vn()]
    return rows, log


def tong_hop_nhieu_file(sources, dm):
    data, log = [], []
    for name, src in sources:
        rows, lr = tong_hop_file(src, dm, name)
        cty, nam, thang = lr[1], lr[2], lr[3]
        data = [r for r in data if not (r[0] == cty and r[1] == nam and r[2] == thang)] + rows
        log = [r for r in log if not (r[1] == cty and r[2] == nam and r[3] == thang)] + [lr]
    return data, log


# ---------------------------------------------------------------- phân tích
def trang_thai(ton, xuat_ngay):
    if ton <= 0 and xuat_ngay <= 0:
        return TT_KHONG, None
    if xuat_ngay <= 0:
        return TT_CHAM, None
    ngay = ton / xuat_ngay
    if ton <= 0:
        return TT_HET, 0.0
    for nguong, ten in NGUONG:
        if ngay < nguong:
            return ten, ngay
    return TT_DU, ngay


def _sp_tp(data):
    """Gom dòng thành phẩm/hàng hóa có quy đổi theo đuôi mã -> {key: {(cty,(nam,thang)): row}}."""
    sp, nhan = {}, {}
    for r in data:
        if r[20] is None:
            continue
        key = r[5][1:]
        sp.setdefault(key, {})[(r[0], (r[1], r[2]))] = r
        if key not in nhan or r[0] == "Vikoda":
            nhan[key] = [r[6], r[19], r[18]]
    return sp, nhan


def _qd(kg, cty, k, ma, qc, i):
    v = kg.get((cty, k, ma)) if kg else None
    return (v[i] / qc) if (v is not None and qc) else 0


def _xuat_ngay(sp_k, cty, ky_list, den):
    """Xuất kế toán (= bán ra, quy đổi) bình quân/ngày của công ty trong SO_THANG_BQ tháng tính tới kỳ `den`."""
    ds = [k for k in ky_list if k <= den][-SO_THANG_BQ:]
    xuat = ngay = 0
    for k in ds:
        r = sp_k.get((cty, k))
        if r is not None:
            xuat += r[22] or 0
            ngay += r[4]
        else:
            ngay += 30
    return max(xuat, 0) / ngay if ngay else 0


HEAD_TTM = ["MaVikoda", "MaVKD", "TenSanPham", "DVTQuyDoi", "QuyCach",
            "TonKhaDungVikoda", "TonKhaDungVKD", "TongTonKhaDung", "TonKyGuiChoGiao", "TongTonVatLy",
            "XuatBanBQNgayVikoda", "XuatBanBQNgayVKD", "NhuCauNgayHeThong",
            "SoNgayTonKhaDung", "TrangThai", "GoiY",
            "NM_NgayBaoCao", "NM_TonThucTe", "NM_ChoQC", "NM_KeHoachXuat", "NM_DaBanGuiKho", "NM_TonDuKien",
            "Ky", "NguonTonKyGui"]
I_NGAY, I_TT = 13, 14
NM_KEYS = ["cuoi_ngay", "cho_qc", "kh_xuat", "da_ban", "du_kien"]


def phan_tich(data, kg=None, nguon_kg="", nha_may=None):
    """Bảng tồn kho thông minh tháng mới nhất + dữ liệu pivot.

    Trạng thái theo TỒN KHẢ DỤNG = tồn kế toán (hàng đã bán chờ giao ở kho ký gửi đã bị trừ),
    so với nhu cầu = xuất kế toán (bán ra) BQ/ngày. Tồn ký gửi & tồn vật lý chỉ để tham khảo.
    nha_may: {(nam, thang): báo cáo} từ nha_may.tong_hop_nha_may.
    """
    from .nha_may import theo_duoi_ma

    ky_list = sorted({(r[1], r[2]) for r in data})
    sp, nhan = _sp_tp(data)
    if not ky_list:
        return HEAD_TTM, [], ky_list, {}, {}, {}
    cuoi = ky_list[-1]
    nm_moi = theo_duoi_ma(nha_may[max(nha_may)]) if nha_may else {}
    ngay_nm = nha_may[max(nha_may)]["ngay"] if nha_may else None
    rows, p_ton, p_ngay, p_kg = [], {}, {}, {}
    for key, d in sp.items():
        qc = nhan[key][2]

        def ton(cty, k):
            r = d.get((cty, k))
            return (r[23] or 0) if r is not None else 0

        def ky_gui(k):
            return sum(_qd(kg, c, k, p + key, qc, 3) for c, p in (("Vikoda", "1"), ("VKD", "2")))

        def he_thong(k):
            tv, tk = ton("Vikoda", k), ton("VKD", k)
            xv, xk = _xuat_ngay(d, "Vikoda", ky_list, k), _xuat_ngay(d, "VKD", ky_list, k)
            # VKD là kênh bán ra thị trường; Vikoda chủ yếu xuất cho VKD -> nhu cầu hệ thống = xuất bán VKD,
            # mã chỉ có ở Vikoda thì dùng xuất Vikoda.
            return tv, tk, ky_gui(k), xv, xk, (xk if xk > 0 else xv)

        for k in ky_list:
            if any((c, k) in d for c in ("Vikoda", "VKD")):
                tv, tk, g, xv, xk, nc = he_thong(k)
                p_ton.setdefault(key, {})[k] = tv + tk
                p_kg.setdefault(key, {})[k] = g
                p_ngay.setdefault(key, {})[k] = (tv + tk) / nc if nc > 0 else None
        if not any((c, cuoi) in d for c in ("Vikoda", "VKD")):
            continue
        tv, tk, g, xv, xk, nc = he_thong(cuoi)
        tt, ngay = trang_thai(tv + tk, nc)
        nm = nm_moi.get(key)
        if tt == TT_KHONG and not (nm and nm.get("cuoi_ngay")):
            continue
        rows.append(["1" + key, "2" + key, nhan[key][0], nhan[key][1], qc,
                     _r(tv), _r(tk), _r(tv + tk), _r(g), _r(tv + tk + g),
                     _r(xv), _r(xk), _r(nc), _r(ngay), tt, GOI_Y[tt],
                     ngay_nm if nm else None] + [(_r(nm[k]) if nm else None) for k in NM_KEYS] +
                    ["T%02d/%d" % (cuoi[1], cuoi[0]), nguon_kg if g else ""])
    rows.sort(key=lambda r: (UU_TIEN[r[I_TT]], r[I_NGAY] if r[I_NGAY] is not None else 1e9, r[2]))
    return HEAD_TTM, rows, ky_list, p_ton, p_ngay, p_kg


def _r(x, n=2):
    return None if x is None else round(x, n)


def kiem_tra(data, log):
    """Danh sách cảnh báo: lệch cân đối, tồn cuối tháng trước khác tồn đầu tháng sau, thiếu quy cách."""
    out = []
    idx = {(r[0], r[1], r[2], r[5]): r for r in data}
    for r in data:
        if abs(r[10] + r[12] - r[14] - r[16]) > 0.001:
            out.append([r[0], "T%02d/%d" % (r[2], r[1]), r[5], r[6], "Đầu + Nhập - Xuất ≠ Cuối",
                        r[10] + r[12] - r[14], r[16]])
        nam, thang = (r[1], r[2] + 1) if r[2] < 12 else (r[1] + 1, 1)
        nx = idx.get((r[0], nam, thang, r[5]))
        if nx is not None and abs(nx[10] - r[16]) > 0.001:
            out.append([r[0], "T%02d/%d" % (thang, nam), r[5], r[6], "Tồn đầu ≠ tồn cuối tháng trước",
                        r[16], nx[10]])
    for lr in log:
        for ma in filter(None, (x.strip() for x in lr[9].split(","))):
            out.append([lr[1], "T%02d/%d" % (lr[3], lr[2]), ma, "", "Thiếu Quy cách trong Danh mục", None, None])
    return ["CongTy", "Ky", "MaVatTu", "TenVatTu", "VanDe", "GiaTriKyVong", "GiaTriThucTe"], out


# ---------------------------------------------------------------- ghi file
def build_workbook(data, log, fingerprint="", kg=None, nguon_kg="", nha_may=None, nm_bo_qua=()):
    thu_tu = {"Vikoda": 0, "VKD": 1}
    data = sorted(data, key=lambda r: (r[1], r[2], thu_tu.get(r[0], 9), r[7], r[5]))
    log = sorted(log, key=lambda r: (r[2], r[3], thu_tu.get(r[1], 9)))
    head, rows, ky_list, p_ton, p_ngay, p_kg = phan_tich(data, kg, nguon_kg, nha_may)
    _, nhan = _sp_tp(data)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ton thong minh"
    _ghi_bang(ws, "tblTonThongMinh", head, rows, "TableStyleMedium2")
    for i in range(2, len(rows) + 2):
        for j in list(range(6, I_TT + 1)) + list(range(18, 23)):
            ws.cell(i, j).number_format = "#,##0.0"
        ws.cell(i, 17).number_format = "dd/mm/yyyy"
        fill = MAU.get(ws.cell(i, I_TT + 1).value)
        if fill:
            ws.cell(i, I_TT + 1).fill = PatternFill("solid", fgColor=fill)
    for j, w in enumerate([11, 11, 44, 9, 8] + [12] * 9 + [16, 26] + [12] * 6 + [9, 18], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "D2"

    keys = sorted(p_ton, key=lambda k: (str(nhan[k][0]), k))
    ky_head = ["T%02d/%d" % (t, n) for n, t in ky_list]
    for title, tbl, src, fmt in (("Pivot Ton kha dung", "tblPivotTon", p_ton, "#,##0.0"),
                                 ("Pivot Ton ky gui", "tblPivotKyGui", p_kg, "#,##0.0"),
                                 ("Pivot So ngay ton", "tblPivotNgayTon", p_ngay, "#,##0.0")):
        wp = wb.create_sheet(title)
        prow = [["1" + k, "2" + k, nhan[k][0], nhan[k][1]] + [_r(src.get(k, {}).get(x)) for x in ky_list]
                for k in keys]
        _ghi_bang(wp, tbl, ["MaVikoda", "MaVKD", "TenSanPham", "DVTQuyDoi"] + ky_head, prow, "TableStyleMedium2")
        for i in range(2, len(prow) + 2):
            for j in range(5, 5 + len(ky_list)):
                wp.cell(i, j).number_format = fmt
        for j, w in enumerate([11, 11, 44, 9], start=1):
            wp.column_dimensions[get_column_letter(j)].width = w
        for j in range(5, 5 + len(ky_list)):
            wp.column_dimensions[get_column_letter(j)].width = 11
        wp.freeze_panes = "D2"
        if tbl == "tblPivotNgayTon" and prow and ky_list:
            rng = "E2:%s%d" % (get_column_letter(4 + len(ky_list)), len(prow) + 1)
            wp.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["7"],
                                                          fill=PatternFill("solid", fgColor=MAU["Thiếu hàng"])))
            wp.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["60"],
                                                          fill=PatternFill("solid", fgColor=MAU[TT_DU])))

    _ghi_nha_may(wb, nha_may or {})

    wkg = wb.create_sheet("Ky gui")
    ten = {}
    for r in data:
        ten.setdefault((r[0], r[5]), (r[6], r[18], r[19]))
    kg_rows = []
    for (cty, k, ma), (dau, vao, ra, cuoi_) in sorted((kg or {}).items(), key=lambda x: (x[0][1], x[0][0], x[0][2])):
        tn, qc, dvt = ten.get((cty, ma), ("", None, None))
        f = (lambda v: _r(v / qc)) if qc else (lambda v: None)
        kg_rows.append([cty, k[0], k[1], ma, tn, dvt, qc, dau, vao, ra, cuoi_, f(dau), f(vao), f(ra), f(cuoi_), nguon_kg])
    _ghi_bang(wkg, "tblKyGui", ["CongTy", "Nam", "Thang", "MaSanPham", "TenSanPham", "DVTQuyDoi", "QuyCach",
                                "DauKySL", "ChuyenVaoSL", "XuatRaSL", "CuoiKySL",
                                "DauKyQD", "ChuyenVaoQD", "XuatRaQD", "CuoiKyQD", "NguonDauKy"], kg_rows, "TableStyleLight9")
    for i in range(2, len(kg_rows) + 2):
        for j in range(8, 16):
            wkg.cell(i, j).number_format = "#,##0" if j < 12 else "#,##0.0"
    for j, w in enumerate([8, 6, 6, 11, 44, 9, 8] + [12] * 8 + [18], start=1):
        wkg.column_dimensions[get_column_letter(j)].width = w
    wkg.freeze_panes = "A2"

    wd = wb.create_sheet("Data")
    _ghi_bang(wd, "tblNXT", DATA_COLS, data, "TableStyleLight9")
    for i in range(2, len(data) + 2):
        wd.cell(i, 4).number_format = "mm/yyyy"
        for j in range(11, 19):
            wd.cell(i, j).number_format = "#,##0"
        for j in range(21, 25):
            wd.cell(i, j).number_format = "#,##0.00"
    for j, w in enumerate([8, 6, 6, 9, 8, 11, 44, 20, 10, 6] + [13] * 8 + [8, 8] + [12] * 4, start=1):
        wd.column_dimensions[get_column_letter(j)].width = w
    wd.freeze_panes = "A2"

    kh, kr = kiem_tra(data, log)
    kr += [["", "", "", name, "File trong 'Ton thuc te' không đúng mẫu BÁO CÁO KHO THÀNH PHẨM -> bỏ qua", None, None]
           for name in nm_bo_qua]
    wk = wb.create_sheet("Kiem tra")
    _ghi_bang(wk, "tblKiemTra", kh, kr, "TableStyleLight9")
    for j, w in enumerate([8, 9, 11, 40, 32, 14, 14], start=1):
        wk.column_dimensions[get_column_letter(j)].width = w

    wl = wb.create_sheet("Log")
    _ghi_bang(wl, "tblLogNXT", LOG_COLS, log, "TableStyleLight9")
    for i in range(2, len(log) + 2):
        wl.cell(i, 5).number_format = wl.cell(i, 6).number_format = "dd/mm/yyyy"
        wl.cell(i, 11).number_format = "dd/mm/yyyy hh:mm"
    for j, w in enumerate([24, 8, 6, 6, 11, 11, 8, 10, 9, 22, 17], start=1):
        wl.column_dimensions[get_column_letter(j)].width = w

    wg = wb.create_sheet("Huong dan")
    for line in [
        ["Tồn kho thông minh – cách tính"],
        ["Đơn vị: số lượng ÷ Quy cách (Danh mục) → Thùng/Két/Bình. Chỉ áp dụng Thành phẩm (Vikoda) / Hàng hóa (VKD)."],
        ["Tồn khả dụng = Tồn kế toán (Vikoda + VKD). Kế toán ghi 'Chuyển qua kho ký gửi' là xuất = hàng ĐÃ BÁN, "
         "nằm ở kho ký gửi chờ giao -> không còn khả dụng (giống nhà máy trừ 'SL hàng đã bán (gửi kho)')."],
        ["Tồn ký gửi chờ giao = Đầu kỳ + Chuyển vào ký gửi − Xuất từ KHO KÝ GỬI (báo cáo Xuất kho); đầu kỳ là ƯỚC TÍNH "
         "TỐI THIỂU. Tồn vật lý = Tồn khả dụng + Tồn ký gửi (tham khảo)."],
        [f"Nhu cầu = Xuất kế toán (bán ra) BQ/ngày {SO_THANG_BQ} tháng gần nhất của VKD; mã chỉ có ở Vikoda dùng xuất Vikoda."],
        ["Số ngày tồn khả dụng = Tồn khả dụng ÷ Nhu cầu/ngày; trạng thái theo số này."],
        ["Cột NM_*: báo cáo kho thành phẩm nhà máy mới nhất (thư mục Ton thuc te): tồn thực tế cuối ngày, chờ QA/QC, "
         "kế hoạch xuất, đã bán gửi kho, tồn dự kiến = thực tế − kế hoạch xuất − đã bán."],
        ["Trạng thái: Hết hàng (tồn ≤ 0) | Thiếu hàng < 7 ngày | Thấp 7–15 | An toàn 15–30 | Cao 30–60 | Dư tồn > 60 | "
         "Không luân chuyển (còn tồn nhưng không xuất)."],
        ["Lưu ý: tháng đang mở (chưa tính giá) có Thành tiền tồn cuối = 0 → dùng số lượng, không dùng giá trị của tháng đó."],
    ]:
        wg.append(line)
    wg.column_dimensions["A"].width = 120

    wm = wb.create_sheet("Meta")
    wm["A1"], wm["B1"] = "fingerprint", fingerprint
    wm.sheet_state = "hidden"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(rows), len(ky_list), len(kr)


NM_HEAD = ["MaVikoda", "MaVKD", "TenSanPham", "DVT", "NgayBaoCao", "NhapTrongThang", "XuatTrongThang",
           "TonDauNgay", "TonThucTe", "ChoQC", "KeHoachXuat", "DaBanGuiKho", "TonDuKien", "SourceFile"]
NM_KEY = ["nhap_thang", "xuat_thang", "dau_ngay", "cuoi_ngay", "cho_qc", "kh_xuat", "da_ban", "du_kien"]


def _ghi_nha_may(wb, nha_may):
    """Sheet 'Ton nha may' (mọi báo cáo theo tháng, mới nhất trước) + 'Pivot Ton nha may' (tồn thực tế cuối ngày)."""
    ws = wb.create_sheet("Ton nha may")
    rows = []
    for k in sorted(nha_may, reverse=True):
        bc = nha_may[k]
        for d in bc["rows"]:
            rows.append([d["ma_vik"], d["ma_vkd"], d["ten"], d["dvt"], bc["ngay"]] +
                        [d[x] for x in NM_KEY] + [bc["source"]])
    _ghi_bang(ws, "tblTonNhaMay", NM_HEAD, rows, "TableStyleMedium2")
    for i in range(2, len(rows) + 2):
        ws.cell(i, 5).number_format = "dd/mm/yyyy"
        for j in range(6, 14):
            ws.cell(i, j).number_format = "#,##0"
    for j, w in enumerate([11, 11, 44, 7, 11] + [12] * 8 + [18], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "D2"

    wp = wb.create_sheet("Pivot Ton nha may")
    ky = sorted(nha_may)
    sp = {}
    for k in ky:
        for d in nha_may[k]["rows"]:
            ma = d["ma_vik"] or d["ma_vkd"]
            sp.setdefault(ma[1:], {"nhan": [d["ma_vik"], d["ma_vkd"], d["ten"], d["dvt"]]})[k] = d["cuoi_ngay"]
    head = ["MaVikoda", "MaVKD", "TenSanPham", "DVT"] + [nha_may[k]["ngay"].strftime("%d/%m/%Y") for k in ky]
    prow = [v["nhan"] + [v.get(k) for k in ky] for _, v in sorted(sp.items(), key=lambda x: (x[1]["nhan"][2], x[0]))]
    _ghi_bang(wp, "tblPivotNhaMay", head, prow, "TableStyleMedium2")
    for i in range(2, len(prow) + 2):
        for j in range(5, 5 + len(ky)):
            wp.cell(i, j).number_format = "#,##0"
    for j, w in enumerate([11, 11, 44, 7] + [12] * len(ky), start=1):
        wp.column_dimensions[get_column_letter(j)].width = w
    wp.freeze_panes = "D2"
