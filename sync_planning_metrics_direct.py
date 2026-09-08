import hashlib
import json
from io import BytesIO

from openpyxl import load_workbook

import sync_planning_metrics as metrics
import sync_planning_metrics_compat as compat


DEBT_SHEET = "No kho"
DEBT_CODE_COL = 1   # A - Mã Sản Phẩm
DEBT_VALUE_COL = 4  # D - Số lượng nợ


def read_debt_from_no_kho(dest_bytes):
    """Đọc Nợ kho trực tiếp từ No kho!D theo Mã Sản Phẩm ở cột A."""
    workbook = load_workbook(
        BytesIO(dest_bytes),
        data_only=True,
        read_only=True,
    )

    try:
        if DEBT_SHEET not in workbook.sheetnames:
            raise RuntimeError(f"Không tìm thấy sheet {DEBT_SHEET!r}.")

        worksheet = workbook[DEBT_SHEET]
        code_header = worksheet.cell(row=1, column=DEBT_CODE_COL).value
        debt_header = worksheet.cell(row=1, column=DEBT_VALUE_COL).value

        if str(code_header or "").strip().casefold() != "mã sản phẩm".casefold():
            raise RuntimeError(
                f"{DEBT_SHEET}!A1 phải là 'Mã Sản Phẩm', hiện là {code_header!r}."
            )
        if str(debt_header or "").strip().casefold() != "số lượng nợ".casefold():
            raise RuntimeError(
                f"{DEBT_SHEET}!D1 phải là 'Số lượng nợ', hiện là {debt_header!r}."
            )

        debts = {}
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, max_col=DEBT_VALUE_COL, values_only=True),
            start=2,
        ):
            code = metrics.normalize_code(values[DEBT_CODE_COL - 1] if values else None)
            if not code:
                continue

            if code in debts:
                raise RuntimeError(
                    f"Mã {code} bị lặp trong {DEBT_SHEET}!A."
                )

            raw_debt = values[DEBT_VALUE_COL - 1]
            debt = 0 if raw_debt in (None, "") else metrics.to_number(
                raw_debt,
                f"{DEBT_SHEET}!D{row_number}",
            )
            debts[code] = metrics.clean_number(debt)

        if not debts:
            raise RuntimeError(
                f"Không đọc được dữ liệu từ {DEBT_SHEET}!A:D."
            )

        print(
            f"[{DEBT_SHEET}] Đọc {len(debts)} mã; Nợ kho lấy trực tiếp từ cột D."
        )
        return debts
    finally:
        workbook.close()


def hash_debt_sheet(dest_bytes):
    """Tính SHA-256 fingerprint của sheet No kho dựa trên mã sản phẩm và số lượng nợ."""
    workbook = load_workbook(
        BytesIO(dest_bytes),
        data_only=True,
        read_only=True,
    )
    try:
        if DEBT_SHEET not in workbook.sheetnames:
            return {}, ""
    finally:
        workbook.close()

    debts = read_debt_from_no_kho(dest_bytes)
    payload = json.dumps(
        {k: debts[k] for k in sorted(debts)},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return debts, hashlib.sha256(payload).hexdigest()


def read_planning_rows_with_no_kho(dest_bytes):
    selector, plan_month, rows = compat.read_planning_rows_robust(dest_bytes)
    debts = read_debt_from_no_kho(dest_bytes)

    missing = sorted(set(rows) - set(debts))
    if missing:
        raise RuntimeError(
            f"Thiếu {len(missing)} mã Ke_hoach_SX trong {DEBT_SHEET}!A: "
            + ", ".join(missing[:10])
        )

    for code, row in rows.items():
        row["current_debt"] = debts[code]

    return selector, plan_month, rows


def clamp_nonnegative_production(output):
    """O và P chỉ nhận giá trị > 0; còn lại trả 0. Q = 0 khi P = 0."""
    output = dict(output)

    required = float(output.get("required_production", 0) or 0)
    rounded = float(output.get("rounded_production", 0) or 0)

    if required <= 0:
        output["required_production"] = 0
        output["rounded_production"] = 0
        output["production_days"] = 0
        return output

    output["required_production"] = metrics.clean_number(required)

    if rounded <= 0:
        output["rounded_production"] = 0
        output["production_days"] = 0
    else:
        output["rounded_production"] = metrics.clean_number(rounded)

    return output


def calculate_metrics_from_no_kho(
    *,
    report_date,
    plan_month,
    planning_rows,
    actual_receipts,
    system_receipts,
    current_consignments,
    state,
):
    """Tính M:R, trong đó N lấy nguyên giá trị No kho!D theo mã sản phẩm."""
    source_key = metrics._month_key(report_date.year, report_date.month)
    plan_year = metrics._resolve_plan_year(report_date, plan_month)
    plan_key = metrics._month_key(plan_year, plan_month)
    state_changed = False

    # M vẫn dùng snapshot gửi kho đầu kỳ như thuật toán hiện tại.
    consign_state = state.setdefault("opening_consignment_by_month", {})
    opening_consignment = consign_state.get(plan_key)
    if opening_consignment is None:
        opening_consignment = {
            code: metrics.clean_number(current_consignments.get(code, 0))
            for code in planning_rows
        }
        consign_state[plan_key] = opening_consignment
        state_changed = True

    if metrics._is_month_end(report_date):
        next_year, next_month = metrics._next_month(
            report_date.year,
            report_date.month,
        )
        next_key = metrics._month_key(next_year, next_month)
        next_consignment = {
            code: metrics.clean_number(current_consignments.get(code, 0))
            for code in planning_rows
        }
        if consign_state.get(next_key) != next_consignment:
            consign_state[next_key] = next_consignment
            state_changed = True

    result = {}
    for code, row in planning_rows.items():
        leadtime = metrics.LEADTIME_BY_CODE.get(code)
        if leadtime is None:
            raise RuntimeError(f"Chưa có Leadtime Danh_muc!J cho mã {code}.")

        warehouse_debt = float(row["current_debt"] or 0)
        calculated = metrics.calculate_row(
            fc=row["fc"],
            actual_stock=row["actual_stock"],
            book_stock=row["book_stock"],
            opening_consignment=float(opening_consignment.get(code, 0) or 0),
            warehouse_debt=warehouse_debt,
            leadtime=leadtime,
            batch=row["batch"],
            per_shift=row["per_shift"],
            shifts_per_day=row["shifts_per_day"],
            classification=row["classification"],
            plan_year=plan_year,
            plan_month=plan_month,
        )
        result[code] = clamp_nonnegative_production(calculated)

    return result, state_changed, source_key, plan_key


# Runtime production: nguồn sự thật cho N là No kho!D.
metrics.read_planning_rows = read_planning_rows_with_no_kho
metrics.calculate_metrics = calculate_metrics_from_no_kho


if __name__ == "__main__":
    compat.main_with_retry()
