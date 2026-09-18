"""Read and reconcile NVL workbook data without SharePoint I/O."""

from __future__ import annotations

from io import BytesIO
from typing import Any
import zipfile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from excel.openpyxl_io import safe_close_workbook
from excel.workbook_xml import column_number, find_sheet_xml_path
from nvl.models import NVLConfig, NVLReconcileResult
from nvl.values import EPS, normalize_nvl_code, parse_nvl_quantity

def check_target_sheet_safety(
    target_bytes: bytes,
    sheet_name: str,
    target_col_letter: str,
    target_start_row: int,
) -> None:
    """Kiểm tra an toàn cấu trúc sheet đích trước khi patch:

    - Kiểm tra bảo vệ sheet (<sheetProtection>).
    - Kiểm tra gộp ô (<mergeCell>): không cho phép cột đích nằm trong vùng gộp ô.
    """
    target_col = column_number(target_col_letter)
    buf = BytesIO(target_bytes)
    z = zipfile.ZipFile(buf, "r")
    try:
        sheet_path = find_sheet_xml_path(z, sheet_name)
        sheet_root = etree.fromstring(z.read(sheet_path))

        # 1. Kiểm tra sheetProtection
        for prot in sheet_root.xpath('//*[local-name()="sheetProtection"]'):
            is_prot = any(
                prot.get(attr) in ("1", "true")
                for attr in ("sheet", "objects", "scenarios")
            )
            if is_prot:
                raise RuntimeError(
                    f"[Đích] Sheet '{sheet_name}' đang được bật bảo vệ (sheetProtection). "
                    "Không thể ghi đè dữ liệu tồn kho."
                )

        # 2. Kiểm tra mergeCell phủ lên cột đích
        for m in sheet_root.xpath('//*[local-name()="mergeCell"]'):
            ref = m.get("ref", "")
            if not ref:
                continue
            try:
                min_c, min_r, max_c, max_r = range_boundaries(ref)
                if min_c <= target_col <= max_c and max_r >= target_start_row:
                    raise RuntimeError(
                        f"[Đích] Ô trong vùng {sheet_name}!{target_col_letter}{max(min_r, target_start_row)} "
                        f"nằm trong dải ô gộp '{ref}'. Cột {target_col_letter} phải là cột đơn lẻ, "
                        "không được gộp ô."
                    )
            except Exception as exc:
                if "nằm trong dải ô gộp" in str(exc):
                    raise
    finally:
        z.close()
        z.fp = None


def read_nvl_source_stock(
    source_bytes: bytes,
    config: NVLConfig,
) -> tuple[dict[str, float], dict[str, Any]]:
    """Đọc dữ liệu số tồn kho từ file nguồn XNT_ketoan_Vikoda.xlsm!Sheet1.

    Sử dụng streaming reader qua iter_rows() sau khi reset_dimensions() để
    hoàn toàn miễn nhiễm với dimension bị thiếu hoặc khai báo hụt.
    """
    wb_val = load_workbook(BytesIO(source_bytes), data_only=True, read_only=True, keep_vba=True)
    wb_raw = load_workbook(BytesIO(source_bytes), data_only=False, read_only=True, keep_vba=True)

    try:
        if config.source_sheet not in wb_val.sheetnames:
            raise RuntimeError(
                f"[Nguồn: {config.source_name}] Không tìm thấy sheet '{config.source_sheet}'."
            )

        ws_val = wb_val[config.source_sheet]
        ws_raw = wb_raw[config.source_sheet]

        # Xóa dimension cache để đọc trọn vẹn dữ liệu từ sheet XML
        ws_val.reset_dimensions()
        ws_raw.reset_dimensions()

        reporting_period = None
        for r_hdr in range(1, min(config.source_start_row, 15)):
            c_val = ws_val.cell(r_hdr, 1).value
            if isinstance(c_val, str):
                c_str = c_val.strip()
                if any(k in c_str.lower() for k in ("từ ngày", "đến ngày", "kỳ", "tháng")):
                    reporting_period = c_str
                    break

        stock: dict[str, float] = {}
        seen_rows: dict[str, int] = {}
        total_scanned = 0

        val_stream = ws_val.iter_rows(min_row=config.source_start_row, values_only=True)
        raw_stream = ws_raw.iter_rows(min_row=config.source_start_row, values_only=True)

        code_idx = config.source_code_col - 1
        val_idx = config.source_value_col - 1

        for r, (val_row, raw_row) in enumerate(zip(val_stream, raw_stream), start=config.source_start_row):
            total_scanned += 1
            if not val_row or len(val_row) <= code_idx:
                continue

            raw_code = val_row[code_idx]
            if isinstance(raw_code, str):
                lower_code = raw_code.strip().lower()
                if any(k in lower_code for k in ("tổng cộng", "tong cong", "total", "cộng")):
                    continue

            code = normalize_nvl_code(raw_code)
            if not code:
                continue

            if config.reject_duplicates and code in stock:
                raise RuntimeError(
                    f"[Nguồn: {config.source_name}] Mã vật tư '{code}' bị lặp tại dòng {r} "
                    f"(đã xuất hiện tại dòng {seen_rows[code]})."
                )

            seen_rows[code] = r
            cell_ref = f"{config.source_sheet}!{config.source_value_col_letter}{r}"

            c_val = val_row[val_idx] if len(val_row) > val_idx else None
            c_raw = raw_row[val_idx] if len(raw_row) > val_idx else None

            if isinstance(c_raw, str) and c_raw.startswith("="):
                if c_val is None:
                    raise RuntimeError(
                        f"[Nguồn: {config.source_name}] Ô {cell_ref} chứa công thức '{c_raw}' "
                        f"nhưng không có giá trị cached (chưa được Excel tính toán và lưu)."
                    )

            qty = parse_nvl_quantity(
                c_val,
                convention=config.number_convention,
                cell_name=cell_ref,
            )

            if not config.allow_negative and qty < 0:
                raise ValueError(f"[Nguồn: {config.source_name}] Ô {cell_ref} có số lượng âm ({qty}) bị từ chối.")
            if not config.allow_zero and abs(qty) < EPS:
                raise ValueError(f"[Nguồn: {config.source_name}] Ô {cell_ref} có số lượng 0 bị từ chối.")

            stock[code] = qty

        if not stock:
            raise RuntimeError(
                f"[Nguồn: {config.source_name}] Không đọc được bất kỳ mã vật tư hợp lệ nào "
                f"trong {config.source_sheet}!{config.source_code_col_letter}."
            )

        metadata = {
            "source_name": config.source_name,
            "sheet_name": config.source_sheet,
            "total_scanned_rows": total_scanned,
            "valid_codes_count": len(stock),
            "seen_rows": seen_rows,
            "reporting_period": reporting_period,
        }
        return stock, metadata
    finally:
        try:
            val_stream.close()
        except Exception:
            pass
        try:
            raw_stream.close()
        except Exception:
            pass
        safe_close_workbook(wb_val)
        safe_close_workbook(wb_raw)

def reconcile_nvl_target(
    target_bytes: bytes,
    source_stock: dict[str, float],
    config: NVLConfig,
) -> NVLReconcileResult:
    """Đối chiếu mã vật tư và lập danh sách thay đổi cho Kế hoạch mua hàng.xlsx!Ton_NVL.

    Sử dụng streaming reader với reset_dimensions() và kiểm tra cấu trúc an toàn
    (bảo vệ sheet, gộp ô, công thức người dùng).
    """
    # 1. Kiểm tra an toàn cấu trúc sheet đích
    check_target_sheet_safety(
        target_bytes,
        config.target_sheet,
        config.target_value_col_letter,
        config.target_start_row,
    )

    # 2. Đọc dữ liệu đích
    wb_val = load_workbook(BytesIO(target_bytes), data_only=True, read_only=True)
    wb_raw = load_workbook(BytesIO(target_bytes), data_only=False, read_only=True)

    try:
        if config.target_sheet not in wb_val.sheetnames:
            raise RuntimeError(
                f"[Đích: {config.target_name}] Không tìm thấy sheet '{config.target_sheet}'."
            )

        ws_val = wb_val[config.target_sheet]
        ws_raw = wb_raw[config.target_sheet]

        ws_val.reset_dimensions()
        ws_raw.reset_dimensions()

        seen_target_codes: dict[str, int] = {}
        changes: list[dict[str, Any]] = []
        unchanged: list[dict[str, Any]] = []
        missing_in_source: list[dict[str, Any]] = []

        val_stream = ws_val.iter_rows(min_row=config.target_start_row, values_only=True)
        raw_stream = ws_raw.iter_rows(min_row=config.target_start_row, values_only=True)

        code_idx = config.target_code_col - 1
        val_idx = config.target_value_col - 1

        for r, (val_row, raw_row) in enumerate(zip(val_stream, raw_stream), start=config.target_start_row):
            if not val_row or len(val_row) <= code_idx:
                continue

            raw_code = val_row[code_idx]
            if isinstance(raw_code, str):
                lower_code = raw_code.strip().lower()
                if any(k in lower_code for k in ("tổng cộng", "tong cong", "total", "cộng")):
                    continue

            code = normalize_nvl_code(raw_code)
            if not code:
                continue

            if config.reject_duplicates and code in seen_target_codes:
                raise RuntimeError(
                    f"[Đích: {config.target_name}] Mã vật tư '{code}' bị lặp trong "
                    f"{config.target_sheet}!{config.target_code_col_letter} tại dòng {r} "
                    f"(đã xuất hiện tại dòng {seen_target_codes[code]})."
                )

            seen_target_codes[code] = r
            cell_ref = f"{config.target_sheet}!{config.target_value_col_letter}{r}"

            target_val = val_row[val_idx] if len(val_row) > val_idx else None
            target_raw = raw_row[val_idx] if len(raw_row) > val_idx else None

            if config.reject_formula_in_target_cell and isinstance(target_raw, str) and target_raw.startswith("="):
                raise RuntimeError(
                    f"[Đích: {config.target_name}] Ô {cell_ref} chứa công thức '{target_raw}'. "
                    f"Cột {config.target_value_col_letter} phải là cột nhập tồn, không được ghi đè lên công thức người dùng."
                )

            if code not in source_stock:
                if not config.preserve_missing_in_source:
                    raise RuntimeError(
                        f"[Đích: {config.target_name}] Mã vật tư '{code}' không có trong file nguồn và "
                        f"cấu hình 'preserve_missing_in_source' đang tắt."
                    )
                missing_in_source.append({
                    "row": r,
                    "code": code,
                    "current_value": target_val,
                })
                continue

            new_qty = source_stock[code]
            needs_change = False

            if target_val is None or (isinstance(target_val, str) and target_val.strip() == ""):
                needs_change = True
            else:
                try:
                    curr_float = parse_nvl_quantity(
                        target_val,
                        convention=config.number_convention,
                        cell_name=cell_ref,
                    )
                    if abs(curr_float - new_qty) > EPS:
                        needs_change = True
                except Exception:
                    # Nếu ô đích hiện tại chứa text không parse được thành số hợp lệ -> cần cập nhật
                    needs_change = True

            if needs_change:
                changes.append({
                    "row": r,
                    "code": code,
                    "before": target_val,
                    "after": new_qty,
                })
            else:
                unchanged.append({
                    "row": r,
                    "code": code,
                    "value": target_val,
                })

        if not seen_target_codes:
            raise RuntimeError(
                f"[Đích: {config.target_name}] Không đọc được bất kỳ mã vật tư hợp lệ nào trong "
                f"{config.target_sheet}!{config.target_code_col_letter}."
            )

        matched_count = len(changes) + len(unchanged)
        if matched_count == 0:
            raise RuntimeError(
                f"[Đích: {config.target_name}] Không có mã vật tư nào trong "
                f"{config.target_sheet}!{config.target_code_col_letter} khớp với nguồn {config.source_name}. Bị chặn cập nhật."
            )

        source_only = [c for c in source_stock if c not in seen_target_codes]

        if missing_in_source:
            status = "completed_with_warnings"
            msg = f"Đối soát hoàn tất: {len(changes)} ô cần cập nhật, còn {len(missing_in_source)} mã đích không có trong nguồn."
        elif changes:
            status = "success"
            msg = f"Đối soát thành công: {len(changes)} ô cần cập nhật, {len(unchanged)} ô không đổi."
        else:
            status = "unchanged"
            msg = "Dữ liệu tồn kho khớp hoàn toàn, không có ô nào cần thay đổi."

        return NVLReconcileResult(
            changes=changes,
            unchanged=unchanged,
            missing_in_source=missing_in_source,
            source_only=source_only,
            target_codes=seen_target_codes,
            status=status,
            message=msg,
        )
    finally:
        try:
            val_stream.close()
        except Exception:
            pass
        try:
            raw_stream.close()
        except Exception:
            pass
        safe_close_workbook(wb_val)
        safe_close_workbook(wb_raw)


__all__ = [
    "check_target_sheet_safety",
    "read_nvl_source_stock",
    "reconcile_nvl_target",
]
