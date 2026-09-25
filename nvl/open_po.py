"""Đồng bộ Tồn Đơn hàng (PO còn phải nhận) vào Ton_NVL!E.

Nguồn: BCTheodoiDMHANG.xlsm / REPORT_DONMUAHANG
  - F: Mã hàng
  - I: Số lượng mua
  - K: Số lượng nhận
  - AB: Trạng thái

Quy tắc nghiệp vụ:
  - Chỉ lấy dòng có AB trống (đơn chưa hoàn tất), tương đương VBA Lay_PO.
  - PO còn phải nhận từng dòng = MAX(I - K, 0).
  - Cộng theo mã hàng F.
  - Ghép Ton_NVL!A và ghi kết quả vào Ton_NVL!E (Tồn Đơn hàng).
  - Mã Ton_NVL không có PO mở được ghi 0 để tránh giữ số cũ.

File đích được patch trực tiếp XML của đúng sheet, không round-trip toàn workbook.
"""

from __future__ import annotations

import argparse
import hashlib
from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
import json
import math
from pathlib import Path
from urllib.parse import quote
import zipfile

from lxml import etree
from openpyxl import load_workbook

from nvl.values import normalize_nvl_code, parse_nvl_quantity
from excel.openpyxl_io import safe_close_workbook
from excel.workbook_xml import find_sheet_xml_path, set_numeric_cell
from sharepoint.client import GRAPH, GraphClient, get_access_token

DEFAULT_SOURCE_CONFIG = "nvl_open_po_config.json"
DEFAULT_TARGET_CONFIG = "nvl_stock_config.json"
DEFAULT_PROPOSAL = "nvl_open_po_proposal.xlsx"
DEFAULT_REPORT = "nvl_open_po_report.json"
EPS = 1e-6

__all__ = [
    "OpenPOConfig",
    "load_open_po_config",
    "patch_target_workbook",
    "read_open_po",
    "reconcile_target",
    "run_online",
    "verify_target",
]


@dataclass(frozen=True)
class OpenPOConfig:
    source_name: str
    source_sourcedoc: str
    source_sheet: str
    source_start_row: int
    source_code_col: int
    source_purchase_col: int
    source_received_col: int
    source_status_col: int

    target_name: str
    target_path: str
    target_sourcedoc: str
    target_sheet: str
    target_start_row: int
    target_code_col: int
    target_value_col: int


def _normalize_guid(value: str | None) -> str:
    return str(value or "").strip().strip("{}").replace("-", "").lower()


def load_open_po_config(
    source_config_path: str | Path = DEFAULT_SOURCE_CONFIG,
    target_config_path: str | Path = DEFAULT_TARGET_CONFIG,
) -> OpenPOConfig:
    with open(source_config_path, "r", encoding="utf-8") as f:
        src_data = json.load(f)
    with open(target_config_path, "r", encoding="utf-8") as f:
        tgt_data = json.load(f)

    src = src_data.get("source", {})
    tgt = tgt_data.get("target", {})

    return OpenPOConfig(
        source_name=str(src.get("name", "BCTheodoiDMHANG.xlsm")),
        source_sourcedoc=str(src.get("sourcedoc", "")),
        source_sheet=str(src.get("sheet_name", "REPORT_DONMUAHANG")),
        source_start_row=int(src.get("start_row", 6)),
        source_code_col=int(src.get("code_column", 6)),
        source_purchase_col=int(src.get("purchase_quantity_column", 9)),
        source_received_col=int(src.get("received_quantity_column", 11)),
        source_status_col=int(src.get("status_column", 28)),
        target_name=str(tgt.get("name", "Kế hoạch mua hàng.xlsx")),
        target_path=str(tgt.get("sharepoint_path", "")),
        target_sourcedoc=str(tgt.get("sourcedoc", "")),
        target_sheet=str(tgt.get("sheet_name", "Ton_NVL")),
        target_start_row=int(tgt.get("start_row", 2)),
        target_code_col=int(tgt.get("code_column", 1)),
        target_value_col=5,
    )


def _qty(value, *, cell_name: str, blank_as_zero: bool = False) -> float:
    if value is None or (isinstance(value, str) and value.strip() == ""):
        if blank_as_zero:
            return 0.0
        raise ValueError(f"[{cell_name}] Số lượng mua đang rỗng.")
    return parse_nvl_quantity(value, convention="strict", cell_name=cell_name)


def read_open_po(source_bytes: bytes, config: OpenPOConfig) -> tuple[dict[str, float], dict]:
    """Đọc PO mở và trả về tổng lượng còn phải nhận theo mã vật tư."""
    wb = load_workbook(BytesIO(source_bytes), data_only=True, read_only=True, keep_vba=True)
    try:
        if config.source_sheet not in wb.sheetnames:
            raise RuntimeError(
                f"Không tìm thấy sheet {config.source_sheet!r} trong {config.source_name}."
            )
        ws = wb[config.source_sheet]
        max_col = max(
            config.source_code_col,
            config.source_purchase_col,
            config.source_received_col,
            config.source_status_col,
        )

        totals: dict[str, float] = {}
        open_rows = 0
        positive_rows = 0
        fully_received_or_over_rows = 0
        over_received_rows = 0
        skipped_no_code = 0

        for row_num, row in enumerate(
            ws.iter_rows(
                min_row=config.source_start_row,
                max_col=max_col,
                values_only=True,
            ),
            start=config.source_start_row,
        ):
            status = row[config.source_status_col - 1]
            if status is not None and str(status).strip() != "":
                continue

            code = normalize_nvl_code(row[config.source_code_col - 1])
            if not code:
                skipped_no_code += 1
                continue

            purchase = _qty(
                row[config.source_purchase_col - 1],
                cell_name=f"{config.source_sheet}!I{row_num}",
            )
            received = _qty(
                row[config.source_received_col - 1],
                cell_name=f"{config.source_sheet}!K{row_num}",
                blank_as_zero=True,
            )
            raw_remaining = purchase - received
            remaining = max(raw_remaining, 0.0)

            open_rows += 1
            if raw_remaining > EPS:
                positive_rows += 1
            else:
                fully_received_or_over_rows += 1
                if raw_remaining < -EPS:
                    over_received_rows += 1

            totals[code] = totals.get(code, 0.0) + remaining

        return totals, {
            "open_rows": open_rows,
            "source_codes": len(totals),
            "positive_rows": positive_rows,
            "fully_received_or_over_rows": fully_received_or_over_rows,
            "over_received_rows": over_received_rows,
            "skipped_no_code": skipped_no_code,
            "total_open_po": sum(totals.values()),
        }
    finally:
        safe_close_workbook(wb)


def _numeric_equal(current, target: float) -> bool:
    if current in (None, ""):
        return False
    if isinstance(current, bool):
        return False
    try:
        current_num = float(current)
    except (TypeError, ValueError):
        return False
    return math.isclose(current_num, float(target), rel_tol=1e-12, abs_tol=EPS)


def reconcile_target(
    target_bytes: bytes,
    open_po: dict[str, float],
    config: OpenPOConfig,
) -> tuple[list[dict], dict[str, float]]:
    """Đối chiếu Ton_NVL và lập patch cho cột E."""
    wb = load_workbook(BytesIO(target_bytes), data_only=False, read_only=False)
    try:
        if config.target_sheet not in wb.sheetnames:
            raise RuntimeError(
                f"Không tìm thấy sheet {config.target_sheet!r} trong {config.target_name}."
            )
        ws = wb[config.target_sheet]

        expected_header = "tồn đơn hàng"
        header = str(ws.cell(1, config.target_value_col).value or "").strip().lower()
        if header != expected_header:
            raise RuntimeError(
                f"Header {config.target_sheet}!E1 không đúng. "
                f"Mong đợi 'Tồn Đơn hàng', nhận {ws.cell(1, config.target_value_col).value!r}."
            )

        for merged_range in ws.merged_cells.ranges:
            if (
                merged_range.min_col <= config.target_value_col <= merged_range.max_col
                and merged_range.max_row >= config.target_start_row
            ):
                raise RuntimeError(
                    f"Cột E có merged cell ({merged_range}); từ chối patch để bảo toàn workbook."
                )

        changes: list[dict] = []
        expected: dict[str, float] = {}
        seen: set[str] = set()

        for row_num in range(config.target_start_row, ws.max_row + 1):
            code = normalize_nvl_code(ws.cell(row_num, config.target_code_col).value)
            if not code:
                continue
            if code in seen:
                raise RuntimeError(
                    f"Mã NVL {code} bị lặp trong {config.target_sheet}!A."
                )
            seen.add(code)

            cell = ws.cell(row_num, config.target_value_col)
            if cell.data_type == "f" or (
                isinstance(cell.value, str) and cell.value.startswith("=")
            ):
                raise RuntimeError(
                    f"{config.target_sheet}!E{row_num} đang chứa công thức; từ chối ghi đè."
                )

            target_value = float(open_po.get(code, 0.0))
            expected[code] = target_value
            if not _numeric_equal(cell.value, target_value):
                changes.append(
                    {
                        "row": row_num,
                        "code": code,
                        "current": cell.value,
                        "target": target_value,
                    }
                )

        if not expected:
            raise RuntimeError(f"Không đọc được mã NVL trong {config.target_sheet}!A.")

        return changes, expected
    finally:
        safe_close_workbook(wb)


def patch_target_workbook(
    target_bytes: bytes,
    changes: list[dict],
    config: OpenPOConfig,
) -> bytes:
    if not changes:
        return target_bytes

    source_buffer = BytesIO(target_bytes)
    output_buffer = BytesIO()
    changes_by_row = {
        int(item["row"]): float(item["target"])
        for item in changes
    }

    with zipfile.ZipFile(source_buffer, "r") as source_zip:
        sheet_path = find_sheet_xml_path(source_zip, config.target_sheet)
        root = etree.fromstring(source_zip.read(sheet_path))
        row_nodes = {
            int(node.get("r")): node
            for node in root.xpath('//*[local-name()="sheetData"]/*[local-name()="row"]')
            if node.get("r")
        }

        for row_num, value in changes_by_row.items():
            row_node = row_nodes.get(row_num)
            if row_node is None:
                raise RuntimeError(
                    f"Không tìm thấy XML row {row_num} trong sheet {config.target_sheet}."
                )
            set_numeric_cell(row_node, row_num, "E", value)

        new_xml = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        with zipfile.ZipFile(
            output_buffer,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as output_zip:
            for item in source_zip.infolist():
                data = new_xml if item.filename == sheet_path else source_zip.read(item.filename)
                output_zip.writestr(item, data)

    return output_buffer.getvalue()


def verify_target(
    target_bytes: bytes,
    expected: dict[str, float],
    config: OpenPOConfig,
) -> None:
    wb = load_workbook(BytesIO(target_bytes), data_only=True, read_only=True)
    try:
        ws = wb[config.target_sheet]
        checked = 0
        for row in ws.iter_rows(
            min_row=config.target_start_row,
            max_col=config.target_value_col,
            values_only=True,
        ):
            code = normalize_nvl_code(row[config.target_code_col - 1])
            if not code or code not in expected:
                continue
            current = row[config.target_value_col - 1]
            if not _numeric_equal(current, expected[code]):
                raise RuntimeError(
                    f"Verify thất bại cho mã {code}: "
                    f"E={current!r}, mong đợi {expected[code]}."
                )
            checked += 1
        if checked != len(expected):
            raise RuntimeError(
                f"Verify chỉ thấy {checked}/{len(expected)} mã trong {config.target_sheet}."
            )
    finally:
        safe_close_workbook(wb)


def _item_metadata(graph: GraphClient, drive_id: str, item_id: str) -> dict:
    return graph.get_json(
        f"{GRAPH}/drives/{drive_id}/items/{item_id}"
        "?$select=id,name,eTag,cTag,size,webUrl,lastModifiedDateTime,sharepointIds,parentReference"
    )


def resolve_source_item(
    graph: GraphClient,
    drive_id: str,
    config: OpenPOConfig,
) -> dict:
    """Tìm file nguồn theo tên rồi khóa identity bằng sourcedoc GUID."""
    encoded_name = quote(config.source_name, safe="")
    result = graph.get_json(
        f"{GRAPH}/drives/{drive_id}/root/search(q='{encoded_name}')"
        "?$select=id,name,eTag,cTag,size,webUrl,lastModifiedDateTime,sharepointIds,parentReference"
    )
    exact = [
        item
        for item in result.get("value", [])
        if str(item.get("name", "")).casefold() == config.source_name.casefold()
    ]
    if not exact:
        raise FileNotFoundError(
            f"Không tìm thấy {config.source_name!r} trong SharePoint drive mặc định."
        )

    expected_guid = _normalize_guid(config.source_sourcedoc)
    hydrated = [_item_metadata(graph, drive_id, item["id"]) for item in exact]
    if expected_guid:
        matched = [
            item
            for item in hydrated
            if _normalize_guid(
                item.get("sharepointIds", {}).get("listItemUniqueId")
            ) == expected_guid
        ]
        if len(matched) != 1:
            found = [
                item.get("sharepointIds", {}).get("listItemUniqueId")
                for item in hydrated
            ]
            raise RuntimeError(
                f"Không xác minh được sourcedoc {config.source_sourcedoc} "
                f"cho {config.source_name}. Các UniqueId tìm thấy: {found}"
            )
        return matched[0]

    if len(hydrated) != 1:
        raise RuntimeError(
            f"Có {len(hydrated)} file cùng tên {config.source_name!r}; "
            "cần sourcedoc để định danh."
        )
    return hydrated[0]


def _validate_target_identity(item: dict, config: OpenPOConfig) -> None:
    if str(item.get("name", "")) != config.target_name:
        raise RuntimeError(
            f"File đích không đúng tên: {item.get('name')!r}; "
            f"mong đợi {config.target_name!r}."
        )
    expected_guid = _normalize_guid(config.target_sourcedoc)
    actual_guid = _normalize_guid(
        item.get("sharepointIds", {}).get("listItemUniqueId")
    )
    if expected_guid and expected_guid != actual_guid:
        raise RuntimeError(
            f"Sai sourcedoc file đích: {actual_guid}; mong đợi {expected_guid}."
        )


def build_report(
    config: OpenPOConfig,
    source_info: dict,
    changes: list[dict],
    expected: dict[str, float],
) -> dict:
    return {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "source": {
            "name": config.source_name,
            "sourcedoc": config.source_sourcedoc,
            "sheet": config.source_sheet,
            **source_info,
        },
        "target": {
            "name": config.target_name,
            "path": config.target_path,
            "sheet": config.target_sheet,
            "column": "E",
            "codes": len(expected),
            "changed_cells": len(changes),
            "positive_codes": sum(1 for value in expected.values() if value > EPS),
            "zero_codes": sum(1 for value in expected.values() if abs(value) <= EPS),
            "total_open_po_for_target_codes": sum(expected.values()),
        },
        "rule": (
            "SUM by code of MAX(purchase_qty(I) - received_qty(K), 0) "
            "for rows with status(AB) blank"
        ),
        "changes_preview": changes[:100],
    }


def run_online(
    config: OpenPOConfig,
    *,
    publish: bool,
    proposal_path: str | Path = DEFAULT_PROPOSAL,
    report_path: str | Path = DEFAULT_REPORT,
) -> dict:
    token = get_access_token()
    graph = GraphClient(token)
    site_id = graph.get_site_id()
    drive_id = graph.get_default_drive_id(site_id)

    source_item_before = resolve_source_item(graph, drive_id, config)
    target_item_before = graph.get_item_by_path(drive_id, config.target_path)
    target_item_before = _item_metadata(graph, drive_id, target_item_before["id"])
    _validate_target_identity(target_item_before, config)

    source_bytes = graph.download_file(drive_id, source_item_before["id"])
    target_bytes = graph.download_file(drive_id, target_item_before["id"])

    source_after_download = _item_metadata(
        graph,
        drive_id,
        source_item_before["id"],
    )
    if source_after_download.get("eTag") != source_item_before.get("eTag"):
        raise RuntimeError(
            "File BCTheodoiDMHANG.xlsm thay đổi trong lúc download; hủy snapshot."
        )

    open_po, source_info = read_open_po(source_bytes, config)
    changes, expected = reconcile_target(target_bytes, open_po, config)
    patched_bytes = patch_target_workbook(target_bytes, changes, config)
    verify_target(patched_bytes, expected, config)

    Path(proposal_path).write_bytes(patched_bytes)
    report = build_report(config, source_info, changes, expected)
    report["source"]["revision"] = source_item_before.get("eTag")
    report["target"]["revision_before"] = target_item_before.get("eTag")
    report["target"]["sha256_before"] = hashlib.sha256(target_bytes).hexdigest()
    report["proposal"] = {
        "sha256": hashlib.sha256(patched_bytes).hexdigest(),
    }
    report["publish_requested"] = publish
    report["published"] = False
    report["publish_evidence"] = {
        "expected_etag": target_item_before.get("eTag"),
        "upload_skipped": not bool(changes),
        "upload_etag": None,
        "target_revision_after": None,
        "post_upload_verified": False,
        "server_sha256": None,
    }

    if publish and changes:
        source_pre_upload = _item_metadata(
            graph,
            drive_id,
            source_item_before["id"],
        )
        if source_pre_upload.get("eTag") != source_item_before.get("eTag"):
            raise RuntimeError(
                "Nguồn PO thay đổi trước upload; hủy publish để tránh snapshot cũ."
            )

        target_pre_upload = graph.get_item_by_path(drive_id, config.target_path)
        target_pre_upload = _item_metadata(
            graph,
            drive_id,
            target_pre_upload["id"],
        )
        _validate_target_identity(target_pre_upload, config)
        if target_pre_upload.get("eTag") != target_item_before.get("eTag"):
            raise RuntimeError(
                "File Kế hoạch mua hàng.xlsx thay đổi trước upload; "
                "hủy publish để tránh ghi đè."
            )

        upload_result = graph.upload_file(
            drive_id,
            target_item_before["id"],
            patched_bytes,
            expected_etag=target_item_before.get("eTag"),
        )
        server_bytes = graph.download_file(
            drive_id,
            target_item_before["id"],
        )
        verify_target(server_bytes, expected, config)
        report["published"] = True
        report["upload_name"] = upload_result.get("name", config.target_name)
        report["target"]["revision_after"] = upload_result.get("eTag")
        report["publish_evidence"] = {
            "expected_etag": target_item_before.get("eTag"),
            "upload_skipped": False,
            "upload_etag": upload_result.get("eTag"),
            "target_revision_after": upload_result.get("eTag"),
            "post_upload_verified": True,
            "server_sha256": hashlib.sha256(server_bytes).hexdigest(),
        }
    elif publish:
        report["published"] = True
        report["message"] = "Cột E đã đúng; không cần upload."
        report["target"]["revision_after"] = target_item_before.get("eTag")
        report["publish_evidence"] = {
            "expected_etag": target_item_before.get("eTag"),
            "upload_skipped": True,
            "upload_etag": None,
            "target_revision_after": target_item_before.get("eTag"),
            "post_upload_verified": True,
            "server_sha256": hashlib.sha256(target_bytes).hexdigest(),
        }

    Path(report_path).write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Đồng bộ Tồn Đơn hàng vào Ton_NVL!E"
    )
    parser.add_argument("--source-config", default=DEFAULT_SOURCE_CONFIG)
    parser.add_argument("--target-config", default=DEFAULT_TARGET_CONFIG)
    parser.add_argument("--publish", action="store_true")
    parser.add_argument("--proposal", default=DEFAULT_PROPOSAL)
    parser.add_argument("--report", default=DEFAULT_REPORT)
    args = parser.parse_args()

    config = load_open_po_config(args.source_config, args.target_config)
    report = run_online(
        config,
        publish=args.publish,
        proposal_path=args.proposal,
        report_path=args.report,
    )
    target = report["target"]
    print(
        f"[Ton_NVL!E] {target['codes']} mã, "
        f"{target['positive_codes']} mã có PO mở, "
        f"{target['changed_cells']} ô cần đổi, "
        f"tổng còn phải nhận={target['total_open_po_for_target_codes']}."
    )
    print("Publish:", "YES" if report.get("published") else "NO")


if __name__ == "__main__":
    main()
