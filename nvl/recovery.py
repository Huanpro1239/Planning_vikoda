"""Fail-closed NVL release recovery.

Recovery restores only Ton_NVL columns D/E from a verified historical release
into the current SharePoint workbook. It never replaces the whole workbook.
"""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO
import hashlib
import json
from pathlib import Path
from typing import Any
import zipfile

from lxml import etree
from openpyxl import load_workbook

from excel.openpyxl_io import safe_close_workbook
from excel.workbook_xml import column_number, find_sheet_xml_path
from nvl.config import load_nvl_config
from nvl.release_audit import audit_nvl_release_ledger
from nvl.values import normalize_nvl_code
from nvl.verify_release import verify_nvl_release
from nvl.workbook import identify_sharepoint_metadata_exemption
from sharepoint.client import (
    GraphClient,
    assert_same_revision,
    download_file_with_retry,
    get_access_token,
    get_item_metadata,
    validate_item_identity,
)


RECOVERY_REPORT = "nvl_recovery_report.json"
RECOVERY_PROPOSAL = "nvl_recovery_proposal.xlsx"
RECOVERY_BACKUP = "nvl_recovery_current_backup.xlsx"


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _load_json(path: Path, label: str) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        raise RuntimeError(f"Không đọc được {label}: {path}") from exc
    if not isinstance(value, dict):
        raise RuntimeError(f"{label} phải là JSON object: {path}")
    return value


def resolve_release_manifest(
    release: str | Path,
    releases_dir: str | Path,
) -> Path:
    candidate = Path(release)
    if candidate.is_file():
        return candidate
    release_id = str(release).strip()
    path = Path(releases_dir) / f"{release_id}.json"
    if not path.is_file():
        raise FileNotFoundError(
            f"Không tìm thấy NVL release manifest cho {release_id!r}: {path}"
        )
    return path


def _code_rows(workbook_bytes: bytes, config) -> dict[str, int]:
    wb = load_workbook(BytesIO(workbook_bytes), data_only=False, read_only=True)
    try:
        if config.target_sheet not in wb.sheetnames:
            raise RuntimeError(
                f"Không tìm thấy sheet {config.target_sheet!r}."
            )
        ws = wb[config.target_sheet]
        result: dict[str, int] = {}
        for row_num in range(config.target_start_row, ws.max_row + 1):
            code = normalize_nvl_code(
                ws.cell(row=row_num, column=config.target_code_col).value
            )
            if not code:
                continue
            if code in result:
                raise RuntimeError(
                    f"Mã NVL {code!r} bị lặp tại dòng "
                    f"{result[code]} và {row_num}."
                )
            result[code] = row_num
        if not result:
            raise RuntimeError(
                f"Không đọc được mã NVL trong {config.target_sheet}."
            )
        return result
    finally:
        safe_close_workbook(wb)


def _values_de(
    workbook_bytes: bytes,
    config,
    code_rows: dict[str, int],
) -> dict[str, dict[str, Any]]:
    wb = load_workbook(BytesIO(workbook_bytes), data_only=False, read_only=True)
    try:
        ws = wb[config.target_sheet]
        return {
            code: {
                "row": row,
                "D": ws.cell(row=row, column=4).value,
                "E": ws.cell(row=row, column=5).value,
            }
            for code, row in code_rows.items()
        }
    finally:
        safe_close_workbook(wb)


def _sheet_cell_map(root) -> tuple[dict[int, Any], dict[str, Any]]:
    rows: dict[int, Any] = {}
    cells: dict[str, Any] = {}
    for row in root.xpath('//*[local-name()="sheetData"]/*[local-name()="row"]'):
        raw = row.get("r")
        if not raw or not raw.isdigit():
            continue
        row_num = int(raw)
        rows[row_num] = row
        for cell in row.xpath('./*[local-name()="c"]'):
            ref = str(cell.get("r") or "")
            if ref:
                cells[ref] = cell
    return rows, cells


def _insert_cell_sorted(row, cell) -> None:
    ref = str(cell.get("r") or "")
    target_col = column_number(ref)
    for existing in row.xpath('./*[local-name()="c"]'):
        if column_number(existing.get("r", "")) > target_col:
            existing.addprevious(cell)
            return
    row.append(cell)


def build_recovery_proposal(
    current_bytes: bytes,
    historical_bytes: bytes,
    config,
) -> tuple[bytes, list[dict[str, Any]]]:
    """Restore only D/E cells from historical workbook into current bytes."""
    current_codes = _code_rows(current_bytes, config)
    historical_codes = _code_rows(historical_bytes, config)

    if current_codes != historical_codes:
        current_only = sorted(set(current_codes) - set(historical_codes))
        historical_only = sorted(set(historical_codes) - set(current_codes))
        moved = sorted(
            code
            for code in set(current_codes).intersection(historical_codes)
            if current_codes[code] != historical_codes[code]
        )
        raise RuntimeError(
            "Ton_NVL code/row layout đã thay đổi; recovery dừng fail-closed. "
            f"current_only={current_only[:10]}, "
            f"historical_only={historical_only[:10]}, moved={moved[:10]}"
        )

    before = _values_de(current_bytes, config, current_codes)
    desired = _values_de(historical_bytes, config, historical_codes)

    current_buf = BytesIO(current_bytes)
    historical_buf = BytesIO(historical_bytes)
    output_buf = BytesIO()

    with zipfile.ZipFile(current_buf, "r") as current_zip, zipfile.ZipFile(
        historical_buf, "r"
    ) as historical_zip:
        current_sheet = find_sheet_xml_path(
            current_zip, config.target_sheet
        )
        historical_sheet = find_sheet_xml_path(
            historical_zip, config.target_sheet
        )
        current_root = etree.fromstring(current_zip.read(current_sheet))
        historical_root = etree.fromstring(
            historical_zip.read(historical_sheet)
        )
        current_rows, current_cells = _sheet_cell_map(current_root)
        _, historical_cells = _sheet_cell_map(historical_root)

        for code, row_num in current_codes.items():
            row = current_rows.get(row_num)
            if row is None:
                raise RuntimeError(
                    f"Thiếu XML row {row_num} cho mã {code!r} trong current workbook."
                )
            for col in ("D", "E"):
                ref = f"{col}{row_num}"
                existing = current_cells.get(ref)
                historical_cell = historical_cells.get(ref)

                if existing is not None:
                    row.remove(existing)
                if historical_cell is not None:
                    replacement = deepcopy(historical_cell)
                    replacement.set("r", ref)
                    _insert_cell_sorted(row, replacement)

        new_sheet_xml = etree.tostring(
            current_root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        with zipfile.ZipFile(
            output_buf,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as output_zip:
            for item in current_zip.infolist():
                payload = (
                    new_sheet_xml
                    if item.filename == current_sheet
                    else current_zip.read(item.filename)
                )
                output_zip.writestr(item, payload)

    proposal = output_buf.getvalue()
    after = _values_de(proposal, config, current_codes)

    changes: list[dict[str, Any]] = []
    for code, row_num in current_codes.items():
        for col in ("D", "E"):
            if before[code][col] != after[code][col]:
                changes.append(
                    {
                        "code": code,
                        "row": row_num,
                        "column": col,
                        "before": before[code][col],
                        "after": after[code][col],
                    }
                )
    return proposal, changes


def verify_recovery_workbook(
    current_bytes: bytes,
    proposal_or_server_bytes: bytes,
    historical_bytes: bytes,
    config,
    *,
    is_server_comparison: bool = False,
) -> dict[str, Any]:
    """Verify D/E equal historical while everything else stays current."""
    current_codes = _code_rows(current_bytes, config)
    historical_codes = _code_rows(historical_bytes, config)
    candidate_codes = _code_rows(proposal_or_server_bytes, config)
    if current_codes != historical_codes or current_codes != candidate_codes:
        raise RuntimeError("Ton_NVL code/row layout không còn đồng nhất.")

    desired = _values_de(historical_bytes, config, historical_codes)
    actual = _values_de(proposal_or_server_bytes, config, candidate_codes)
    for code, row_num in current_codes.items():
        for col in ("D", "E"):
            if actual[code][col] != desired[code][col]:
                raise RuntimeError(
                    f"Recovery verify thất bại {code} {col}{row_num}: "
                    f"expected={desired[code][col]!r}, actual={actual[code][col]!r}"
                )

    exempted: list[dict[str, str]] = []
    with zipfile.ZipFile(BytesIO(current_bytes), "r") as current_zip, zipfile.ZipFile(
        BytesIO(proposal_or_server_bytes), "r"
    ) as candidate_zip:
        current_sheet = find_sheet_xml_path(
            current_zip, config.target_sheet
        )
        candidate_sheet = find_sheet_xml_path(
            candidate_zip, config.target_sheet
        )

        for name in current_zip.namelist():
            if name == current_sheet:
                continue
            if name not in candidate_zip.namelist():
                raise RuntimeError(
                    f"Recovery làm mất ZIP part không liên quan: {name}"
                )
            left = current_zip.read(name)
            right = candidate_zip.read(name)
            if left == right:
                continue
            if is_server_comparison:
                exemption = identify_sharepoint_metadata_exemption(
                    name,
                    left,
                    right,
                    current_zip,
                    candidate_zip,
                )
                if exemption is not None:
                    exempted.append(exemption)
                    continue
            raise RuntimeError(
                f"Recovery làm thay đổi ZIP part không liên quan: {name}"
            )

        for name in candidate_zip.namelist():
            if name in current_zip.namelist() or name == candidate_sheet:
                continue
            if is_server_comparison:
                exemption = identify_sharepoint_metadata_exemption(
                    name,
                    b"",
                    candidate_zip.read(name),
                    current_zip,
                    candidate_zip,
                )
                if exemption is not None:
                    exempted.append(exemption)
                    continue
            raise RuntimeError(
                f"Recovery sinh ZIP part lạ: {name}"
            )

        current_root = etree.fromstring(current_zip.read(current_sheet))
        candidate_root = etree.fromstring(candidate_zip.read(candidate_sheet))
        _, current_cells = _sheet_cell_map(current_root)
        _, candidate_cells = _sheet_cell_map(candidate_root)

        all_refs = set(current_cells).union(candidate_cells)
        for ref in sorted(all_refs):
            col = "".join(ch for ch in ref if ch.isalpha()).upper()
            if col in {"D", "E"}:
                continue
            left = current_cells.get(ref)
            right = candidate_cells.get(ref)
            if left is None or right is None:
                raise RuntimeError(
                    f"Ô ngoài D/E bị thêm/xóa trong recovery: {ref}"
                )
            if etree.tostring(left) != etree.tostring(right):
                raise RuntimeError(
                    f"Ô ngoài D/E bị thay đổi trong recovery: {ref}"
                )

    return {
        "ok": True,
        "exempted_server_metadata_parts": exempted,
    }


def _write_report(path: Path, report: dict[str, Any]) -> None:
    path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )


def run_recovery(
    release: str | Path,
    *,
    historical_workbook: str | Path,
    config_path: str | Path = "nvl_stock_config.json",
    releases_dir: str | Path = "runtime-state/nvl/releases",
    latest_path: str | Path | None = None,
    out_dir: str | Path = ".",
    publish: bool = False,
    approve_release_id: str | None = None,
    graph: GraphClient | None = None,
    verify_git: bool = False,
) -> dict[str, Any]:
    """Build or publish a controlled NVL recovery proposal."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    report_path = out_dir / RECOVERY_REPORT
    proposal_path = out_dir / RECOVERY_PROPOSAL
    backup_path = out_dir / RECOVERY_BACKUP

    manifest_path = resolve_release_manifest(release, releases_dir)
    manifest = _load_json(manifest_path, "NVL release manifest")
    release_id = str(manifest.get("release_id") or "").strip()
    if not release_id:
        raise RuntimeError("Release manifest thiếu release_id.")

    ledger = audit_nvl_release_ledger(
        releases_dir,
        latest_path=latest_path,
        verify_commit=verify_git,
    )
    if ledger.get("status") != "passed":
        raise RuntimeError(
            "NVL release ledger không PASS; từ chối recovery. "
            f"status={ledger.get('status')!r}"
        )
    ledger_ids = {
        row.get("release_id")
        for row in ledger.get("releases", [])
    }
    if release_id not in ledger_ids:
        raise RuntimeError(
            f"Release {release_id!r} không thuộc ledger đã audit."
        )

    historical_path = Path(historical_workbook)
    if not historical_path.is_file():
        raise FileNotFoundError(
            f"Không tìm thấy historical workbook: {historical_path}"
        )
    verify_nvl_release(
        manifest_path,
        final_workbook_path=historical_path,
        strict=False,
        verify_commit=verify_git,
    )
    historical_bytes = historical_path.read_bytes()

    config = load_nvl_config(config_path)
    if graph is None:
        graph = GraphClient(get_access_token())
    site_id = graph.get_site_id()
    drive_id = graph.get_default_drive_id(site_id)

    target_before = graph.get_item_by_path(
        drive_id,
        config.target_path,
    )
    target_before = validate_item_identity(
        target_before,
        expected_name=config.target_name,
        expected_sourcedoc=config.target_sourcedoc,
        expected_drive_id=drive_id,
        label="NVL recovery target",
    )
    current_bytes = download_file_with_retry(
        graph,
        drive_id,
        target_before["id"],
    )
    backup_path.write_bytes(current_bytes)

    proposal_bytes, changes = build_recovery_proposal(
        current_bytes,
        historical_bytes,
        config,
    )
    verify_recovery_workbook(
        current_bytes,
        proposal_bytes,
        historical_bytes,
        config,
    )
    proposal_path.write_bytes(proposal_bytes)

    report: dict[str, Any] = {
        "schema": "nvl_recovery_report_v1",
        "status": "proposal_ready",
        "mode": "publish" if publish else "dry_run",
        "release_id": release_id,
        "release_manifest": str(manifest_path),
        "target": {
            "name": target_before.get("name"),
            "path": config.target_path,
            "item_id": target_before.get("id"),
            "revision_before": target_before.get("eTag"),
            "sha256_before": _sha256(current_bytes),
        },
        "historical_workbook": {
            "path": str(historical_path),
            "sha256": _sha256(historical_bytes),
        },
        "proposal": {
            "path": str(proposal_path),
            "sha256": _sha256(proposal_bytes),
            "changed_cells": len(changes),
            "changed_D": sum(
                1 for item in changes if item["column"] == "D"
            ),
            "changed_E": sum(
                1 for item in changes if item["column"] == "E"
            ),
        },
        "changes": changes,
        "publish_evidence": {
            "approved_release_id": approve_release_id,
            "expected_etag": target_before.get("eTag"),
            "upload_result": None,
            "post_upload_verified": False,
            "server_sha256": None,
        },
    }
    _write_report(report_path, report)

    if not publish:
        return report

    if str(approve_release_id or "").strip() != release_id:
        raise RuntimeError(
            "Publish recovery yêu cầu --approve-release-id khớp chính xác "
            f"release {release_id!r}."
        )

    target_pre_upload = graph.get_item_by_path(
        drive_id,
        config.target_path,
    )
    target_pre_upload = validate_item_identity(
        target_pre_upload,
        expected_name=config.target_name,
        expected_sourcedoc=config.target_sourcedoc,
        expected_drive_id=drive_id,
        label="NVL recovery target before upload",
    )
    assert_same_revision(
        target_before,
        target_pre_upload,
        label="NVL recovery target",
    )

    if not changes:
        report["status"] = "published_no_change"
        report["publish_evidence"]["post_upload_verified"] = True
        report["publish_evidence"]["server_sha256"] = _sha256(
            current_bytes
        )
        _write_report(report_path, report)
        return report

    upload_result = graph.upload_file(
        drive_id,
        target_before["id"],
        proposal_bytes,
        expected_etag=target_before.get("eTag"),
    )
    server_bytes = download_file_with_retry(
        graph,
        drive_id,
        target_before["id"],
    )
    verify_result = verify_recovery_workbook(
        current_bytes,
        server_bytes,
        historical_bytes,
        config,
        is_server_comparison=True,
    )

    report["status"] = "published"
    report["publish_evidence"] = {
        "approved_release_id": approve_release_id,
        "expected_etag": target_before.get("eTag"),
        "upload_result": upload_result,
        "post_upload_verified": True,
        "server_sha256": _sha256(server_bytes),
        "exempted_server_metadata_parts": verify_result.get(
            "exempted_server_metadata_parts", []
        ),
    }
    _write_report(report_path, report)
    return report
