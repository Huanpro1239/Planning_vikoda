from io import BytesIO

from openpyxl import load_workbook

import sync_planning_metrics as metrics
from sync_stock_compat import read_conversion_factors_robust


def read_planning_rows_robust(dest_bytes):
    workbook = load_workbook(
        BytesIO(dest_bytes),
        data_only=True,
        read_only=True,
    )

    try:
        if metrics.PLANNING_SHEET not in workbook.sheetnames:
            raise RuntimeError(
                f"Không tìm thấy sheet {metrics.PLANNING_SHEET!r}."
            )
        if metrics.FC_SHEET not in workbook.sheetnames:
            raise RuntimeError(
                f"Không tìm thấy sheet {metrics.FC_SHEET!r}."
            )

        worksheet = workbook[metrics.PLANNING_SHEET]
        selector = workbook[metrics.FC_SHEET][metrics.FC_SELECTOR_CELL].value
        plan_month = metrics._parse_plan_month(selector)
        rows = {}

        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, max_col=18, values_only=True),
            start=2,
        ):
            code = metrics.normalize_code(values[0] if values else None)
            if not code:
                continue

            if code in rows:
                raise RuntimeError(
                    f"Mã {code} bị lặp trong {metrics.PLANNING_SHEET}!A."
                )

            rows[code] = {
                "row": row_number,
                "batch": metrics.to_number(
                    values[3], f"{metrics.PLANNING_SHEET}!D{row_number}"
                ),
                "per_shift": metrics.to_number(
                    values[4], f"{metrics.PLANNING_SHEET}!E{row_number}"
                ),
                "classification": str(values[7] or "").strip(),
                "shifts_per_day": metrics.to_number(
                    values[8], f"{metrics.PLANNING_SHEET}!I{row_number}"
                ),
                "actual_stock": metrics.to_number(
                    values[9], f"{metrics.PLANNING_SHEET}!J{row_number}"
                ),
                "book_stock": metrics.to_number(
                    values[10], f"{metrics.PLANNING_SHEET}!K{row_number}"
                ),
                "fc": metrics.to_number(
                    values[11], f"{metrics.PLANNING_SHEET}!L{row_number}"
                ),
                "current_debt": metrics.to_number(
                    values[13], f"{metrics.PLANNING_SHEET}!N{row_number}"
                ),
                "current_outputs": {
                    "expected_end_stock": values[12],
                    "warehouse_debt": values[13],
                    "required_production": values[14],
                    "rounded_production": values[15],
                    "production_days": values[16],
                    "production_start": values[17],
                },
            }

        if not rows:
            raise RuntimeError(
                f"Không đọc được dữ liệu trong {metrics.PLANNING_SHEET}."
            )

        print(
            f"[{metrics.PLANNING_SHEET}] Đọc {len(rows)} dòng bằng "
            "chế độ tương thích sheet không có dimension."
        )
        return selector, plan_month, rows
    finally:
        workbook.close()


metrics.read_planning_rows = read_planning_rows_robust
metrics.read_conversion_factors = read_conversion_factors_robust


if __name__ == "__main__":
    metrics.main()
