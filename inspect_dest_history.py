from io import BytesIO

from openpyxl import load_workbook

from sync_stock import DEST_PATH, GRAPH, GraphClient, get_access_token


TARGET_SHEET = "Ke_hoach_SX"
TARGET_MIN_COL = 13  # M
TARGET_MAX_COL = 18  # R


def main():
    token = get_access_token()
    graph = GraphClient(token)
    site_id = graph.get_site_id()
    drive_id = graph.get_default_drive_id(site_id)
    item = graph.get_item_by_path(drive_id, DEST_PATH)

    versions_url = f"{GRAPH}/drives/{drive_id}/items/{item['id']}/versions"
    versions = graph.get_json(versions_url).get("value", [])

    print(f"Destination: {DEST_PATH}")
    print(f"Found {len(versions)} versions")

    for version in versions:
        version_id = version["id"]
        content_url = (
            f"{GRAPH}/drives/{drive_id}/items/{item['id']}"
            f"/versions/{version_id}/content"
        )
        response = graph.session.get(
            content_url,
            timeout=120,
            allow_redirects=True,
        )
        graph._raise(response)

        try:
            workbook = load_workbook(
                BytesIO(response.content),
                data_only=False,
                read_only=True,
            )
        except Exception as exc:
            print(f"VERSION {version_id}: cannot open: {exc}")
            continue

        try:
            sheet_names = workbook.sheetnames
            target_formulas = []
            if TARGET_SHEET in sheet_names:
                ws = workbook[TARGET_SHEET]
                for row in range(1, ws.max_row + 1):
                    for col in range(TARGET_MIN_COL, TARGET_MAX_COL + 1):
                        value = ws.cell(row=row, column=col).value
                        if isinstance(value, str) and value.startswith("="):
                            target_formulas.append(
                                f"{ws.cell(row=row, column=col).coordinate}:{value}"
                            )

            weekly_present = "Ke hoach SX tuan" in sheet_names
            print(
                f"VERSION {version_id} | modified={version.get('lastModifiedDateTime')} "
                f"| sheets={sheet_names} | weekly={weekly_present} "
                f"| M:R formulas={len(target_formulas)}"
            )
            for formula in target_formulas[:18]:
                print(f"  {formula}")
        finally:
            workbook.close()


if __name__ == "__main__":
    main()
