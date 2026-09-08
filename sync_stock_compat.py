import hashlib
import json
from io import BytesIO

from openpyxl import load_workbook

import sync_planning_weekly_model as weekly_model
import sync_stock


def read_conversion_factors_robust(dest_bytes):
    workbook = load_workbook(
        BytesIO(dest_bytes),
        data_only=True,
        read_only=True,
    )

    try:
        if sync_stock.MASTER_SHEET not in workbook.sheetnames:
            raise RuntimeError(
                f"Không tìm thấy sheet {sync_stock.MASTER_SHEET!r} trong file đích."
            )

        worksheet = workbook[sync_stock.MASTER_SHEET]
        headers = [str(c or "").strip() for c in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())]
        debt_col = weekly_model.find_header_col(headers, weekly_model.DEBT_HEADER_NAMES)
        profile_col = weekly_model.find_header_col(headers, weekly_model.PROFILE_HEADER_NAMES)

        factors = {}
        master_meta = {}

        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if not values:
                continue
            code = sync_stock.normalize_code(values[0])
            if not code:
                continue

            if code in factors:
                raise RuntimeError(
                    f"[{sync_stock.MASTER_SHEET}] Mã {code} bị lặp trong cột A."
                )

            factor = sync_stock.to_number(
                values[8] if len(values) >= 9 else None,
                f"{sync_stock.MASTER_SHEET}!I{row_number}",
            )
            if factor <= 0:
                raise RuntimeError(
                    f"[{sync_stock.MASTER_SHEET}] Quy cách của mã {code} "
                    f"phải > 0, hiện là {factor!r}."
                )

            factors[code] = factor
            raw_debt_mode = values[debt_col] if debt_col is not None and debt_col < len(values) else None
            raw_profile = values[profile_col] if profile_col is not None and profile_col < len(values) else None
            master_meta[code] = {
                "batch": sync_stock.to_number(values[3] if len(values) >= 4 else None, default=0.0),
                "per_shift": sync_stock.to_number(values[4] if len(values) >= 5 else None, default=0.0),
                "line": str(values[5] or "").strip() if len(values) >= 6 else "",
                "group": str(values[6] or "").strip() if len(values) >= 7 else "",
                "classification": str(values[7] or "").strip() if len(values) >= 8 else "",
                "mold": factor,
                "leadtime": sync_stock.to_number(values[9] if len(values) >= 10 else None, default=0.0),
                "debt_mode": weekly_model.normalize_debt_mode(raw_debt_mode) or "",
                "profile": weekly_model.normalize_profile(raw_profile) or "",
            }

        if not factors:
            raise RuntimeError(
                f"[{sync_stock.MASTER_SHEET}] Không đọc được mã/quy cách từ A:I."
            )

        payload = json.dumps(
            {k: master_meta[k] for k in sorted(master_meta)},
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
        conversion_hash = hashlib.sha256(payload).hexdigest()

        print(
            f"[{sync_stock.MASTER_SHEET}] Đọc {len(factors)} quy cách "
            "bằng chế độ tương thích sheet không có dimension."
        )
        return factors, conversion_hash
    finally:
        workbook.close()


sync_stock.read_conversion_factors = read_conversion_factors_robust


if __name__ == "__main__":
    sync_stock.main()
